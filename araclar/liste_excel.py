"""İE-13 PR-2 · Fiyatlı listenin Excel sayfası: teklifin (`araclar/teklif.py`) ve `fiyatla.py --excel`'in tek yazıcısı.

Düzen Patron'un yükleme listesinin düzeni (28 Ekim 06:40–06:41): grup adı sütun değil, grubun üstünde başlık satırı ve altında ara toplam;
tek "Adet" sütunu; ölçüsü birim sözcüğü olan satırda (adet, metre, kova, paket, takım) ölçü boş, birim o. Genel toplam ve fiyatı olmayan
satır sayısı `fiyatla.fiyatla`'nın özetinden; sayfa saymaz, yazar. Motor'u yüklemez (yalnız openpyxl).
"""
from __future__ import annotations

BIRIM = {"m": "metre"}
BIRIM_SOZCUK = {"adet", "metre", "kova", "paket", "takım", "m²"}    # Motor'un ölçü yerine yazdığı birimler
TL = "#,##0.00"
RENK = {"sari": "FFF2A8", "gri": "E7E7E7", "baslik": "D9E1F2", "cizgi": "8EA9DB",
        "öneri onaylı": "FCE4D6", "ikinci kaynak": "DDEBF7", "öneri": "FFF2A8"}
GENEL = "GENEL TOPLAM (fiyatı olan satırlar)"


def buyuk(s: str) -> str:
    """Türkçe büyük harf: "Elektrik" → "ELEKTRİK", "Sıhhi" → "SIHHİ" (str.upper i'yi I yapar)."""
    return s.replace("i", "İ").replace("ı", "I").upper()


def gorunen(olcu: str | None, birim: str | None) -> tuple[str, str]:
    """(ölçü, birim): ölçü bir birim sözcüğüyse ölçü boş kalır, birim o olur; öteki ölçüler ("—", "roll form çelik") yerinde kalır."""
    olcu = olcu or ""
    if olcu in BIRIM_SOZCUK:
        return "", olcu
    return olcu, BIRIM.get(birim, birim or "")


def eksik(s: dict) -> str | None:
    """Satırın miktarı yoksa nedeni (Not sütunu, gri satır, Özet'in Neden'i hep buradan): adet ya da kg'lı kalemde kg."""
    if s.get("birim") == "kg":
        return "kg yok" if s.get("kg") is None else None
    return "Motor'da kuralı yok, adet yok" if s.get("adet") is None else None


def _stil():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    dolgu = {k: PatternFill("solid", fgColor=v) for k, v in RENK.items()}
    return Alignment, Border, Font, dolgu, Side


def liste_sayfasi(ws, R: dict, baslik: str, aciklama: str, renk=None, kaynak_sutunu: bool = False) -> None:
    """`fiyatla.fiyatla`'nın sonucu (R) bir sayfaya: grup başlık satırı + satırlar + ara toplam; en altta genel toplam ve fiyatı olmayan
    satır sayısı (R'nin özetinden). `renk(satır)` birim fiyat hücresinin dolgusunu (RENK anahtarı) verir. `kaynak_sutunu`: fiyatla.py'nin
    defter görünümü için Durum ve Kaynak sütunları. Ara toplamların toplamı özetin toplamını tutmazsa durur."""
    Alignment, Border, Font, dolgu, Side = _stil()
    satirlar, ozet = R["satirlar"], R["ozet"]
    kg = any(s.get("kg") is not None for s in satirlar)
    sutun = ["Ölçü", "Birim", "Adet"] + (["kg"] if kg else []) + ["Birim fiyat", "Tutar", "Not"] + (["Durum", "Kaynak"] if kaynak_sutunu else [])
    n = 1 + len(sutun)
    i_bf, i_tut, i_not = (sutun.index(x) + 2 for x in ("Birim fiyat", "Tutar", "Not"))    # sayfadaki sütun (1'den; A Kalem)
    ws.append([baslik]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([aciklama]); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n); ws.row_dimensions[2].height = 32
    gruplar: dict[str, list[dict]] = {}                     # grup ilk göründüğü sırada, grup içinde Motor'un sırası
    for s in satirlar:
        gruplar.setdefault(s.get("grup_adi") or str(s.get("grup") or ""), []).append(s)
    toplam = 0.0
    for ad, ss in gruplar.items():
        ws.append([]); ws.append([buyuk(ad), *sutun])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True); c.fill = dolgu["baslik"]
        ara = round(sum(s["tutar"] or 0 for s in ss), 2)
        for s in ss:
            olcu, birim = gorunen(s.get("olcu"), s.get("birim"))
            neden = eksik(s)
            notu = neden or ("fiyat yok" if s["birim_fiyat"] is None else "")
            ws.append([s["kalem"], olcu, birim, s.get("adet")] + ([s.get("kg")] if kg else []) + [s["birim_fiyat"], s["tutar"], notu]
                      + ([s["durum"], s.get("kaynak_adi") or s.get("kaynak")] if kaynak_sutunu else []))
            r = ws[ws.max_row]
            r[i_bf - 1].number_format = r[i_tut - 1].number_format = TL
            k = renk(s) if renk else None
            if neden:
                for c in r: c.fill = dolgu["gri"]
            elif k:
                r[i_bf - 1].fill = dolgu[k]
        toplam += ara
        ws.append([f"{ad} toplamı"] + [""] * (i_tut - 2) + [ara])
        r = ws[ws.max_row]; r[0].font = r[i_tut - 1].font = Font(bold=True); r[i_tut - 1].number_format = TL
        for c in r[:n]: c.border = Border(top=Side(style="thin", color=RENK["cizgi"]))
    if abs(toplam - ozet["toplam"]) > 0.005:
        raise ValueError(f"liste sayfası: ara toplamların toplamı {toplam:.2f}, özetin toplamı {ozet['toplam']:.2f}")
    ws.append([])
    ws.append([GENEL] + [""] * (i_tut - 2) + [ozet["toplam"], f"fiyatı olmayan satır: {len(ozet['fiyatsiz'])}"])
    r = ws[ws.max_row]; r[0].font = r[i_tut - 1].font = Font(bold=True, size=12); r[i_tut - 1].number_format = TL
    for col, w in zip("ABCDEFGHIJ", [62, 20, 9, 8] + ([9] if kg else []) + [13, 14, 30] + ([16, 30] if kaynak_sutunu else [])):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if c.column in (3, 4): c.alignment = Alignment(horizontal="center")
            elif c.column >= i_not: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

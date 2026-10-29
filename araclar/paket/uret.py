"""İE-18 · Her model için düzenli paket: çizim, liste, kimlik kartlı model, montaj paftası, rehber, kıyas ve `bilgi.md`.

    python -m araclar.paket <KOD|hepsi> <çıktı klasörü> [--glb] [--gecici <klasör>]
    python -m araclar.paket <KOD> <çıktı klasörü> --is-emri <İş Emri.json> [--tavan asma,ters] [--glb] [--gecici <klasör>]

Çıktı klasörü PF `Portman_Proje_Dosyalari/Modeller/` düzenindedir (`OKU_BENI.md`, CIO onayı 27 Ekim 2026 22:00 UTC):
`prefabrik/<KOD>/<KOD>_<kalem>[_<seçenek>].<uzantı>` ve `hepsi` ile kökte `ozet.xlsx`. Üreteç canlıdaki işlevleri çağırır, kendi çizim
kodu yoktur: `motor.dxf_out.uret` (ön onay PDF'i ve DXF), `motor.bayi_gorseli.plan_gorseli` (WhatsApp plan görseli), liste üreteci
(`motor.prefabrik_liste`), kart modeli (`motor.kart_model.uret`: IFC ve yerleşim planları). Rehber 62 m² rehberinin şablonundan
(D-28; `araclar/paket/rehber/`: motorun planı, kart modeli ve liste → şablonun verisi), kıyas depodaki hakem verisinden (hakemin pafta
katmanı ve kör test sonucu; gerçek liste açılmaz) ve kart başına model–liste tablosundan.

- Motor bir şeyi yapamıyorsa (doğrulama freni, beşik dışı çatıda liste ya da kart modeli, rehberin kuramadığı ev) dosya yazılmaz;
  `bilgi.md`'de "Yok: neden" yazar. Her dosya için Var ya da Yok. Doğrulama geçmezse (K23) ya da Motor çizemezse teslim çıktısı
  (bayi PDF'i, plan görseli, liste, model, montaj paftası, rehber) hiç yazılmaz; DXF denetim için yazılır.
- Kör model: kıyas yok (gerçek liste kör protokolüyle açılmadan girmez). Kaynak dosyalar (paftalar, DWG, gerçek liste) kopyalanmaz;
  `bilgi.md` yalnız mühürdeki kimliklerini (bayt, SHA-256, Drive kimliği) taşır. Kişi ve bayi adı yazılmaz.
- Yeniden üretimde klasördeki eski dosyalar `_eski/<eski üretimin tarihi>/`'ne taşınır; hiçbir şey silinmez.
- GLB (`--glb`): `rehber/glb.js` Playwright ile; three.js 0.160.0 `THREE_DIR` ortam değişkeninin gösterdiği klasörden (npm `three`).
- Kayıtta olmayan ev (`--is-emri`; çizim masasının yeni talepleri, 28 Ekim 2026): model satırı İş Emri'nden kurulur, küme "teklif";
  gövde, m² ve verandalar Motor'un planından (`kimlik`), çatı İş Emri'nden. Kayıttaki bir teklif projesine (TK-01) yeni revizyon da
  verilebilir, kimliği yine o İş Emri'nden; kayıttaki öteki modellerin İş Emri'si kayıttan gelir (kıyasları ona bağlı). Tavan
  seçenekleri teklifle aynı işlevden (`araclar/teklif.secenekler`): `--tavan`, yoksa İş Emri'nin `teklif.tavan`'ı, yoksa kayıttaki
  satırınki, yoksa İş Emri'nin kendi tavanı. Çıktıdan önce İş Emri'nin kopyasında kişi ve bayi adı yerine kod yazılır,
  metinlerdeki telefon numaraları maskelenir (`beyin/numara.py`); serbest metindeki adlar tanınmaz, İş Emri'ne ad yazılmaz.

LLM yok, ağ yok (GLB'de three.js yerelden).
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

KOK = Path(__file__).resolve().parents[2]
if str(KOK) not in sys.path:
    sys.path.insert(0, str(KOK))

KAYIT = Path(__file__).with_name("modeller.json")
MUHUR = KOK / "docs" / "notlar" / "2026-10-18_AIPrefabrik_muhur.md"
REHBER = Path(__file__).with_name("rehber")
KALEMLER = [("bilgi", "md", "kimlik, kaynaklar ve dosyaların durumu"), ("bayi_paftasi", "pdf", "canlıda bayiye giden ön onay PDF'i"),
            ("plan", "png", "WhatsApp plan görseli"), ("gorunusler", "pdf", "dört görünüş"), ("cizim", "dxf", "düzenlenebilir çizim"),
            ("montaj_paftasi", "pdf", "yerleşim planları (kart modeli)"), ("yukleme_listesi", "xlsx", "liste üreteci"),
            ("model", "ifc", "kimlik kartlı model"), ("montaj_rehberi", "html", "adım adım 3B rehber"), ("model", "glb", "rehberin 3B modeli"),
            ("kiyas", "xlsx", "mühendisin paftası ve gerçek listesiyle kıyas")]




def kayit() -> dict:
    return json.loads(KAYIT.read_text(encoding="utf-8"))


def muhur(kod: str) -> dict:
    """Mühür notundaki satırlar: dosyalar (tür, uzantı, bayt, SHA-256) ve DWG (bayt, Drive kimliği)."""
    dosya, dwg = [], []
    if not kod.startswith("AP-"):
        return {"dosya": dosya, "dwg": dwg}
    for satir in MUHUR.read_text(encoding="utf-8").splitlines():
        c = [x.strip() for x in satir.strip().strip("|").split("|")]
        if not c or c[0] != kod:
            continue
        if len(c) == 7 and c[1].isdigit():
            dosya.append({"tur": c[2], "uzanti": c[3], "bayt": c[4], "sha256": c[5].strip("`")})
        elif len(c) == 6 and "DWG" in c[1]:
            dwg.append({"tur": c[1], "bayt": c[2], "yukleme": c[3], "drive": c[4].strip("`"), "sha256": c[5]})
    return {"dosya": dosya, "dwg": dwg}


def depo_durumu() -> dict:
    def git(*a):
        r = subprocess.run(["git", "-C", str(KOK), *a], capture_output=True, text=True)
        return r.stdout.strip()
    return {"commit": git("rev-parse", "--short", "HEAD"), "dal": git("rev-parse", "--abbrev-ref", "HEAD"),
            "temiz": not git("status", "--porcelain", "--", "motor", "kurallar", "araclar/paket")}


KOD_DESENI = re.compile(r"[A-Z]{2}-\d\d")            # kayıttaki kodlar gibi (AP-19, TK-01); dosya adları ve kart kodu bundan


def _maskele(ie: dict, kod: str) -> dict:
    """Dışarı giden çıktılar için İş Emri kopyası: talepte kişi ve bayi adı yerine kod (antet, IFC), bütün metinlerde telefon numarası
    yerine "…" (AGENTS.md: kişisel veri maskelenir, bayi adları dahil). Kayıttaki İş Emri'leri zaten böyle yazılmıştır."""
    from beyin import numara

    def m(o):
        if isinstance(o, str):
            return numara.maskele(o)
        if isinstance(o, list):
            return [m(x) for x in o]
        if isinstance(o, dict):
            return {k: m(v) for k, v in o.items()}
        return o
    ie = m(copy.deepcopy(ie))
    ie["talep"] = {"bayi": {"kod": kod.replace("-", ""), "ad": f"Bayi {kod}"}, "musteri_adi": kod}   # talep metni çıktıya girmez
    return ie


def kimlik(ie: dict) -> dict:
    """Model satırının kimlik alanları İş Emri'nden ve Motor'un planından (D-28 madde 1: ölçü ve m² Motor'un; `dis_olcu_cm` isteğe
    bağlı, ızgara İş Emri'sinde yok): gövde kutusu (`geometry.govde_kutusu`, taşan veranda hariç), taşan veranda (cephe, en, derinlik),
    gömme açık mahaller (m²), çatı (tip, eğim, kaplama). Motor kuramazsa gövde yerine nedeni yazar; paket Motor'un hatasını ayrıca yazar."""
    from motor import geometry as G
    from motor.cati_prefabrik import TIP_ADI
    c, oz = ie.get("cati") or {}, ie.get("ozellikler") or {}
    kap = ", ".join(str(oz[k]).lower() for k in ("Çatı Kaplaması", "Çatı Kaplama Rengi") if oz.get(k))
    kap = kap.replace(", ", " (", 1) + ")" if ", " in kap else kap
    tip_adi = {"besik": "beşik"} | TIP_ADI
    tip, eg = c.get("tip", "besik"), c.get("egim_yuzde")                  # Motor'un varsayılanı beşik, eğim kitaptan
    out = {"cati": f"{tip_adi.get(tip, tip)}" + (f" %{eg:g}" if eg is not None else " (eğim kitaptan)") + (f", {kap}" if kap else "")}
    try:
        plan = G.kur(copy.deepcopy(ie))
    except Exception as e:                                              # noqa: BLE001 — neden satıra yazılır
        return out | {"m2": "—", "govde": f"Motor İş Emri'ni kuramadı ({type(e).__name__}: {str(e)[:120]})", "veranda": "—"}
    x0, y0, x1, y1 = G.govde_kutusu(plan)
    gx, gy = round(x1 - x0, 2), round(y1 - y0, 2)
    yon = {"on": "önde", "arka": "arkada", "sol": "solda", "sag": "sağda"}
    tasan = [f"{yon.get(v['cephe'], v['cephe'])} {v['en']:g} × {v['derinlik']:g}" for v in plan.dis_verandalar]
    dis = {v["ad"] for v in plan.dis_verandalar}
    gomme = [f"{m.alan_m2:.1f}".replace(".", ",") for m in plan.mahaller.values() if m.acik and m.ad not in dis]
    ver = (["taşan: " + ", ".join(tasan)] if tasan else []) + (["gömme " + ", ".join(gomme) + " m²"] if gomme else [])
    return out | {"m2": f"{gx * gy / 1e4:.0f}", "govde": f"dış ölçü {gx:g} × {gy:g}", "veranda": "; ".join(ver) or "yok"}


def talep_modeli(kod: str, yol: Path, tavan: list[str] | None, kayit_: dict) -> tuple[dict, dict]:
    """`--is-emri`: İş Emri dosyasından model satırı ve maskeli İş Emri. Kimlik (gövde, m², veranda, çatı) her zaman bu İş Emri'nden
    (`kimlik`); kayıttaki teklif projesinden yalnız takma ad ve, İş Emri'nde ya da `--tavan`'da yoksa, tavan seçenekleri gelir.
    Dönen: (model, İş Emri)."""
    import hashlib
    from araclar import teklif
    if not KOD_DESENI.fullmatch(kod):
        raise SystemExit(f"DUR: kod {kod!r}; iki büyük harf, tire, iki rakam (TK-02 gibi)")
    yol = Path(yol)
    ham = yol.read_bytes()
    ie = json.loads(ham.decode("utf-8"))
    kayitta = next((x for x in kayit_["modeller"] if x["kod"] == kod), None)
    if kayitta is not None and kayitta["kume"] != "teklif":
        raise SystemExit(f"DUR: {kod} kayıtta {kayitta['kume']} model; İş Emri'si ve kıyası kayıttan gelir (`--is-emri` olmadan koşun)")
    secim = [o["Tavan"].lower() for _, o in teklif.secenekler(ie, tavan) if o] or (kayitta or {}).get("tavan_secenekleri")
    try:
        goster = str(yol.resolve().relative_to(KOK))
    except ValueError:
        goster = str(yol)
    komut = f"python -m araclar.paket {kod} <Modeller klasörü> --is-emri {goster}" + (f" --tavan {','.join(tavan)}" if tavan else "")
    m = {"kod": kod, "takma": (kayitta or {}).get("takma"), "kume": "teklif", **kimlik(ie), "pencere_ic_duvar": "—", "engel": "—",
         "adim": "teklif", "kor_test": None, "hakem": None,
         "is_emri": goster, "tavan_secenekleri": secim, "komut": komut, "girdi_sha256": hashlib.sha256(ham).hexdigest()}
    return m, _maskele(ie, kod)


# ----------------------------------------------------------------------------- parçalar
def cizim(ie: dict, gec: Path) -> dict:
    """Ön onay PDF'i, DXF ve plan görseli (canlıdaki işlevler). Dönen: {'pdf','dxf','png','gorunusler','neden'}."""
    from motor import bayi_gorseli, dxf_out
    m = dxf_out.uret(copy.deepcopy(ie), gec / "motor")
    out = {"dxf": gec / "motor" / m["dxf"], "pdf": gec / "motor" / m["pdf"] if m.get("pdf") else None, "uyarilar": m.get("warnings") or []}
    if out["pdf"] is None:
        out["neden"] = f"Motor'un doğrulaması geçmedi; teslim PDF'i üretilmedi ({m.get('pdf_skipped_reason')})"
    v = json.loads((gec / "motor" / "validation.json").read_text(encoding="utf-8")) if (gec / "motor" / "validation.json").exists() else None
    out["png"] = bayi_gorseli.plan_gorseli(copy.deepcopy(ie), gec / "plan.png", v)
    import ezdxf
    d = ezdxf.readfile(out["dxf"])
    yazi = [e.dxf.text if e.dxftype() == "TEXT" else e.text for e in d.modelspace() if e.dxftype() in ("TEXT", "MTEXT")]
    out["gorunusler"] = [re.sub(r" GÖRÜNÜŞ.*", "", t) for t in yazi if "GÖRÜNÜŞ" in t]
    return out


def listeler(ie: dict, secenekler: list[str] | None) -> dict:
    """Liste üreteci (canlıdaki kurallar), komut satırı ve teklifle aynı yoldan (`prefabrik_liste.is_emrinden`, D-28 madde 1; kırma ve cumba
    çatının freni orada). Dönen: {seçenek ya da '': L} ya da {'neden': ...}."""
    from motor import prefabrik_liste as PL
    try:
        return {s: PL.is_emrinden(copy.deepcopy(ie), ozellikler={"Tavan": s.capitalize()} if s else None) for s in secenekler or [""]}
    except PL.ListeYazilmaz as e:
        return {"neden": e.neden}


def kart_modeli(ie: dict, gec: Path, kart_kodu: str) -> dict:
    from motor import kart_model
    from motor.cati_prefabrik import TIP_ADI
    tip = (ie.get("cati") or {}).get("tip", "besik")
    if tip != "besik":
        return {"neden": f"kart modeli {TIP_ADI.get(tip, tip)} çatıyı kurmuyor (yerleşimi yalnız düz beşik; İE-15 K5)"}
    say = kart_model.uret(copy.deepcopy(ie), gec / "kart", kart_kodu)
    return {"klasor": gec / "kart", "ifc": gec / "kart" / f"{kart_kodu}_model.ifc", "pdf": gec / "kart" / f"{kart_kodu}_yerlesim_planlari.pdf",
            "kural_yok": say["kural_yok"], "sema_hatasi": say["sema_hatasi"], "uyarilar": say.get("yerlesim", {}).get("uyarilar", [])}


def rehber(km: dict, kart_kodu: str, ie: dict, L: dict, kod: str, gec: Path, glb: bool) -> dict:
    from araclar.paket.rehber import sayfa, veri
    try:
        d = veri.kur(copy.deepcopy(ie), km["klasor"], kart_kodu, L["satirlar"], kod)
    except veri.Desteklenmiyor as e:
        return {"neden": f"rehber üreteci bu evi kurmuyor: {e}"}
    html = sayfa.yaz(d, gec / "rehber.html")
    out = {"html": html, "veri": d, "kiyas": veri.kiyas(km["klasor"], kart_kodu, L["satirlar"])}
    if not glb:
        out["glb_neden"] = "GLB bu koşuda istenmedi (`--glb` yok)"
    elif not os.environ.get("THREE_DIR"):
        out["glb_neden"] = "GLB yazılmadı: THREE_DIR (three.js 0.160.0) bu ortamda yok"
    else:
        r = subprocess.run(["node", str(REHBER / "glb.js"), str(html.resolve()), str((gec / "model.glb").resolve())],
                           capture_output=True, text=True, timeout=600)
        if r.returncode == 0 and (gec / "model.glb").exists():
            out["glb"] = gec / "model.glb"
        else:
            out["glb_neden"] = f"GLB yazılamadı: {(r.stderr or r.stdout).strip()[-200:]}"
    return out


def kiyas(m: dict, km: dict | None, rh: dict | None, yol: Path) -> str | None:
    """Kıyas Excel'i (açık model, depoda hakem verisi varsa). Dönen: None (yazıldı) ya da Yok nedeni."""
    if m["kume"] == "kör":
        return "kör model; gerçek liste kör protokolüyle (bizim liste commit'li, Patron'un açılış kaydı) açılmadan kıyas yok"
    if m["kume"] == "teklif":
        return "teklif projesi; mühendis paftası ve gerçek yükleme listesi yok (modelin adetleri rehberde Motor'un listesiyle kıyaslı)"
    if m["kume"] == "kapalı":
        return "listesi kapalı model"
    hk, kt = m.get("hakem"), m.get("kor_test")
    if not (hk or kt):
        return "depoda bu evin hakem verisi (pafta katmanı) ve kör test sonucu yok"
    if km is None or "ifc" not in km:
        return "kart modeli yok; paftayla kıyas modelin adetleriyle yapılır"
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    sys.path.insert(0, str(KOK / "testler" / "kimlik" / "hakem"))
    import kiyas as KY
    wb = Workbook()
    bas = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="2F3B4C")) for c in ws[2]]
    renk = {"geçer": "E2EFDA", "geçmez": "FCE4D6", "aynı": "E2EFDA", "farklı": "FCE4D6"}
    ws = wb.active; ws.title = "Pafta"
    if hk:
        cevap = json.loads((KOK / hk["cevap"]).read_text(encoding="utf-8"))
        proje = json.loads((KOK / hk["proje"]).read_text(encoding="utf-8"))
        satirlar, gorulen, metre = [], set(), []
        for adim in [None] + sorted(proje.get("kiyas", {}).get("adimlar", {})):
            k = KY.kiyasla(km["ifc"], cevap, proje, adim)
            for s in k["satirlar"]:
                if s["anahtar"] not in gorulen:
                    gorulen.add(s["anahtar"]); satirlar.append(s)
            metre += [x for x in k.get("metre", []) if x["anahtar"] not in {y["anahtar"] for y in metre}]
        ws.append([f"{m['kod']} · mühendisin paftası (hakemin okuması, {hk['cevap']}) ile kimlik kartlı model · pafta SHA {cevap.get('pafta_sha256', '')[:12]}…"])
        ws.append(["Anahtar", "Paftadan", "Modelden", "Fark", "Sonuç", "Kapıda", "Sayım yolu"]); bas(ws)
        for s in satirlar:
            ws.append([s["anahtar"], s["paftadan"], s["modelden"], s["fark"], s["sonuc"], "evet" if s["kapida"] else "bilgi", s["sayim_yolu"]])
            if s["sonuc"] in renk:
                for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor=renk[s["sonuc"]])
        if metre:
            ws.append([]); ws.append(["Metre (±%3)", "Paftadan (m)", "Modelden (m)", "Fark (%)", "Sonuç", "Parça"])
            for x in metre:
                ws.append([x["anahtar"], x["paftadan_m"], x["modelden_m"], x["fark_yuzde"], x["sonuc"], x["parca"]])
        for i, w in enumerate([44, 11, 11, 8, 14, 8, 30], 1): ws.column_dimensions[chr(64 + i)].width = w
    else:
        ws.append([f"{m['kod']} · depoda hakemin pafta okuması yok"])
    if kt:
        s = json.loads((KOK / kt / "sonuc.json").read_text(encoding="utf-8"))
        w2 = wb.create_sheet("Gerçek liste (kör test)")
        w2.append([f"{s['kod']} ({s['ap']}) · kör test sonucu {s['tarih']} · bizim liste commit {s['bizim_liste']['commit'][:8]} · kıyas commit {s['kiyas_commit'][:8]}. "
                   "Yalnız tutan satırlar ve özet; tutmayan satırların gerçek değeri depoya girmez."])
        w2.append(["Özet", "Satır"]); bas(w2)
        for k_, v_ in s["sayilar"].items():
            w2.append([k_, v_])
        if s.get("not"):
            w2.append(["Not", s["not"]])
        w2.append([]); w2.append(["Tutan satır: bizim kalem", "Bizim ölçü", "Adet", "Eşleme", "Sonuç"])
        for c in w2[w2.max_row]: c.font = Font(bold=True)
        for t in s["tutanlar"]:
            for b in t["bizim"]:
                w2.append([b["kalem"], b["olcu"], t["deger"], t["eslesme"], t["sonuc"]])
        for i, w in enumerate([52, 22, 8, 12, 10], 1): w2.column_dimensions[chr(64 + i)].width = w
    if rh and "veri" in rh:
        w3 = wb.create_sheet("Model ve Motor'un listesi")
        w3.append([f"{m['kod']} · kart başına adet: kimlik kartlı model ve Motor'un yükleme listesi"])
        w3.append(["Kart", "Parça", "Modelde", "Listede", "Not"]); bas(w3)
        for k_ in rh["kiyas"]:
            w3.append([k_["kod"], k_["ad"], k_["model"], k_["liste"], k_["not"]])
            if k_["fark"]:
                for c in w3[w3.max_row]: c.fill = PatternFill("solid", fgColor="FCE4D6")
        for i, w in enumerate([12, 46, 10, 10, 60], 1): w3.column_dimensions[chr(64 + i)].width = w
    for w_ in wb.worksheets:
        for row in w_.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(yol)
    return None


# ----------------------------------------------------------------------------- paket
def _eskiyi_tasi(klasor: Path) -> str | None:
    eski = [p for p in klasor.iterdir() if p.name != "_eski"] if klasor.exists() else []
    if not eski:
        return None
    # eski üretimin tarihi: en eski dosyasınınki (sonradan elle güncellenen bilgi.md tarihi kaydırmasın)
    tarih = dt.datetime.fromtimestamp(min(p.stat().st_mtime for p in eski), dt.timezone.utc).date().isoformat()
    hedef = klasor / "_eski" / tarih
    hedef.mkdir(parents=True, exist_ok=True)
    for p in eski:
        h = hedef / p.name
        n = 1
        while h.exists():                                # aynı gün ikinci üretim: üzerine yazılmaz
            h = hedef / f"{p.stem}_{dt.datetime.now(dt.timezone.utc):%H%M%S}_{n}{p.suffix}"
            n += 1
        shutil.move(str(p), str(h))
    return tarih


def paket(m: dict, kok: Path, gecici: Path, glb: bool = False, kayit_: dict | None = None, ie: dict | None = None) -> dict:
    """Bir modelin paketi; `ie` verilirse (`talep_modeli`'nin maskeli kopyası) İş Emri dosyadan okunmaz.
    Dönen: ozet satırı ({'kod', ..., 'dosyalar': {ad: 'Var' | 'Yok: ...'}})."""
    kayit_ = kayit_ or kayit()
    kod = m["kod"]
    tur = "prefabrik"
    klasor = kok / tur / kod
    gecici.mkdir(parents=True, exist_ok=True)
    gec = Path(tempfile.mkdtemp(prefix=f"{kod}_", dir=gecici))     # her koşuya yeni klasör; eskisi silinmez
    durum: dict[str, str] = {}
    notlar: list[str] = []
    yaz: dict[str, Path] = {}
    ad = lambda kalem, uz, s="": f"{kod}_{kalem}{'_' + s if s else ''}.{uz}"
    if ie is None and m.get("is_emri"):
        ie = json.loads((KOK / m["is_emri"]).read_text(encoding="utf-8"))
    kart_kodu = kod.replace("-", "")
    km = rh = None
    L = {}
    teslim_yok = None                                   # K23: doğrulamadan geçmeyen İş Emri'nden teslim çıktısı (liste, model, pafta, rehber) yok
    if ie is None:
        neden = ("İş Emri depoda yok; motor bu evi bugün çizebilir, ölçüm İş Emri'si yazılınca üretilir" if m["adim"] == "0" else
                 f"çizilemiyor, yol haritası adım {m['adim']} ({kayit_['adimlar'].get(m['adim'], '')})")
        for kalem, uz, _ in KALEMLER[1:]:
            durum[ad(kalem, uz)] = f"Yok: {neden}"
    else:
        # çizim
        try:
            cz = cizim(ie, gec)
            # K23 ve canlı akış (`kanal/worker.py`): doğrulama geçmezse bayiye ne PDF ne plan görseli gider; DXF denetim için yazılır.
            if cz["pdf"]:
                yaz[ad("bayi_paftasi", "pdf")] = cz["pdf"]
                yaz[ad("plan", "png")] = cz["png"]
            else:
                durum[ad("bayi_paftasi", "pdf")] = f"Yok: {cz['neden']}"
                durum[ad("plan", "png")] = "Yok: Motor'un doğrulaması geçmedi; canlıdaki gibi plan görseli de yazılmaz"
                teslim_yok = "Motor'un doğrulaması geçmedi (K23); İş Emri düzeltilip yeniden üretilir, DXF denetim için yazıldı"
            yaz[ad("cizim", "dxf")] = cz["dxf"]
            g = cz["gorunusler"]
            if len(g) >= 4:
                durum[ad("gorunusler", "pdf")] = f"Yok: ayrı dosya değil; dört görünüş ({', '.join(x.lower() for x in g)}) `{ad('bayi_paftasi', 'pdf')}`'in içinde"
            else:
                durum[ad("gorunusler", "pdf")] = (f"Yok: canlıdaki ön onay PDF'inde beşik çatıda yalnız {', '.join(x.lower() for x in g) or 'hiçbir'} görünüş ve "
                                                  "çatı profili var; dört görünüşü bugün yalnız beşik dışı çatılar alıyor (İE-17 PR-1)")
            if cz["uyarilar"]:
                notlar.append("Motor'un uyarıları: " + "; ".join(map(str, cz["uyarilar"])))
        except Exception as e:                                          # noqa: BLE001 — Motor'un hatası nedenle yazılır
            for kalem, uz in (("bayi_paftasi", "pdf"), ("plan", "png"), ("cizim", "dxf"), ("gorunusler", "pdf")):
                durum[ad(kalem, uz)] = f"Yok: Motor çizemedi ({type(e).__name__}: {str(e)[:160]})"
            teslim_yok = "Motor çizemedi, doğrulama yapılamadı (K23)"
        if teslim_yok:
            # K23 ve canlı akış: doğrulamadan geçmeyen (ya da Motor'un çizemediği) İş Emri'nden teslim çıktısı yok; DXF denetim için
            for kalem, uz in (("yukleme_listesi", "xlsx"), ("model", "ifc"), ("montaj_paftasi", "pdf"), ("montaj_rehberi", "html"), ("model", "glb")):
                durum[ad(kalem, uz)] = f"Yok: {teslim_yok}"
        else:
            # liste
            L = listeler(ie, m.get("tavan_secenekleri"))
            if "neden" in L:
                durum[ad("yukleme_listesi", "xlsx")] = f"Yok: {L['neden']}"
                L = {}
            else:
                from motor import prefabrik_liste as PL
                for s, Ls in L.items():
                    PL.xlsx(Ls, gec / f"liste_{s or 'tek'}.xlsx")
                    yaz[ad("yukleme_listesi", "xlsx", s)] = gec / f"liste_{s or 'tek'}.xlsx"
                    ks = Ls.get("kural_seti") or {}
                    notlar.append(f"Liste{' (' + s + ' tavan)' if s else ''}: {Ls['ozet']['satir']} satır, kural yok {Ls['ozet']['kural_yok']}; "
                                  f"üretici {Ls['uretici_surum'].split(' ')[0]}, kural seti {ks.get('kitap_seti', '')} ({ks.get('ortak', '')})")
            # kart modeli
            try:
                km = kart_modeli(ie, gec, kart_kodu)
            except Exception as e:                                          # noqa: BLE001 — hata nedenle yazılır, öteki modeller sürer
                km = {"neden": f"kart modeli hata verdi ({type(e).__name__}: {str(e)[:160]})"}
            if "neden" in km:
                durum[ad("model", "ifc")] = durum[ad("montaj_paftasi", "pdf")] = f"Yok: {km['neden']}"
                km = None
            else:
                yaz[ad("model", "ifc")] = km["ifc"]
                yaz[ad("montaj_paftasi", "pdf")] = km["pdf"]
                if km["kural_yok"]:
                    notlar.append(f"Kart modeli: {len(km['kural_yok'])} parça \"kural yok\" (montaj paftası bu parçalarda tam değil): " +
                                  "; ".join(f"{x['parca']}: {x['neden']}" for x in km["kural_yok"]))
            # rehber
            if km is None:
                durum[ad("montaj_rehberi", "html")] = durum[ad("model", "glb")] = durum[ad("model", "ifc")]
            elif not L:
                durum[ad("montaj_rehberi", "html")] = durum[ad("model", "glb")] = "Yok: liste yok (rehberin listeden gelen parçaları ve adetleri)"
            else:
                s_r = "asma" if "asma" in L else next(iter(L))
                ie_r = copy.deepcopy(ie)
                if s_r:                                     # rehberin tavanı adetlerini aldığı listeninki (listeler() gibi)
                    ie_r.setdefault("ozellikler", {})["Tavan"] = s_r.capitalize()
                try:
                    rh = rehber(km, kart_kodu, ie_r, L[s_r], kod, gec, glb)
                except Exception as e:                                      # noqa: BLE001 — hata nedenle yazılır, öteki modeller sürer
                    rh = {"neden": f"rehber üreteci hata verdi ({type(e).__name__}: {str(e)[:160]})"}
                if "neden" in rh:
                    durum[ad("montaj_rehberi", "html")] = durum[ad("model", "glb")] = f"Yok: {rh['neden']}"
                    rh = None
                else:
                    yaz[ad("montaj_rehberi", "html")] = rh["html"]
                    if "glb" in rh:
                        yaz[ad("model", "glb")] = rh["glb"]
                    else:
                        durum[ad("model", "glb")] = f"Yok: {rh['glb_neden']}"
                    d = rh["veri"]
                    notlar.append(f"Rehber (62 m² rehberinin şablonu, D-28): {d['N']} adım, {len(d['rooms'])} oda, {len(d['panels'])} pano, {len(d['hs'])} H, "
                                  f"{len(d['posts'])} köşe direği, {len(d['makas'])} makas; motorun modelinde olmayanlar kuraldan (kartlarında yazar); adetler " +
                                  ("asma tavanlı listeyle" if "asma" in L else "İş Emri'nin tavanıyla yazılan listeyle"))
    # kıyas
    kiyas_yolu = gec / "kiyas.xlsx"
    neden = kiyas(m, km, rh, kiyas_yolu) if ie is not None and not teslim_yok else durum[ad("model", "ifc")][5:]
    if neden:
        durum[ad("kiyas", "xlsx")] = f"Yok: {neden}"
    else:
        yaz[ad("kiyas", "xlsx")] = kiyas_yolu
    # yaz
    eski = _eskiyi_tasi(klasor)
    klasor.mkdir(parents=True, exist_ok=True)
    for a, p in yaz.items():
        shutil.copy2(p, klasor / a)
        durum[a] = "Var"
    dd = depo_durumu()
    tarih = dt.datetime.now(dt.timezone.utc)
    bilgi = bilgi_md(m, kayit_, durum, notlar, dd, tarih, eski, ie, yaz)
    (klasor / ad("bilgi", "md")).write_text(bilgi, encoding="utf-8")
    durum[ad("bilgi", "md")] = "Var"
    return {"kod": kod, "takma": m.get("takma"), "kume": m["kume"], "m2": m["m2"], "govde": m["govde"], "cati": m["cati"], "adim": m["adim"],
            "dosyalar": durum, "uretim": f"{tarih:%Y-%m-%d} · {dd['dal']} {dd['commit']}" + ("" if dd["temiz"] else " (yerel değişiklik)"),
            "eski": eski}


def _dosya_sirasi(kod: str, durum: dict) -> list[str]:
    out = []
    for kalem, uz, _ in KALEMLER:
        ilk = f"{kod}_{kalem}.{uz}"
        sec = sorted(a for a in durum if a.startswith(f"{kod}_{kalem}_") and a.endswith("." + uz) and kalem != "model")
        out += sec or [ilk]
    return out


def bilgi_md(m, kayit_, durum, notlar, dd, tarih, eski, ie, yaz) -> str:
    kod = m["kod"]
    baslik = f"# {kod}" + (f" ({m['takma']})" if m.get("takma") else "") + f" · tek katlı prefabrik · {m['m2']} m²"
    komut = m.get("komut") or f"python -m araclar.paket {kod} <Modeller klasörü>"
    s = [baslik, "", f"Küme: **{m['kume']}**. Paket: İE-18 (`{komut}`). Kişi ve bayi adı yazılmaz; yalnız kod ve m².", "",
         "## Kimlik", "", "| | |", "|---|---|", f"| Gövde | {m['govde']} |", f"| Veranda | {m['veranda']} |", f"| Çatı (paftada) | {m['cati']} |",
         f"| Pencere / iç duvar | {m['pencere_ic_duvar']} |", f"| Motor'un engeli | {m['engel']} |",
         f"| Yol haritası adımı | {m['adim']}: {kayit_['adimlar'].get(m['adim'], '')} |"]
    if ie is not None:
        y = ie["yapi"]; c = ie.get("cati") or {}
        s += [f"| İş Emri | `{m['is_emri']}` ({ie.get('is_emri_no')} r{int(ie.get('revizyon', 1)):02d}) |",
              f"| Dış ölçü (İş Emri) | " + (f"{y['dis_olcu_cm']['x']} × {y['dis_olcu_cm']['y']} cm" if y.get("dis_olcu_cm") else
                                           f"yazılı değil (ızgaradan; Motor: {m['govde']})") +
              f", duvar {y.get('duvar_yuksekligi_cm', '250 (varsayılan)')} cm |",
              f"| Çatı (İş Emri) | {c.get('tip')}, %{c.get('egim_yuzde')}, saçak {c.get('sacak_cm')}, topuk {c.get('topuk_cm')} |"]
        oz = ie.get("ozellikler") or {}
        if oz:
            s.append("| Özellikler | " + "; ".join(f"{k}: {v}" for k, v in oz.items()) + " |")
    mh = muhur(kod)
    s += ["", "## Kaynak dosyalar (kopyalanmaz)", ""]
    if mh["dosya"] or mh["dwg"]:
        s += ["Mühür: `docs/notlar/2026-10-18_AIPrefabrik_muhur.md` (18 Ekim 2026). Dosyalar Drive'da \"AIPrefabrik\" altında; antetlerinde ad var, buraya gelmez.", "",
              "| Tür | Uzantı | Bayt | SHA-256 / Drive kimliği |", "|---|---|---:|---|"]
        s += [f"| {d['tur']} | {d['uzanti']} | {d['bayt']} | `{d['sha256']}` |" for d in mh["dosya"]]
        s += [f"| {d['tur']} | .dwg | {d['bayt']} | Drive `{d['drive']}` (SHA {d['sha256']}) |" for d in mh["dwg"]]
        if m["kume"] == "kör":
            s += ["", "Kör model: yükleme listesinin yalnız kimliği burada; içeriği kör protokolüyle (bizim liste commit'li, Patron'un açılış kaydı) açılmadan hiçbir dosyaya girmez."]
    elif m["kume"] == "teklif" and m.get("girdi_sha256"):
        s.append(f"Teklif projesi: İş Emri `{m['is_emri']}` (SHA-256 `{m['girdi_sha256']}`); çıktılarda kişi ve bayi adı yerine kod, "
                 "telefon numaraları maskeli. Mühendis paftası ve gerçek liste yok.")
    elif m["kume"] == "teklif":
        s.append("Teklif projesi: bayi çizimi `Teklif_Calisma/" + kod + "/girdi/` altında; mühendis paftası ve gerçek liste yok.")
    else:
        s.append("Mühürde satırı yok.")
    s += ["", "## Üretim", "", f"- Tarih: {tarih:%Y-%m-%d %H:%M} UTC.",
          f"- Depo: dal `{dd['dal']}`, commit `{dd['commit']}`" + ("" if dd["temiz"] else " (motor, kurallar ya da paket üretecinde commit'lenmemiş değişiklik var)") + ".",
          "- Çizim `motor/dxf_out.uret` ve `motor/bayi_gorseli.plan_gorseli`; liste `motor/prefabrik_liste`; kart modeli `motor/kart_model.uret`; rehber `araclar/paket/rehber/`; kıyas hakem aracı (`testler/kimlik/hakem/kiyas.py`) ve kör test sonucu."]
    if eski:
        s.append(f"- Önceki üretim `_eski/{eski}/`'e taşındı.")
    s += ["", "## Dosyalar", "", "| Dosya | Durum |", "|---|---|"]
    for a in _dosya_sirasi(kod, durum | {f"{kod}_bilgi.md": "Var"}):
        s.append(f"| `{a}` | {durum.get(a, 'Var') if a != f'{kod}_bilgi.md' else 'Var'} |")
    if notlar:
        s += ["", "## Notlar", ""] + [f"- {n}" for n in notlar]
    return "\n".join(s) + "\n"


def ozet(satirlar: list[dict], yol: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Modeller"
    ws.append([f"Modeller · her modelin paketi (İE-18) · {dt.datetime.now(dt.timezone.utc):%Y-%m-%d %H:%M} UTC · Var ya da Yok: neden"])
    kalemler = [f"{k}.{u}" for k, u, _ in KALEMLER]
    ws.append(["Kod", "Takma", "Küme", "m²", "Gövde", "Çatı (paftada)", "Adım"] + kalemler + ["Üretim"])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F3B4C")
    for r in satirlar:
        d = r["dosyalar"]
        hucre = []
        for k, u, _ in KALEMLER:
            ilgili = [v for a, v in d.items() if a == f"{r['kod']}_{k}.{u}" or (a.startswith(f"{r['kod']}_{k}_") and a.endswith("." + u) and k != "model")]
            hucre.append("Var" if ilgili and all(v == "Var" for v in ilgili) else (ilgili[0] if ilgili else "Yok"))
        ws.append([r["kod"], r["takma"] or "", r["kume"], r["m2"], r["govde"], r["cati"], r["adim"]] + hucre + [r["uretim"]])
        for c, v in zip(ws[ws.max_row][7:7 + len(kalemler)], hucre):
            c.fill = PatternFill("solid", fgColor="E2EFDA" if v == "Var" else "F2F2F2")
    for i, w in enumerate([8, 8, 8, 10, 26, 26, 10] + [16] * len(kalemler) + [30], 1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B3"
    wb.save(yol)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("kod", help="model kodu (AP-19, TK-01) ya da 'hepsi'")
    ap.add_argument("cikti", type=Path, help="Modeller klasörü (prefabrik/<KOD>/ altına yazar)")
    ap.add_argument("--glb", action="store_true", help="rehberin GLB'sini de yaz (Playwright + THREE_DIR)")
    ap.add_argument("--gecici", type=Path, default=None, help="ara çıktıların klasörü (varsayılan: geçici klasör)")
    ap.add_argument("--is-emri", type=Path, default=None, help="kayıtta olmayan evin (ya da teklif projesinin yeni revizyonunun) İş Emri dosyası")
    ap.add_argument("--tavan", default=None, help="--is-emri ile: tavan seçenekleri, virgülle (asma,ters); teklifle aynı kural")
    a = ap.parse_args(argv)
    k = kayit()
    ieler: dict[str, dict] = {}
    if a.is_emri is not None:
        if a.kod == "hepsi":
            ap.error("--is-emri tek ev içindir; kodu yazın (TK-02 gibi)")
        m, ieler[a.kod] = talep_modeli(a.kod, a.is_emri, [t.strip() for t in a.tavan.split(",") if t.strip()] if a.tavan else None, k)
        modeller = [m]
    elif a.tavan:
        ap.error("--tavan yalnız --is-emri ile; kayıttaki modelin seçenekleri kayıttan gelir")
    else:
        modeller = k["modeller"] if a.kod == "hepsi" else [x for x in k["modeller"] if x["kod"] == a.kod]
    if not modeller:
        ap.error(f"kayıtta yok: {a.kod} (araclar/paket/modeller.json); kayıtta olmayan ev için --is-emri <İş Emri.json>")
    gecici = a.gecici or Path(tempfile.mkdtemp(prefix="paket_"))
    satirlar = []
    for m in modeller:
        r = paket(m, a.cikti, gecici, a.glb, k, ieler.get(m["kod"]))
        var = sum(1 for v in r["dosyalar"].values() if v == "Var")
        print(f"{r['kod']:6} {r['kume']:7} Var {var:2} / {len(r['dosyalar'])}" + (f"  (eskisi _eski/{r['eski']}/)" if r["eski"] else ""), flush=True)
        satirlar.append(r)
    if a.kod == "hepsi":
        oz = a.cikti / "ozet.xlsx"
        if oz.exists():
            t = dt.datetime.fromtimestamp(oz.stat().st_mtime, dt.timezone.utc).date().isoformat()
            (a.cikti / "_eski" / t).mkdir(parents=True, exist_ok=True)
            shutil.move(str(oz), str(a.cikti / "_eski" / t / f"ozet_{dt.datetime.now(dt.timezone.utc):%H%M%S}.xlsx"))
        ozet(satirlar, oz)
        print("ozet:", oz)
    return 0

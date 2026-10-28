"""İE-13 PR-2 · Teklif: İş Emri → Motor'un yükleme listesi → fiyat defteri → teklif Excel ve JSON. D-28: teklif yalnız bu komutla üretilir.

    python3 araclar/teklif.py <İş Emri.json> <çıktı klasörü> [--tavan asma,ters] [--kosullar <json>] [--defter <json>] [--oneri]

Liste dosyası almaz: listeyi İş Emri'nden Motor'la kendisi üretir (`motor/prefabrik_liste.is_emrinden`, paket üreteciyle ortak). Elle
düzeltilmiş liste giremez, adet düzeltmesi yoktur: düzeltme kural kitabına ya da İş Emri'ne işlenir (D-28). Fiyat `araclar/fiyatla.py`
ile defterden; fiyat uydurulmaz, fiyatı olmayan satır boş kalır ve sayılır. LLM yok, ağ yok.

- Proje kararları İş Emri'ndedir (D-28): isteğe bağlı `teklif` bölümü, `{"tavan": ["asma", "ters"], "kosullar": {...}}`. Motor bu
  bölümü okumaz.
- Seçenekler: İş Emri'nin `teklif.tavan`'ı ya da `--tavan asma,ters` (komut satırı İş Emri'nin üstüne yazar) her seçenekte İş Emri'nin
  tavanını değiştirir (TK-01: Patron 27 Ekim 13:49, asma tavanla fiyat, ters tavanla da); ikisi de yoksa İş Emri'nin kendi tavanıyla tek
  seçenek.
- Teklif koşulları (teslim süresi, ödeme şekli): genel değer `kurallar/teklif/kosullar.json`, üstüne İş Emri'nin `teklif.kosullar`'ı, onun
  üstüne `--kosullar` dosyası. Değer yoksa hücre sarı kalır ve "Patron'dan" yazar; değer uydurulmaz (Patron 28 Ekim 06:51: bayiye süre ve
  ödeme "teklifle birlikte bildirilir").
- Excel düzeni Patron'un yükleme listesinin düzeni (28 Ekim 06:40–06:41): grup adı sütun değil, grubun üstünde başlık satırı ve altında
  ara toplam; tek "Adet" sütunu; ölçüsü birim olan satırda (adet, metre, kova …) ölçü boş, birim o. Sayfalar: Özet, her seçenek bir
  sayfa, Fiyat kaynağı. Liste sayfasını `fiyatla.py --excel` de bu modülün yazıcısıyla yazar (tek yazıcı).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parents[1]
if str(KOK) not in sys.path:
    sys.path.insert(0, str(KOK))
from araclar import fiyatla as F  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402

KOSULLAR = KOK / "kurallar" / "teklif" / "kosullar.json"
TAVAN = {"asma": "Asma", "ters": "Ters"}
BIRIM = {"m": "metre"}
TL = "#,##0.00"
RENK = {"sari": "FFF2A8", "gri": "E7E7E7", "baslik": "D9E1F2", "cizgi": "8EA9DB",
        "öneri onaylı": "FCE4D6", "ikinci kaynak": "DDEBF7", "öneri": "FFF2A8"}


def buyuk(s: str) -> str:
    """Türkçe büyük harf: "Elektrik" → "ELEKTRİK", "Sıhhi" → "SIHHİ" (str.upper i'yi I yapar)."""
    return s.replace("i", "İ").replace("ı", "I").upper()


def gorunen(olcu: str | None, birim: str | None) -> tuple[str, str]:
    """(ölçü, birim): ölçü yalnız birim sözcüğüyse (adet, metre, kova, paket, takım) ölçü boş kalır, birim o olur."""
    olcu = olcu or ""
    if re.search(r"\d", olcu):
        return olcu, BIRIM.get(birim, birim or "")
    return "", olcu or BIRIM.get(birim, birim or "")


def kaynak_durumu(s: dict, defter: dict, hat: str) -> tuple[str, dict | None]:
    """Satırın fiyatının defterdeki yeri: (durum, kayıt). Durum defterin `durum`'u (eşleşti, bant, öneri onaylı, ikinci kaynak, kaynakta
    yok, …), defterin bandıyla bulunduysa "bant", hiç yoksa "defterde yok"."""
    H = defter["hatlar"][hat]
    g = H["kalemler"].get(F.anahtar(s, hat))
    if g is not None:
        return g["durum"], g
    if hat == "prefabrik" and (b := F.bant_bul(s, H.get("bantlar") or [])):
        alt, ust = b[1]["en_mm"]
        return "bant", {"fiyat": b[1]["fiyat"], "kaynak": b[0]["kaynak"], "kaynak_satir": b[0].get("kaynak_satir"),
                        "kaynak_kalem": b[0].get("kaynak_kalem"), "onay": b[1]["onay"], "not": f"pano eni bandı {alt:g}–{ust:g} mm"}
    return "defterde yok", None


def secenekler(ie: dict, tavan: list[str] | None) -> list[tuple[str, dict]]:
    tavan = tavan or (ie.get("teklif") or {}).get("tavan")
    if not tavan:
        t = (ie.get("ozellikler") or {}).get("Tavan")
        return [(f"{t} tavan" if t else "Teklif", {})]
    bilinmeyen = [t for t in tavan if t not in TAVAN]
    if bilinmeyen or len(set(tavan)) != len(tavan):
        raise SystemExit(f"DUR: tavan seçeneği {tavan}; bilinenler {sorted(TAVAN)}, her biri bir kez")
    return [(f"{TAVAN[t]} tavan", {"Tavan": TAVAN[t]}) for t in tavan]


def kosullari_oku(ek: Path | str | None = None, ie: dict | None = None) -> dict:
    """Genel koşullar, üstüne İş Emri'nin `teklif.kosullar`'ı, onun üstüne `ek` dosyası. Genel dosyada olmayan koşul adı durur (yazım
    hatası yeni satır açmasın)."""
    K = dict(json.loads(KOSULLAR.read_text(encoding="utf-8"))["kosullar"])
    katmanlar = [((ie or {}).get("teklif") or {}).get("kosullar") or {}]
    if ek:
        katmanlar.append(json.loads(Path(ek).read_text(encoding="utf-8")).get("kosullar") or {})
    for k in katmanlar:
        if bilinmeyen := sorted(set(k) - set(K)):
            raise SystemExit(f"DUR: bilinmeyen teklif koşulu {bilinmeyen}; bilinenler {sorted(K)}")
        K.update(k)
    return K


def uret(ie: dict, defter: dict, secim: list[tuple[str, dict]], kosullar: dict, oneri: bool = False) -> dict:
    """Her seçenek için Motor'un listesi ve fiyatı; teklifin JSON'u (zaman damgası yok: aynı girdi aynı çıktıyı verir)."""
    sec, kod = [], None
    for ad, oz in secim:
        L = PL.is_emrinden(ie, ozellikler=oz)
        kod = kod or L.get("kod")
        R = F.fiyatla(L, defter, "prefabrik", oneri)
        sec.append({"ad": ad, "ozellikler": oz, "liste": {k: L[k] for k in ("uretici_surum", "kural_seti", "girdi_ozeti", "ozet")}, "fiyat": R})
    return {"sema": "portman.teklif/1", "kod": kod or ie.get("is_emri_no"),
            "is_emri_no": ie.get("is_emri_no"), "revizyon": ie.get("revizyon"), "defter_surumu": defter.get("surum"),
            "oneri_dahil": oneri, "kosullar": kosullar, "secenekler": sec}


# ── Excel ──────────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
def _stil():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    dolgu = {k: PatternFill("solid", fgColor=v) for k, v in RENK.items()}
    return Alignment, Border, Font, dolgu, Side


def liste_sayfasi(ws, satirlar: list[dict], baslik: str, aciklama: str, renk=None, kaynak_sutunu: bool = False) -> float:
    """Fiyatlı liste sayfası (`fiyatla.fiyatla`'nın satırları). Grup başlık satırı + satırlar + ara toplam; en altta genel toplam.
    `renk(satır)` birim fiyat hücresinin dolgusunu (RENK anahtarı) verir. `kaynak_sutunu`: fiyatla.py'nin defter görünümü için
    Durum ve Kaynak sütunları. Döner: genel toplam."""
    Alignment, Border, Font, dolgu, Side = _stil()
    kg = any(s.get("kg") is not None for s in satirlar)
    sutun = ["Ölçü", "Birim", "Adet"] + (["kg"] if kg else []) + ["Birim fiyat", "Tutar", "Not"] + (["Durum", "Kaynak"] if kaynak_sutunu else [])
    n = 1 + len(sutun)
    i_bf, i_tut, i_not = (sutun.index(x) + 2 for x in ("Birim fiyat", "Tutar", "Not"))    # sayfadaki sütun (1'den; A Kalem)
    ws.append([baslik]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([aciklama]); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n); ws.row_dimensions[2].height = 32
    gruplar: dict[str, list[dict]] = {}                     # grup ilk göründüğü sırada, grup içinde Motor'un sırası
    for s in satirlar:
        gruplar.setdefault(s.get("grup_adi") or str(s.get("grup") or ""), []).append(s)
    toplam, fiyatsiz = 0.0, 0
    for ad, ss in gruplar.items():
        ws.append([]); ws.append([buyuk(ad), *sutun])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True); c.fill = dolgu["baslik"]
        ara = 0.0
        for s in ss:
            olcu, birim = gorunen(s.get("olcu"), s.get("birim"))
            miktar = s.get("kg") if s.get("birim") == "kg" else s.get("adet")
            if s["tutar"] is not None:
                ara += s["tutar"]
            elif miktar != 0:
                fiyatsiz += 1
            notu = "Motor'da kuralı yok, adet yok" if s.get("adet") is None and s.get("birim") != "kg" else ("fiyat yok" if s["birim_fiyat"] is None else "")
            ws.append([s["kalem"], olcu, birim, s.get("adet")] + ([s.get("kg")] if kg else []) + [s["birim_fiyat"], s["tutar"], notu]
                      + ([s["durum"], s.get("kaynak_adi") or s.get("kaynak")] if kaynak_sutunu else []))
            r = ws[ws.max_row]
            r[i_bf - 1].number_format = r[i_tut - 1].number_format = TL
            k = renk(s) if renk else None
            if s.get("adet") is None and s.get("birim") != "kg":
                for c in r: c.fill = dolgu["gri"]
            elif k:
                r[i_bf - 1].fill = dolgu[k]
        toplam += ara
        ws.append([f"{ad} toplamı"] + [""] * (i_tut - 2) + [round(ara, 2)])
        r = ws[ws.max_row]; r[0].font = r[i_tut - 1].font = Font(bold=True); r[i_tut - 1].number_format = TL
        for c in r[:n]: c.border = Border(top=Side(style="thin", color=RENK["cizgi"]))
    ws.append([])
    ws.append(["GENEL TOPLAM (fiyatı olan satırlar)"] + [""] * (i_tut - 2) + [round(toplam, 2), f"fiyatı olmayan satır: {fiyatsiz}"])
    r = ws[ws.max_row]; r[0].font = r[i_tut - 1].font = Font(bold=True, size=12); r[i_tut - 1].number_format = TL
    for col, w in zip("ABCDEFGHIJ", [62, 20, 9, 8] + ([9] if kg else []) + [13, 14, 30] + ([16, 30] if kaynak_sutunu else [])):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if c.column in (3, 4): c.alignment = Alignment(horizontal="center")
            elif c.column >= i_not: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    return round(toplam, 2)


def excel_yaz(T: dict, defter: dict, yol: Path) -> None:
    from openpyxl import Workbook
    Alignment, Border, Font, dolgu, Side = _stil()
    B = Font(bold=True)
    kaynak_kisa = {k: v.get("kisa") for k, v in defter["kaynaklar"].items()}
    def durum_renk(s: dict) -> str | None:
        if s["birim_fiyat"] is None:
            return "gri"
        if s["durum"].startswith("öneri"):                  # onaysız öneri (yalnız --oneri ile)
            return "öneri"
        d = kaynak_durumu(s, defter, "prefabrik")[0]
        return d if d in ("öneri onaylı", "ikinci kaynak") else None
    wb = Workbook(); oz = wb.active; oz.title = "Özet"
    kimlik = f"{T['is_emri_no']} r{T['revizyon']:02d}" if isinstance(T.get("revizyon"), int) else str(T.get("is_emri_no"))
    acik = (f"İş Emri {kimlik}; liste Motor'dan, fiyat defterden ({T['defter_surumu']}"
            + ("; onaysız öneriler dahil" if T["oneri_dahil"] else "") + "). Fiyatı boş satır toplama girmez. "
            "Birim fiyat rengi: turuncu onaylı öneri, mavi Haziran malzeme listesi, gri fiyat yok. Gri satır: Motor'da kuralı yok.")
    toplamlar = []
    for sec in T["secenekler"]:
        ws = wb.create_sheet(sec["ad"])
        toplamlar.append(liste_sayfasi(ws, sec["fiyat"]["satirlar"], f"{T['kod']} · teklif · {sec['ad']} · TASLAK", acik, durum_renk))
        assert abs(toplamlar[-1] - sec["fiyat"]["ozet"]["toplam"]) < 0.01, (sec["ad"], toplamlar[-1], sec["fiyat"]["ozet"]["toplam"])

    # Özet
    oz.append([f"{T['kod']} · teklif · TASLAK"]); oz["A1"].font = Font(bold=True, size=13)
    oz.append([acik]); oz["A2"].alignment = Alignment(wrap_text=True, vertical="top"); oz.merge_cells("A2:D2"); oz.row_dimensions[2].height = 32
    oz.append([])
    oz.append(["Seçenek", "Toplam", "Fiyatsız satır", "Satır"])
    for c in oz[oz.max_row]: c.font = B
    for sec in T["secenekler"]:
        o = sec["fiyat"]["ozet"]
        oz.append([sec["ad"], o["toplam"], len(o["fiyatsiz"]), f"{o['satir']} satır, {o['fiyatli']} fiyatlı"])
        oz.cell(oz.max_row, 2).number_format = TL; oz.cell(oz.max_row, 2).font = B
    if len(T["secenekler"]) == 2:
        a, b = (s["fiyat"]["ozet"]["toplam"] for s in T["secenekler"])
        oz.append([f"Fark ({T['secenekler'][1]['ad']} − {T['secenekler'][0]['ad']})", round(b - a, 2)]); oz.cell(oz.max_row, 2).number_format = TL
    oz.append(["Toplam yalnız fiyatı olan satırların toplamıdır; fiyatı olmayan satırlar aşağıda."])
    oz.append([])
    oz.append(["Teklif koşulları", "Değer", "", "Açıklama"])
    for c in oz[oz.max_row]: c.font = B
    for k, v in T["kosullar"].items():
        oz.append([k, v or "", "", "" if v else "Patron'dan"])
        if not v: oz.cell(oz.max_row, 2).fill = dolgu["sari"]
    oz.append([])
    oz.append(["Fiyatın kaynağı"] + [s["ad"] for s in T["secenekler"]])
    for c in oz[oz.max_row]: c.font = B
    sayim = [{} for _ in T["secenekler"]]
    for i, sec in enumerate(T["secenekler"]):
        for s in sec["fiyat"]["satirlar"]:
            d = kaynak_durumu(s, defter, "prefabrik")[0]
            sayim[i][d] = sayim[i].get(d, 0) + 1
    for d in sorted({d for c in sayim for d in c}, key=lambda d: (d in ("kaynakta yok", "defterde yok"), d)):
        oz.append([d] + [c.get(d, 0) for c in sayim])
        if d in RENK: oz.cell(oz.max_row, 1).fill = dolgu[d]
    oz.append([])
    bos = [""] * (3 - len(T["secenekler"]))                 # Neden hep E sütununda (en çok üç seçenek)
    oz.append(["Fiyatı olmayan satırlar"] + [f"{s['ad']} adedi" for s in T["secenekler"]] + bos + ["Neden"])
    for c in oz[oz.max_row]: c.font = B
    adet = [{F.anahtar(s, "prefabrik"): s["adet"] for s in sec["fiyat"]["satirlar"]} for sec in T["secenekler"]]
    goruldu = []
    for sec in T["secenekler"]:
        for f in sec["fiyat"]["ozet"]["fiyatsiz"]:
            if f["anahtar"] in goruldu: continue
            goruldu.append(f["anahtar"])
            kayit = defter["hatlar"]["prefabrik"]["kalemler"].get(f["anahtar"]) or {}
            neden = "Motor'da kuralı yok, adet yok" if f["miktar"] is None else (kayit.get("not") or f["durum"])
            oz.append([f["anahtar"].replace(" | ", " · ")] + [a.get(f["anahtar"]) for a in adet] + bos + [neden])
            oz.cell(oz.max_row, 1).fill = dolgu["gri"]
    oz.append([])
    oz.append(["Kaynaklar: " + "; ".join(v["kisa"] for v in defter["kaynaklar"].values() if v.get("patron_verdi")) + " (Patron'un verdiği listeler)."])
    for col, w in zip("ABCDE", (58, 16, 16, 22, 90)): oz.column_dimensions[col].width = w
    for row in oz.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (1, 5))

    # Fiyat kaynağı: seçeneklerin anahtarları bir arada, defterdeki kaydıyla
    fk = wb.create_sheet("Fiyat kaynağı")
    fk.append(["Fiyatın defterdeki kaydı (kurallar/fiyat). Satır numarası kaynak listedeki Excel satırıdır."]); fk["A1"].font = B
    fk.append(["Kalem", "Ölçü"] + [f"{s['ad']} adedi" for s in T["secenekler"]] + ["Durum", "Birim fiyat", "Kaynak", "Kaynaktaki satır", "Kaynaktaki kalem", "Not"])
    for c in fk[2]: c.font = B
    sira = []
    for sec in T["secenekler"]:
        for s in sec["fiyat"]["satirlar"]:
            if F.anahtar(s, "prefabrik") not in [F.anahtar(x, "prefabrik") for x in sira]: sira.append(s)
    for s in sira:
        a = F.anahtar(s, "prefabrik"); d, g = kaynak_durumu(s, defter, "prefabrik"); g = g or {}
        fk.append([s["kalem"], s["olcu"]] + [x.get(a) for x in adet] + [d, g.get("fiyat"), kaynak_kisa.get(g.get("kaynak")) or "",
                   ", ".join(map(str, g.get("kaynak_satir") or [])), g.get("kaynak_kalem") or "", g.get("not") or ""])
        if d in RENK: fk.cell(fk.max_row, len(T["secenekler"]) + 3).fill = dolgu[d]
        fk.cell(fk.max_row, len(T["secenekler"]) + 4).number_format = TL
    for col, w in zip("ABCDEFGHIJK", [52, 20] + [10] * len(T["secenekler"]) + [14, 12, 30, 12, 50, 60]): fk.column_dimensions[col].width = w
    fk.freeze_panes = "C3"
    oz.page_setup.orientation = fk.page_setup.orientation = "landscape"
    oz.sheet_properties.pageSetUpPr.fitToPage = True; oz.page_setup.fitToWidth, oz.page_setup.fitToHeight = 1, 0
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)


def main(argv: list[str] | None = None) -> int:
    a = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    a.add_argument("is_emri", type=Path); a.add_argument("cikti", type=Path)
    a.add_argument("--tavan", help="seçenekler, virgülle: asma,ters")
    a.add_argument("--kosullar", type=Path, help="projeye özel teklif koşulları (kurallar/teklif/kosullar.json biçimi)")
    a.add_argument("--defter", type=Path, default=F.DEFTER); a.add_argument("--oneri", action="store_true")
    o = a.parse_args(argv)
    ie = json.loads(o.is_emri.read_text(encoding="utf-8"))
    secim = secenekler(ie, [t.strip() for t in o.tavan.split(",")] if o.tavan else None)
    defter = F.oku(o.defter)
    try:
        T = uret(ie, defter, secim, kosullari_oku(o.kosullar, ie), o.oneri)
    except PL.ListeYazilmaz as e:
        print(f"teklif: {e.neden}", file=sys.stderr)
        return 1
    kod = str(T["kod"]).replace("-", "")
    o.cikti.mkdir(parents=True, exist_ok=True)
    (o.cikti / f"{kod}_teklif.json").write_text(json.dumps(T, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    excel_yaz(T, defter, o.cikti / f"{kod}_teklif.xlsx")
    for sec in T["secenekler"]:
        oz = sec["fiyat"]["ozet"]
        print(f"{T['kod']} · {sec['ad']}: {oz['toplam']:,.2f} TL · {oz['fiyatli']}/{oz['satir']} satır fiyatlı, {len(oz['fiyatsiz'])} fiyatsız")
    print("koşullar:", ", ".join(f"{k} {v or 'yok (Patron’dan)'}" for k, v in T["kosullar"].items()))
    print("→", o.cikti / f"{kod}_teklif.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())

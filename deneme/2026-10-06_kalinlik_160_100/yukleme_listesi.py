"""6 Ekim 2026 denemesi: İş Emri → yükleme listesi (Excel).

    python deneme/2026-10-06_kalinlik_160_100/yukleme_listesi.py [is_emri.json]

Örnek: Astra dönemindeki `Cumbali-6x5-Yukleme-Listesi.xlsx` (YÜKLEME + Blok Sayımı; adetler sayımdan SUMIFS ile).
O liste yalnız çizilmiş blokları sayar; çatı, profil ve tavanı "çizimde tanımlı değil" diye dışarıda bırakır. Burada:
  - Sayım: her pano, doğrama ve bağlantı tek satır; cephesi, sıra numarası (görünüş paftasındaki kod) ve IFC adıyla.
  - YÜKLEME: adetler Sayım'dan formülle; türetilen kalemler (pano eki H'ları, alt çerçeve, makas, çatı, tavan, dere) Metraj'dan.
  - Metraj: her türetme kuralı girdileriyle birlikte formül olarak; aynı kural P6001'in girdileriyle de hesaplanır ve
    P6001'in gerçek yükleme listesindeki adetle karşılaştırılır (docs/notlar/2026-09-21_Yukleme_Listesi_Analizi.md).
  - Açık konular: gerçek listede olup kuralı bilinmeyen kalemler ve teyit bekleyenler.
Hiçbir adet elle yazılmaz: sayılanlar modelden, türetilenler formülden gelir. Canlı Motor'a dokunmaz.
"""
from __future__ import annotations
import json, os, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BURASI = Path(__file__).resolve().parent
KOK = BURASI.parents[1]
CIKTI = BURASI / "cikti"
sys.path.insert(0, str(KOK)); sys.path.insert(0, str(BURASI))
os.environ.setdefault("PORTMAN_REF_DIR", str(KOK / "kanal"))

F = "Arial"
YAZI = Font(name=F, size=10); KALIN = Font(name=F, size=10, bold=True); BASLIK = Font(name=F, size=14, bold=True)
BEYAZ_KALIN = Font(name=F, size=10, bold=True, color="FFFFFF"); GIRDI = Font(name=F, size=10, color="0000FF"); KUCUK = Font(name=F, size=9, color="555555")
D_BASLIK = PatternFill("solid", fgColor="2F3B4C"); D_GRUP = PatternFill("solid", fgColor="E7E2D6")
D_LISTE = PatternFill("solid", fgColor="FFF2CC"); D_CIKARIM = PatternFill("solid", fgColor="FCE4D6"); D_TOPLAM = PatternFill("solid", fgColor="D9E1F2")
INCE = Side(style="thin", color="BFBFBF"); KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
ORTA = Alignment(horizontal="center", vertical="center"); SOL = Alignment(horizontal="left", vertical="center", wrap_text=True)

# P6001'in gerçek yükleme listesindeki adetler ve Motor'un P6001 planından ölçülenler (21 Eylül analizi; 6 Ekim'de yeniden ölçüldü)
P6001 = {"x": 773, "y": 763, "aciklik": 773, "bina_boyu": 763, "dis_pano": 25, "dis_kosu": 6, "ic_pano": 11, "ic_kosu": 4, "uclu_h": 2,
         "dis_boy": 30.02, "ic_boy": 11.155, "aks": 7, "p02": 2, "kapatma": 0.87}
P6001_LISTE = {"h_dis": 19, "h_ic": 5, "alt_dis": 9, "alt_ic": 3, "kor_kasa": 2, "makas_bas": 2, "makas_orta": 5, "kaplama": 20,
               "mahya": 3, "sacak_saci": 7, "alin_saci": 8, "asik": 37, "alcipan": 20, "silte": 7}


def model(ie: dict) -> dict:
    """Plan + malzeme kodları: Motor'un deneme kopyasından (DXF ve IFC ile aynı sayım)."""
    import yama
    yama.uret(CIKTI)
    sys.path.insert(0, str(CIKTI))
    from motor_k import geometry as G, ifc_uret
    from motor_k.dxf_out import Cizici
    from motor_k.katalog import Katalog

    plan = G.kur(json.loads(json.dumps(ie)))
    c = Cizici(plan, Katalog()); c.ciz_plan(0.0, 0.0, "ANA", "", "")
    H0 = float(ie["yapi"].get("duvar_yuksekligi_cm", 250))
    nx2, ny2 = int(round(plan.nx * 2)), int(round(plan.ny * 2))
    cephe = {("x", 0): ("Ön", "Ö", False), ("x", ny2): ("Arka", "A", True), ("y", 0): ("Sol", "SL", True), ("y", nx2): ("Sağ", "SĞ", False)}
    satirlar, pid = [], 0
    ic_no = 0
    for d in plan.duvarlar:
        dis = d.tur == "dis"
        if dis:
            ad, on_ek, ayna = cephe[(d.yon, d.hat)]
            sira = sorted(d.panolar, key=lambda p: -p.a0 if ayna else p.a0)       # görünüş paftasındaki sıra: bakana göre soldan sağa
            duvar = f"{ad} cephe"
        else:
            ic_no += 1
            on_ek = f"İ{ic_no}-"; sira = list(d.panolar)
            duvar = f"İç duvar {ic_no} ({'x' if d.yon == 'y' else 'y'} = {d.c:g} cm, {'soldan sağa' if d.yon == 'x' else 'önden arkaya'})"
        no = {id(p): i + 1 for i, p in enumerate(sira)}
        for p in d.panolar:
            _, _, kod = c.pano_blok(p)
            r = p.rect(); L = p.boy
            ifc_ad = f"PANO_{'D' if dis else 'I'}_{ifc_uret.pano_sinifi(L, 2 * d.h).upper()}_{L:.1f}_{pid:02d}"; pid += 1
            ac = p.acilim
            aciklik = "" if ac is None else (f"pencere {ac.olcu.replace('x', '×')}" if ac.tur == "pencere" else f"{'dış' if ac.tur == 'dis_kapi' else 'iç'} kapı {ac.genislik:g}")
            satirlar.append({"kod": kod, "tur": "Dış pano" if dis else "İç pano", "duvar": duvar, "sira": f"{on_ek}{no[id(p)]}", "boy": round(L * 10, 1),
                             "kal": round(2 * d.h * 10, 1), "yuk": round(H0 * 10), "aciklik": aciklik, "x": round((r[0] + r[2]) / 2, 2), "y": round((r[1] + r[3]) / 2, 2), "ifc": ifc_ad})
            if ac is not None:
                g = p.gap()
                if ac.tur == "pencere":
                    dk, tur2 = c.dograma[ac.olcu][1], "Pencere"; en, yuk = (int(v) * 10 for v in ac.olcu.split("x"))
                elif ac.tur == "ic_kapi":
                    dk, tur2, en, yuk = c.dograma["ic82"][1], "Kapı", round(ac.genislik * 10), None
                else:
                    dk, tur2, en, yuk = c._dis_kapi_blok(ac.genislik)[1], "Kapı", round(ac.genislik * 10), None
                satirlar.append({"kod": dk, "tur": tur2, "duvar": duvar, "sira": f"{on_ek}{no[id(p)]}", "boy": en, "kal": None, "yuk": yuk, "aciklik": aciklik,
                                 "x": round((g[0] + g[2]) / 2, 2), "y": round((g[1] + g[3]) / 2, 2), "ifc": ""})
    for i, k in enumerate(plan.koseler, 1):
        _, kod = c.kose_blok(k); x0, y0, x1, y1 = k.rect
        satirlar.append({"kod": kod, "tur": "Bağlantı", "duvar": "Dış köşe" if k.tur == "dis" else "İç köşe", "sira": f"K{i}", "boy": round(max(x1 - x0, y1 - y0) * 10, 1),
                         "kal": round(min(x1 - x0, y1 - y0) * 10, 1), "yuk": round(H0 * 10), "aciklik": "", "x": round((x0 + x1) / 2, 2), "y": round((y0 + y1) / 2, 2), "ifc": ""})
    for i, h in enumerate(plan.h_baglantilar, 1):
        satirlar.append({"kod": c.h_kod, "tur": "Bağlantı", "duvar": "İç duvar T birleşimi", "sira": f"H{i}", "boy": round(G.T_IC * 10, 1), "kal": None, "yuk": round(H0 * 10),
                         "aciklik": "", "x": round(h.x, 2), "y": round(h.y, 2), "ifc": ""})
    dis = [d for d in plan.duvarlar if d.tur == "dis"]; ic = [d for d in plan.duvarlar if d.tur == "ic"]
    cati = ie.get("cati", {})
    makas_y = cati.get("makas_ekseni", "x") == "y"
    kalkan = [d for d in dis if d.yon == ("x" if makas_y else "y")]           # mahyaya paralel dış duvar: makaslar bunun eklerine oturur
    kodlar = {}
    for s in satirlar:
        m = c.kat.malzemeler.get(s["kod"], {})
        kodlar.setdefault(s["kod"], {"label": m.get("label", ""), "dims": m.get("dimensions_mm"), "tur": s["tur"], "note": m.get("note", "")})
    return {"plan": plan, "satirlar": satirlar, "kodlar": kodlar, "H0": H0, "t_dis": G.T_DIS, "t_ic": G.T_IC,
            "girdi": {"x": plan.W, "y": plan.H, "aciklik": plan.H if makas_y else plan.W, "bina_boyu": plan.W if makas_y else plan.H,
                      "dis_kosu": len(dis), "ic_kosu": len(ic), "aks": len(kalkan[0].panolar) + 1,
                      "sacak": float(cati.get("sacak_cm", 30)), "egim": float(cati.get("egim_yuzde", 33)) / 100.0,
                      "net_kapali": round(sum(m.alan_m2 for m in plan.mahaller.values() if not m.acik), 2)},
            "kontrol": {"dis_pano": sum(len(d.panolar) for d in dis), "ic_pano": sum(len(d.panolar) for d in ic), "uclu_h": len(plan.h_baglantilar),
                        "dis_boy": round(sum(p.boy for d in dis for p in d.panolar) / 100, 3), "ic_boy": round(sum(p.boy for d in ic for p in d.panolar) / 100, 3)}}


def _yaz(ws, r, c, v, font=YAZI, fill=None, hiza=None, fmt=None, kenar=True):
    h = ws.cell(row=r, column=c, value=v); h.font = font
    if fill is not None: h.fill = fill
    h.alignment = hiza or (SOL if isinstance(v, str) and not str(v).startswith("=") else ORTA)
    if fmt: h.number_format = fmt
    if kenar: h.border = KENAR
    return h


def olcu_yazi(dims) -> str:
    return " × ".join(f"{v:g}".replace(".", ",") for v in (dims or []) if v is not None)


def kitap(ie: dict, m: dict) -> tuple[Workbook, dict]:
    wb = Workbook()
    # ------------------------------------------------------------------ Sayım
    ws = wb.active; ws.title = "Sayım"
    bas = ["Kod", "Eleman", "Tür", "Duvar", "Sıra (pafta kodu)", "Boy / en (mm)", "Kalınlık (mm)", "Yükseklik (mm)", "Açıklık", "Merkez X (cm)", "Merkez Y (cm)", "IFC adı"]
    _yaz(ws, 1, 1, "MODELDEN ALINAN SAYIM — her satır bir eleman", KALIN, kenar=False)
    _yaz(ws, 2, 1, f"Kaynak: İş Emri {ie['is_emri_no']} r{int(ie.get('revizyon', 1)):02d}; Motor deneme kopyasının planı (DXF ve IFC ile aynı). Koordinatlar sol-ön dış köşeden, cm. "
                   "Sıra kodu görünüş paftasındaki pano numarasıdır (cephede bakana göre soldan sağa).", KUCUK, kenar=False)
    for j, b in enumerate(bas, 1):
        _yaz(ws, 4, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    r0 = 5
    for i, s in enumerate(m["satirlar"]):
        r = r0 + i
        for j, v in enumerate([s["kod"], m["kodlar"][s["kod"]]["label"], s["tur"], s["duvar"], s["sira"], s["boy"], s["kal"], s["yuk"], s["aciklik"], s["x"], s["y"], s["ifc"]], 1):
            _yaz(ws, r, j, v, hiza=SOL if j in (2, 4, 9, 12) else ORTA, fmt="0.0" if j == 6 and isinstance(v, float) and v != int(v) else None)
    r1 = r0 + len(m["satirlar"]) - 1
    for j, w in enumerate([8, 30, 11, 46, 16, 13, 13, 14, 20, 13, 13, 30], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    KOD = f"Sayım!$A${r0}:$A${r1}"; TUR = f"Sayım!$C${r0}:$C${r1}"; BOY = f"Sayım!$F${r0}:$F${r1}"

    # ------------------------------------------------------------------ Metraj
    wm = wb.create_sheet("Metraj")
    _yaz(wm, 1, 1, "TÜRETİLEN KALEMLER — kural, girdi ve formül", KALIN, kenar=False)
    _yaz(wm, 2, 1, "Mavi: girdi (modelden ya da kuraldan). Siyah: formül. P6001 sütunu aynı formülü P6001'in girdileriyle hesaplar; yanındaki sütun P6001'in gerçek "
                   "yükleme listesindeki adettir (60 m², İkranur Prefabrik; 21 Eylül 2026 analizi). Tutmayan kural kırmızı ✗ ile durur.", KUCUK, kenar=False)
    for j, b in enumerate(["Kalem", "Bu yapı", "P6001 (aynı formül)", "P6001 gerçek liste", "Tutuyor mu", "Kural ve kaynağı"], 1):
        _yaz(wm, 4, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    g = m["girdi"]; A = {}
    r = 5

    def grup(ad):
        nonlocal r
        for j in range(1, 7):
            _yaz(wm, r, j, ad if j == 1 else None, KALIN, D_GRUP)
        r += 1

    def girdi(anahtar, ad, b, c_, kaynak, fmt=None, formul_b=False):
        nonlocal r
        _yaz(wm, r, 1, ad)
        _yaz(wm, r, 2, b, YAZI if formul_b else GIRDI, fmt=fmt); _yaz(wm, r, 3, c_, GIRDI, fmt=fmt)
        _yaz(wm, r, 4, None); _yaz(wm, r, 5, None); _yaz(wm, r, 6, kaynak)
        A[anahtar] = r; r += 1

    def hesap(anahtar, ad, formul, kaynak, liste=None, fmt="0"):
        """formul: '{x}' yer tutucuları girdi/hesap anahtarlarıdır; B ve C sütunları için ayrı ayrı açılır."""
        nonlocal r
        _yaz(wm, r, 1, ad)
        for kol in ("B", "C"):
            _yaz(wm, r, 2 if kol == "B" else 3, "=" + formul.format(**{k: f"{kol}{v}" for k, v in A.items()}), fmt=fmt)
        if liste is not None:
            _yaz(wm, r, 4, liste, GIRDI)
            h = _yaz(wm, r, 5, f'=IF(C{r}=D{r},"✓","✗")', KALIN)
        else:
            _yaz(wm, r, 4, None); _yaz(wm, r, 5, "—")
        _yaz(wm, r, 6, kaynak)
        A[anahtar] = r; r += 1

    grup("Girdiler · yapı")
    girdi("x", "Dış ölçü, genişlik (cm)", g["x"], P6001["x"], "İş Emri: dış ölçü (Lokman, 6 Ekim 18:50). P6001: 6 pano × 125,5 + 20.")
    girdi("y", "Dış ölçü, derinlik (cm)", g["y"], P6001["y"], "İş Emri: dış ölçü. P6001: 6 pano × 125,5 + 10.")
    girdi("aciklik", "Makas açıklığı (cm)", g["aciklik"], P6001["aciklik"], "Makaslar kısa kenarı geçer (Lokman, 20–21 Eylül). Bu yapıda derinlik; P6001'de genişlik.")
    girdi("bina_boyu", "Mahya boyunca bina boyu (cm)", g["bina_boyu"], P6001["bina_boyu"], "Makas açıklığına dik kenar.")
    girdi("sacak", "Saçak (cm)", g["sacak"], 30, "P6001 çizim kuralı: saçak 30 cm.")
    girdi("egim", "Çatı eğimi", g["egim"], 0.33, "P6001 çizim kuralı: %33.", fmt="0%")
    grup("Girdiler · modelden sayılanlar")
    girdi("dis_pano", "Dış pano adedi", f'=COUNTIF({TUR},"Dış pano")', P6001["dis_pano"], "Sayım sayfasından.", formul_b=True)
    girdi("dis_kosu", "Dış duvar koşusu (kesintisiz duvar)", g["dis_kosu"], P6001["dis_kosu"], "Plandan: dört cephe. P6001'de veranda girintisiyle 6.")
    girdi("ic_pano", "İç pano adedi", f'=COUNTIF({TUR},"İç pano")', P6001["ic_pano"], "Sayım sayfasından.", formul_b=True)
    girdi("ic_kosu", "İç duvar koşusu", g["ic_kosu"], P6001["ic_kosu"], "Plandan.")
    girdi("uclu_h", "Üçlü H (iç–iç T birleşimi)", f'=COUNTIFS({TUR},"Bağlantı",Sayım!$D${r0}:$D${r1},"İç duvar T birleşimi")', P6001["uclu_h"], "Sayım sayfasından.", formul_b=True)
    girdi("dis_boy", "Dış pano boyları toplamı (m)", f'=SUMIFS({BOY},{TUR},"Dış pano")/1000', P6001["dis_boy"], "Sayım sayfasından.", fmt="0.00", formul_b=True)
    girdi("ic_boy", "İç pano boyları toplamı (m)", f'=SUMIFS({BOY},{TUR},"İç pano")/1000', P6001["ic_boy"], "Sayım sayfasından.", fmt="0.00", formul_b=True)
    girdi("aks", "Mahya boyunca pano aksı sayısı", g["aks"], P6001["aks"], "Mahyaya paralel dış duvarın pano adedi + 1.")
    p02 = next((k for k, v in m["kodlar"].items() if v["tur"] == "Pencere" and (v["dims"] or [0, 0])[:2] == [1200, 1800]), "P02")
    girdi("p02", "120×180 pencere adedi", f'=COUNTIF({KOD},"{p02}")', P6001["p02"], "Sayım sayfasından.", formul_b=True)
    grup("Girdiler · kural katsayıları")
    girdi("alt_parca", "Alt çerçeve parça boyu (m)", 3.75, 3.75, "P6001 gerçek listesi: 100'lük ve 60'lık alt çerçeve L 3750.", fmt="0.00")
    girdi("kapatma", "Çatı kaplamasının kapatma genişliği (m)", 1.0, P6001["kapatma"], "Lokman, 4 Ekim 15:18: tek kat panel 100 cm kapatır. 2021 ürününde 87 cm.", fmt="0.00")
    girdi("tasma", "Kaplamanın saçaktan taşması (m)", 0.05, 0.05, "Mühendis, 4 Ekim 12:14: kaplama saçaktan 5 cm taşar.", fmt="0.00")
    girdi("sac_net", "Sac parçasının bindirmeden sonraki boyu (m)", 2.5, 2.5, "Lokman, 3 Ekim 13:27: saçak, mahya ve alın sacında bindirme 30 cm; L 2800 parçada uzunluk ÷ 2,5.", fmt="0.00")
    girdi("asik_aralik", "Aşık aralığı (m)", 0.6, 0.6, "Montaj paftası notu: aşıklar 60 cm aralıklı.", fmt="0.00")
    girdi("asik_parca", "Aşık parça boyu (m)", 3.75, 3.75, "P6001 gerçek listesi: omega aşık L 3750. (Çelik listelerinde L 4200.)", fmt="0.00")
    girdi("asik_zayiat", "Aşık zayiatı", 0.05, 0.05, "Gerçek listelere uydurulan pay; teyit bekliyor.", fmt="0%")
    girdi("alcipan_m2", "Alçıpan levhası (m²)", 3.0, 3.0, "1,2 × 2,5 m.", fmt="0.00")
    girdi("silte_m2", "Şilte (m²/adet)", 9.6, 9.6, "Çelik listeleri: çatı şiltesi 9,6 m²/adet.", fmt="0.00")
    girdi("dere_net", "Dere parçasının bindirmeden sonraki boyu (m)", 3.8, 3.8, "Excel hesaplama motoru (CTKBF): L 4000 − 20 cm bindirme.", fmt="0.00")
    grup("Bağlantı profilleri")
    hesap("h_dis", "2'li H, dış pano ekleri", "{dis_pano}-{dis_kosu}", "Her koşuda pano adedi − 1 ek. P6001 gerçek listesi: H 100 = 19.", P6001_LISTE["h_dis"])
    hesap("h_ic", "2'li H, iç pano ekleri", "{ic_pano}-{ic_kosu}-{uclu_h}", "İç ekler − üçlü H'nin yerini aldığı ekler. P6001 gerçek listesi: H 60 = 5.", P6001_LISTE["h_ic"])
    hesap("alt_dis", "Alt çerçeve, dış duvar (adet)", "ROUNDUP({dis_boy}/{alt_parca},0)", "Dış pano boyları toplamı ÷ parça boyu, yukarı yuvarlanır. Kapı boşluğu düşülmez (P6001 listesine göre).", P6001_LISTE["alt_dis"])
    hesap("alt_ic", "Alt çerçeve, iç duvar (adet)", "ROUNDUP({ic_boy}/{alt_parca},0)", "İç pano boyları toplamı ÷ parça boyu.", P6001_LISTE["alt_ic"])
    hesap("kor_kasa", "Pencere kör kasası 120/180", "{p02}", "P6001 gerçek listesi: yalnız 120/180 pencerelerde kör kasa (2 pencere, 2 kasa). Teyit bekliyor.", P6001_LISTE["kor_kasa"])
    grup("Çatı (beşik)")
    hesap("makas_bas", "Baş makas", "2", "İki kalkan duvarın üstünde.", P6001_LISTE["makas_bas"])
    hesap("makas_orta", "Orta makas", "{aks}-2", "Makaslar pano akslarında. P6001 gerçek listesi: 6 pano boyunca 2 baş + 5 orta.", P6001_LISTE["makas_orta"])
    hesap("egik", "Eğimli boy, mahyadan saçak ucuna (m)", "({aciklik}/2+{sacak})/100*SQRT(1+{egim}^2)", "Geometri.", fmt="0.000")
    hesap("mahya_boy", "Mahya boyu (m)", "({bina_boyu}+2*{sacak})/100", "Bina boyu + iki uçta saçak.", fmt="0.00")
    hesap("cati_m2", "Çatı yüzeyi (m²)", "2*{egik}*{mahya_boy}", "İki yüz. IFC modelindeki örtü alanıyla aynı olmalı.", fmt="0.00")
    hesap("kaplama", "Çatı kaplaması, tek kat panel (adet)", "2*ROUNDUP({mahya_boy}/{kapatma},0)", "Her yüzde mahya boyu ÷ kapatma genişliği. P6001 gerçek listesi: 20 adet.", P6001_LISTE["kaplama"])
    hesap("kaplama_boy", "Kaplama boyu (m)", "{egik}+{tasma}", "Eğimli boy + taşma. P6001 gerçek listesinde L 4350 (bu formül 4,44 verir; liste 9 cm kısa).", fmt="0.000")
    hesap("mahya", "Mahya parçası L 2800 (adet)", "ROUNDUP({mahya_boy}/{sac_net},0)", "Mahya boyu ÷ 2,5. P6001 gerçek listesinde 3: bindirmesiz hesapla (8,23 ÷ 2,8) tutuyor, 30 cm bindirmeyle tutmuyor.", P6001_LISTE["mahya"])
    hesap("sacak_saci", "Saçak sacı L 2800 (adet)", "ROUNDUP(2*{mahya_boy}/{sac_net},0)", "İki saçak kenarı ÷ 2,5. P6001 gerçek listesi: 7.", P6001_LISTE["sacak_saci"])
    hesap("alin_saci", "Alın sacı L 2800 (adet)", "ROUNDUP(4*{egik}/{sac_net},0)", "Dört kalkan kenarı ÷ 2,5. P6001 gerçek listesi: 8.", P6001_LISTE["alin_saci"])
    hesap("asik_sira", "Aşık sırası, yüz başına", "ROUNDDOWN({egik}/{asik_aralik},0)+1", "Eğimli boy ÷ aralık + 1.")
    hesap("asik", "Omega aşık (adet)", "ROUNDUP(2*{asik_sira}*{mahya_boy}*(1+{asik_zayiat})/{asik_parca},0)", "İki yüz × sıra × mahya boyu × (1 + zayiat) ÷ parça boyu. P6001 gerçek listesi: 37.", P6001_LISTE["asik"])
    grup("Tavan ve izolasyon")
    hesap("brut", "Brüt alan (m²)", "{x}*{y}/10000", "Dış ölçülerden.", fmt="0.00")
    hesap("alcipan", "Tavan alçıpanı (adet)", "ROUNDUP({brut}/{alcipan_m2},0)", "Brüt alan ÷ levha alanı. P6001 gerçek listesi: 15 beyaz + 5 yeşil = 20. Beyaz–yeşil ayrımının kuralı yok.", P6001_LISTE["alcipan"])
    hesap("silte", "Çatı şiltesi (adet)", "ROUNDUP({brut}/{silte_m2},0)", "Brüt alan ÷ 9,6. P6001 gerçek listesi: 7.", P6001_LISTE["silte"])
    grup("Dere (çelik listelerinin kuralıyla; prefabrik listede sağlaması yok)")
    hesap("dere", "Metal dere L 4000 (adet)", "ROUNDUP(2*{mahya_boy}/{dere_net},0)", "İki saçak kenarı ÷ 3,8 (Excel hesaplama motoru; üç gerçek çelik listede tutuyor).")
    hesap("inis", "Dere inişi (adet)", "4+IF({mahya_boy}>=10,2,0)", "Lokman, 3 Ekim 13:27: her köşeye bir iniş; 10 m ve daha uzun cephenin ortasına bir tane daha.")
    hesap("dere_kelepce", "Dere kelepçesi (adet)", "{dere}*8", "Excel hesaplama motoru: dere adedi × 8.")
    hesap("acik_dirsek", "Açık dirsek (adet)", "{inis}*2", "Excel hesaplama motoru: iniş başına 2.")
    hesap("inis_kelepce", "İniş kelepçesi (adet)", "{inis}*2", "Excel hesaplama motoru: iniş başına 2.")
    for j, w in enumerate([44, 12, 18, 18, 12, 110], 1):
        wm.column_dimensions[get_column_letter(j)].width = w
    wm.freeze_panes = "B5"

    # ------------------------------------------------------------------ YÜKLEME
    wy = wb.create_sheet("YÜKLEME", 0)
    yapi = ie["yapi"]; kal = f"dış {m['t_dis'] * 10:g} mm, iç {m['t_ic'] * 10:g} mm"
    _yaz(wy, 1, 1, "PREFABRİK BİNA MALZEMELERİ", BASLIK, kenar=False); _yaz(wy, 1, 6, "YÜKLEME LİSTESİ", BASLIK, kenar=False)
    _yaz(wy, 3, 1, f"{ie.get('talep', {}).get('musteri_adi', '')}  /  dış ölçü {g['x']:g} × {g['y']:g} cm  /  duvar yüksekliği {m['H0']:g} cm  /  {kal}  /  verandasız, beşik çatı", KALIN, kenar=False)
    _yaz(wy, 4, 1, f"06.10.2026 — İş Emri {ie['is_emri_no']} r{int(ie.get('revizyon', 1)):02d}. Tek bina. Kapılı ve pencereli panolar tek pano; doğramalar ayrıca sayılır. "
                   "TASLAK: üretim ve sevkiyat listesi değildir.", KUCUK, kenar=False)
    for j, b in enumerate(["Sıra", "Kod", "Malzeme", "Ölçü (mm)", "Adet", "Birim", "Dayanak", "Açıklama"], 1):
        _yaz(wy, 6, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    r = 7; sira = 0; adet_hucreleri = {"Dış pano": [], "İç pano": []}

    def ygrup(ad):
        nonlocal r
        for j in range(1, 9):
            _yaz(wy, r, j, ad if j == 1 else None, KALIN, D_GRUP)
        r += 1

    def ysatir(kod, malzeme, olcu, adet, birim, dayanak, aciklama, tur=None):
        nonlocal r, sira
        sira += 1
        dolgu = {"Sayım": None, "Kural": None, "Liste kuralı": D_LISTE, "Çıkarım": D_CIKARIM}[dayanak]
        for j, v in enumerate([sira, kod, malzeme, olcu, adet, birim, dayanak, aciklama], 1):
            _yaz(wy, r, j, v, fill=dolgu, hiza=SOL if j in (3, 8) else ORTA, fmt="0" if j == 5 else None)
        if tur in adet_hucreleri:
            adet_hucreleri[tur].append(f"E{r}")
        r += 1

    def M(anahtar):
        return f"=Metraj!B{A[anahtar]}"

    lento = m["H0"] - 200.0
    pen_not = {(1200, 1200): f"Pencere ayrı kalem. Gerçek listelerde üst panel (h {lento * 10:g}) + alt panel (h 800) diye yazılır.",
               (1200, 1800): f"Pencere ayrı kalem. Gerçek listelerde üst panel (h {lento * 10:g}) + alt panel (h 200) diye yazılır.",
               (600, 400): "Vasistas ayrı kalem. Gerçek listelerde oyma panel diye yazılır."}

    def pano_satirlari(tur):
        for kod in sorted(k for k, v in m["kodlar"].items() if v["tur"] == tur):
            v = m["kodlar"][kod]; lab = v["label"]; boy = (v["dims"] or [0])[0]
            ack = ""
            for (en, yuk), n in pen_not.items():
                if f"{en // 10}×{yuk // 10}" in lab:
                    ack = n
            if "kapılı" in lab.lower():
                ack = "Kapı ayrı kalem; pano tek parça sayıldı."
            if abs(boy - 1255) > 0.5 and not ack:
                ack = ("Ölçüye kesim: cephenin son panosu." if tur == "Dış pano" else
                       "Birleşimde duvar yüzüne dayanan ya da ölçüye kesilen pano; boy çizimdeki net boydur.")
            if boy < 400:
                ack += " DAR PARÇA: böyle mi üretilecek, komşu panoyla mı paylaştırılacak (açık konu 1)."
            ysatir(kod, lab, olcu_yazi(v["dims"]), f'=COUNTIF({KOD},"{kod}")', "adet", "Sayım", ack.strip(), tur)

    ygrup(f"Dış duvar panoları ({m['t_dis'] * 10:g} mm)"); pano_satirlari("Dış pano")
    ygrup(f"İç duvar panoları ({m['t_ic'] * 10:g} mm)"); pano_satirlari("İç pano")
    ygrup("Kapı ve pencereler")
    for tur in ("Pencere", "Kapı"):
        for kod in sorted(k for k, v in m["kodlar"].items() if v["tur"] == tur):
            v = m["kodlar"][kod]
            olc = olcu_yazi(v["dims"]) if tur == "Pencere" else f"Açıklık {v['dims'][0]:g}"
            ysatir(kod, v["label"], olc, f'=COUNTIF({KOD},"{kod}")', "adet", "Sayım", v["note"] if "otomatik" not in v["note"] else "Plandaki açıklıktır; ürün ölçüsü siparişte teyit edilmeli.")
    ysatir("—", "Pencere kör kasası 120/180", "1200 × 1800", M("kor_kasa"), "adet", "Liste kuralı", "P6001 gerçek listesinde yalnız 120/180 pencerelerde. Teyit bekliyor.")
    ygrup("Bağlantı profilleri ve alt çerçeve")
    for kod in sorted(k for k, v in m["kodlar"].items() if v["tur"] == "Bağlantı"):
        v = m["kodlar"][kod]
        ysatir(kod, v["label"], olcu_yazi(v["dims"]), f'=COUNTIF({KOD},"{kod}")', "adet", "Çıkarım",
               f"Adet plandan sayıldı. Kesit {m['t_dis'] * 10:g} / {m['t_ic'] * 10:g} mm duvar için çıkarım; profil ve kesim boyu mühendis teyidi bekliyor.")
    ysatir("—", f"2'li H profil, dış pano eki ({m['t_dis'] * 10:g}'lık)", f"L {m['H0'] * 10:g}", M("h_dis"), "adet", "Liste kuralı", "Her cephede pano adedi − 1. P6001 gerçek listesinde tutuyor (19). Kulaklı / dübelli ayrımının kuralı yok.")
    ysatir("—", f"2'li H profil, iç pano eki ({m['t_ic'] * 10:g}'lük)", f"L {m['H0'] * 10:g}", M("h_ic"), "adet", "Liste kuralı", "İç ekler − üçlü H'nin yerini aldığı ekler. P6001 gerçek listesinde tutuyor (5).")
    ysatir("—", f"Alt çerçeve, dış duvar ({m['t_dis'] * 10:g}'lık)", "L 3750", M("alt_dis"), "adet", "Liste kuralı", "Dış duvar boyu ÷ 3,75 m. P6001 gerçek listesinde tutuyor (9).")
    ysatir("—", f"Alt çerçeve, iç duvar ({m['t_ic'] * 10:g}'lük)", "L 3750", M("alt_ic"), "adet", "Liste kuralı", "İç duvar boyu ÷ 3,75 m. P6001 gerçek listesinde tutuyor (3).")
    ygrup("Çatı (beşik, %33; makaslar kısa kenarı geçer)")
    ysatir("—", "Baş makas", f"açıklık {g['aciklik'] * 10:g}", M("makas_bas"), "adet", "Liste kuralı", "Kalkan duvarların üstünde. Makas boyu ve kesiti mühendisten.")
    ysatir("—", "Orta makas", f"açıklık {g['aciklik'] * 10:g}", M("makas_orta"), "adet", "Liste kuralı", "Pano akslarında; son aralık 114 cm. P6001 gerçek listesinde tutuyor (5).")
    ysatir("—", "Çatı kaplaması, tek kat panel", f'=CONCATENATE("L ",ROUNDUP(Metraj!B{A["kaplama_boy"]}*100,0)*10)', M("kaplama"), "adet", "Kural", "Kapatma 100 cm (Lokman, 4 Ekim); boy eğimli boy + 5 cm taşma. Her yüzde mahya boyu ÷ kapatma genişliği.")
    ysatir("—", "Mahya", "L 2800", M("mahya"), "adet", "Kural", "30 cm bindirme (Lokman, 3 Ekim). P6001 gerçek listesinde bu kural 4 verir, liste 3: bindirme teyit edilmeli.")
    ysatir("—", "Saçak sacı", "L 2800", M("sacak_saci"), "adet", "Kural", "İki saçak kenarı, 30 cm bindirme. P6001 gerçek listesinde tutuyor (7).")
    ysatir("—", "Alın sacı", "L 2800", M("alin_saci"), "adet", "Kural", "Dört kalkan kenarı, 30 cm bindirme. P6001 gerçek listesinde tutuyor (8).")
    ysatir("—", "Omega aşık", "L 3750", M("asik"), "adet", "Liste kuralı", "60 cm aralık, %5 zayiat. P6001 gerçek listesinde tutuyor (37). Zayiat payı teyit bekliyor.")
    ysatir("—", "Çatı şiltesi", "9,6 m²/adet", M("silte"), "adet", "Liste kuralı", "Brüt alan ÷ 9,6. P6001 gerçek listesinde tutuyor (7).")
    ygrup("Tavan")
    ysatir("—", "Tavan alçıpanı", "12 × 1200 × 2500", M("alcipan"), "adet", "Liste kuralı", "Brüt alan ÷ 3 m². P6001 gerçek listesinde tutuyor (20). Beyaz–yeşil ayrımı ve tavan H'ı için kural yok.")
    ygrup("Dere")
    ysatir("—", "Metal dere", "L 4000", M("dere"), "adet", "Çıkarım", "Çelik listelerinin kuralı (÷ 3,8). Prefabrik pakette dere var mı, teyit.")
    ysatir("—", "Dere kelepçesi", "", M("dere_kelepce"), "adet", "Çıkarım", "Dere adedi × 8.")
    ysatir("—", "İniş borusu", "L 3000", M("inis"), "adet", "Çıkarım", "Dört köşe + 10 m'den uzun iki saçak cephesinin ortası (Lokman, 3 Ekim; çelik).")
    ysatir("—", "İniş aparatı", "", M("inis"), "adet", "Çıkarım", "İniş başına 1.")
    ysatir("—", "Kapalı dirsek", "", M("inis"), "adet", "Çıkarım", "İniş başına 1.")
    ysatir("—", "Açık dirsek", "", M("acik_dirsek"), "adet", "Çıkarım", "İniş başına 2.")
    ysatir("—", "İniş kelepçesi", "", M("inis_kelepce"), "adet", "Çıkarım", "İniş başına 2.")
    r += 1
    for j in range(1, 9):
        _yaz(wy, r, j, None, KALIN, D_TOPLAM)
    _yaz(wy, r, 1, "TOPLAM DUVAR PANOSU", KALIN, D_TOPLAM)
    toplam_r = r
    _yaz(wy, r, 5, "=" + "+".join(adet_hucreleri["Dış pano"] + adet_hucreleri["İç pano"]), KALIN, D_TOPLAM, ORTA, "0")
    _yaz(wy, r, 6, "adet", KALIN, D_TOPLAM, ORTA)
    _yaz(wy, r, 8, f'=IF(E{r}=COUNTIF({TUR},"Dış pano")+COUNTIF({TUR},"İç pano"),"Sayım sayfasındaki pano satırlarıyla aynı.","UYARI: Sayım sayfasıyla tutmuyor.")', KALIN, D_TOPLAM, SOL)
    r += 2
    notlar = [
        "Pano ölçüsü: boy × kalınlık × yükseklik. Pencereler: en × yükseklik. Ölçüler mm.",
        "Dayanak — Sayım: modelden sayıldı (DXF ve IFC ile aynı). Kural: Lokman'ın ya da mühendisin söylediği kural. "
        "Liste kuralı (sarı): P6001'in gerçek yükleme listesinden çıkarıldı, o listede tutuyor, teyit bekliyor. Çıkarım (turuncu): başka sistemin kuralı ya da kesit varsayımı.",
        "Adetlerin hepsi formüldür: sayılanlar Sayım sayfasından, türetilenler Metraj sayfasından gelir. Metraj'da her kuralın P6001 gerçek listesiyle sağlaması var.",
        "Bu listede olmayanlar Açık konular sayfasında: çektirme U'su, duvar ve saçak omegası, pencere U'su, tavan H'ı, dış–iç üçlü H, vida, elektrik, tesisat, boya.",
    ]
    for n in notlar:
        _yaz(wy, r, 1, n, KUCUK, kenar=False); r += 1
    for j, w in enumerate([6, 7, 40, 20, 8, 8, 13, 100], 1):
        wy.column_dimensions[get_column_letter(j)].width = w
    wy.freeze_panes = "A7"

    # ------------------------------------------------------------------ Açık konular
    wa = wb.create_sheet("Açık konular")
    _yaz(wa, 1, 1, "AÇIK KONULAR — listeye girmeyenler ve teyit bekleyenler", KALIN, kenar=False)
    for j, b in enumerate(["No", "Konu", "Bu yapıda durum", "P6001 gerçek listesinde (60 m²)", "Gereken"], 1):
        _yaz(wa, 3, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    konular = [
        ("Dar kalan son parça", "Arkada 30 cm dış pano (2 adet), arka yatak odalarının duvarlarında 22 cm iç pano (2 adet).", "Yok (tam modül).", "Böyle mi üretilecek, bir önceki panoyla mı paylaştırılacak? En az pano boyu."),
        ("160 mm dış ve 100 mm iç duvarın bağlantı profilleri", "Köşe 160×80, iç köşe 100×50, üçlü H 100, 2'li H'lar: kesitler çıkarım.", "Köşe direği 100×100 ×6, 60×60 ×1; H 100 ×19; H 60 ×5; 3'lü H 60 ×2.", "Mühendisten profil listesi ve kesim boyları."),
        ("Dış–iç üçlü H (iç duvarın dış duvara dayandığı yer)", "6 birleşim: 3'ü dış pano ekinde, 3'ü pano ortasında. Listeye yazılmadı.", "3'lü H 100×100×60 ×2 (kulaklı 1 + dübelli 1).", "Pano ortasına gelen birleşim nasıl bağlanıyor?"),
        ("Kulaklı / dübelli H ayrımı", "Ayrılmadı; 2'li H tek kalem.", "H 100: 8 kulaklı + 11 dübelli.", "Hangi eklerde kulaklı?"),
        ("Çektirme U'su", "Listeye yazılmadı.", "100'lük ×5, 60'lık ×3 (L 2440).", "Nerede sayılır?"),
        ("Duvar omegası ve saçak omegası", "Listeye yazılmadı.", "100'lük duvar omegası ×3, 60'lık ×15, saçak omegası ×7.", "Neyi taşıyor, nasıl sayılır?"),
        ("Pencere U'su", "Listeye yazılmadı.", "×6 (5 pencere).", "Pencere başına kural."),
        ("Tavan H'ı; alçıpanda beyaz–yeşil ayrımı", "Alçıpan toplam adet verildi; H ve ayrım yok.", "Tavan H'ı 1185 ×70; alçıpan 15 beyaz + 5 yeşil.", "H adedi kuralı; yeşil levha hangi mahallerde?"),
        ("Makas boyu ve kesiti", "Adet verildi (2 baş + 8 orta); boy yazılmadı.", "Baş ve orta makas L 8230 (açıklık 7730).", "10,50 m açıklık için makas ölçüsü ve statik."),
        ("Mahya bindirmesi", "30 cm bindirmeyle 5 adet.", "3 adet (bindirmesiz hesapla tutuyor).", "Prefabrikte mahya bindirmesi kaç cm?"),
        ("Pencereli panonun yazılışı", "Tek pano + ayrı doğrama (Astra listesindeki gibi).", "Üst panel + alt panel + oyma panel ayrı kalemler.", "Üretim listesi hangi biçimi istiyor?"),
        ("Vida, elektrik, sıhhi tesisat, boya", "Listeye yazılmadı.", "Dört ayrı grup.", "Çelik listelerindeki kurallar prefabrikte de geçerli mi?"),
        ("Pencere ölçüleri", "El çizimindeki 240/240, 150/240, 100/240, 180/120, 60/60 standart ölçülere çevrildi.", "—", "Standart mı, özel doğrama mı?"),
    ]
    for i, (a, b, c_, d) in enumerate(konular, 1):
        for j, v in enumerate([i, a, b, c_, d], 1):
            _yaz(wa, 3 + i, j, v, hiza=ORTA if j == 1 else SOL)
        wa.row_dimensions[3 + i].height = 32
    for j, w in enumerate([5, 42, 62, 52, 58], 1):
        wa.column_dimensions[get_column_letter(j)].width = w
    wa.freeze_panes = "A4"
    return wb, {"metraj_satir": A, "yukleme_toplam_satir": toplam_r, "sayim_ilk": r0, "sayim_son": r1}


def uret(is_emri_yolu: Path) -> dict:
    ie = json.loads(Path(is_emri_yolu).read_text(encoding="utf-8"))
    m = model(ie)
    wb, yer = kitap(ie, m)
    hedef = CIKTI / ie["is_emri_no"] / "yukleme"; hedef.mkdir(parents=True, exist_ok=True)
    yol = hedef / f"Portman_{ie['is_emri_no']}_r{int(ie.get('revizyon', 1)):02d}_Yukleme_Listesi.xlsx"
    wb.save(yol)
    from collections import Counter
    return {"yol": str(yol), "yer": yer, "kod_sayim": dict(sorted(Counter(s["kod"] for s in m["satirlar"]).items())), "kontrol": m["kontrol"], "girdi": m["girdi"]}


if __name__ == "__main__":
    yol = Path(sys.argv[1]) if len(sys.argv) > 1 else BURASI / "is_emri_3arti1_1150x1050.json"
    if not yol.is_absolute() and not yol.exists():
        yol = BURASI / yol
    print(json.dumps(uret(yol), ensure_ascii=False, indent=1))

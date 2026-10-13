"""Prefabrik kör test · kendi yükleme listemiz (İE-11 PR-2, M6; protokol `docs/sartname/prefabrik_kor_test.md` adım 3).

    python3 deneme/2026-10-13_pk_01/yukleme_listesi.py <İş Emri.json> <çıktı klasörü> [--pafta <pafta_katmani.json>]

Gerçek liste AÇILMADAN yazılır ve depoya kaydedilir; kıyas ondan sonra. Hiçbir adet elle yazılmaz:
  - **Sayım**: canlı pano Motor'unun planından (`motor.geometry.kur`; panolar, doğramalar, köşe, dikme, pano ekleri / üçlü H).
  - **Kural**: kural kitabından (`kurallar/liste/kurallar.json`; yöntem B formülleri genel motorla aynı: saçak ve alın sacı 2,5 − 0,2
    bindirme, aşık sırası eğimli boy ÷ 0,60 + saçak + mahya, tavan omegası alan ÷ 0,60 ÷ 3,75, baş makas Z'si kenar kenar ÷ 2,8, dere ve iniş)
    ve 6 Ekim prefabrik denemesinin P6001'e (Uğur Başarır 60 m² listesi, 21 Eylül) uydurulmuş kurallarından ("liste kuralı").
  - **Kural yok**: gerçek prefabrik listelerinde bulunan ama kuralı olmayan kalemler adetsiz yazılır (K21: uydurulmaz).
  - **Pafta katmanı** (isteğe bağlı): o projenin montaj paftasından okunan adetler ayrı sütunda; varsayılan hesabı değiştirmez (protokol §5).
Gruplar gerçek listelerin gruplarıdır (1001 panel … 1012 boya; `docs/notlar/2026-09-21_Yukleme_Listesi_Analizi.md`). LLM yok, ağ yok.
"""
from __future__ import annotations

import argparse
import json
import math
import sys
import time
from collections import Counter
from pathlib import Path

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))

from motor import geometry as G  # noqa: E402

GRUPLAR = {"1001": "Panel", "1002": "Metal aksam", "1003": "Kapı ve pencere", "1004": "Asma tavan", "1005": "Makas", "1006": "Çatı",
           "1007": "Elektrik", "1008": "Vida", "1009": "Depo", "1010": "Sıhhi tesisat", "1011": "Dere", "1012": "Boya"}
TAVAN = lambda x: int(math.ceil(x - 1e-9))
PENCERE_UST_KOTU_CM = 200.0           # Motor: pencere üst kotu / kapı üstü 200 (lento = h − 200)
PENCERE_ESIK_CM = {"120x120": 80.0, "120x180": 20.0, "60x40": 160.0}


def _kural(k: dict, *yol):
    v = k
    for a in yol:
        v = v[a]
        if isinstance(v, dict) and "deger" in v and a == yol[-1]:
            return v["deger"]
        if isinstance(v, dict) and "deger" in v:
            v = v["deger"]
    return v


def nominal(cm: float) -> str:
    """Listedeki yazılış: aks boyu 125,5 → 1250, T'de kısalan 122,5 → 1220 (P6001 listesi; 5 mm H payı). Öteki boylar Motor'un net boyu."""
    if abs(cm - G.MOD) < 0.05 or abs(cm - (G.MOD - 3)) < 0.05:
        return mm(cm - 0.5)
    return mm(cm)


def mm(cm: float) -> str:
    v = round(cm * 10, 1)
    return f"{v:g}".replace(".", ",")


def pano_ekleri(plan) -> list:
    """Bir duvardaki ardışık iki panonun her eki; ekte üçlü H varsa (`plan.h_baglantilar`) "uclu", yoksa "pano_eki". Yalnız okur.
    (İE-12 dalındaki `G.pano_ekleri` ile aynı kural; PK-01 paftasının H planıyla 19 + 10 / 5 + 4.)"""
    cik = []
    for d in plan.duvarlar:
        for a, _b in zip(d.panolar, d.panolar[1:]):
            x, y = (a.a1, d.c) if d.yon == "x" else (d.c, a.a1)
            uclu = any(abs(h.x - x) <= d.h + 3.5 and abs(h.y - y) <= d.h + 3.5 for h in plan.h_baglantilar)
            cik.append({"duvar": d.tur, "cins": "uclu" if uclu else "pano_eki"})
    return cik


def model(ie: dict) -> dict:
    """Motor planından sayım girdileri (cm)."""
    plan = G.kur(json.loads(json.dumps(ie)))
    y = ie["yapi"]
    H0 = float(y.get("duvar_yuksekligi_cm", 250))
    t_dis, t_ic = 2 * next(d.h for d in plan.duvarlar if d.tur == "dis"), 2 * next(d.h for d in plan.duvarlar if d.tur == "ic")
    dis = [d for d in plan.duvarlar if d.tur == "dis"]; ic = [d for d in plan.duvarlar if d.tur == "ic"]
    ekler = pano_ekleri(plan)
    from shapely.geometry import box
    from shapely.ops import unary_union
    kapali = [m for m in plan.mahaller.values() if not (m.acik or m.tur in G.ACIK_MAHAL_TURLERI)]
    acik = [m for m in plan.mahaller.values() if (m.acik or m.tur in G.ACIK_MAHAL_TURLERI)]
    g = unary_union([plan.footprint()] + [m.poly for m in plan.mahaller.values() if m.poly is not None] + [box(*r) for r in plan.dikmeler])
    bx0, by0, bx1, by1 = g.bounds
    return {"plan": plan, "H0": H0, "t_dis": t_dis, "t_ic": t_ic, "dis": dis, "ic": ic, "ekler": ekler, "kapali": kapali, "acik": acik,
            "W": bx1 - bx0, "D": by1 - by0, "brut_m2": round((bx1 - bx0) * (by1 - by0) / 1e4, 3)}


def liste(ie: dict, k: dict, pafta: dict | None = None) -> dict:
    m = model(ie)
    plan, H0 = m["plan"], m["H0"]
    kucuk = lambda x: str(x).replace("İ", "i").replace("I", "ı").lower()
    oz = {kucuk(kk): kucuk(v) for kk, v in (ie.get("ozellikler") or {}).items()}
    alici = [kucuk(x) for x in ((ie.get("alici_ait") or {}).get("kalemler") or [])]
    satirlar = []

    def ekle(grup, kalem, olcu, adet, dayanak, formul, kaynak="", birim="adet", pafta_adet=None, pafta_notu=""):
        satirlar.append({"grup": grup, "grup_adi": GRUPLAR[grup], "kalem": kalem, "olcu": olcu, "adet": adet, "birim": birim, "dayanak": dayanak,
                         "formul": formul, "kaynak": kaynak, "pafta": pafta_adet, "pafta_notu": pafta_notu})

    P = pafta or {}
    pf = lambda anahtar: (P.get(anahtar) or {}).get("adet")
    pn = lambda anahtar: (P.get(anahtar) or {}).get("not", "")

    # ── 1001 Panel ──────────────────────────────────────────────────────────
    # Gerçek liste biçimi (P6001 / Uğur Başarır 60 m² listesi, 21 Eylül): tam pano 1250 nominal (aks 1255); pencereli pano üst + alt panel;
    # vasistas oyma panel; kapılı karkaslı panel. İç duvarda T'ye gelen 122,5 / 120,5'lik pano listede standart 1250 sayılıyordu (liste kuralı).
    sayac = Counter(); notlar = {}
    for d in plan.duvarlar:
        kal = "100'lük" if d.tur == "dis" else "60'lık"
        kal_mm = round(2 * d.h * 10)
        for p in d.panolar:
            L = p.boy; a = p.acilim
            if a is None:
                if abs(L - G.MOD) < 0.05:
                    anahtar = (f"{kal} yapıştırma dolu pano", f"{kal_mm} × 1250 × {mm(H0)}", "Sayım", "tam pano (aks 1255, nominal 1250)")
                elif d.tur == "ic" and G.MOD - 5.01 <= L < G.MOD - 0.05:
                    anahtar = (f"{kal} yapıştırma dolu pano", f"{kal_mm} × 1250 × {mm(H0)}", "Liste kuralı",
                               "iç duvarın T'ye gelen panosu (122,5 / 120,5) listede standart 1250 (P6001 listesi)")
                else:
                    anahtar = (f"{kal} dolu pano (özel boy)", f"{kal_mm} × {mm(L)} × {mm(H0)}", "Sayım",
                               "ölçüye kesilmiş / özel pano; boy çizimdeki net boy (sipariş boyu kuralı yok)")
                sayac[anahtar] += 1
            elif a.tur == "pencere":
                en, boy = G.PENCERE_OLCULERI[a.olcu]
                if a.olcu == "60x40":
                    sayac[(f"Vasistas oyma panel {kal}", f"60/40 · pano {nominal(L)}", "Liste kuralı", "vasistaslı pano listede oyma panel")] += 1
                else:
                    ust = H0 - PENCERE_UST_KOTU_CM
                    sayac[(f"Pencere üst paneli {kal}", f"{kal_mm} × {nominal(L)} × {mm(ust)}", "Liste kuralı", "pencereli pano listede üst + alt panel (P6001 listesi)")] += 1
                    sayac[(f"Pencere alt paneli {kal} ({en}/{boy})", f"{kal_mm} × {nominal(L)} × {mm(PENCERE_ESIK_CM[a.olcu])}", "Liste kuralı",
                           "alt panel yüksekliği pencere eşiği")] += 1
            elif a.tur in ("ic_kapi", "dis_kapi"):
                sayac[(f"{kal} karkaslı {'dış ' if a.tur == 'dis_kapi' else ''}kapılı panel", f"{kal_mm} × {nominal(L)} × {mm(H0)}", "Sayım",
                       "kapı boşluklu pano" + (" (kapı alıcıya ait; pano bizim)" if a.dograma == "alıcıya ait" else ""))] += 1
            else:
                sayac[(f"{kal} geçişli pano", f"{kal_mm} × {mm(L)} × {mm(H0)}", "Sayım", "kapısız geçiş")] += 1
    for (kalem, olcu, dayanak, kaynak), n in sorted(sayac.items(), key=lambda x: (x[0][0], x[0][1])):
        ekle("1001", kalem, olcu, n, dayanak, "Motor planından sayım", kaynak)

    # ── 1002 Metal aksam ────────────────────────────────────────────────────
    kc = Counter("dis" if kk.tur == "dis" else "ic" for kk in plan.koseler)
    ekle("1002", "Köşe direği 100 × 100", f"L {mm(H0)}", kc["dis"], "Sayım", "Motor köşe parçası (B02)", "P6001 listesinde 'köşe direği 100×100'")
    if kc["ic"]:
        ekle("1002", "Köşe direği 60 × 60 (iç köşe)", f"L {mm(H0)}", kc["ic"], "Sayım", "Motor iç köşe (B01)")
    if plan.dikmeler:
        ekle("1002", "Flanşlı kutu profil 100 × 100 (veranda dikmesi)", f"L {mm(H0)}", len(plan.dikmeler), "Sayım", "İş Emri `dikmeler`")
    ec = Counter((e["duvar"], e["cins"]) for e in m["ekler"])
    ekle("1002", "H 100 (dış pano eki)", f"100 × 100 × {mm(H0)}", ec[("dis", "pano_eki")], "Sayım", "her dış ekte bir H",
         "kulaklı / dübelli ayrımının kuralı yok", pafta_adet=pf("h_dis"), pafta_notu=pn("h_dis"))
    ekle("1002", "H 60 standart (iç pano eki)", f"60 × 60 × {mm(H0)}", ec[("ic", "pano_eki")], "Sayım", "her iç ekte bir H", "", pafta_adet=pf("h_ic"), pafta_notu=pn("h_ic"))
    ekle("1002", "3'lü H 100 × 100 × 60 (dış duvar T'si)", f"L {mm(H0)}", ec[("dis", "uclu")], "Sayım", "iç duvarın dış duvara geldiği T", "kulaklı / dübelli ayrımının kuralı yok",
         pafta_adet=pf("uclu_dis"), pafta_notu=pn("uclu_dis"))
    ekle("1002", "3'lü H 60 standart (iç T)", f"L {mm(H0)}", ec[("ic", "uclu")], "Sayım", "iç–iç T birleşimi", "", pafta_adet=pf("uclu_ic"), pafta_notu=pn("uclu_ic"))
    dis_boy = sum(p.boy for d in m["dis"] for p in d.panolar) / 100; ic_boy = sum(p.boy for d in m["ic"] for p in d.panolar) / 100
    ekle("1002", "100'lük alt çerçeve", "L 3750", TAVAN(dis_boy / 3.75), "Liste kuralı", f"dış pano boyları {dis_boy:.2f} m ÷ 3,75, yukarı", "P6001 listesinde tutuyor (9)")
    ekle("1002", "60'lık alt çerçeve", "L 3750", TAVAN(ic_boy / 3.75), "Liste kuralı", f"iç pano boyları {ic_boy:.2f} m ÷ 3,75, yukarı", "P6001 listesinde tutuyor (3)")
    for kalem, olcu in (("Çektirme U'su 100", "L 2440"), ("Çektirme U'su 60", "L 2440"), ("100'lük duvar omegası", "L 2500"), ("60'lık duvar omegası", "L 3750")):
        ekle("1002", kalem, olcu, None, "Kural yok", "—", "P6001 listesinde var; nerede sayıldığı bilinmiyor (21 Eylül sorusu açık)")
    # veranda kirişi: açık kenarın dış sınırı − dikme (P6001: 4392,5 → kiriş 4292,5)
    if plan.dikmeler and m["acik"]:
        for v in m["acik"]:
            vx0, vy0, vx1, vy1 = v.poly.bounds
            yanlar = set()
            for (x0, y0, x1, y1, *_r) in plan.acik_kenarlar:
                if v.poly.buffer(1.0).contains(__import__("shapely.geometry", fromlist=["LineString"]).LineString([(x0, y0), (x1, y1)])):
                    yanlar.add("x" if abs(y1 - y0) < 1e-6 else "y")
            for eks in sorted(yanlar):
                dis_sinir = (vx1 - vx0) if eks == "x" else (vy1 - vy0)
                ekle("1002", "Kutu profil 100 × 100 (veranda kirişi)", f"L {mm(dis_sinir - 10)}", 1, "Liste kuralı",
                     f"verandanın açık kenarı boyunca dış ölçüsü {mm(dis_sinir)} − 100 (dikme)", "P6001 listesinde 4392,5 → 4292,5 ve 2510 → 2410")
    # saçak omegası (yöntem B; prefabrikte kullanılır: Lokman 7 Ekim)
    mahya_boy = None
    c = ie.get("cati") or {}
    sac = float(c.get("sacak_cm", 30)) / 100; eg = float(c.get("egim_yuzde", 33)) / 100
    mx = c.get("makas_ekseni", "x") == "x"
    aciklik, bina_boyu = (m["W"], m["D"]) if mx else (m["D"], m["W"])
    egik = (aciklik / 200 + sac) * math.sqrt(1 + eg * eg)
    mahya_boy = bina_boyu / 100 + 2 * sac
    ps, bs = _kural(k, "parca_boyu_m", "sacak_saci"), _kural(k, "bindirme_m", "sacak_saci")
    ekle("1002", "Saçak omegası", "L 2500", 2 * TAVAN(mahya_boy / ps), "Kural (uydurma)", f"her saçak kenarı {mahya_boy:.2f} m ÷ 2,5, ayrı ayrı yukarı",
         "kurallar.json › sacak_parcasi_kenar_kenar; prefabrikte saçak omegası kullanılır (Lokman, 7 Ekim)")
    ekle("1002", "Saçak sacı (30'luk)", "L 2500", TAVAN(2 * mahya_boy / (ps - bs)), "Kural", f"saçak boyu {2 * mahya_boy:.2f} m ÷ (2,5 − 0,2 bindirme)",
         "kurallar.json › bindirme_m (Lokman, 7 Ekim: 20 cm)")
    alin = 4 * egik
    ekle("1002", "Alın sacı (30'luk)", "300 × 2500", TAVAN(alin / (_kural(k, "parca_boyu_m", "alin_saci") - _kural(k, "bindirme_m", "alin_saci"))), "Kural",
         f"alın boyu 4 × {egik:.3f} = {alin:.2f} m ÷ (2,5 − 0,2)", "beşikte dört kalkan kenarı")
    bz = _kural(k, "bas_makas_z")
    ekle("1002", "Baş makas Z'si", f"L {round(bz['parca_boyu_m'] * 1000)}", 4 * TAVAN(egik / bz["parca_boyu_m"]), "Kural (uydurma)",
         f"her alın kenarı {egik:.3f} m ÷ 2,8, ayrı ayrı yukarı (4 kenar)", "kurallar.json › bas_makas_z")
    tavan_m2 = sum(mm_.alan_m2 for mm_ in m["kapali"]) + sum(mm_.alan_m2 for mm_ in m["acik"])
    ekle("1002", "Tavan omegası (60'lık)", "L 3750", TAVAN(tavan_m2 / _kural(k, "tavan_omegasi_aralik_m") / _kural(k, "parca_boyu_m", "tavan_omegasi")),
         "Kural (uydurma)", f"tavan alanı {tavan_m2:.2f} m² (kapalı + veranda) ÷ 0,60 ÷ 3,75", "kurallar.json › tavan_omegasi_aralik_m",
         pafta_adet=pf("tavan_omegasi"), pafta_notu=pn("tavan_omegasi"))
    sira = math.floor(egik / _kural(k, "asik", "aralik_m") + 1e-9) + 2
    asik_boy = 2 * sira * mahya_boy
    ekle("1002", "Omega aşık", "L 4200", TAVAN(asik_boy * (1 + _kural(k, "asik", "zayiat")) / _kural(k, "parca_boyu_m", "asik")), "Kural (uydurma)",
         f"yüz başına {sira} sıra (eğimli boy {egik:.3f} ÷ 0,60 + saçak + mahya aşığı) × 2 yüz × {mahya_boy:.2f} m = {asik_boy:.1f} m × 1,05 ÷ 4,2",
         "kurallar.json › asik, parca_boyu_m.asik", pafta_adet=pf("asik"), pafta_notu=pn("asik"))

    # ── 1003 Kapı ve pencere ────────────────────────────────────────────────
    acl = [p.acilim for p in plan.panolar if p.acilim]
    for olcu, n in sorted(Counter(a.olcu for a in acl if a.tur == "pencere").items()):
        en, boy = G.PENCERE_OLCULERI[olcu]
        ekle("1003", "PVC vasistas" if olcu == "60x40" else "PVC pencere", f"{en}/{boy}", n, "Sayım", "Motor planı")
    n_ic = sum(a.tur == "ic_kapi" for a in acl)
    ekle("1003", "Amerikan panel kapı (iç)", "80 × 198", n_ic, "Sayım", "Motor planı; ölçü P6001 listesinden", "kanat yönü D-13 (sol / sağ ayrımı sayılmadı)")
    for a in (a for a in acl if a.tur == "dis_kapi"):
        kapi_alici = a.dograma == "alıcıya ait" or "alıcı" in oz.get("dış kapı", "")
        ekle("1003", "Çelik dış kapı", f"{mm(a.genislik)} × 2050", 0 if kapi_alici else 1, "Sayım", "Motor planı" + ("; alıcıya ait → listeye girmez" if kapi_alici else ""))
    ekle("1003", "Pencere kör kasası 120/180", "1200 × 1800", sum(a.olcu == "120x180" for a in acl if a.tur == "pencere"), "Liste kuralı",
         "120/180 pencere adedi", "P6001 listesinde yalnız 120/180'de; teyit bekliyor")
    ekle("1003", "Pencere U'su", "—", None, "Kural yok", "—", "P6001 listesinde 5 pencereye 6")

    # ── 1004 Asma tavan ─────────────────────────────────────────────────────
    lev = _kural(k, "levha_m2", "alcipan")
    toplam = TAVAN(m["brut_m2"] / lev)
    yesil = TAVAN(sum(mm_.alan_m2 for mm_ in m["acik"]) / lev) if m["acik"] else 0
    ekle("1004", "Alçıpan 12 × 1200 × 2500 (beyaz)", "12 × 1200 × 2500", toplam - yesil, "Liste kuralı",
         f"brüt {m['brut_m2']:.2f} m² ÷ 3 = {toplam}, yeşil düşülür", "P6001: brüt ÷ 3 = 20 (15 beyaz + 5 yeşil); beyaz / yeşil ayrımının kuralı yok")
    ekle("1004", "Alçıpan 12 × 1200 × 2500 (yeşil)", "12 × 1200 × 2500", yesil, "Kural",
         "veranda tavanı ÷ 3", "kurallar.json › veranda_tavani (prefabrikte yeşil alçıpan, Lokman 7 Ekim); ıslak hacimde yeşil mi: kural yok")
    ekle("1004", "Tavan H'ı", "L 1185", None, "Kural yok", "—", "P6001 listesinde 60 m²'ye 70")
    kb = _kural(k, "kosebent")
    cevre = sum(mm_.poly.length for mm_ in m["kapali"]) / 100
    ic_kose = 0
    for mm_ in m["kapali"]:
        pts = list(mm_.poly.simplify(0.01).exterior.coords)[:-1]
        ic_kose += len(pts)              # dikdörtgen olmayan odada iç köşe sayısı = köşe sayısı (her köşede bir düşey köşebent)
    ekle("1004", "Köşebent", f"L {round(kb['parca_boyu_m'] * 1000)}", TAVAN((cevre + ic_kose * kb["kose_boyu_m"]) / kb["parca_boyu_m"]), "Çıkarım",
         f"(oda tavan çevreleri {cevre:.2f} m + {ic_kose} köşe × 2,7) ÷ 2,5", "kurallar.json › kosebent (Ubeydullah, 3 Ekim; çelik kuralı, prefabrikte çıkarım)")
    ekle("1004", "Çatı şiltesi (80 mm)", "9,6 m²/adet", TAVAN(m["brut_m2"] / _kural(k, "silte_adet_m2")), "Liste kuralı", f"brüt {m['brut_m2']:.2f} ÷ 9,6",
         "P6001 listesinde tutuyor (7)")

    # ── 1005 Makas ──────────────────────────────────────────────────────────
    n_makas = TAVAN((bina_boyu - 10) / G.MOD) + 1
    ekle("1005", "Baş makas", f"açıklık {mm(aciklik)}", 2, "Kural", "iki kalkan duvarın üstünde", "makas boyu ve kesiti mühendisten", pafta_adet=pf("makas_bas"), pafta_notu=pn("makas_bas"))
    ekle("1005", "Orta makas", f"açıklık {mm(aciklik)}", n_makas - 2, "Liste kuralı",
         f"makaslar pano akslarında: ⌈({mm(bina_boyu)} − 100) ÷ 1255⌉ + 1 = {n_makas}, − 2 baş", "P6001 listesinde 2 + 5; Motor IFC'si aynı kural", pafta_adet=pf("makas_orta"), pafta_notu=pn("makas_orta"))
    ekle("1005", "Baş makas betopanı 125 × 250", "125 × 250", None, "Kural yok", "—", "", pafta_adet=pf("bas_makas_betopan"), pafta_notu=pn("bas_makas_betopan"))

    # ── 1006 Çatı ───────────────────────────────────────────────────────────
    kap = oz.get("çatı kaplaması", "")
    sandvic = "sandvi" in kap
    kap_m = _kural(k, "sandvic_panel_kapatma_m" if sandvic else "tek_kat_panel_kapatma_m")
    ekle("1006", "Çatı kaplaması · " + ("sandviç panel" if sandvic else "tek kat panel"), f"L {round((egik + 0.05) * 1000)}", 2 * TAVAN(mahya_boy / kap_m), "Kural",
         f"her yüzde mahya boyu {mahya_boy:.2f} ÷ {kap_m:g} kapatma; boy eğimli {egik:.3f} + 0,05 taşma",
         "kapatma: Lokman 4 Ekim; taşma: mühendis 4 Ekim" + ("" if kap else "; kaplama İş Emri'nde yok, tek kat panel varsayıldı"))
    pm, bm = _kural(k, "parca_boyu_m", "mahya"), _kural(k, "bindirme_m", "mahya")
    ekle("1006", "Mahya", "L 2800", TAVAN(mahya_boy / (pm - bm)), "Kural", f"mahya boyu {mahya_boy:.2f} ÷ (2,8 − 0,2)", "P6001'de bindirmesiz hesap tutuyordu (3); teyit")
    ekle("1006", "Nem bariyeri / şilte (çatı)", "—", None, "Kural yok", "—", "prefabrik listelerinde şilte 1004'te sayıldı")

    # ── 1007 Elektrik · 1008 Vida · 1009 Depo · 1010 Tesisat · 1011 Dere · 1012 Boya ──
    ekle("1007", "Elektrik grubu (duy, priz, anahtar, sigorta, kablo)", "—", None, "Kural yok",
         "kural montaj paftasındaki elektrik planından sayım ister", "kurallar.json › elektrik ('AutoCAD montaj paftasından sayılır'); bu projede elektrik planı yok")
    vd = _kural(k, "vida")
    ekle("1008", "Alçıpan vidası", "—", toplam * vd["alcipan_vidasi_levha_basina"], "Çıkarım", f"{toplam} levha × {vd['alcipan_vidasi_levha_basina']}", "çelik / Excel motoru kuralı")
    ekle("1008", "Pencere vidası", "—", sum(a.tur == "pencere" for a in acl) * vd["pencere_vidasi_pencere_basina"], "Çıkarım",
         f"pencere × {vd['pencere_vidasi_pencere_basina']}", "çelik / Excel motoru kuralı")
    ekle("1008", "Öteki vidalar (trapez, çatı paneli, dübel)", "—", None, "Kural yok", "—", "prefabrik için kural yok")
    ekle("1009", "Depo", "—", None, "Kural yok", "—", "")
    tes_alici = any("tesisat" in x for x in alici)
    ekle("1010", "Sıhhi tesisat grubu", "—", 0 if tes_alici else None, "İş Emri" if tes_alici else "Kural yok",
         "alıcıya ait (İş Emri › alici_ait)" if tes_alici else "—", "paftanın açıklaması" if tes_alici else "")
    dere_hat = [mahya_boy, mahya_boy]
    pd_, bd = _kural(k, "parca_boyu_m", "dere"), _kural(k, "bindirme_m", "dere")
    dere = sum(TAVAN(x / (pd_ - bd)) for x in dere_hat)
    ekle("1011", "Metal dere", "L 4000", dere, "Kural", f"iki saçak kenarı {mahya_boy:.2f} m ÷ (4,0 − 0,2), ayrı ayrı yukarı", "çelik kuralı; prefabrikte dere var mı teyit")
    ekle("1011", "Dere kelepçesi", "—", dere * _kural(k, "dere_kelepcesi_dere_basina"), "Kural", "dere × 8", "Excel motoru")
    ins = _kural(k, "inis")
    inis = 2 * 2 * ins["kose_basina"] + sum(1 for x in dere_hat if x > ins["orta_inis_esigi_m"])
    ekle("1011", "İniş borusu", "L 3000", inis, "Kural (Lokman)", f"her dere hattının iki ucu + {ins['orta_inis_esigi_m']:g} m'yi geçen hattın ortası", "kurallar.json › inis (Lokman, 8 Ekim)")
    ib = _kural(k, "inis_basina")
    for kalem, anahtar in (("İniş aparatı", "inis_aparati"), ("Kapalı dirsek", "kapali_dirsek"), ("Açık dirsek", "acik_dirsek"), ("İniş kelepçesi", "inis_kelepcesi")):
        ekle("1011", kalem, "—", inis * ib[anahtar], "Kural", f"iniş × {ib[anahtar]}", "Excel motoru")
    ekle("1011", "Dere kapağı", "—", _kural(k, "dere_kapagi").get("besik", 0), "Kural", "beşikte 5", "Excel motoru")
    boya_alici = any("boya" in x for x in alici) or ("alıcı" in oz.get("dış cephe boya", "") and "alıcı" in oz.get("iç cephe boya", ""))
    ekle("1012", "Boya grubu (dış, iç, tavan, astar, mastik)", "—", 0 if boya_alici else None, "İş Emri" if boya_alici else "Kural yok",
         "alıcıya ait (İş Emri › alici_ait / özellikler)" if boya_alici else "—", "paftanın açıklaması: mastikler dahil" if boya_alici else "")

    say = {"pano_dis": sum(len(d.panolar) for d in m["dis"]), "pano_ic": sum(len(d.panolar) for d in m["ic"]), "brut_m2": m["brut_m2"],
           "dis_olcu_cm": [round(m["W"], 2), round(m["D"], 2)], "egimli_boy_m": round(egik, 3), "mahya_boyu_m": round(mahya_boy, 3),
           "cati_m2": round(2 * egik * mahya_boy, 2), "tavan_m2": round(tavan_m2, 2), "makas": n_makas}
    return {"is_emri_no": ie.get("is_emri_no"), "kod": ie.get("yapi", {}).get("referans_model"), "uretildi": time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime()),
            "kural_seti": k.get("surum") or k.get("sema"), "girdi_ozeti": say, "satirlar": satirlar,
            "ozet": {"satir": len(satirlar), "adetli": sum(s["adet"] is not None for s in satirlar), "kural_yok": sum(s["dayanak"] == "Kural yok" for s in satirlar),
                     "pafta_katmanli": sum(s["pafta"] is not None for s in satirlar)}}


def xlsx(L: dict, yol: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Liste"
    ws.append([f"{L['kod']} · kendi yükleme listemiz (kör test; gerçek liste açılmadan) · {L['is_emri_no']} · {L['uretildi']}"])
    ws.append(["Grup", "Kalem", "Ölçü", "Adet", "Birim", "Dayanak", "Formül", "Kaynak / not", "Pafta katmanı", "Pafta notu"])
    for c_ in ws[2]:
        c_.font = Font(bold=True, color="FFFFFF"); c_.fill = PatternFill("solid", fgColor="2F3B4C")
    renk = {"Kural yok": "F2F2F2", "Liste kuralı": "FFF2CC", "Kural (uydurma)": "FCE4D6", "Çıkarım": "FCE4D6"}
    for s in L["satirlar"]:
        ws.append([f"{s['grup']} {s['grup_adi']}", s["kalem"], s["olcu"], s["adet"], s["birim"], s["dayanak"], s["formul"], s["kaynak"], s["pafta"], s["pafta_notu"]])
        if s["dayanak"] in renk:
            for c_ in ws[ws.max_row]:
                c_.fill = PatternFill("solid", fgColor=renk[s["dayanak"]])
    for i, w in enumerate([16, 46, 22, 7, 6, 15, 70, 70, 10, 50], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=3):
        for c_ in row:
            c_.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C3"
    wb.save(yol)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("is_emri", type=Path); ap.add_argument("cikti", type=Path); ap.add_argument("--pafta", type=Path, default=None)
    a = ap.parse_args()
    ie = json.loads(a.is_emri.read_text(encoding="utf-8"))
    k = json.loads((KOK / "kurallar" / "liste" / "kurallar.json").read_text(encoding="utf-8"))
    pafta = json.loads(a.pafta.read_text(encoding="utf-8")) if a.pafta else None
    L = liste(ie, k, pafta)
    a.cikti.mkdir(parents=True, exist_ok=True)
    kod = (L["kod"] or "PK").replace("-", "")
    (a.cikti / f"{kod}_liste.json").write_text(json.dumps(L, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    xlsx(L, a.cikti / f"{kod}_liste.xlsx")
    print(json.dumps(L["ozet"], ensure_ascii=False), "→", a.cikti / f"{kod}_liste.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())

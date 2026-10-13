#!/usr/bin/env python3
"""Yükleme listesinin Excel'i (İE-08; D-11, Lokman 13 Ekim 2026): `<kod>_yukleme.json` → `<kod>_yukleme.xlsx` (openpyxl).

Düzen `deneme/2026-10-06_kalinlik_160_100/yukleme_listesi.py`'deki YÜKLEME sayfası gibi: başlık, model ve yöntem satırı, TASLAK notu;
grup satırları (açık bej); her kalem bir satır: sıra, kalem, boyut, adet, ölçü, birim, ağırlık, kural (formül metni), dayanak, not.
Sayılar sayı hücresidir (düzenlenebilir); dayanağı varsayım, uydurma ya da açık soru olan satır renklidir (teyit bekliyor). İkinci
sayfa "Ara değerler" (listenin hesap girdileri: dış net alan, oda çevreleri, iniş sayısı…), üçüncüsü "Açıklama" (dayanakların anlamı,
renkler, kural).

openpyxl sunucu paket listesinde yok (D-11): bu modül yalnız kabul paketinde (`araclar/kabul_paketi.py`, yerel) çağrılır, `uret`
çağırmaz; openpyxl çağrı anında yüklenir (sunucudaki içe aktarma denetimini etkilemez). Excel çıktıdır, kaynak değil: düzeltme kural
kitabına (`kurallar/`) ya da tanıma (`modeller/*.json`, D-07) işlenir, liste yeniden üretilir.

    python -m motor.genel.excel_liste <kod>_yukleme.json [<çıktı>.xlsx]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

DAYANAK = {   # docs/sartname/yukleme_listesi_yontemi.md › Durum
    "lokman": "Lokman söyledi", "muhendis": "mühendis söyledi", "excel": "yükleme listesi çalışma kitabındaki formül",
    "motor": "Excel motorundaki formül ya da not", "uydurma": "gerçek listelerdeki adedi tutturan kural (teyit bekliyor)",
    "varsayim": "Claude'un varsayımı (teyit bekliyor)", "soru": "kural bilinmiyor", "bekliyor": "değer henüz yok (sorulacak ya da paftadan gelecek)",
    "model": "modeldeki açıklıklardan ya da plandan", "liste": "örnekte gerçek listeden alınan sayı", "sablon": "Eylül 2026 şablonunun formülü (deneme)"}
KURAL_METNI = ("Excel çıktıdır, kaynak değil (D-11, Lokman 13 Ekim 2026): bu dosyadaki düzeltme dosyada kalmaz; kural kitabına (kurallar/) ya da "
               "tanıma (modeller/*.json; yeni dosya ya da sürüm, D-07) işlenir, liste yeniden üretilir. Bir sonraki üretim elle düzeltmeyi ezer.")


def xlsx_yaz(liste: dict, hedef: Path, kaynak: str = "", set_adi: str | None = None) -> Path:
    """`liste`: `yukleme.hesapla` sonucu (`<kod>_yukleme.json`). Dönen: yazılan dosyanın yolu."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    YAZI = Font(name=F, size=10); KALIN = Font(name=F, size=10, bold=True); BASLIK = Font(name=F, size=14, bold=True)
    BEYAZ_KALIN = Font(name=F, size=10, bold=True, color="FFFFFF"); KUCUK = Font(name=F, size=9, color="555555")
    D_BASLIK = PatternFill("solid", fgColor="2F3B4C"); D_GRUP = PatternFill("solid", fgColor="E7E2D6"); D_TOPLAM = PatternFill("solid", fgColor="D9E1F2")
    D_UYDURMA = PatternFill("solid", fgColor="FFF2CC"); D_VARSAYIM = PatternFill("solid", fgColor="FCE4D6"); D_SORU = PatternFill("solid", fgColor="F8CBAD")
    DOLGU = {"uydurma": D_UYDURMA, "varsayim": D_VARSAYIM, "soru": D_SORU, "bekliyor": D_SORU}
    INCE = Side(style="thin", color="BFBFBF"); KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
    ORTA = Alignment(horizontal="center", vertical="center"); SOL = Alignment(horizontal="left", vertical="center", wrap_text=True)
    DUZ = Alignment(horizontal="left", vertical="center")

    def yaz(ws, r, c, v, font=YAZI, fill=None, hiza=None, fmt=None, kenar=True):
        h = ws.cell(row=r, column=c, value=v); h.font = font
        if fill is not None:
            h.fill = fill
        h.alignment = hiza or (DUZ if not kenar else (SOL if isinstance(v, str) else ORTA))
        if fmt:
            h.number_format = fmt
        if kenar:
            h.border = KENAR
        return h

    wb = Workbook()
    wb.properties.creator = "Portman (genel motor)"
    ws = wb.active; ws.title = "YÜKLEME"
    yaz(ws, 1, 1, "ÇELİK BİNA MALZEMELERİ", BASLIK, kenar=False); yaz(ws, 1, 6, "YÜKLEME LİSTESİ", BASLIK, kenar=False)
    yaz(ws, 3, 1, str(liste.get("ad") or ""), KALIN, kenar=False)
    yaz(ws, 4, 1, f"Çatı yöntemi: {liste.get('yontem_adi') or liste.get('yontem') or '...'} · Çatı kaplaması: {liste.get('cati_kaplamasi') or '...'}"
        + (f" · Kapalı gruplar: {', '.join(liste['kapali_gruplar'])}" if liste.get("kapali_gruplar") else "")
        + (f" · Kural seti: {set_adi}" if set_adi else "") + (f" · Kaynak: {kaynak}" if kaynak else ""), KUCUK, kenar=False)
    yaz(ws, 5, 1, "TASLAK: üretim ve sevkiyat listesi değildir. " + KURAL_METNI, KUCUK, kenar=False)
    basliklar = ["Sıra", "Kalem", "Boyut", "Adet", "Ölçü", "Birim", "Ağırlık (kg)", "Kural (formül)", "Dayanak", "Not"]
    for j, b in enumerate(basliklar, 1):
        yaz(ws, 7, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    r, sira, grup = 8, 0, None
    for x in liste.get("satirlar") or []:
        if x.get("grup") != grup:
            grup = x.get("grup")
            for j in range(1, len(basliklar) + 1):
                yaz(ws, r, j, grup if j == 1 else None, KALIN, D_GRUP, DUZ)
            r += 1
        toplam = str(x.get("kalem") or "").startswith("Toplam")
        kalem = (x.get("liste_adi") or x.get("kalem") or "") + (f" ({x['renk']})" if x.get("renk") else "")
        not_ = (x.get("not") or "") + (((". " if x.get("not") else "") + f"Bugünkü kurallarla: {x['bugunku_adet']}") if "bugunku_adet" in x else "")
        if not toplam:
            sira += 1
        dolgu = D_TOPLAM if toplam else DOLGU.get(x.get("durum"))
        font = KALIN if toplam else YAZI
        degerler = [None if toplam else sira, kalem, x.get("boyut") or "", x.get("adet"), x.get("olcu"), x.get("birim") or "", x.get("kg"),
                    x.get("formul") or "", x.get("durum") or "", not_]
        for j, v in enumerate(degerler, 1):
            yaz(ws, r, j, v, font, dolgu, SOL if j in (2, 3, 8, 10) else ORTA, fmt={4: "0", 5: "0.0", 7: "0"}.get(j) if isinstance(v, (int, float)) else None)
        r += 1
    r += 1
    for metin in ["Ölçüler: Ölçü sütunu birimindedir (m, m², adet…); Boyut mm ya da ürün ölçüsü. Ağırlık kg.",
                  "Renkler: sarı = uydurma (gerçek listelerdeki adedi tutturan kural), turuncu = varsayım, kırmızı = kural bilinmiyor ya da değer bekleniyor; "
                  "üçü de teyit bekliyor. Dayanakların anlamı Açıklama sayfasında.",
                  KURAL_METNI]:
        yaz(ws, r, 1, metin, KUCUK, kenar=False); r += 1
    for j, w in enumerate([6, 46, 22, 8, 10, 8, 12, 60, 11, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A8"

    wa = wb.create_sheet("Ara değerler")
    yaz(wa, 1, 1, "ARA DEĞERLER — listenin hesap girdileri (genel motorun metrajından)", KALIN, kenar=False)
    for j, b in enumerate(["Değer", "Sayı / metin"], 1):
        yaz(wa, 3, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    ra = 4
    for k, v in (liste.get("ara") or {}).items():
        for kk, vv in ((v.items() if isinstance(v, dict) else [(None, v)])):
            ad = k if kk is None else f"{k} › {kk}"
            deger = vv if isinstance(vv, (int, float, str)) or vv is None else json.dumps(vv, ensure_ascii=False)
            yaz(wa, ra, 1, ad); yaz(wa, ra, 2, deger, hiza=SOL if isinstance(deger, str) else ORTA); ra += 1
    wa.column_dimensions["A"].width = 44; wa.column_dimensions["B"].width = 70
    wa.freeze_panes = "A4"

    wc = wb.create_sheet("Açıklama")
    yaz(wc, 1, 1, "DAYANAK — her satırın kuralı nereden geliyor", KALIN, kenar=False)
    for j, b in enumerate(["Dayanak", "Anlamı"], 1):
        yaz(wc, 3, j, b, BEYAZ_KALIN, D_BASLIK, ORTA)
    for i, (k, v) in enumerate(DAYANAK.items(), 4):
        yaz(wc, i, 1, k, fill=DOLGU.get(k), hiza=ORTA); yaz(wc, i, 2, v)
    son = 4 + len(DAYANAK) + 1
    for i, metin in enumerate(["Katsayılar ve kaynakları: kurallar/liste/kurallar.json (kural kitabı). Yöntem: docs/sartname/yukleme_listesi_yontemi.md.",
                               KURAL_METNI], son):
        yaz(wc, i, 1, metin, KUCUK, kenar=False)
    wc.column_dimensions["A"].width = 14; wc.column_dimensions["B"].width = 90
    for w_ in (ws, wa, wc):                                   # yazdırma: yatay, sayfa genişliğine sığdır
        w_.page_setup.orientation = "landscape"; w_.page_setup.fitToWidth = 1; w_.page_setup.fitToHeight = 0
        w_.sheet_properties.pageSetUpPr.fitToPage = True
    hedef = Path(hedef); hedef.parent.mkdir(parents=True, exist_ok=True)
    wb.save(hedef)
    return hedef


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        print(__doc__); return 2
    kaynak = Path(argv[0])
    hedef = Path(argv[1]) if len(argv) > 1 else kaynak.with_suffix(".xlsx")
    yol = xlsx_yaz(json.loads(kaynak.read_text(encoding="utf-8")), hedef, kaynak=kaynak.name)
    print("yazıldı:", yol)
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""İE-15 · K1 · Kimlik kartlı model (D-24): bugünkü IFC'nin parçaları kartlarına bağlanır, adetler modelden sayılır.

    python -m motor.kart_model <İş Emri.json> <çıktı klasörü> [--kod PK01]

Çıktılar: `<kod>.ifc` (bugünkü IFC, `motor.ifc_uret`; dokunulmaz), `<kod>_model.ifc` (kart modeli; Bonsai'de açılır),
`<kod>_kart_sayim.json` ve `<kod>_kart_katalogu.xlsx` (kartlar, sayım, H türleri, kartsız parçalar).

- **Kart** bir üründür; modelde IFC tür nesnesidir (`IfcWallType`, `IfcMemberType`…). Kartlar `kurallar/urun/kartlar.json`'da; alanları türde
  `PortatifEv_Kart` setinde durur. Parçaya özgü alanlar örnekte `PortatifEv_Parca` setindedir: kart kodu, paftadaki anahtarı
  (`PaftaAnahtari`, `;` ile; yerleşim kapısının hakemi bu sözleşmeyle sayar), bugünkü katalog kodu ve uygulanan kural.
- **Kart modeli bugünkü IFC'nin kopyasıdır** (plan §4): parça eklenmez, silinmez, yeri ve geometrisi değişmez; yalnız türleri kartlara geçer.
  Duvar kartı eski pano tipinin katmanlarını ve `PortatifEv_PanoTipi` setini taşır, eski pano tipi kalkar. Ortak parçada çelişki çıkarsa
  bugünkü IFC esastır (test). İkisinin tek dosyada birleşmesi K4'ün işidir.
- **H türü bir kuraldır** (`kurallar/yerlesim/prefabrik.json › h_turu`, PK-07 lejantı): dış duvarda makasa denk gelen H dübelli & kulaklı,
  öteki dış H dübelli. Girdiler bugünkü IFC'de: dış duvarların ve H'lerin yeri, makasların yeri (makas parçalarından). Denk gelme payı makas
  aralığının yarısıdır; aralık `kurallar/liste/prefabrik.json`'dan okunur. İki makas aynı H'ye denk gelirse kayıt en yakınını yazar.
- **Çatı tipi İş Emri'nden gelir:** bugünkü IFC çatıyı her zaman beşik çizer. Eşlemesi ya da kuralı `cati_tipleri` taşıyan parça (makaslar ve
  makasa bağlı dış H'ler) İş Emri'nin çatısı o tiplerden değilse "kural yok" yazar (ikisinde de yazılıysa ikisi de sağlanır); model makası
  olmayan bir çatıdan tür türetmez. Nedende çatının kodu ve adı (K5: İE-17'nin cumbalı beşiği ve kırması); çatı düzlemleri İE-17 PR-1b'de.
- **Kartı olmayan parça sayılmaz:** hiçbir karta uymayan parça "kural yok" yazar; iki karta uyan parçada model üretilmez. Görünüm (üçlü H'nin
  yüz görünümü), bileşen (makasın çubukları) ve kartsız olduğu bildirilen parçalar (çatı yüzü, alın, zemin betonu, boşluk) ayrı listelenir.
  Kart dosyasının `kural_yok` listesindeki parça (K5: taşan verandanın çatı yüzü ve alın dolgusu) yine "kural yok" yazar, nedeni oradan.
- **K2, K3 ve K3b katmanları** (`motor/yerlesim/`): komut İş Emri'yle çatı ve tavan katmanlarını (K2), elektriği (K3) ve montaj rehberinin
  kartsız yedi parçasını (K3b: alt çerçeve ve çatı kenarları) de kurar; beşinci çıktı `<kod>_yerlesim_planlari.pdf` (beş çatı ve tavan planı,
  altıncı sayfa elektrik planı, yedinci sayfa alt çerçeve ve çatı kenarları).

LLM yok, ağ yok. Sayım kodda; Bonsai bakma yeridir.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib
import json
import math
import re
import sys
import uuid
from collections import Counter
from pathlib import Path

import ifcopenshell
import ifcopenshell.api
import ifcopenshell.geom
import ifcopenshell.guid
import ifcopenshell.util.element as ue

KOK = Path(__file__).resolve().parents[1]
KARTLAR = KOK / "kurallar" / "urun" / "kartlar.json"
YERLESIM = KOK / "kurallar" / "yerlesim" / "prefabrik.json"
LISTE_KURALLARI = KOK / "kurallar" / "liste" / "prefabrik.json"
LISTE_YAZILMAZ = "yazılmaz"     # K5: liste üreteci bu İş Emri'nde liste yazmıyor (beşik dışı ana çatı); nedeni `kenar.liste_yazilmaz`
KART_SETI, PARCA_SETI = "PortatifEv_Kart", "PortatifEv_Parca"
TUR_AD_ALANI = uuid.UUID("5d1f6a52-6c2b-4b8e-9a51-0f1e2d3c4b5a")     # kart türünün GlobalId'si koddan türer (her üretimde aynı)


def _oku(yol: Path) -> dict:
    return json.loads(Path(yol).read_text(encoding="utf-8"))


def _goreli(yol: Path) -> str:
    try:
        return Path(yol).resolve().relative_to(KOK).as_posix()
    except ValueError:
        return str(yol)


def _api(ad: str, *a, **k):
    """`ifcopenshell.api.<modül>.<işlev>`; modül burada yüklenir (`import ifcopenshell.api` alt modülleri yüklemez)."""
    modul, islev = ad.split(".")
    return getattr(importlib.import_module(f"ifcopenshell.api.{modul}"), islev)(*a, **k)


# ----------------------------------------------------------------------------- geometri (dünya koordinatı, cm)
class _Kutular:
    def __init__(self):
        self.s = ifcopenshell.geom.settings()
        self.s.set("use-world-coords", True)
        self._bellek: dict[int, tuple[list[float], list[float]]] = {}

    def __call__(self, e) -> tuple[list[float], list[float]]:
        if e.id() not in self._bellek:
            parcalar = [p for r in e.IsDecomposedBy for p in r.RelatedObjects] if e.is_a("IfcElementAssembly") else [e]
            xs, ys, zs = [], [], []
            for p in parcalar:
                v = ifcopenshell.geom.create_shape(self.s, p).geometry.verts
                xs += v[0::3]; ys += v[1::3]; zs += v[2::3]
            self._bellek[e.id()] = ([min(xs) * 100, min(ys) * 100, min(zs) * 100], [max(xs) * 100, max(ys) * 100, max(zs) * 100])
        return self._bellek[e.id()]

    def merkez(self, e) -> tuple[float, float]:
        a, b = self(e)
        return (a[0] + b[0]) / 2, (a[1] + b[1]) / 2


# ----------------------------------------------------------------------------- eşleme
def _bosluk(e) -> str:
    """Duvarın boşluk türleri (`PortatifEv_Bosluk.Tur`), sıralı ve virgülle; boşluksuz duvarda boş."""
    return ",".join(sorted(str(ue.get_psets(r.RelatedOpeningElement).get("PortatifEv_Bosluk", {}).get("Tur", "?"))
                           for r in getattr(e, "HasOpenings", ()) or ()))


def ic_duvar_cm(f: ifcopenshell.file) -> str:
    """Evin iç duvar kalınlığı: bugünkü IFC'deki iç panoların (`IsExternal` değil) kalınlıkları, sıralı ve virgülle; iç duvarsız evde boş.
    İş Emri'nde tek değerdir (`yapi.duvar_kalinligi_cm.ic`); BK-18'den beri 100'lük de olabilir."""
    return ",".join(f"{k:g}" for k in sorted({float(ue.get_psets(w).get("PortatifEv_Pano", {}).get("Kalinlik_cm"))
                                               for w in f.by_type("IfcWall")
                                               if ue.get_psets(w).get("Pset_WallCommon", {}).get("IsExternal") is False
                                               and ue.get_psets(w).get("PortatifEv_Pano", {}).get("Kalinlik_cm") is not None}))


def _uyar(e, m: dict, ev: dict | None = None) -> bool:
    """Eşlemenin sınıf, ad, özellik, boşluk ve ev (`ic_duvar_cm`: evin iç duvar kalınlığı) koşulları (yerleşim kuralı ayrı)."""
    if not e.is_a(m["sinif"]):
        return False
    if "ic_duvar_cm" in m and not re.search(m["ic_duvar_cm"], (ev or {}).get("ic_duvar_cm", "")):
        return False
    if m.get("ad") and not re.search(m["ad"], e.Name or ""):
        return False
    ps = ue.get_psets(e)
    for pset, oz, desen in m.get("kosul", []):
        deger = ps.get(pset, {}).get(oz)
        if deger is None or not re.search(desen, str(deger)):
            return False
    if "bosluk" in m and not re.search(m["bosluk"], _bosluk(e)):
        return False
    return True


# ----------------------------------------------------------------------------- H türü kuralı
def makaslar(f: ifcopenshell.file, kutu: _Kutular) -> list[dict]:
    """Bugünkü IFC'nin makasları: {ad, rol (bas | orta; `PortatifEv_Kapsam.BOMKalem`), eksen (uzandığı), konum (dizildiği eksende),
    aralik (uzandığı eksende)}."""
    out = []
    for e in sorted(f.by_type("IfcElementAssembly"), key=lambda e: e.Name or ""):
        if ue.get_predefined_type(e) != "TRUSS":
            continue
        a, b = kutu(e)
        eksen = 0 if b[0] - a[0] > b[1] - a[1] else 1                 # makasın uzandığı eksen
        kalem = str(ue.get_psets(e).get("PortatifEv_Kapsam", {}).get("BOMKalem", ""))
        out.append({"ad": e.Name, "rol": "bas" if "(baş)" in kalem else "orta", "eksen": eksen,
                    "konum": (a[1 - eksen] + b[1 - eksen]) / 2, "aralik": (a[eksen], b[eksen])})
    return out


def h_turleri(f: ifcopenshell.file, adaylar: list, kural: dict, makas_aralik_cm: float, kutu: _Kutular) -> dict[int, dict]:
    """Dış H adaylarının türü (PK-07 lejantı). Dönen: parça id → {tur, makas, uzaklik_cm}."""
    pay_cizgi = kural["duvar_cizgisi_payi_cm"]["deger"]
    pay_denk = makas_aralik_cm * kural["denk_gelme_payi"]["carpan"]
    makaslar_ = makaslar(f, kutu)
    duvarlar = []                                                      # dış duvar parçaları: (uzandığı eksen, orta çizgi, uç, uç)
    for w in f.by_type("IfcWall"):
        if ue.get_psets(w).get("Pset_WallCommon", {}).get("IsExternal") is not True:
            continue
        a, b = kutu(w)
        yon = 0 if b[0] - a[0] > b[1] - a[1] else 1
        duvarlar.append((yon, (a[1 - yon] + b[1 - yon]) / 2, a[yon], b[yon]))
    H = [(e, kutu.merkez(e)) for e in adaylar]
    sonuc = {e.id(): {"tur": kural["turler"]["denk_degil"], "makas": None, "uzaklik_cm": None} for e, _ in H}
    en_yakin: dict[int, float] = {}                                    # iki makas aynı H'ye denk gelirse kayıt en yakınını yazar
    for m in makaslar_:
        yon = 1 - m["eksen"]                                           # oturduğu duvar makasın dizildiği eksende uzanır
        cizgiler = sorted({round(c, 1) for y, c, u0, u1 in duvarlar
                           if y == yon and m["aralik"][0] <= c <= m["aralik"][1] and u0 <= m["konum"] <= u1})
        for c in cizgiler:
            ustteki = [(e, xy) for e, xy in H if abs(xy[1 - yon] - c) <= pay_cizgi]
            if not ustteki:
                continue
            e, xy = min(ustteki, key=lambda t: (abs(t[1][yon] - m["konum"]), t[0].Name or ""))
            d = abs(xy[yon] - m["konum"])
            if d <= pay_denk and d < en_yakin.get(e.id(), math.inf):
                en_yakin[e.id()] = d
                sonuc[e.id()] = {"tur": kural["turler"]["denk"], "makas": m["ad"], "uzaklik_cm": round(d, 2)}
    return sonuc


# ----------------------------------------------------------------------------- sınıflandırma
def _cati_sinirlari(m: dict, yerlesim: dict) -> list[list[str]]:
    """Eşlemenin çatı tipi sınırları: eşlemede ve uyguladığı yerleşim kuralında yazılı olanlar. Parça hepsini sağlamalı (biri ötekini
    genişletmez); hiçbiri yoksa her çatıda geçerli."""
    sinirlar = [m.get("cati_tipleri")]
    if m.get("kural"):
        sinirlar.append(yerlesim["kurallar"][m["kural"]["ad"]].get("cati_tipleri"))
    return [x for x in sinirlar if x]


def siniflandir(f: ifcopenshell.file, kartlar: dict, yerlesim: dict, makas_aralik_cm: float, cati_tipi: str) -> list[dict]:
    """Her `IfcElement` için rol: kart, görünüm, bileşen, kartsız ya da kural yok. `cati_tipi` İş Emri'ndendir (bugünkü IFC hep beşik çizer)."""
    from motor.yerlesim import girdi
    kutu = _Kutular()
    ev = {"ic_duvar_cm": ic_duvar_cm(f)}
    elemanlar = sorted(f.by_type("IfcElement"), key=lambda e: (e.is_a(), e.Name or "", e.GlobalId))
    h_adaylari = [e for e in elemanlar if any(_uyar(e, m, ev) for k in kartlar["kartlar"] for m in k.get("bugunku_ifc", [])
                                               if m.get("kural", {}).get("ad") == "h_turu")]
    hk = h_turleri(f, h_adaylari, yerlesim["kurallar"]["h_turu"], makas_aralik_cm, kutu) if h_adaylari else {}
    kural_sonucu = {"h_turu": lambda e: hk.get(e.id(), {}).get("tur")}
    out = []
    for e in elemanlar:
        ust = [r.RelatingObject for r in getattr(e, "Decomposes", ()) or () if r.is_a("IfcRelAggregates")]
        if any(u.is_a("IfcElementAssembly") for u in ust):
            out.append({"e": e, "rol": "bilesen", "asil": ust[0]})
            continue
        g = next((g for g in kartlar.get("gorunum", []) if _uyar(e, g)), None)
        if g:
            xy = kutu.merkez(e)
            asillar = [(a, kutu.merkez(a)) for a in elemanlar if _uyar(a, g["asil"])]
            a, axy = min(asillar, key=lambda t: (abs(t[1][0] - xy[0]) + abs(t[1][1] - xy[1]), t[0].Name or ""), default=(None, None))
            if a is not None and max(abs(axy[0] - xy[0]), abs(axy[1] - xy[1])) <= g["pay_cm"]:
                out.append({"e": e, "rol": "gorunum", "asil": a, "neden": g["neden"]})
                continue
            out.append({"e": e, "rol": "kural_yok", "neden": "görünümün sayılan parçası bulunamadı"})
            continue
        uyan = []
        for k in kartlar["kartlar"]:
            for m in k.get("bugunku_ifc", []):
                if _uyar(e, m, ev) and (not m.get("kural") or kural_sonucu[m["kural"]["ad"]](e) == m["kural"]["deger"]):
                    uyan.append((k, m))
                    break
        if len(uyan) > 1:
            raise ValueError(f"{e.Name}: iki karta uyuyor ({', '.join(k['kod'] for k, _ in uyan)}); kart dosyası düzeltilmeli")
        if uyan:
            k, m = uyan[0]
            izin = next((x for x in _cati_sinirlari(m, yerlesim) if cati_tipi not in x), None)
            if izin:
                out.append({"e": e, "rol": "kural_yok", "neden": girdi.cati_disi(cati_tipi, f"{k['kod']} eşlemesi", izin), "cati": cati_tipi})
                continue
            kayit = {"e": e, "rol": "kart", "kart": k}
            if e.id() in hk:
                kayit["h"] = hk[e.id()]
            out.append(kayit)
            continue
        ks = next((k for k in kartlar.get("kartsiz", []) if _uyar(e, k)), None)
        if ks:
            out.append({"e": e, "rol": "kartsiz", "neden": ks["neden"]})
            continue
        kb = next((k for k in kartlar.get("kural_yok", []) if _uyar(e, k)), None)      # K5: bilinen, kuralı sonraki adımda (nedeniyle)
        out.append({"e": e, "rol": "kural_yok", "neden": kb["neden"] if kb else "hiçbir karta uymuyor"})
    return out


def _katalog_kodu(e) -> str:
    ps = ue.get_psets(e)
    return next((str(ps[s]["BOMKod"]) for s in ("PortatifEv_Pano", "PortatifEv_Dograma", "PortatifEv_Kapsam") if ps.get(s, {}).get("BOMKod")), "")


def pafta_anahtarlari(e, kart: dict) -> list[str]:
    anahtarlar = list(kart.get("pafta", []))
    if kart.get("pafta_dis_ic"):
        dis = ue.get_psets(e).get("Pset_WallCommon", {}).get("IsExternal")
        anahtarlar.append(kart["pafta_dis_ic"]["dis" if dis else "ic"])
    return list(dict.fromkeys(anahtarlar))


H_TURU_AD = {"dubelli_kulakli": "dübelli & kulaklı", "dubelli": "dübelli"}


# ----------------------------------------------------------------------------- K2: çatı ve tavan katmanları
def yerlesim_kartlari(p: dict, kartlar: dict) -> list[dict]:
    """Yerleşim kuralının kurduğu parçanın kartları: kuralı, çeşidi ve (yazılıysa) bilgileri kartın `yerlesim` eşlemesiyle aynı olan."""
    return [k for k in kartlar["kartlar"] if k.get("yerlesim") and k["yerlesim"]["kural"] == p["kural"] and k["yerlesim"]["parca"] == p["parca"]
            and all(p["bilgi"].get(a) == v for a, v in k["yerlesim"].get("bilgi", {}).items())]


def _yerlesim_anahtarlari(kart: dict, p: dict, tavan: str) -> list[str]:
    anahtarlar = list(kart.get("pafta", []))
    if kart.get("pafta_tavan"):
        anahtarlar.append(kart["pafta_tavan"][tavan])
    if kart.get("pafta_bilgi") and kart["pafta_bilgi"]["alan"] in p["bilgi"]:
        anahtarlar.append(f"{kart['pafta_bilgi']['onek']}.{p['bilgi'][kart['pafta_bilgi']['alan']]}")
    return list(dict.fromkeys(anahtarlar))


def _k2(f, kartlar: dict, yerlesim: dict, liste_kurallari: dict, is_emri: dict, plan: dict, surum: str, kaynak: str,
        elektrik: bool = False, kenar: bool = False) -> dict:
    """Yerleşim katmanlarını (`motor/yerlesim/`) kart modeline ekler. Dönen: sınıflandırma kayıtları (rol "kart"), kural yok, kesim planı,
    sabit paketler ve yerleşimin özeti. `elektrik` ise K3'ün elektrik katmanı da eklenir (`motor/yerlesim/elektrik.py`); `kenar` ise K3b'nin
    alt çerçeve ve çatı kenarları (`motor/yerlesim/kenar.py`); özetleri ayrıdır. K3b'nin özetinde liste üretecinin aynı satırları da durur
    (bilgi: model ile liste yan yana; liste değişmez). Beşik dışı ana çatıda liste üreteci liste yazmaz: satırların yerinde
    `LISTE_YAZILMAZ`, nedeni `liste_yazilmaz`'da (K5)."""
    from motor import yerlesim as Y
    from motor.yerlesim import kesim, yaz

    sonuc = Y.katmanlar(is_emri, plan, makaslar(f, _Kutular()), yerlesim, liste_kurallari)
    g, D = sonuc["girdi"], sonuc["degerler"]
    kural_yok = list(sonuc["kural_yok"])
    sira: Counter = Counter()

    def esle(parcalar: list[dict]) -> list[dict]:
        out = []
        for p in parcalar:
            uyan = yerlesim_kartlari(p, kartlar)
            if len(uyan) > 1:
                raise ValueError(f"{p['kural']} / {p['parca']}: iki karta uyuyor ({', '.join(k['kod'] for k in uyan)}); kart dosyası düzeltilmeli")
            sira[p["parca"]] += 1
            ad = f"{Y.PARCA_ADI[p['parca']]}_{sira[p['parca']]:03d}"
            if not uyan:
                kural_yok.append({"parca": ad, "sinif": "—", "neden": f"{p['kural']} kuralının \"{p['parca']}\" parçasına uyan kart yok"})
                continue
            k = uyan[0]
            out.append(dict(p, ad=ad, kart=k, sinif=k["yerlesim"]["sinif"], pafta=_yerlesim_anahtarlari(k, p, g["tavan"])))
        return out
    eslesen = esle(sonuc["parcalar"])
    E, el_eslesen = None, []
    if elektrik:
        from motor.yerlesim import elektrik as EL
        E = EL.katman(is_emri, plan, D[Y.ELEKTRIK], yerlesim["kurallar"][Y.ELEKTRIK])
        kural_yok += E["kural_yok"]
        el_eslesen = esle(E["parcalar"])
    KN, kn_eslesen = None, []
    if kenar:
        from motor.yerlesim import kenar as KE
        KN = KE.katman(is_emri, g, D, D["gosterim"]["kesit_cm"], yerlesim["kurallar"])
        kural_yok += KN["kural_yok"]
        kn_eslesen = esle(KN["parcalar"])
    hepsi = eslesen + el_eslesen + kn_eslesen
    turler = {k["kod"]: _tur_olustur(f, k, surum) for k in kartlar["kartlar"] if any(p["kart"] is k for p in hepsi)}
    elemanlar = yaz.yaz(f, hepsi, turler, kaynak, PARCA_SETI, str(is_emri.get("is_emri_no", "")))
    kayitlar = [{"e": e, "rol": "kart", "kart": p["kart"], "pafta": p["pafta"], "k2": p} for e, p in zip(elemanlar, hepsi)]

    kitap = Y.kitap_oku(liste_kurallari)
    kesim_plani, sabit = {}, {}
    for k in kartlar["kartlar"]:
        kp = [p for p in eslesen + kn_eslesen if p["kart"] is k and "boy_cm" in p]
        boylar = [p["boy_cm"] for p in kp]
        if k["sayim"] == "boy_kesim" and boylar:
            R = D[k["yerlesim"]["kural"]]
            if k["yerlesim"].get("kesim"):                             # kartın parça boyu ve payı kuralın başka değerlerinden (OM-02: D-32 ek)
                R = {**R, **{a: R[v] for a, v in k["yerlesim"]["kesim"].items()}}
            cm = lambda m: round(m * 100, 6)
            if R.get("yontem") == "toplam":                            # K3b: liste üretecinin yolu, hatların toplamı (kesik hat bir kez)
                hat = {}
                for p in kp:
                    hat.setdefault(p["bilgi"].get("Hat", p["ad"]), p["bilgi"].get("HatBoyu_cm", p["boy_cm"]))
                kesim_plani[k["kod"]] = kesim.toplam_karti(list(hat.values()), cm(R["parca_boyu"]), cm(R.get("bindirme", 0.0)), len(kp))
            elif R.get("yontem") == "kenar":                           # K3b: kenar başına
                kesim_plani[k["kod"]] = kesim.kenar_karti(boylar, cm(R["parca_boyu"]))
            else:
                kesim_plani[k["kod"]] = (kesim.dere_karti(boylar, cm(R["parca_boyu"]), cm(R["bindirme"])) if "bindirme" in R
                                         else kesim.boy_karti(boylar, cm(R["parca_boyu"]), cm(R["pay"])))
        if k.get("sabit"):
            sabit[k["kod"]] = {"adet": Y.oku(k["sabit"]["okunan"], liste_kurallari, kitap), "okunan": k["sabit"]["okunan"]}
    ozet = {"tavan": g["tavan"], "cati_tipi": g["cati_tipi"], "baslangic": g["baslangic"], "uyarilar": sonuc["uyarilar"],
            "omega_hatlari": [{"no": h["no"], "yer": h["yer"], "eksen": h["eksen"], "s_cm": h["s"], "r0_cm": round(h["r0"], 4), "r1_cm": round(h["r1"], 4),
                               "mahal": h["mahal"], "parcalar_cm": [round(b - a, 4) for a, b in h["parcalar"]], "bosluk_cm": round(h["bosluk_cm"], 4)}
                              for h in sonuc["hatlar"]],
            "parcalar": [{"parca": p["ad"], "kart": p["kart"]["kod"], "kural": p["kural"], **({"boy_cm": p["boy_cm"]} if "boy_cm" in p else {}),
                          **({"alan_m2": p["alan_m2"]} if "alan_m2" in p else {}), "bilgi": p["bilgi"], "aciklama": p["aciklama"]} for p in eslesen],
            "degerler": {ad: D[ad] for ad in Y.KURALLAR}}
    out = {"kayitlar": kayitlar, "kural_yok": kural_yok, "kesim_plani": kesim_plani, "sabit": sabit, "yerlesim": ozet}
    if E is not None:
        out["elektrik"] = {"toplam": E["toplam"], "kablo": E["kablo"], "uyarilar": E["uyarilar"], "odalar": E["odalar"],
                           "noktalar": [{"parca": p["ad"], "kart": p["kart"]["kod"], "pafta": p["pafta"], **({"boy_cm": p["boy_cm"]} if "boy_cm" in p else {}),
                                         "bilgi": p["bilgi"], "aciklama": p["aciklama"]} for p in el_eslesen],
                           "degerler": D[Y.ELEKTRIK]}
    if KN is not None:
        from motor import prefabrik_liste as PL
        kalem = {k["kod"]: (k["liste"]["satirlar"] or [None])[0] for k in kartlar["kartlar"] if k.get("yerlesim", {}).get("kural") in KE.KENAR}
        satir: dict[str, int] = {}
        try:                                                            # listeyle aynı yol (D-28); beşik dışı ana çatıda liste yazılmaz (K5)
            L = PL.is_emrinden(copy.deepcopy(is_emri))
        except PL.ListeYazilmaz as hata:
            L, yazilmaz = None, hata.neden
        for x in (L or {"satirlar": []})["satirlar"]:
            if x["kalem"] in kalem.values():                          # yalnız bu kartların satırları (başka satırın adedi boş olabilir)
                satir[x["kalem"]] = satir.get(x["kalem"], 0) + x["adet"]
        yok = {x["parca"] for x in KN["kural_yok"]}
        kurulmadi = [k["kod"] for k in kartlar["kartlar"] if KE.CATI_ADI.get(k.get("yerlesim", {}).get("kural")) in yok]
        out["kenar"] = {"uyarilar": KN["uyarilar"], "hatlar": KN["hatlar"],
                        "liste": {kod: (LISTE_YAZILMAZ if L is None else satir.get(ad, 0) if ad else None) for kod, ad in kalem.items()},
                        **({"liste_yazilmaz": yazilmaz} if L is None else {}), **({"kurulmadi": kurulmadi} if kurulmadi else {}),
                        "parcalar": [{"parca": p["ad"], "kart": p["kart"]["kod"], "kural": p["kural"], "boy_cm": p["boy_cm"], "bilgi": p["bilgi"],
                                      "aciklama": p["aciklama"]} for p in kn_eslesen],
                        "degerler": {ad: D[ad] for ad in KE.KENAR + (KE.GOSTERIM,)}}
    return out


# ----------------------------------------------------------------------------- kart modeli
def _tur_olustur(f, kart: dict, surum: str):
    sinif, alt = kart["ifc"]["sinif"], kart["ifc"]["alt_tur"]
    t = _api("root.create_entity", f, ifc_class=sinif, name=f"{kart['kod']} {kart['ad']}", predefined_type=alt)
    t.GlobalId = ifcopenshell.guid.compress(uuid.uuid5(TUR_AD_ALANI, kart["kod"]).hex)
    t.Tag = kart["kod"]
    if sinif == "IfcDoorType":
        t.OperationType = "NOTDEFINED"
    elif sinif == "IfcWindowType":
        t.PartitioningType = "NOTDEFINED"
    ps = _api("pset.add_pset", f, product=t, name=KART_SETI)
    _api("pset.edit_pset", f, pset=ps, properties={
        "Kod": kart["kod"], "Ad": kart["ad"], "ListeGrubu": kart["liste"]["grup"], "ListeSatirlari": "; ".join(kart["liste"]["satirlar"]) or "(listede yok)",
        "Olcu": kart["olcu"], "Birim": kart["birim"], "Nerede": kart["nerede"], "Sayim": kart["sayim"],
        "KaynakDurum": kart["kaynak"]["durum"], "KaynakTarih": kart["kaynak"]["tarih"], "Dayanak": kart["kaynak"]["dayanak"],
        "KartSurumu": surum})
    return t


def _pano_tipi_icerigi(tip) -> tuple:
    """Pano tipinin ürün içeriği: katmanları (malzeme adı, kalınlık) ve `PortatifEv_PanoTipi` seti; tipin ve katman dizisinin adı hariç."""
    m = ue.get_material(tip)
    katman = tuple((x.Material.Name, x.LayerThickness) for x in m.MaterialLayers) if m is not None and m.is_a("IfcMaterialLayerSet") else m
    return katman, tuple(sorted((k, v) for k, v in ue.get_psets(tip).get("PortatifEv_PanoTipi", {}).items() if k != "id"))


def _duvar_turu_tasi(f, t, duvarlar: list) -> None:
    """Duvar kartı eski pano tipinin katmanlarını ve `PortatifEv_PanoTipi` setini alır (parçanın özellikleri aynı kalır).
    İç duvarı 100'lük evde (BK-18) D kartının parçaları iki tipten gelir (`PE-DIS-100`, `PE-IC-100`); katmanları ve seti aynıysa kart
    adı sıradaki ilk tipinkini alır. Öteki tipin parçalarının katman kullanımı (`IfcMaterialLayerSetUsage`) da kartın katman dizisini
    gösterir (IFC4: kullanım türün dizisine bağlanır; yön ve kaçıklık aynı kalır, katmanlar aynı). İçerik farklıysa kart iki ürünü birden
    taşıyamaz."""
    eski = sorted({ue.get_type(w) for w in duvarlar} - {None}, key=lambda x: x.Name or "")
    if not eski or len({_pano_tipi_icerigi(x) for x in eski}) != 1:
        raise ValueError(f"{t.Tag}: parçaları {len(eski)} ayrı pano tipinden geliyor ({', '.join(x.Name or '?' for x in eski)}); "
                         "katmanları ya da pano tipi bilgileri farklı, kart iki katman dizisini taşıyamaz")
    eski = eski[0]
    malzeme = ue.get_material(eski)
    if malzeme is not None:
        _api("material.assign_material", f, products=[t], type=malzeme.is_a(), material=malzeme)
    if malzeme is not None and malzeme.is_a("IfcMaterialLayerSet"):
        kume = set(duvarlar)
        for w in duvarlar:
            u = ue.get_material(w, should_inherit=False)
            if u is None or not u.is_a("IfcMaterialLayerSetUsage") or u.ForLayerSet == malzeme:
                continue
            kullanan = {o for r in f.get_inverse(u) if r.is_a("IfcRelAssociatesMaterial") for o in r.RelatedObjects}
            if not kullanan <= kume:                                   # bugünkü IFC'de her duvarın kendi kullanımı var (ifc_uret)
                raise ValueError(f"{t.Tag}: {w.Name} parçasının katman kullanımı bu kartın dışındaki parçalarla ortak; kart katman dizisini taşıyamaz")
            u.ForLayerSet = malzeme
    pt = ue.get_psets(eski).get("PortatifEv_PanoTipi")
    if pt:
        ps = _api("pset.add_pset", f, product=t, name="PortatifEv_PanoTipi")
        _api("pset.edit_pset", f, pset=ps, properties={k: v for k, v in pt.items() if k != "id"})


def _eski_turu_kaldir(f, t) -> None:
    """Önce malzeme bağı elle çözülür: API'nin `unassign_material`'ı türün katman dizisini kullanan bütün parçaların katman kullanımını
    (`IfcMaterialLayerSetUsage`) da siler (ifcopenshell 0.9.0); kart modelinde duvarlar bugünkü IFC'deki kullanımı korur (test)."""
    for rel in list(t.HasAssociations):
        if rel.is_a("IfcRelAssociatesMaterial"):
            kalan = [o for o in rel.RelatedObjects if o != t]
            if kalan:
                rel.RelatedObjects = kalan
            else:
                f.remove(rel)
    _api("root.remove_product", f, product=t)


def kart_modeli(ifc_yolu: Path, cikti_yolu: Path, cati_tipi: str, kartlar: dict | None = None, yerlesim: dict | None = None,
                liste_kurallari: dict | None = None, is_emri: dict | None = None, plan: dict | None = None, elektrik: bool = False,
                kenar: bool = False) -> dict:
    """Bugünkü IFC → kart modeli. Bugünkü IFC'ye yazılmaz. `cati_tipi` İş Emri'nin `cati.tip`'i. Dönen: sayım (JSON'a yazılacak biçimde).
    `is_emri` ve `plan` (`motor.ifc_uret.hazirla(is_emri)`) verilirse K2'nin çatı ve tavan katmanları da eklenir; verilmezse model K1'inkidir.
    `elektrik` ise K3'ün elektrik katmanı, `kenar` ise K3b'nin alt çerçeve ve çatı kenarları da eklenir (ikisi de İş Emri ister); komut
    (`uret`) her zaman ekler."""
    kartlar = kartlar or _oku(KARTLAR)
    yerlesim = yerlesim or _oku(YERLESIM)
    liste_kurallari = liste_kurallari or _oku(LISTE_KURALLARI)
    okunan = yerlesim["kurallar"]["h_turu"]["denk_gelme_payi"]["okunan"]
    makas_aralik_cm = liste_kurallari["kurallar"][okunan.rsplit(".", 1)[1]]["deger"]
    f = ifcopenshell.open(str(ifc_yolu))
    sinif = siniflandir(f, kartlar, yerlesim, makas_aralik_cm, cati_tipi)
    surum = kartlar["surum"]
    kaynak = f"{_goreli(KARTLAR)} sürüm {surum}"

    gruplar: dict[str, list] = {}
    for s in sinif:
        if s["rol"] == "kart":
            gruplar.setdefault(s["kart"]["kod"], []).append(s)
    turler, eski_turler = {}, set()
    for kart in kartlar["kartlar"]:
        parcalar = [s["e"] for s in gruplar.get(kart["kod"], [])]
        if not parcalar:
            continue
        t = turler[kart["kod"]] = _tur_olustur(f, kart, surum)
        if kart["ifc"]["sinif"] == "IfcWallType":
            _duvar_turu_tasi(f, t, parcalar)
        eski_turler |= {ue.get_type(p) for p in parcalar} - {None}
        _api("type.assign_type", f, related_objects=parcalar, relating_type=t, should_map_representations=False)
    for eski in sorted(eski_turler, key=lambda t: t.Name or ""):
        if not getattr(eski, "Types", None):                           # kullanılmayan eski pano tipi kalkar; katmanlar kart türünde
            _eski_turu_kaldir(f, eski)

    for s in sinif:
        e = s["e"]
        if s["rol"] == "kart":
            kural = ""
            if "h" in s:
                h = s["h"]
                kural = f"h_turu: {H_TURU_AD[h['tur']]}" + (f" ({h['makas']} {h['uzaklik_cm']:.2f} cm)" if h["makas"] else "")
            ps = _api("pset.add_pset", f, product=e, name=PARCA_SETI)
            _api("pset.edit_pset", f, pset=ps, properties={"KartKodu": s["kart"]["kod"], "PaftaAnahtari": ";".join(pafta_anahtarlari(e, s["kart"])),
                                                        "KatalogKodu": _katalog_kodu(e), "Kural": kural, "Kaynak": kaynak})
        elif s["rol"] == "gorunum":
            ps = _api("pset.add_pset", f, product=e, name=PARCA_SETI)
            _api("pset.edit_pset", f, pset=ps, properties={"Rol": "görünüm", "Asil": s["asil"].Name, "AsilGlobalId": s["asil"].GlobalId,
                                                        "Not": s["neden"], "Kaynak": kaynak})
    k2 = None
    if elektrik and is_emri is None:
        raise ValueError("K3'ün elektrik katmanı İş Emri ister (`is_emri` ve `plan`)")
    if kenar and is_emri is None:
        raise ValueError("K3b'nin alt çerçeve ve çatı kenarları İş Emri ister (`is_emri` ve `plan`)")
    if is_emri is not None:
        if plan is None:
            raise ValueError("K2 katmanları için `plan` gerekir (`motor.ifc_uret.hazirla(is_emri)`)")
        k2 = _k2(f, kartlar, yerlesim, liste_kurallari, is_emri, plan, surum, kaynak, elektrik, kenar)
        sinif = sinif + k2["kayitlar"]
    cikti_yolu = Path(cikti_yolu)
    cikti_yolu.parent.mkdir(parents=True, exist_ok=True)
    f.write(str(cikti_yolu))
    return sayim(sinif, kartlar, yerlesim, makas_aralik_cm, cati_tipi, ifc_yolu, cikti_yolu, k2)


def sayim(sinif: list[dict], kartlar: dict, yerlesim: dict, makas_aralik_cm: float, cati_tipi: str, ifc_yolu: Path, model_yolu: Path,
          k2: dict | None = None) -> dict:
    adet = Counter(s["kart"]["kod"] for s in sinif if s["rol"] == "kart")
    metre: dict[str, float] = {}
    for s in sinif:
        if "k2" in s and "boy_cm" in s["k2"]:
            metre[s["kart"]["kod"]] = metre.get(s["kart"]["kod"], 0.0) + s["k2"]["boy_cm"] / 100
    katalog = {}
    for s in sinif:
        if s["rol"] == "kart":
            katalog.setdefault(s["kart"]["kod"], Counter())[_katalog_kodu(s["e"]) or "—"] += 1
    pafta = Counter(a for s in sinif if s["rol"] == "kart" for a in (s["pafta"] if "pafta" in s else pafta_anahtarlari(s["e"], s["kart"])))
    h = [{"parca": s["e"].Name, "kart": s["kart"]["kod"], "tur": s["h"]["tur"], "makas": s["h"]["makas"], "uzaklik_cm": s["h"]["uzaklik_cm"]}
         for s in sinif if s["rol"] == "kart" and "h" in s]
    sabit = (k2 or {}).get("sabit", {})
    out = {
        "aciklama": "İE-15 kart sayımı (K1; K2 ile çatı ve tavan; K3 ile elektrik; K3b ile alt çerçeve ve çatı kenarları): her parça kartına bağlı; kartsız parça sayılmaz. Boyla alınan kartta metre ve kesim "
                    "planı, sabit pakette liste kuralının adedi. Paftayla kıyas hakemin işidir (`testler/kimlik/hakem/kiyas.py`).",
        "bugunku_ifc": Path(ifc_yolu).name, "bugunku_ifc_sha256": hashlib.sha256(Path(ifc_yolu).read_bytes()).hexdigest(),
        "kart_modeli": Path(model_yolu).name,
        "kartlar_surumu": kartlar["surum"], "yerlesim_surumu": yerlesim["surum"], "makas_aralik_cm": makas_aralik_cm, "cati_tipi": cati_tipi,
        "kartlar": [{"kod": k["kod"], "ad": k["ad"], "adet": adet[k["kod"]], "katalog_kodlari": dict(sorted(katalog.get(k["kod"], {}).items())),
                     **({"metre": round(metre[k["kod"]], 4)} if k["kod"] in metre else {}),
                     **({"sabit": sabit[k["kod"]]["adet"]} if k["kod"] in sabit else {})}
                    for k in kartlar["kartlar"]],
        "pafta_anahtarlari": dict(sorted(pafta.items())),
        "h_turu": sorted(h, key=lambda x: x["parca"]),
        "gorunum": [{"parca": s["e"].Name, "asil": s["asil"].Name} for s in sinif if s["rol"] == "gorunum"],
        "bilesen": dict(sorted(Counter(s["asil"].Name for s in sinif if s["rol"] == "bilesen").items())),
        "kartsiz": [{"parca": s["e"].Name, "sinif": s["e"].is_a(), "neden": s["neden"]} for s in sinif if s["rol"] == "kartsiz"],
        "kural_yok": [{"parca": s["e"].Name, "sinif": s["e"].is_a(), "neden": s["neden"], **({"cati": s["cati"]} if "cati" in s else {})}
                      for s in sinif if s["rol"] == "kural_yok"]
                     + (k2 or {}).get("kural_yok", []),
    }
    if k2 is not None:
        out |= {"kesim_plani": k2["kesim_plani"], "sabit": sabit, "yerlesim": k2["yerlesim"]}
        if "elektrik" in k2:
            out["elektrik"] = k2["elektrik"]
        if "kenar" in k2:
            out["kenar"] = k2["kenar"]
    return out


# ----------------------------------------------------------------------------- kart kataloğu (Excel)
def katalog_excel(kartlar: dict, say: dict | None, yol: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    kalin, dolgu = Font(bold=True), PatternFill("solid", fgColor="DDE6F0")

    def sayfa(ws, basliklar, satirlar, genislik):
        ws.append(basliklar)
        for c in ws[1]:
            c.font, c.fill = kalin, dolgu
        for s in satirlar:
            ws.append(s)
        for i, g in enumerate(genislik):
            ws.column_dimensions[chr(65 + i)].width = g
        for satir in ws.iter_rows(min_row=2):
            for c in satir:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"

    yollar = kartlar["sayim_yollari"]
    sayfa(wb.active, ["Kod", "Ad", "Liste grubu", "Listedeki satırları", "Ölçü", "Birim", "Nerede durur", "Nasıl sayılır",
                      "IFC tür sınıfı (alt tür)", "Paftadaki anahtarı", "Kaynak", "Tarih", "Dayanak"],
          [[k["kod"], k["ad"], k["liste"]["grup"], "; ".join(k["liste"]["satirlar"]) or "(listede yok)", k["olcu"], k["birim"], k["nerede"], yollar[k["sayim"]],
            f"{k['ifc']['sinif']} ({k['ifc']['alt_tur']})",
            "; ".join(k.get("pafta", []) + ([f"{k['pafta_dis_ic']['dis']} / {k['pafta_dis_ic']['ic']}"] if k.get("pafta_dis_ic") else [])) or "—",
            k["kaynak"]["durum"], k["kaynak"]["tarih"], k["kaynak"]["dayanak"]] for k in kartlar["kartlar"]],
          [8, 34, 10, 34, 34, 8, 40, 22, 26, 40, 11, 11, 50])
    wb.active.title = "Kartlar"
    if say:
        sayfa(wb.create_sheet("Sayım"), ["Kod", "Ad", "Adet (modelde)", "Boy toplamı (m)", "Sabit paket", "Bugünkü katalog kodları"],
              [[k["kod"], k["ad"], k["adet"], k.get("metre", "—"), k.get("sabit", "—"),
                ", ".join(f"{a} × {n}" for a, n in k["katalog_kodlari"].items()) or "—"] for k in say["kartlar"]],
              [8, 40, 10, 12, 10, 50])
        sayfa(wb.create_sheet("H türü"), ["Parça", "Kart", "Tür", "Denk geldiği makas", "Uzaklık (cm)"],
              [[x["parca"], x["kart"], H_TURU_AD[x["tur"]], x["makas"] or "—", x["uzaklik_cm"] if x["uzaklik_cm"] is not None else "—"]
               for x in say["h_turu"]], [24, 8, 18, 18, 12])
        kural_yok = {x["parca"] for x in say["kural_yok"]}
        sayfa(wb.create_sheet("Kartsız"), ["Parça", "IFC sınıfı", "Rol", "Neden"],
              [[x["parca"], x["sinif"], "kartsız", x["neden"]] for x in say["kartsiz"]]
              + [[x["parca"], x["sinif"], "kural yok", x["neden"]] for x in say["kural_yok"]]
              + [[x["parca"], "IfcMember", "görünüm", f"sayılan parça {x['asil']}"] for x in say["gorunum"]]
              + [[a, "IfcElementAssembly", "bileşen", f"{n} çubuk; birleşim kural yok, sayılmadı" if a in kural_yok else f"{n} çubuk birleşimin kartında"]
                 for a, n in say["bilesen"].items()],
              [26, 20, 12, 60])
    if say and "yerlesim" in say:
        sayfa(wb.create_sheet("Çatı ve tavan"), ["Parça", "Kart", "Kural", "Boy (cm)", "Alan (m²)", "Bilgi", "Açıklama"],
              [[p["parca"], p["kart"], p["kural"], p.get("boy_cm", "—"), p.get("alan_m2", "—"),
                "; ".join(f"{a}: {v}" for a, v in p["bilgi"].items()), p["aciklama"]] for p in say["yerlesim"]["parcalar"]],
              [14, 8, 28, 10, 10, 50, 40])
        sayfa(wb.create_sheet("Kesim planı"), ["Kart", "Yöntem", "Stok boyu (cm)", "Parça", "Toplam (m)", "Stok adedi", "Dizilim / hat başına"],
              [[kod, x["yontem"], x["parca_boyu_cm"], x["parca"], x["toplam_m"], x["stok"],
                " | ".join(" + ".join(f"{b:g}" for b in d) for d in x["dizilim"]) if "dizilim" in x
                else ", ".join(map(str, x["hat_basina"])) if "hat_basina" in x else "hatlar " + " + ".join(f"{h:g}" for h in x["hatlar_cm"])]
               for kod, x in say["kesim_plani"].items()]
              + [[kod, f"sabit paket ({x['okunan']})", "—", "—", "—", x["adet"], "—"] for kod, x in say["sabit"].items()],
              [8, 40, 12, 8, 10, 10, 80])
    if say and "kenar" in say:
        liste, kurulmadi = say["kenar"]["liste"], set(say["kenar"].get("kurulmadi", []))
        sayfa(wb.create_sheet("Alt çerçeve ve çatı kenarı"), ["Parça", "Kart", "Kural", "Boy (cm)", "Bilgi", "Açıklama"],
              [[p["parca"], p["kart"], p["kural"], p["boy_cm"], "; ".join(f"{a}: {v}" for a, v in p["bilgi"].items()), p["aciklama"]]
               for p in say["kenar"]["parcalar"]]
              + [[]] + [["Kart", "Model (kesim planı)", "Liste üreteci", "", "", ""]]
              + [[kod, ("kurulmadı" if kod in kurulmadi else say["kesim_plani"][kod]["stok"] if kod in say["kesim_plani"] else
                        f"{next(k.get('metre', 0) for k in say['kartlar'] if k['kod'] == kod):g} m"),
                  "listede yok" if n is None else "liste yazılmaz" if n == LISTE_YAZILMAZ else n, "", "", ""] for kod, n in liste.items()]
              + ([["Liste yazılmaz", "", say["kenar"]["liste_yazilmaz"], "", "", ""]] if "liste_yazilmaz" in say["kenar"] else []),
              [16, 8, 26, 10, 50, 50])
    yol = Path(yol)
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(yol))
    return yol


# ----------------------------------------------------------------------------- komut
def uret(ie: dict, klasor: Path, kod: str) -> dict:
    from motor import ifc_uret
    from motor.yerlesim import cizim

    klasor = Path(klasor)
    bugunku = klasor / f"{kod}.ifc"
    ozet = ifc_uret.uret(ie, bugunku)
    say = kart_modeli(bugunku, klasor / f"{kod}_model.ifc", ie.get("cati", {}).get("tip", "besik"),    # ifc_uret'in varsayılanı
                      is_emri=ie, plan=ifc_uret.hazirla(ie), elektrik=True, kenar=True)
    say["sema_hatasi"] = {"bugunku_ifc": ozet["sema_hatasi"], "kart_modeli": ifc_uret.ozet(klasor / f"{kod}_model.ifc")["sema_hatasi"]}
    (klasor / f"{kod}_kart_sayim.json").write_text(json.dumps(say, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    katalog_excel(_oku(KARTLAR), say, klasor / f"{kod}_kart_katalogu.xlsx")
    cizim.ciz(klasor / f"{kod}_model.ifc", say, klasor / f"{kod}_yerlesim_planlari.pdf", kod)
    return say


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("is_emri", type=Path)
    ap.add_argument("klasor", type=Path)
    ap.add_argument("--kod", help="çıktı adlarının öneki (varsayılan: İş Emri dosyasının adı)")
    a = ap.parse_args(argv)
    say = uret(_oku(a.is_emri), a.klasor, a.kod or a.is_emri.stem)
    for k in say["kartlar"]:
        if k["adet"] or "sabit" in k:
            ek = f"  {k['metre']:.2f} m" if "metre" in k else (f"  sabit {k['sabit']}" if "sabit" in k else "")
            print(f"{k['kod']:6} {k['adet']:4}  {k['ad']}{ek}")
    print(f"görünüm {len(say['gorunum'])}, bileşen {sum(say['bilesen'].values())}, kartsız {len(say['kartsiz'])}, kural yok {len(say['kural_yok'])}; "
          f"şema hatası {say['sema_hatasi']}")
    for x in say["kural_yok"]:
        print(f"kural yok: {x['parca']}: {x['neden']}")
    for u in say.get("yerlesim", {}).get("uyarilar", []) + say.get("elektrik", {}).get("uyarilar", []) + say.get("kenar", {}).get("uyarilar", []):
        print(f"uyarı ({u['katman']}): {u['metin']}")
    return 1 if say["kural_yok"] or any(say["sema_hatasi"].values()) else 0


if __name__ == "__main__":
    sys.exit(main())

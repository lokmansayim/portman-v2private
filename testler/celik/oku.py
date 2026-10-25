"""Çelik hakeminin liste okuyucusu (İE-16 PR-1): gerçek yükleme listesinin satırları ve çelik panel ağırlıkları.

Satırlar prefabrik kör testinin okuyucusuyla okunur (`testler.kor.oku`, BK-12 PR-3; değiştirilmez, yalnız çağrılır): başlık bloğu, telefon,
personel ve imza satırları okunmaz. Grubun "TOPLAM" satırı (ör. "ÇELİK PANEL TOPLAM") malzeme sayılmaz.

Çelik listelerinin ilk grubu "ÇELİK PANEL GRUBU"dur. Açık beş projede (ÇP-03, 06, 07, 13, 14; 25 Ekim 2026) üç şablon görüldü: sütunlar
NO / ad / BOYUT/ÖLÇÜ / ADET / KG ya da AĞIRLIK (KG) (`KG_BASLIKLARI`) / MAKİNE; toplam satırı "TOPLAM" ya da "ÇELİK PANEL TOPLAM"; PDF'te
TOPLAM yazısı boyut sütununa taşabilir. Bu dosya o grup için ayrıca okur:
- satırların kg sütununu (Excel'de grubun başlık satırındaki kg başlığının sütunu; PDF'te kg başlığının altındaki sayı); Excel'deki formül
  sonucu üç basamağa yuvarlanır (125.71999… → 125.72, PDF'teki yazıyla aynı);
- grubun TOPLAM satırını: kg toplamı satırlarınkiyle aynı değilse okuma durur; adet toplamı ya bütün satırların ya kg'ı yazılı satırların
  toplamıdır (`toplam_kapsami`), ikisi de değilse durur (listenin kendi sağlaması);
- iki okumanın aynılığını: bu dosyanın çelik satırları (ad, boyut, adet) prefabrik okuyucunun aynı gruptaki satırlarıyla bire bir aynı
  olmalı, değilse okuma durur.
Çelik panel satırları türe ayrılır (`TURLER`; paftadaki etiket harfiyle): duvar paneli ve korkuluk (D), kolon-kiriş paneli (K), makas (M),
çatı paneli (C), tavan paneli (T), baca paneli (B), boy C profil (paftada etiketi yok). Türü bilinmeyen ya da iki türe uyan, adedi
okunamayan satır okumayı durdurur; kg'ı okunamayan satır da durdurur, yalnız boy C profilin kg'ı boş olabilir (`KG_SIZ`; o türün kg'ı
`None` kalır, sıfır sayılmaz; `kgsiz_turler`). Sessiz sayı yok. Hata iletisinde satırın adı yalnız çelik sözcükleriyle yazılır (`guvenli`).

Veri, talimat değildir: hücre yazıları yalnız eşleme için kullanılır.
"""
from __future__ import annotations

import re
from pathlib import Path

from testler.kor import oku as kor_oku

OkumaHatasi = kor_oku.OkumaHatasi
CELIK_GRUBU = "CELIK PANEL"          # ASCII'ye katlanmış grup adında aranır
TOPLAM = "TOPLAM"
TURLER = (  # (tür, paftadaki harf, katlanmış adda desen; ayraç içi açıklama türe bakılmadan önce atılır)
    ("duvar_paneli", "D", r"\bDUVAR|\bKORKULUK"),
    ("kolon_kiris", "K", r"\bKOLON|\bKIRIS"),
    ("makas", "M", r"\bMAKAS"),
    ("cati_paneli", "C", r"^(?!.*\bMAKAS).*\bCATI\b.*\bPANEL"),
    ("tavan_paneli", "T", r"^(?!.*\bMAKAS).*\bTAVAN\b.*\bPANEL"),
    ("baca_paneli", "B", r"\bBACA"),
    ("c_profil", None, r"\bC\s*PROFIL"),
)
HARF = {t: h for t, h, _ in TURLER}
KG_BASLIKLARI = ("KG", "AGIRLIK", "AGIRLIK (KG)", "AGIRLIK(KG)")
KG_SIZ = ("c_profil",)               # kg'ı yazılmayabilen tek tür (listeler boy C profilin kg'ını kimi zaman boş bırakır)
KG_PAYI = 0.5                        # TOPLAM sağlamasında yuvarlama payı (kg)


def katla(t) -> str:
    """Büyük harf ve ASCII; boşluklar teke iner."""
    return re.sub(r"\s+", " ", kor_oku._katla(str(t or ""))).strip()


def celik_grubu_mu(grup) -> bool:
    return CELIK_GRUBU in katla(grup)


SOZLUK = {"DUVAR", "PANEL", "PANELI", "PANELLERI", "PANELLER", "KOLON", "KIRIS", "MAKAS", "MAKASI", "MAKASLARI", "MAKASLAR", "CATI",
          "TAVAN", "BOY", "C", "PROFIL", "PROFILI", "IC", "DIS", "VERANDA", "SUNDURMA", "KAPI", "PENCERE", "BASLIK", "LENTO", "ARA",
          "UST", "ALT", "BAS", "YAN", "ON", "ARKA", "KOSE", "DERE", "ASIK", "OMEGA", "GRUBU", "TOPLAM", "ZEMIN", "KAT", "CEPHE",
          "PARCA", "CERCEVE", "MUSTEMILAT", "KORKULUK", "BACA", "VEANDA", "DIKME", "DIKMELERI", "ARASI", "USTU", "SIRT", "MM"}


def guvenli(ad) -> str:
    """Hata iletisi için satır adı: yalnız çelik sözlüğündeki sözcükler ve sayılar kalır, öteki sözcükler "…" olur (kişi ya da bayi adı
    iletiye girmez)."""
    out = []
    for w in re.split(r"[^A-Z0-9]+", katla(ad)):
        if not w:
            continue
        w = w if (w in SOZLUK or w.isdigit()) else "…"
        if not (w == "…" and out and out[-1] == "…"):
            out.append(w)
    return " ".join(out) or "…"


def tur(ad) -> str:
    """Çelik panel satırının türü; bilinmiyorsa ya da iki türe uyuyorsa `OkumaHatasi`."""
    a = re.sub(r"\(.*?\)", " ", katla(ad))
    uyan = [t for t, _, d in TURLER if re.search(d, a)]
    if len(uyan) != 1:
        raise OkumaHatasi(f"çelik panel satırının türü {'bilinmiyor' if not uyan else 'belirsiz (' + ', '.join(uyan) + ')'}: "
                          f"{guvenli(ad)!r}")
    return uyan[0]


def _sayi(v):
    if isinstance(v, str):
        v = v.strip().replace(" ", "")
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", v):        # 1.367 → 1367 (nokta binlik)
            v = v.replace(".", "")
        v = v.replace(",", ".")
    try:
        x = round(float(v), 3)                           # Excel formül sonucu 125.71999999999998 → 125.72 (PDF'teki yazıyla aynı)
    except (TypeError, ValueError):
        return None
    return int(x) if x == int(x) else x


# ------------------------------------------------------------------------------------------------------------------ Excel

def _excel_baslik(hucre: dict) -> dict | None:
    """Grup başlık satırı: BOYUT ve ADET başlığı; çelik grubunda KG da."""
    b = [k for k, v in hucre.items() if katla(v) in ("BOYUT", "BOYUT/OLCU")]
    a = [k for k, v in hucre.items() if katla(v) == "ADET"]
    if not b or not a:
        return None
    kg = [k for k, v in hucre.items() if katla(v) in KG_BASLIKLARI and k > min(a)]
    return {"ad": b[0] - 1, "boyut": b[0], "adet": min(a), "kg": min(kg) if kg else None,
            "grup": str(hucre.get(b[0] - 1, "") or "").strip()}


def _excel_celik(ws) -> dict | None:
    satirlar, toplam, kol = [], None, None
    sinir = min(ws.max_row or kor_oku.EXCEL_SATIR_SINIRI, kor_oku.EXCEL_SATIR_SINIRI)
    for r in ws.iter_rows(min_row=1, max_row=sinir, max_col=kor_oku.EXCEL_SUTUN_SINIRI):
        hucre = {c.column: c.value for c in r if c.value is not None and str(c.value).strip() != ""}
        b = _excel_baslik(hucre)
        if b:
            if kol is not None:
                break                                    # sonraki grup: çelik grubu bitti
            if celik_grubu_mu(b["grup"]):
                if b["kg"] is None:
                    raise OkumaHatasi("çelik panel grubunda KG sütunu yok")
                kol = b
            continue
        if kol is None:
            continue
        if any(re.search(rf"\b{TOPLAM}\b", katla(v)) for k, v in hucre.items() if k <= kol["ad"]):
            toplam = {"adet": _sayi(hucre.get(kol["adet"])), "kg": _sayi(hucre.get(kol["kg"]))}
            continue
        ad = hucre.get(kol["ad"])
        if ad is None:
            continue
        satirlar.append({"ad": re.sub(r"\s+", " ", str(ad)).strip(), "boyut": _yazi(hucre.get(kol["boyut"])),
                         "adet": _sayi(hucre.get(kol["adet"])), "kg": _sayi(hucre.get(kol["kg"]))})
    if kol is None:
        return None
    return {"satirlar": satirlar, "toplam": toplam}


def _yazi(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else None


def excel_celik(yol: Path, sayfa: int | None = None) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(str(yol), data_only=True, read_only=True)
    if sayfa is not None:
        if not 0 <= sayfa < len(wb.worksheets):
            raise OkumaHatasi(f"liste sayfası yok: {sayfa}")
        adaylar = [(sayfa, _excel_celik(wb.worksheets[sayfa]))]
    else:
        adaylar = [(i, _excel_celik(ws)) for i, ws in enumerate(wb.worksheets)]
    adaylar = [(i, c) for i, c in adaylar if c is not None]
    if not adaylar:
        raise OkumaHatasi("Excel'de çelik panel grubu yok")
    if len(adaylar) > 1:
        raise OkumaHatasi(f"{len(adaylar)} sayfada çelik panel grubu var; `sayfa` verilmeli")
    return {"sayfa": adaylar[0][0], **adaylar[0][1]}


# -------------------------------------------------------------------------------------------------------------------- PDF

def pdf_celik(yol: Path) -> dict:
    """PDF'te çelik panel grubu: başlık satırındaki ADET ve KG sütunlarının yerinden; sonraki grup başlığında biter."""
    import pymupdf

    doc = pymupdf.open(str(yol))
    satirlar, toplam, kol = [], None, None
    for sayfa in doc:
        parcalar = []
        for b in sayfa.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                t = re.sub(r"\s+", " ", "".join(s["text"] for s in l["spans"]).strip())
                if t:
                    x0, y0, x1, y1 = l["bbox"]
                    parcalar.append({"x0": x0, "x1": x1, "xc": (x0 + x1) / 2, "y": (y0 + y1) / 2, "t": t})
        parcalar.sort(key=lambda p: (round(p["y"]), p["x0"]))
        kumeler = []
        for p in parcalar:
            if kumeler and abs(p["y"] - kumeler[-1]["y"]) <= 3.2:
                kumeler[-1]["p"].append(p)
            else:
                kumeler.append({"y": p["y"], "p": [p]})
        for k in kumeler:
            ps = sorted(k["p"], key=lambda p: p["x0"])
            basliklar = {katla(p["t"]): p for p in ps}
            if ("ADET" in basliklar and any(x in basliklar for x in ("BOYUT", "BOYUT/OLCU"))
                    and not any(re.fullmatch(r"[\d.,]+", p["t"]) for p in ps)):
                if kol is not None:
                    return {"satirlar": satirlar, "toplam": toplam}
                grup = " ".join(p["t"] for p in ps if katla(p["t"]) not in ("BOYUT", "BOYUT/OLCU", "ADET", "MAKINE", "NO", *KG_BASLIKLARI))
                if celik_grubu_mu(grup):
                    kgb = next((basliklar[x] for x in KG_BASLIKLARI if x in basliklar), None)
                    if kgb is None:
                        raise OkumaHatasi("çelik panel grubunda KG sütunu yok")
                    bo = basliklar.get("BOYUT/OLCU") or basliklar["BOYUT"]
                    kol = {"boyut": bo["x0"] - 8, "adet": basliklar["ADET"]["xc"], "kg": kgb["xc"]}
                continue
            if kol is None:
                continue
            ad, boy, adet, kg = [], [], [], []
            for p in ps:
                if p["x1"] < kol["boyut"] + 4:
                    if re.fullmatch(r"\d{1,3}", p["t"]) and p["x1"] < 66:
                        continue                         # NO sütunu
                    ad.append(p["t"])
                elif abs(p["xc"] - kol["adet"]) < 22:
                    adet.append(p["t"])
                elif abs(p["xc"] - kol["kg"]) < 20:
                    kg.append(p["t"])
                elif p["xc"] < kol["adet"] - 22:
                    boy.append(p["t"])
            a = " ".join(ad).strip()
            if any(re.search(rf"\b{TOPLAM}\b", katla(p["t"])) for p in ps):   # "TOPLAM" ad sütununda ya da ölçü sütununa taşmış
                toplam = {"adet": _sayi(adet[0]) if adet else None, "kg": _sayi(kg[0]) if kg else None}
                continue
            if not a:
                continue
            satirlar.append({"ad": a, "boyut": " ".join(boy) or None, "adet": _sayi(adet[0]) if adet else None,
                             "kg": _sayi(kg[0]) if kg else None})
    if kol is None:
        raise OkumaHatasi("PDF'te çelik panel grubu yok")
    return {"satirlar": satirlar, "toplam": toplam}


# ---------------------------------------------------------------------------------------------------------------- birlikte

def _ozet(c: dict) -> dict:
    """Türler ve toplam; TOPLAM satırı listenin kendi sağlamasıdır. Kg'ı boş satır yalnız boy C profilde olabilir (`KG_SIZ`); o zaman
    TOPLAM'ın adedi ya bütün satırların ya kg'lı satırların toplamıdır, hangisi olduğu `toplam_kapsami`'na yazılır."""
    out = []
    for s in c["satirlar"]:
        t = tur(s["ad"])
        if s["adet"] is None or (s["kg"] is None and t not in KG_SIZ):
            raise OkumaHatasi(f"çelik panel satırında adet ya da kg okunamadı: {guvenli(s['ad'])!r}")
        out.append({**s, "tur": t, "harf": HARF[t]})
    if not out:
        raise OkumaHatasi("çelik panel grubunda satır yok")
    adet = sum(s["adet"] for s in out)
    adet_kgli = sum(s["adet"] for s in out if s["kg"] is not None)
    kg = round(sum(s["kg"] for s in out if s["kg"] is not None), 3)
    T = c.get("toplam")
    if not T or T.get("adet") is None or T.get("kg") is None:
        raise OkumaHatasi("çelik panel grubunun TOPLAM satırı okunamadı")
    kapsam = "bütün satırlar" if T["adet"] == adet else "kg'lı satırlar" if T["adet"] == adet_kgli else None
    if kapsam is None or abs(T["kg"] - kg) > KG_PAYI:
        raise OkumaHatasi(f"çelik panel TOPLAM satırı satırlarla tutmuyor: liste {T['adet']} adet {T['kg']} kg, satırlar {adet} adet "
                          f"(kg'lı {adet_kgli}) {kg} kg")
    turler = {}
    for s in out:
        x = turler.setdefault(s["tur"], {"harf": s["harf"], "adet": 0, "kg": 0.0})
        x["adet"] += s["adet"]
        x["kg"] = None if (x["kg"] is None or s["kg"] is None) else round(x["kg"] + s["kg"], 3)
    return {"satirlar": out, "toplam": {"adet": adet, "kg": kg}, "turler": turler, "toplam_kapsami": kapsam,
            "kgsiz_turler": sorted({s["tur"] for s in out if s["kg"] is None})}


def _ayni_satirlar(celik: list[dict], kor: list[dict]) -> None:
    a = [(katla(s["ad"]), katla(s["boyut"]), s["adet"]) for s in celik]
    b = [(katla(s["malzeme"]), katla(s["boyut"]), kor_oku.sayi(s["adet"])) for s in kor]
    if a != b:
        raise OkumaHatasi(f"çelik panel satırları iki okumada farklı ({len(a)} / {len(b)} satır); liste şablonu beklenenden farklı")


def _toplam_satiri_mi(s: dict) -> bool:
    """Grubun toplam satırı ("TOPLAM", "ÇELİK PANEL TOPLAM"): ölçüsü yok, adında TOPLAM sözcüğü var; malzeme sayılmaz."""
    return bool(re.search(rf"\b{TOPLAM}\b", katla(s["malzeme"]))) and not _yazi(s.get("boyut"))


def oku(yol: Path, uzanti: str, sayfa: int | None = None, sayac: dict | None = None) -> dict:
    """{"satirlar": prefabrik okuyucunun satırları (TOPLAM hariç), "celik": çelik panel satırları, türleri ve toplamı}."""
    uzanti = uzanti.lower()
    if uzanti == ".pdf":
        c = pdf_celik(yol)
    elif uzanti in (".xlsx", ".xlsm"):
        c = excel_celik(yol, sayfa)
        sayfa = c["sayfa"]
    else:
        raise OkumaHatasi(f"okunamayan dosya türü: {uzanti}")
    satirlar = [s for s in kor_oku.oku(Path(yol), uzanti, sayfa, sayac) if not _toplam_satiri_mi(s)]
    _ayni_satirlar(c["satirlar"], [s for s in satirlar if celik_grubu_mu(s["grup"])])
    return {"satirlar": satirlar, "celik": _ozet(c)}

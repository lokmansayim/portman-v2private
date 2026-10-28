"""İE-16 PR-1 · çelik hakemi: mühür, açık seçimi, liste ve pafta okuyucuları, kıyasın sırası ve çelik kg kapısı.

Her koşuda: gerçek mühür kaydının yapısı (ad yok, durumlar, çift katlar kapalı), açık dört projenin yazılı kuraldan yeniden kurulması, kör ve
kapalı listelerin açılış kaydı olmadan reddedilmesi; uydurma çelik listesi (Excel ve PDF) ve uydurma pafta ile okuyucular; geçici bir depoda
aracın sırası (bizim yanımız kayıtlı → mühür → açılış) ve kapının sınırı. Kaynak dosya gerekmez. İE-16 PR-6a: açılan listenin yeniden
ölçümü (`acik/` klasörü, durum `acildi`) yalnız kör sonuç açılış kaydına bağlı yazılmışsa; kör sonuç yeniden yazılmaz.

Açık beş projenin (hedef ÇP-06 ve açık dört) hakem cevapları `veri/CPnn_hakem_cevabi.json`'da: her koşuda mühürle ve kendi içinde tutarlı.
Tam koşu (`PORTMAN_CELIK_KAYNAK=<klasör>`; dosyalar Drive kimliğiyle adlı, ör. `<drive_id>.xlsx`): beş projenin kıyas listesi (Excel) ve
indirilmiş PDF kopyaları (ÇP-06, ÇP-13) paftalarıyla yeniden okunur, cevap kayıtlıyla aynı olmalı; bir baytı değiştirilmiş kopya okunmadan
durmalı.
Uydurma listedeki ad ve numara uydurmadır; numara çalışma anında kurulur (dosyada düz yazılmaz).
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
from testler.celik import kiyas as K, oku, pafta  # noqa: E402
from testler.kor.kiyas import Durdu  # noqa: E402

MUHUR = KOK / "testler" / "celik" / "muhur.json"
VERI = KOK / "testler" / "celik" / "veri"
FONT = str(KOK / "motor" / "genel" / "font" / "LiberationSans-Regular.ttf")
UYDURMA_AD = "Deneme Kişisi Uydurma"
UYDURMA_TEL = "0" + "5" + "00 000 07 77"
ACIK = ["ÇP-03", "ÇP-07", "ÇP-13", "ÇP-14"]
CIFT = ["ÇP-01", "ÇP-05", "ÇP-10", "ÇP-15", "ÇP-16", "ÇP-21"]
CELIK = [  # ad, boyut, adet, kg
    ("ÇATI MAKASLARI", "80X45X10X0,8", 6, 240),
    ("KOLON - KİRİŞ PANELLERİ", "80X45X10X1,2", 2, 44),
    ("ÇATI PANELLERİ", "80X45X10X1,2", 1, 60),
    ("DUVAR PANELLERİ", "80X45X10X1,2", 10, 520),
    ("Boy C Profil", "80X45X10X0,8", 2, 10),
]
OTEKI = [  # grup, ad, boyut, adet
    ("METAL AKSAM GRUBU", "KÖŞEBENT", "30*30*2500", 20),
    ("METAL AKSAM GRUBU", "AŞIK", "L:4200", 15),
]


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _git(kok: Path, *args):
    subprocess.run(["git", "-C", str(kok), "-c", "user.name=t", "-c", "user.email=t@t", "-c", "commit.gpgsign=false", *args],
                   check=True, capture_output=True)


# ------------------------------------------------------------------------------------------------------- uydurma dosyalar

def _liste_xlsx(yol: Path, celik=CELIK, toplam: tuple | None = None) -> None:
    """Çelik şablonu: başlık bloğunda uydurma ad ve numara, 11. satırda çelik grubunun başlığı (KG ile), TOPLAM, sonra öteki grup."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ÇELİK BİNA"
    ws["B1"] = "PORTATİFEV YÜKLEME LİSTESİ"
    ws["B3"], ws["C3"] = "MÜŞTERİ ADI", UYDURMA_AD
    ws["B4"], ws["C4"] = "TELEFON", UYDURMA_TEL
    for c, v in enumerate(("NO", "ÇELİK PANEL GRUBU", "BOYUT/ÖLÇÜ", "ADET", "KG", "MAKİNE"), 1):
        ws.cell(11, c, v)
    n = 12
    for i, (a, b, adet, kg) in enumerate(celik, 1):
        for c, v in enumerate((i, a, b, adet, kg, "DENEME MAKİNE"), 1):
            ws.cell(n, c, v)
        n += 1
    t = toplam or (sum(x[2] for x in celik), sum(x[3] for x in celik))
    ws.cell(n, 1, "TOPLAM"), ws.cell(n, 4, t[0]), ws.cell(n, 5, t[1])
    n += 1
    for c, v in enumerate(("NO", "METAL AKSAM GRUBU", "BOYUT/ÖLÇÜ", "ADET"), 1):
        ws.cell(n, c, v)
    for i, (_, a, b, adet) in enumerate(OTEKI, 1):
        n += 1
        for c, v in enumerate((i, a, b, adet), 1):
            ws.cell(n, c, v)
    ws.cell(n + 2, 2, "SORUMLU ÜRETİM PERSONELİ AD SOYAD"), ws.cell(n + 2, 4, UYDURMA_AD)
    wb.save(str(yol))


def _liste_xlsx_b(yol: Path, celik, kgsiz_c=False, toplam_kgli=False, kg_baslik="AĞIRLIK", toplam_ad="ÇELİK PANEL TOPLAM") -> None:
    """Açık listelerdeki öteki iki şablon: NO sütunu yok, ad A'da; kg başlığı "AĞIRLIK" ya da "KG"; toplam "ÇELİK PANEL TOPLAM" ya da
    "TOPLAM"; sağda proje bilgi sütunları (uydurma ad ve numara). `kgsiz_c`: boy C profilin kg'ı boş; `toplam_kgli`: TOPLAM'ın adedi yalnız
    kg'lı satırların."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 KATLI"
    ws["A1"] = "YÜKLEME LİSTESİ"
    for c, v in enumerate(("ÇELİK PANEL GRUBU", "BOYUT", "ADET", kg_baslik, "", "KABA HESAP"), 1):
        ws.cell(10, c, v or None)
    ws["J10"], ws["K10"] = "PROJE İSMİ", UYDURMA_AD
    ws["J11"], ws["K11"] = "İLETİŞİM NUMARASI", UYDURMA_TEL
    n = 11
    for a, b, adet, kg in celik:
        bos = kgsiz_c and oku.tur(a) == "c_profil"
        ws.cell(n, 1, a), ws.cell(n, 2, b), ws.cell(n, 3, adet), ws.cell(n, 4, None if bos else kg), ws.cell(n, 7, "AUTOCAD")
        n += 1
    kgli = [x for x in celik if not (kgsiz_c and oku.tur(x[0]) == "c_profil")]
    ws.cell(n, 1, toplam_ad), ws.cell(n, 3, sum(x[2] for x in (kgli if toplam_kgli else celik)))
    ws.cell(n, 4, sum(x[3] or 0 for x in kgli))
    n += 1
    for c, v in enumerate(("METAL AKSAM GRUBU", "BOYUT", "ADET"), 1):
        ws.cell(n, c, v)
    for _, a, b, adet in OTEKI:
        n += 1
        ws.cell(n, 1, a), ws.cell(n, 2, b), ws.cell(n, 3, adet)
    wb.save(str(yol))


def _liste_pdf(yol: Path, celik=CELIK, toplam: tuple | None = None) -> None:
    """Aynı liste PDF olarak (sütun yerleri gerçek şablondaki gibi)."""
    import pymupdf

    doc = pymupdf.open()
    s = doc.new_page(width=595, height=842)

    def yaz(x, y, t):
        s.insert_text((x, y), str(t), fontsize=8, fontname="ls", fontfile=FONT)

    yaz(60, 60, "MÜŞTERİ ADI"), yaz(200, 60, UYDURMA_AD), yaz(60, 72, "TELEFON"), yaz(200, 72, UYDURMA_TEL)
    y = 170
    for x, t in ((54, "NO"), (127, "ÇELİK PANEL GRUBU"), (284, "BOYUT/ÖLÇÜ"), (366, "ADET"), (427, "KG"), (480, "MAKİNE")):
        yaz(x, y, t)
    for i, (a, b, adet, kg) in enumerate(celik, 1):
        y += 11
        yaz(57, y, i), yaz(68, y, a), yaz(285, y, b), yaz(373, y, adet), yaz(426, y, kg), yaz(473, y, "DENEME MAKİNE")
    y += 11
    t = toplam or (sum(x[2] for x in celik), sum(x[3] for x in celik))
    yaz(185, y, "TOPLAM"), yaz(372, y, t[0]), yaz(424, y, t[1])
    y += 11
    for x, t in ((54, "NO"), (125, "METAL AKSAM GRUBU"), (284, "BOYUT/ÖLÇÜ"), (366, "ADET")):
        yaz(x, y, t)
    for i, (_, a, b, adet) in enumerate(OTEKI, 1):
        y += 11
        yaz(57, y, i), yaz(68, y, a), yaz(292, y, b), yaz(373, y, adet)
    doc.save(str(yol))


def _pafta_pdf(yol: Path, sayfalar: list[list[str]]) -> None:
    import pymupdf

    doc = pymupdf.open()
    for sozcukler in sayfalar:
        s = doc.new_page(width=842, height=595)
        s.insert_text((40, 40), "DUVAR PANELLERİ YERLEŞİM PLANI", fontsize=10, fontname="ls", fontfile=FONT)
        for i, w in enumerate(sozcukler):
            s.insert_text((60 + (i % 10) * 70, 100 + (i // 10) * 30), w, fontsize=9, fontname="ls", fontfile=FONT)
    doc.save(str(yol))


# ------------------------------------------------------------------------------------------------------------ gerçek mühür

@pytest.fixture(scope="module")
def muhur():
    return K.muhur_oku(MUHUR)


IZINLI_ALAN = {"kod", "tur", "uzanti", "bayt", "drive_id", "drive_degisme", "sha256", "durum"}


def test_muhur_kaydinda_ad_yok(muhur):
    """Kayıtlarda yalnız izinli alanlar ve biçimli değerler: dosya ya da klasör adı, kişi ya da bayi adı giremez."""
    for r in muhur["dosyalar"]:
        assert set(r) == IZINLI_ALAN
        assert re.fullmatch(r"ÇP-\d{2}", r["kod"]) and r["kod"] in muhur["projeler"]
        assert re.fullmatch(r"[A-Za-z0-9_-]{25,60}", r["drive_id"])
        assert re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?Z", r["drive_degisme"])
        assert r["sha256"] is None or re.fullmatch(r"[0-9a-f]{64}", r["sha256"])
        assert re.fullmatch(r"\.[a-z]{3,4}", r["uzanti"]) and isinstance(r["bayt"], int) and r["bayt"] > 0
    for k, p in muhur["projeler"].items():
        assert set(p) == {"m2", "kat", "liste_tarihi", "drive_klasor", "durum", "kiyas_listesi"}
        assert re.fullmatch(r"[A-Za-z0-9_-]{25,60}", p["drive_klasor"])
    assert len({r["drive_id"] for r in muhur["dosyalar"]}) == len(muhur["dosyalar"]) == 118


def test_projeler_ve_durumlar(muhur):
    P = muhur["projeler"]
    assert list(P) == [f"ÇP-{i:02d}" for i in range(1, 22)]
    assert sorted(k for k, p in P.items() if p["durum"] == "acik") == ACIK
    assert [k for k, p in P.items() if p["durum"] == "hedef"] == ["ÇP-06"]
    assert sorted(k for k, p in P.items() if p["kat"] == "çift") == CIFT
    assert all(P[k]["durum"] == "kapali" for k in CIFT), "çift katlılar kapalı (D-26)"
    assert sum(p["durum"] == "kor" for p in P.values()) == 10
    for r in muhur["dosyalar"]:
        pd = P[r["kod"]]["durum"]
        if r["tur"] in K.LISTELER:
            assert r["durum"] == ("acik" if pd in ("hedef", "acik") else pd), r["drive_id"]
        else:
            assert r["durum"] == "kaynak"
    assert "ÇP-06" in muhur["acilislar"]
    say = muhur["sayilar"]
    assert say["dosyalar"] == {"acik": 15, "kapali": 14, "kaynak": 65, "kor": 24}
    assert say["sha256_dolu"] == sum(1 for r in muhur["dosyalar"] if r["sha256"]) == 17


def test_muhurun_dondurulmus_bolumleri_degismez(muhur):
    """Mühür birleştikten sonra yalnız `acilislar`, `ilk_indirmeler` ve `kat_kayitlari`'na ekleme yapılır (her biri ayrı commit'le);
    öteki bölümlerden biri değişirse iz tutmaz ve araç durur: hakem değişikliğidir, beyanla ve Patron'un onayıyla gelir."""
    assert set(muhur) == {"aciklama", "kaynaklar", "kat_ayrimi", "secim_kurali", "secim", "kiyas_listesi_kurali", "sayilar", "projeler",
                          "dosyalar", *K.EKLENEBILIR}
    assert K.donmus_iz(muhur) == K.MUHUR_IZI == "311e56af37811219e2419cf884f74ade0f6f696537b4a7f94a28d5de3182704c"
    K.muhur_dogrula(MUHUR, KOK)                    # depodaki mühür: iz ve yalnız-ekleme geçmişi tutuyor
    for degistir in (lambda m: m["projeler"]["ÇP-17"].update(durum="acik"), lambda m: m["dosyalar"][0].update(durum="acik"),
                     lambda m: m["projeler"]["ÇP-17"].update(kiyas_listesi=m["dosyalar"][0]["drive_id"])):
        m = json.loads(json.dumps(muhur))
        degistir(m)
        assert K.donmus_iz(m) != K.MUHUR_IZI


def test_muhur_izi_is_emrinde_de_yazili():
    """İzin ikinci kaydı iş emrinin hakem kuralında: mühürle izi birlikte değiştiren commit iş emrini de değiştirmek zorunda (beyan)."""
    assert K.MUHUR_IZI in (KOK / "docs" / "is_emirleri" / "IE-16.md").read_text(encoding="utf-8")


def test_yalniz_ekleme_kurali():
    a = {"acilislar": {"X": {"karar": "k", "sha256": "1" * 64, "liste_commit": "abcdef1"}}, "ilk_indirmeler": {}, "kat_kayitlari": {}}
    K.yalniz_ekleme([a, {**a, "ilk_indirmeler": {"d": {"sha256": "2" * 64}}}])
    for b in ({**a, "acilislar": {"X": {**a["acilislar"]["X"], "liste_commit": "1234567"}}}, {**a, "acilislar": {}}):
        with pytest.raises(Durdu) as e:
            K.yalniz_ekleme([a, b])
        assert e.value.adim == 2 and "yalnız ekleme" in e.value.neden


def test_eklenen_kayitlar_kurala_uygun(muhur):
    """İlk indirme yalnız SHA'sı bekleyen kayda, mühürdeki bayt ve Drive saatiyle; açılış yalnız kör ya da kapalı listeye, o dosyanın
    özetiyle; kat kaydı yalnız kat sayısı belirsiz projeye."""
    D = {r["drive_id"]: r for r in muhur["dosyalar"]}
    P = muhur["projeler"]
    for did, k in muhur["ilk_indirmeler"].items():
        r = D[did]
        assert set(k) == {"sha256", "bayt", "drive_degisme", "tarih"} and re.fullmatch(r"[0-9a-f]{64}", k["sha256"])
        assert r["sha256"] is None and (k["bayt"], k["drive_degisme"]) == (r["bayt"], r["drive_degisme"])
    assert set(muhur["acilislar"]["ÇP-06"]) == {"tarih", "dayanak"}
    for kod, a in muhur["acilislar"].items():
        if kod == "ÇP-06":
            continue
        assert P[kod]["durum"] in ("kor", "kapali") and set(a) == {"karar", "tarih", "sha256", "liste_commit"}
        assert re.fullmatch(r"[0-9a-f]{40}", a["liste_commit"]), "tam commit"
        assert any(K.kayit_ozeti(r, muhur) == a["sha256"] for r in muhur["dosyalar"] if r["kod"] == kod and r["tur"] in K.LISTELER)
    for kod, k in muhur["kat_kayitlari"].items():
        assert P[kod]["kat"] == "belirsiz" and set(k) == {"kat", "tarih", "dayanak", "pafta_sha256"} and k["kat"] in ("tek", "çift")


def test_kiyas_listesi_yazili_kuraldan(muhur):
    """`kiyas_listesi_kurali` mühürdeki alanlarla yeniden kurulur: projenin en yeni Excel yükleme listesi, yoksa en yeni PDF'i."""
    for kod, p in muhur["projeler"].items():
        L = [r for r in muhur["dosyalar"] if r["kod"] == kod and r["tur"] == K.LISTE]
        aday = [r for r in L if r["uzanti"] == ".xlsx"] or [r for r in L if r["uzanti"] == ".pdf"]
        assert p["kiyas_listesi"] == max(aday, key=lambda r: (r["drive_degisme"], r["drive_id"]))["drive_id"], kod


def test_kor_ve_kapali_listelerin_sha_si_yok(muhur):
    """Kör ve kapalı listeler indirilmedi: SHA'ları ilk indirmede (açılışta) alınır."""
    for r in muhur["dosyalar"]:
        if r["durum"] in ("kor", "kapali"):
            assert r["sha256"] is None, r["drive_id"]


def test_acik_secimi_yazili_kuraldan(muhur):
    """`secim_kurali` mühürdeki alanlarla yeniden kurulur: aynı dört proje."""
    P = muhur["projeler"]

    def bant(m2):
        return "A" if m2 <= 70 else "B" if m2 < 90 else "H" if m2 < 100 else "C" if m2 <= 130 else "D"

    uygun = {k: p for k, p in P.items() if p["durum"] != "hedef" and p["kat"] == "tek"}
    sira = lambda b: sorted((k for k, p in uygun.items() if bant(p["m2"]) == b), key=lambda k: (uygun[k]["liste_tarihi"], k))  # noqa: E731
    secim = {b: (sira(b) or [None])[-1] for b in "ABCD"}
    for b in "ABCD":
        if secim[b] is None:
            secim[b] = next(k for k in reversed(sira("H")) if k not in secim.values())
    assert secim == muhur["secim"]
    assert sorted(secim.values()) == ACIK


@pytest.mark.parametrize("liste", [None, {"commit": "0" * 40}])
def test_kor_ve_kapali_acilis_kaydi_olmadan_reddedilir(muhur, liste):
    n = 0
    for r in muhur["dosyalar"]:
        if r["durum"] in ("kor", "kapali"):
            with pytest.raises(Durdu) as e:
                K.acilabilir_mi(r, muhur, r["kod"], liste)
            assert e.value.adim == 3
            n += 1
    assert n == 38


def _cevap_yolu(kod: str) -> Path:
    return VERI / f"CP{kod[3:]}_hakem_cevabi.json"


AYNI = "aynı"
ACIK_CEVAPLAR = {  # kod: (çelik toplamı, paftadaki etiket sayıları, pafta ↔ liste, pafta uyarıları); 25 Ekim 2026, açık listeler
    "ÇP-06": ({"adet": 51, "kg": 2122}, {"D": 26, "K": 4, "M": 13, "C": 2},
              {"duvar_paneli": AYNI, "kolon_kiris": AYNI, "makas": AYNI, "cati_paneli": AYNI}, []),
    "ÇP-03": ({"adet": 101, "kg": 2522.58}, {"D": 12, "K": 7, "M": 27, "C": 20, "T": 26},
              {"duvar_paneli": AYNI, "kolon_kiris": "farklı", "makas": AYNI, "cati_paneli": AYNI, "tavan_paneli": AYNI}, []),
    "ÇP-07": ({"adet": 98, "kg": 2847}, {"D": 26, "K": 7, "M": 34, "C": 17, "B": 8},
              {"duvar_paneli": AYNI, "kolon_kiris": AYNI, "makas": AYNI, "cati_paneli": AYNI, "baca_paneli": AYNI},
              ["MK: eksik numara MK01, MK14", "Ç: eksik numara Ç10"]),
    "ÇP-13": ({"adet": 52, "kg": 1936.87}, {"D": 21, "K": 4},
              {"duvar_paneli": AYNI, "kolon_kiris": AYNI, "makas": "paftada etiket yok", "cati_paneli": "paftada etiket yok",
               "baca_paneli": "paftada etiket yok"}, []),
    "ÇP-14": ({"adet": 80, "kg": 3059.29}, {"D": 37, "K": 6, "M": 26, "C": 5},
              {"duvar_paneli": AYNI, "kolon_kiris": AYNI, "makas": AYNI, "cati_paneli": AYNI},
              ["MK: eksik numara MK06", "Ç: eksik numara Ç02, Ç03, Ç04, Ç05"]),
}


def test_acik_cevaplar_yalniz_hedef_ve_acik_dort():
    assert sorted(ACIK_CEVAPLAR) == sorted(["ÇP-06", *ACIK])
    assert sorted(p.name for p in VERI.glob("*_hakem_cevabi.json")) == sorted(_cevap_yolu(k).name for k in ACIK_CEVAPLAR)


@pytest.mark.parametrize("kod", list(ACIK_CEVAPLAR))
def test_acik_cevap_muhurle_ve_kendi_icinde_tutarli(muhur, kod):
    toplam, pafta_adet, pafta_liste, uyarilar = ACIK_CEVAPLAR[kod]
    c = _oku(_cevap_yolu(kod))
    kayit = {K.kayit_ozeti(r, muhur): r for r in muhur["dosyalar"] if K.kayit_ozeti(r, muhur)}
    liste, pf = kayit[c["liste_sha256"]], kayit[c["pafta_sha256"]]
    assert (c["kod"], liste["kod"], liste["tur"], pf["kod"], pf["tur"]) == (kod, kod, K.LISTE, kod, K.PAFTA)
    assert liste["drive_id"] == muhur["projeler"][kod]["kiyas_listesi"], "cevap projenin kıyas listesinden"
    assert muhur["projeler"][kod]["durum"] == ("hedef" if kod == "ÇP-06" else "acik") and c["m2"] == muhur["projeler"][kod]["m2"]
    s = c["celik"]["satirlar"]
    assert all(x["kg"] is not None or x["tur"] in oku.KG_SIZ for x in s)
    assert c["celik"]["toplam"] == {"adet": sum(x["adet"] for x in s), "kg": round(sum(x["kg"] for x in s if x["kg"] is not None), 3)}
    assert c["celik"]["toplam"] == toplam
    for t, x in c["celik"]["turler"].items():
        assert x["adet"] == sum(y["adet"] for y in s if y["tur"] == t)
    assert {h: x["adet"] for h, x in c["pafta"].items()} == pafta_adet
    assert {h: sum(o["adet"] for o in x["onekler"].values()) for h, x in c["pafta"].items()} == pafta_adet
    assert c["pafta_liste"] == pafta_liste and c["pafta_uyarilari"] == uyarilar


# ----------------------------------------------------------------------------------------------------------- okuyucular

@pytest.mark.parametrize("ad,beklenen", [
    ("DUVAR PANELLERİ", "duvar_paneli"), ("KOLON - KİRİŞ PANELLERİ", "kolon_kiris"), ("ÇATI MAKASLARI", "makas"),
    ("ÇATI PANELLERİ", "cati_paneli"), ("TAVAN PANELLERİ", "tavan_paneli"), ("Boy C Profil", "c_profil"), ("Duvar paneli", "duvar_paneli"),
    # açık beş projenin listelerindeki adlar (25 Ekim)
    ("ZEMİN KAT DIŞ CEPHE DUVAR PANELLERİ", "duvar_paneli"), ("ZEMİN KAT PENCERE ALTI DUVAR", "duvar_paneli"),
    ("VEANDA DUVAR", "duvar_paneli"), ("VERANDA KORKULUK", "duvar_paneli"), ("MÜŞTEMİLAT DUVAR", "duvar_paneli"),
    ("KOLON", "kolon_kiris"), ("KİRİŞ", "kolon_kiris"), ("ÇERÇEVE KİRİŞ", "kolon_kiris"), ("KİRİŞ PANELİ", "kolon_kiris"),
    ("TAVAN PANELİ", "tavan_paneli"), ("ÇATI MAKAS", "makas"), ("ÇATI MAKAS PANELİ", "makas"), ("ÇATI MAKAS PANELLERİ", "makas"),
    ("MÜŞTEMİLAT ÇATI PANELİ", "cati_paneli"), ("BACA PANELLERİ", "baca_paneli"), ("BACA", "baca_paneli"),
    ("PARÇA C PROFİL", "c_profil"),
    # genel motorun adları (ayraç içi türe bakılmadan atılır)
    ("Çatı makas panelleri", "makas"), ("Tavan paneli (makas arası)", "tavan_paneli"),
    ("Çatı paneli (makas üstü, sırt 80 mm)", "cati_paneli"), ("Veranda kirişleri (kiriş paneli)", "kolon_kiris"),
])
def test_tur(ad, beklenen):
    assert oku.tur(ad) == beklenen


@pytest.mark.parametrize("ad", ["ÇELİK KUTU", "DUVAR MAKAS PANELİ"])
def test_tur_bilinmiyor_ya_da_belirsizse_durur(ad):
    with pytest.raises(oku.OkumaHatasi):
        oku.tur(ad)


def test_hata_iletisinde_ad_yazilmaz():
    with pytest.raises(oku.OkumaHatasi) as e:
        oku.tur(f"{UYDURMA_AD} KUTU")
    assert "Deneme" not in str(e.value) and "Kişisi" not in str(e.value) and "…" in str(e.value)
    assert oku.guvenli("ZEMİN KAT DIŞ CEPHE DUVAR PANELLERİ 140") == "ZEMIN KAT DIS CEPHE DUVAR PANELLERI 140"
    for ad in (f"ÇATI {UYDURMA_TEL}", f"ÇATI {UYDURMA_TEL.replace(' ', '')}", f"DUVAR +9{UYDURMA_TEL}"):
        g = oku.guvenli(ad)
        assert not re.search(r"\d{4}", g) and sum(c.isdigit() for c in g) == 0, "telefon gibi sayı dizisi iletiye girmez"


def _beklenen_celik():
    return {"duvar_paneli": (10, 520), "kolon_kiris": (2, 44), "makas": (6, 240), "cati_paneli": (1, 60), "c_profil": (2, 10)}


@pytest.mark.parametrize("bicim", ["xlsx", "pdf"])
def test_liste_okuyucu(tmp_path, bicim):
    yol = tmp_path / f"liste.{bicim}"
    (_liste_xlsx if bicim == "xlsx" else _liste_pdf)(yol)
    sayac = {}
    r = oku.oku(yol, f".{bicim}", sayac=sayac)
    assert {t: (x["adet"], x["kg"]) for t, x in r["celik"]["turler"].items()} == _beklenen_celik()
    assert r["celik"]["toplam"] == {"adet": 21, "kg": 874}
    metin = json.dumps(r, ensure_ascii=False)
    assert UYDURMA_AD not in metin and UYDURMA_TEL not in metin, "başlık bloğu ve personel satırı okunmaz"
    assert [(s["malzeme"], s["adet"]) for s in r["satirlar"] if not oku.celik_grubu_mu(s["grup"])] == [("KÖŞEBENT", 20), ("AŞIK", 15)]
    assert all(oku.katla(s["malzeme"]) != "TOPLAM" for s in r["satirlar"])


def test_toplam_satiri_tutmazsa_durur(tmp_path):
    yol = tmp_path / "liste.xlsx"
    _liste_xlsx(yol, toplam=(21, 900))
    with pytest.raises(oku.OkumaHatasi, match="TOPLAM"):
        oku.oku(yol, ".xlsx")


@pytest.mark.parametrize("bicim", ["xlsx", "pdf"])
def test_pdf_ve_excel_toplam_tutmazsa_durur(tmp_path, bicim):
    yol = tmp_path / f"liste.{bicim}"
    (_liste_xlsx if bicim == "xlsx" else _liste_pdf)(yol, toplam=(21, 900))
    with pytest.raises(oku.OkumaHatasi, match="TOPLAM"):
        oku.oku(yol, f".{bicim}")


@pytest.mark.parametrize("kg_baslik,toplam_ad", [("AĞIRLIK", "ÇELİK PANEL TOPLAM"), ("KG", "TOPLAM")])
def test_oteki_sablonlar(tmp_path, kg_baslik, toplam_ad):
    yol = tmp_path / "liste.xlsx"
    _liste_xlsx_b(yol, CELIK, kg_baslik=kg_baslik, toplam_ad=toplam_ad)
    r = oku.oku(yol, ".xlsx")
    assert {t: (x["adet"], x["kg"]) for t, x in r["celik"]["turler"].items()} == _beklenen_celik()
    assert r["celik"]["toplam"] == {"adet": 21, "kg": 874} and r["celik"]["toplam_kapsami"] == "bütün satırlar"
    metin = json.dumps(r, ensure_ascii=False)
    assert UYDURMA_AD not in metin and UYDURMA_TEL not in metin, "sağdaki proje bilgi sütunları okunmaz"
    assert not any("TOPLAM" in oku.katla(s["malzeme"]) for s in r["satirlar"]), "grubun toplam satırı malzeme sayılmaz"


@pytest.mark.parametrize("toplam_kgli,kapsam", [(True, "kg'lı satırlar"), (False, "bütün satırlar")])
def test_boy_c_profilin_kg_i_bos_olabilir(tmp_path, toplam_kgli, kapsam):
    yol = tmp_path / "liste.xlsx"
    _liste_xlsx_b(yol, CELIK, kgsiz_c=True, toplam_kgli=toplam_kgli, kg_baslik="KG", toplam_ad="TOPLAM")
    c = oku.oku(yol, ".xlsx")["celik"]
    assert c["toplam"] == {"adet": 21, "kg": 864} and c["toplam_kapsami"] == kapsam and c["kgsiz_turler"] == ["c_profil"]
    assert c["turler"]["c_profil"] == {"harf": None, "adet": 2, "kg": None}


def test_baska_turde_kg_bossa_durur(tmp_path):
    yol = tmp_path / "liste.xlsx"
    _liste_xlsx_b(yol, [(a, b, n, None if a == "ÇATI PANELLERİ" else kg) for a, b, n, kg in CELIK])
    with pytest.raises(oku.OkumaHatasi, match="adet ya da kg"):
        oku.oku(yol, ".xlsx")


def test_bilinmeyen_celik_satiri_durur(tmp_path):
    yol = tmp_path / "liste.xlsx"
    _liste_xlsx(yol, celik=[*CELIK, ("ÇELİK KUTU", "40X40", 1, 5)])
    with pytest.raises(oku.OkumaHatasi, match="türü bilinmiyor"):
        oku.oku(yol, ".xlsx")


# İE-16 PR-5 (Lokman, 28 Ekim 2026, "Üç öneri gibi"; beyan İE-16 › Değişiklik (PR-5)): açılan dört kör listenin üç biçimi.

def _liste_satirli(yol: Path, satirlar) -> None:
    """`_liste_xlsx_b` şablonu, çelik grubunun satırları elle (ad, boyut, adet, kg; ad None ise hücre boş), ardından öteki grup."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for c, v in enumerate(("ÇELİK PANEL GRUBU", "BOYUT", "ADET", "KG"), 1):
        ws.cell(10, c, v)
    n = 11
    for x in satirlar:
        for c, v in enumerate(x, 1):
            if v is not None:
                ws.cell(n, c, v)
        n += 1
    for c, v in enumerate(("METAL AKSAM GRUBU", "BOYUT", "ADET"), 1):
        ws.cell(n, c, v)
    ws.cell(n + 1, 1, OTEKI[0][1]), ws.cell(n + 1, 2, OTEKI[0][2]), ws.cell(n + 1, 3, OTEKI[0][3])
    wb.save(str(yol))


def test_adsiz_toplam_satiri_grubun_sonundaysa_toplamdir(tmp_path):
    """ÇP-20'nin biçimi: grubun son satırından sonra adı boş, yalnız adet ve kg dolu satır; sağlama aynen."""
    yol = tmp_path / "liste.xlsx"
    _liste_satirli(yol, [*CELIK, (None, None, 21, 874)])
    c = oku.oku(yol, ".xlsx")["celik"]
    assert c["toplam"] == {"adet": 21, "kg": 874} and c["toplam_kapsami"] == "bütün satırlar" and c["liste_uyarilari"] == []
    assert {t: (x["adet"], x["kg"]) for t, x in c["turler"].items()} == _beklenen_celik()


@pytest.mark.parametrize("satirlar,neden", [
    ([*CELIK, (None, None, 21, 900)], "TOPLAM satırı satırlarla tutmuyor"),
    ([*CELIK[:2], (None, None, 8, 284), *CELIK[2:]], "satırların arasında"),
    ([*CELIK, ("TOPLAM", None, 21, 874), (None, None, 21, 874)], "adlı bir TOPLAM"),
    ([*CELIK, (None, None, 21, 874), (None, None, 21, 874)], "birden çok"),
    ([*CELIK, (None, "80X45", 21, 874)], "TOPLAM satırı okunamadı"),
])
def test_adsiz_satir_toplam_degilse_durur(tmp_path, satirlar, neden):
    """Adsız satır yalnız grubun sonunda, tek ve adlı TOPLAM yokken toplamdır; tutmuyorsa, arada ise, adlı TOPLAM'la ya da iki kez
    geliyorsa okuma durur. Adet ve kg dışında hücresi olan adsız satır toplam sayılmaz (eskisi gibi atlanır): başka TOPLAM yoksa durur."""
    yol = tmp_path / "liste.xlsx"
    _liste_satirli(yol, satirlar)
    with pytest.raises(oku.OkumaHatasi, match=neden):
        oku.oku(yol, ".xlsx")


def test_toplam_tek_satirin_kg_ini_disarida_birakiyorsa_satirlar_esas(tmp_path):
    """ÇP-18'in biçimi: TOPLAM'ın adedi satırlarınki, kg'ı tam bir satırın (kolon-kiriş, 44) kg'ı kadar eksik: satırlar esas, uyarı yazılır."""
    yol = tmp_path / "liste.xlsx"
    _liste_satirli(yol, [*CELIK, ("TOPLAM", None, 21, 874 - 44)])
    c = oku.oku(yol, ".xlsx")["celik"]
    assert c["toplam"] == {"adet": 21, "kg": 874} and c["liste_uyarilari"] == [oku.TOPLAM_BIR_SATIR_EKSIK]
    assert {t: (x["adet"], x["kg"]) for t, x in c["turler"].items()} == _beklenen_celik()


@pytest.mark.parametrize("celik,adet,kg", [(CELIK, 21, 874 - 44 - 60), (CELIK, 21, 874 + 44), (CELIK, 21, 874 - 43),
                                           (CELIK, 20, 874 - 44), (CELIK[:1], 6, 0)])
def test_toplamin_baska_uyusmazligi_yine_durur(tmp_path, celik, adet, kg):
    """İki satır eksik, bir satır fazla, hiçbir satıra denk gelmeyen eksik, adedi de tutmayan ya da kg'ı sıfır TOPLAM (tek satırlı grupta
    "bir satır eksik" sayılmaz): okuma durur."""
    yol = tmp_path / "liste.xlsx"
    _liste_satirli(yol, [*celik, ("TOPLAM", None, adet, kg)])
    with pytest.raises(oku.OkumaHatasi, match="TOPLAM satırı satırlarla tutmuyor"):
        oku.oku(yol, ".xlsx")


def test_veranda_destegi_kolon_kiristir():
    """ÇP-17'nin satırı: veranda desteği kolon-kiriş (K); veranda duvarı duvar paneli kalır."""
    assert oku.tur("VERANDA DESTEK") == oku.tur("Veranda desteği") == oku.tur("VERANDA DESTEK PANELİ") == "kolon_kiris"
    assert oku.tur("VERANDA DUVAR") == "duvar_paneli"
    with pytest.raises(oku.OkumaHatasi, match="türü bilinmiyor: 'DESTEK'"):
        oku.tur("DESTEK")


def test_pafta_etiketleri(tmp_path):
    yol = tmp_path / "pafta.pdf"
    _pafta_pdf(yol, [["D01", "D02", "D03", "D05", "K01", "DN01", "MD1", "AD01", "D01A", "M01", "M2", "M3"], ["M01", "M02", "C01", "C02"],
                     ["MK01", "MK02", "MKP01", "KLN01", "KRS01", "KRS02", "Ç01", "TP01", "B01", "S280", "VD01"], ["(D06)", "d07,"]])
    e = pafta.etiketler(yol)
    T = e["turler"]
    assert {h: x["adet"] for h, x in T.items()} == {"D": 8, "K": 4, "M": 5, "C": 3, "T": 1, "B": 1}, \
        "tek basamaklı etiket sayılır, noktalama atılır, küçük harf büyütülür, M2/M3 birimdir"
    assert {h: sorted(x["onekler"]) for h, x in T.items()} == {"D": ["D", "MD", "VD"], "K": ["K", "KLN", "KRS"], "M": ["M", "MK", "MKP"],
                                                                "C": ["C", "Ç"], "T": ["TP"], "B": ["B"]}
    assert T["D"]["onekler"]["D"]["eksik"] == [4] and T["D"]["onekler"]["D"]["en_buyuk"] == 7
    assert T["M"]["onekler"]["M"]["yinelenen"] == ["M01"] and T["M"]["sayfalar"] == [1, 2, 3]
    assert e["uyarilar"] == ["D: eksik numara D04", "M: birden çok yerde geçen etiket M01"]
    assert e["taninmayan"] == {"AD#": 1, "D#A": 1, "DN#": 1, "S#": 1}, "önek dışındaki ve harf ekli etiket sayılmaz, bilgi olarak durur"
    assert e["birimler"] == {"M2": 1, "M3": 1}


def test_paftada_etiketi_olmayan_tur_ayri_yazilir():
    gercek = {"turler": {"makas": {"harf": "M", "adet": 16, "kg": 471.3}, "duvar_paneli": {"harf": "D", "adet": 2, "kg": 9.0}}}
    etiket = {"turler": {"D": {"adet": 3}}}
    T = {r["tur"]: r["pafta_liste"] for r in K.celik_tablosu(gercek, None, etiket)}
    assert T == {"duvar_paneli": "farklı", "makas": "paftada etiket yok"}


# ------------------------------------------------------------------------------------------------------- sıra ve kapı

BIZIM = [
    {"grup": "Çelik panel", "malzeme": "Duvar paneli", "boyut": "80X45X10X1,2", "adet": 10, "kg": 530.0},
    {"grup": "Çelik panel", "malzeme": "Kolon-kiriş paneli", "boyut": "80X45X10X1,2", "adet": 2, "kg": 44.0},
    {"grup": "Çelik panel", "malzeme": "Makas", "boyut": "80X45X10X0,8", "adet": 7, "kg": 250.0},
    {"grup": "Çelik panel", "malzeme": "Çatı paneli", "boyut": "80X45X10X1,2", "adet": 1, "kg": 60.0},
    {"grup": "Çelik panel", "malzeme": "Boy C profil", "boyut": "80X45X10X0,8", "adet": 2, "kg": 10.0},
    {"grup": "Çelik panel", "malzeme": "Toplam (adedi ve kg'ı bilinenler)", "adet": 22, "kg": 894.0},
    {"grup": "Metal aksam", "malzeme": "Köşebent", "boyut": "30 × 30 × 2500", "adet": 20},
    {"grup": "Metal aksam", "malzeme": "Aşık", "boyut": "L 4200", "adet": 14},
    {"grup": "Metal aksam", "malzeme": "Alın sacı", "boyut": "2500", "adet": 3},
]


@pytest.fixture
def depo(tmp_path):
    """Geçici depo: uydurma mühür ve kendi listelerimiz; kaynaklar depo dışında. DN-01 açık (kıyas listesi Excel; ayrıca PDF kopyası,
    sundurma listesi ve pafta), DN-02 kör (pafta ile), DN-03 açık (SHA'sı ilk indirmede alınacak liste)."""
    kok = tmp_path / "depo"
    cel = kok / "testler" / "celik"
    cel.mkdir(parents=True)
    dis = tmp_path / "disari"
    dis.mkdir()
    Y = {a: dis / a for a in ("acik.xlsx", "acik_pdf.pdf", "sundurma.xlsx", "pafta.pdf", "kor.xlsx", "pafta2.pdf", "bekleyen.xlsx")}
    _liste_xlsx(Y["acik.xlsx"])
    _liste_pdf(Y["acik_pdf.pdf"])
    _liste_xlsx(Y["sundurma.xlsx"], celik=CELIK[:2])
    _pafta_pdf(Y["pafta.pdf"], [["D01", "D02", "K01", "M01"]])
    _liste_xlsx(Y["kor.xlsx"], celik=CELIK[:-1])
    _pafta_pdf(Y["pafta2.pdf"], [["D01", "D02", "D03"]])
    _liste_xlsx(Y["bekleyen.xlsx"], celik=CELIK[1:])

    def did(ad):
        return "x" * 20 + Path(ad).stem.ljust(13, "_")

    def kayit(kod, tur, ad, durum, sha=True):
        yol = Y[ad]
        return {"kod": kod, "tur": tur, "uzanti": yol.suffix, "bayt": yol.stat().st_size, "drive_id": did(ad),
                "drive_degisme": "2026-09-17T05:54:42.192Z", "sha256": _sha(yol) if sha else None, "durum": durum}

    def proje(durum, kiyas):
        return {"m2": 50, "kat": "tek", "liste_tarihi": "2026-09-17", "drive_klasor": "y" * 33, "durum": durum, "kiyas_listesi": did(kiyas)}

    muhur = {"aciklama": "uydurma",
             "projeler": {"DN-01": proje("acik", "acik.xlsx"), "DN-02": proje("kor", "kor.xlsx"), "DN-03": proje("acik", "bekleyen.xlsx")},
             "acilislar": {}, "ilk_indirmeler": {}, "kat_kayitlari": {},
             "dosyalar": [kayit("DN-01", "yükleme listesi", "acik.xlsx", "acik"), kayit("DN-01", "yükleme listesi", "acik_pdf.pdf", "acik"),
                          kayit("DN-01", "sundurma listesi", "sundurma.xlsx", "acik"), kayit("DN-01", "montaj paftası", "pafta.pdf", "kaynak"),
                          kayit("DN-02", "yükleme listesi", "kor.xlsx", "kor"), kayit("DN-02", "montaj paftası", "pafta2.pdf", "kaynak"),
                          kayit("DN-03", "yükleme listesi", "bekleyen.xlsx", "acik", sha=False)]}
    (cel / "muhur.json").write_text(json.dumps(muhur, ensure_ascii=False), encoding="utf-8")
    for kod, ad in (("DN-01", "DN01"), ("DN-02", "DN02"), ("DN-03", "DN03")):
        (cel / "projeler" / ad).mkdir(parents=True)
        (cel / "projeler" / ad / "bizim_liste.json").write_text(json.dumps({"kod": kod, "satirlar": BIZIM}, ensure_ascii=False),
                                                               encoding="utf-8")
    _git(kok, "init", "-q")
    _git(kok, "add", ".")
    _git(kok, "commit", "-q", "-m", "uydurma")
    return {"kok": kok, "cel": cel, "muhur": cel / "muhur.json", "acik": Y["acik.xlsx"], "acik_pdf": Y["acik_pdf.pdf"],
            "sundurma": Y["sundurma.xlsx"], "kor": Y["kor.xlsx"], "pafta": Y["pafta.pdf"], "pafta2": Y["pafta2.pdf"],
            "bekleyen": Y["bekleyen.xlsx"], "bizim": lambda ad: cel / "projeler" / ad / "bizim_liste.json"}


def _muhur_degistir(depo, f):
    m = _oku(depo["muhur"])
    f(m)
    depo["muhur"].write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
    _git(depo["kok"], "commit", "-q", "-am", "mühür")


def _liste_commit(depo, ad="DN02"):
    return subprocess.run(["git", "-C", str(depo["kok"]), "log", "-1", "--format=%H", "--", f"testler/celik/projeler/{ad}"],
                          capture_output=True, text=True).stdout.strip()


def _ac(depo, sha=None, commit=None):
    """DN-02'nin açılış kaydı (Patron kararının yerine uydurma karar)."""
    a = {"karar": "uydurma karar", "tarih": "2026-10-25", "sha256": sha or _sha(depo["kor"]), "liste_commit": commit or _liste_commit(depo)}
    _muhur_degistir(depo, lambda m: m["acilislar"].update({"DN-02": a}))


@pytest.fixture
def okunmaz(monkeypatch):
    """Okuyucu çağrılırsa test düşer: mühür ya da açılış tutmadığında dosya içeriği okunmamalı."""
    def dur(*a, **k):
        raise AssertionError("dosya okundu")
    monkeypatch.setattr(K.oku, "oku", dur)
    monkeypatch.setattr(K.pafta, "etiketler", dur)


def test_kiyas_acik_liste(depo, tmp_path):
    k = K.kiyasla(depo["bizim"]("DN01"), depo["acik"], depo["pafta"], kok=depo["kok"])
    assert k["kapi"] == {"sonuc": "geçti", "fark_yuzde": round(100 * (894 - 874) / 874, 2), "pay_yuzde": 3.0, "kapi_disi": [],
                         "kapi_disi_bizim_kg": None, "bizim_kapi_kg": 894.0, "gercek_kapi_kg": 874.0}
    T = {r["tur"]: r for r in k["tablo"]}
    assert (T["makas"]["gercek_adet"], T["makas"]["bizim_adet"], T["makas"]["pafta_etiket"]) == (6, 7, 1)
    assert T["duvar_paneli"]["pafta_liste"] == "farklı"
    assert k["satir"]["sayilar"] == {"aynı": 1, "farklı": 1, "adet okunamadı": 0, "bizde yok": 0, "bizde fazla": 1}
    s = K.sonuc_kaydi(k)
    assert s["tutanlar"] == [{"grup": "METAL AKSAM", "ad": "KOSEBENT", "olcu": [30.0, 30.0, 2500.0], "adet": 20}]
    assert {x["tur"] for x in s["celik_turleri_tutan"]} == {"duvar_paneli", "kolon_kiris", "cati_paneli", "c_profil"}
    assert "15" not in json.dumps(s["tutanlar"]), "tutmayan satırın gerçek adedi kayda girmez"
    K.rapor_md(k, tmp_path / "r.md")
    metin = (tmp_path / "r.md").read_text(encoding="utf-8")
    assert "geçti" in metin and UYDURMA_AD not in metin and "acik.xlsx" not in metin


def test_adedi_okunamayan_tur_tutan_sayilmaz(depo):
    """Bizim yanda adedi okunamayan tür gerçek listede hiç yoksa None == None tutan değildir."""
    satirlar = BIZIM + [{"grup": "Çelik panel", "tur": "tavan_paneli", "malzeme": "Tavan paneli", "adet": None, "kg": 5.0}]
    depo["bizim"]("DN01").write_text(json.dumps({"kod": "DN-01", "satirlar": satirlar}, ensure_ascii=False), encoding="utf-8")
    _git(depo["kok"], "commit", "-q", "-am", "bizim")
    s = K.sonuc_kaydi(K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"]))
    assert "tavan_paneli" not in {x["tur"] for x in s["celik_turleri_tutan"]}
    assert all(x["adet"] is not None for x in s["celik_turleri_tutan"])


def test_liste_uyarisi_rapora_ve_sonuc_kaydina_yalniz_varsa_girer(depo, tmp_path):
    """İE-16 PR-5: okuyucunun liste uyarısı (TOPLAM bir satırı dışarıda bırakıyor) sonuç kaydına, rapora ve açık listenin cevabına yazılır;
    uyarı yoksa kayıtta alan yok (PR-4b'ye kadar yazılmış sonuçlar ve açık beşin cevabı aynı kalır)."""
    k = K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"])
    assert k["celik"]["liste_uyarilari"] == [] and "liste_uyarilari" not in K.sonuc_kaydi(k)
    K.rapor_md(k, tmp_path / "r0.md")
    assert "Liste uyarıları" not in (tmp_path / "r0.md").read_text(encoding="utf-8")
    k["celik"]["liste_uyarilari"] = [oku.TOPLAM_BIR_SATIR_EKSIK]
    s = K.sonuc_kaydi(k)
    assert s["liste_uyarilari"] == [oku.TOPLAM_BIR_SATIR_EKSIK] and list(s).index("liste_uyarilari") == list(s).index("celik_kg_kapisi") + 1
    K.rapor_md(k, tmp_path / "r1.md")
    assert f"| Liste uyarıları | {oku.TOPLAM_BIR_SATIR_EKSIK} |" in (tmp_path / "r1.md").read_text(encoding="utf-8")
    c = K.cevap("DN-01", depo["acik"], kok=depo["kok"])          # açık listenin cevabı da (yalnız varsa)
    assert "liste_uyarilari" not in K.cevap_json(c)
    c["celik"]["liste_uyarilari"] = [oku.TOPLAM_BIR_SATIR_EKSIK]
    assert K.cevap_json(c)["liste_uyarilari"] == [oku.TOPLAM_BIR_SATIR_EKSIK]
    K.cevap_md(c, tmp_path / "c.md")
    assert f"Liste uyarıları: {oku.TOPLAM_BIR_SATIR_EKSIK}." in (tmp_path / "c.md").read_text(encoding="utf-8")


def test_sonuc_kaydi_okunan_sayfayi_yazar(depo):
    assert K.sonuc_kaydi(K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"], sayfa=0))["muhur"]["liste_sayfasi"] == 0
    assert K.sonuc_kaydi(K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"]))["muhur"]["liste_sayfasi"] == 0


def test_kor_listede_sayfa_secilmez(depo, okunmaz):
    _ac(depo)
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], depo["pafta2"], kok=depo["kok"], sayfa=0)   # pafta da okunmadan
    assert e.value.adim == 3 and "sayfa seçilmez" in e.value.neden


def test_pdf_te_sayfa_secilmez(depo):
    with pytest.raises(oku.OkumaHatasi):
        oku.oku(depo["acik_pdf"], ".pdf", 3)
    assert oku.oku(depo["acik_pdf"], ".pdf")["sayfa"] is None
    assert K.sonuc_kaydi(K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"]))["muhur"]["liste_sayfasi"] == 0


@pytest.mark.parametrize("ek", [{"durum": "acik"}, {"sayfa": 0}])
def test_acilis_kaydinda_fazla_alan_reddedilir(depo, okunmaz, ek):
    a = {"karar": "uydurma karar", "tarih": "2026-10-25", "sha256": _sha(depo["kor"]), "liste_commit": _liste_commit(depo), **ek}
    _muhur_degistir(depo, lambda m: m["acilislar"].update({"DN-02": a}))
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "alanları" in e.value.neden


def test_kendi_listemiz_yalniz_projenin_yolundan(depo, okunmaz):
    """Aynı commit'te ikinci bir aday liste (başka klasörde ya da başka adla) açılış kaydıyla kıyaslanamaz: projenin tek listesi."""
    _ac(depo)
    for yol in (depo["cel"] / "projeler" / "DN02b" / "bizim_liste.json", depo["kok"] / "docs" / "aday.json"):
        yol.parent.mkdir(parents=True, exist_ok=True)
        yol.write_text(depo["bizim"]("DN02").read_text(encoding="utf-8"), encoding="utf-8")
        _git(depo["kok"], "add", ".")
        _git(depo["kok"], "commit", "-q", "-m", "aday")
        with pytest.raises(Durdu) as e:
            K.kiyasla(yol, depo["kor"], kok=depo["kok"])
        assert e.value.adim == 1 and "yolundan okunur" in e.value.neden


def test_acilis_kaydi_tam_commit_ister(depo, okunmaz):
    _ac(depo, commit=_liste_commit(depo)[:7])
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "liste commit" in e.value.neden


@pytest.mark.parametrize("hangisi", ["bizim", "klasor", "muhur"])
def test_sembolik_bagla_baska_dosya_okunmaz(depo, okunmaz, hangisi):
    """Kanonik yola (ya da üst klasörüne, ya da mühüre) commit'siz sembolik bağ konup denetim başka dosyada yapılamaz."""
    _ac(depo)
    kok = depo["kok"]
    (kok / "docs").mkdir()
    if hangisi == "muhur":
        kopya = kok / "docs" / "muhur_kopya.json"
        kopya.write_text(depo["muhur"].read_text(encoding="utf-8"), encoding="utf-8")
        _git(kok, "add", ".")
        _git(kok, "commit", "-q", "-m", "kopya")
        depo["muhur"].unlink()
        depo["muhur"].symlink_to(kopya)
        beklenen = 2
    elif hangisi == "bizim":
        b = kok / "docs" / "b.json"
        b.write_text(depo["bizim"]("DN02").read_text(encoding="utf-8"), encoding="utf-8")
        _git(kok, "add", ".")
        _git(kok, "commit", "-q", "-m", "b")
        depo["bizim"]("DN02").unlink()
        depo["bizim"]("DN02").symlink_to(b)
        beklenen = 1
    else:
        klasor = depo["cel"] / "projeler" / "DN02"
        yedek = kok / "docs" / "DN02"
        klasor.rename(yedek)
        klasor.symlink_to(yedek, target_is_directory=True)
        beklenen = 1
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=kok)
    assert e.value.adim == beklenen and "sembolik bağ" in e.value.neden


@pytest.mark.parametrize("isaret", ["--assume-unchanged", "--skip-worktree"])
def test_git_isaretiyle_gizlenen_degisiklik_durur(depo, okunmaz, isaret):
    """`git status`'u susturan işaret konup commit'siz açılış kaydı yazılırsa araç HEAD'le baytı kıyaslar ve durur."""
    _git(depo["kok"], "update-index", isaret, "testler/celik/muhur.json")
    m = _oku(depo["muhur"])
    m["acilislar"]["DN-02"] = {"karar": "commit'siz", "tarih": "2026-10-25", "sha256": _sha(depo["kor"]), "liste_commit": _liste_commit(depo)}
    depo["muhur"].write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 2 and "HEAD" in e.value.neden
    _git(depo["kok"], "update-index", isaret.replace("--", "--no-"), "testler/celik/muhur.json")
    _git(depo["kok"], "checkout", "--", "testler/celik/muhur.json")
    _git(depo["kok"], "update-index", isaret, "testler/celik/projeler/DN01/bizim_liste.json")
    depo["bizim"]("DN01").write_text(json.dumps({"kod": "DN-01", "satirlar": BIZIM[:1]}, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"])
    assert e.value.adim == 1 and "HEAD" in e.value.neden


def test_gecmeyen_kapinin_farki_kayda_girmez(depo):
    k = K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"])
    k["kapi"] = K.kg_kapisi(1000, 874) | {"kapi_disi": []}
    s = K.sonuc_kaydi(k)["celik_kg_kapisi"]
    assert s == {"sonuc": "geçmedi", "pay_yuzde": 3.0, "kapi_disi": [], "kapi_disi_bizim_kg": None, "fark_yuzde": None}


@pytest.mark.parametrize("bizim_kg,sonuc", [(874 * 1.03, "geçti"), (874 * 0.97, "geçti"), (874 * 1.0301, "geçmedi"),
                                            (874 * 0.9699, "geçmedi"), (None, "koşulmadı")])
def test_kg_kapisi_siniri(bizim_kg, sonuc):
    assert K.kg_kapisi(bizim_kg, 874)["sonuc"] == sonuc


def test_bizim_kg_ya_da_adet_eksikse_sifir_sayilmaz():
    b = K.bizim_celik([{**BIZIM[0], "kg": None}, {**BIZIM[1], "adet": None}, *BIZIM[2:5]])
    assert b["toplam"]["kg"] is None and K.kg_kapisi(b["toplam"]["kg"], 874)["sonuc"] == "koşulmadı"
    assert b["turler"]["kolon_kiris"]["adet"] is None and b["toplam"]["adet"] is None


def test_bizim_satirin_turu_alanindan():
    """Genel motorun adları ("dış duvar (DIS-80)", "DK") türü söylemez: öneren satıra `tur` yazar."""
    b = K.bizim_celik([{"grup": "Çelik panel", "kalem": "dış duvar (DIS-80)", "olcu": "80 × 45", "tur": "duvar_paneli", "adet": 3,
                        "kg": 30.0}])
    assert b["turler"] == {"duvar_paneli": {"harf": "D", "adet": 3, "kg": 30.0}}
    for satir in ({"grup": "Çelik panel", "kalem": "DK", "adet": 1, "kg": 1.0},
                  {"grup": "Çelik panel", "kalem": "DK", "tur": "kapi", "adet": 1, "kg": 1.0}):
        with pytest.raises(oku.OkumaHatasi):
            K.bizim_celik([satir])


def test_kapi_disi_tur_iki_yanda_da_disarida():
    bizim = K.bizim_celik(BIZIM)
    gercek = {"kgsiz_turler": ["c_profil"], "turler": {"duvar_paneli": {"kg": 500.0}, "c_profil": {"kg": None}}}
    assert K.kapi_kg(bizim, gercek) == (884.0, 500.0, ["c_profil"])


def test_kapi_disi_turun_kgli_satiri_da_disarida():
    """Gerçek listede boy C profilin bir satırı kg'lı, öteki boş: türün kg'lı satırı da kapıya girmez (iki yan aynı kapsamda)."""
    g = oku._ozet({"satirlar": [{"ad": "DUVAR PANELİ", "boyut": "80", "adet": 10, "kg": 1000},
                                {"ad": "BOY C PROFİL", "boyut": "80", "adet": 3, "kg": 30}, {"ad": "BOY C PROFİL", "boyut": "80", "adet": 3, "kg": None}],
                   "toplam": {"adet": 16, "kg": 1030}})
    assert g["kgsiz_turler"] == ["c_profil"] and g["turler"]["c_profil"]["kg"] is None and g["toplam"]["kg"] == 1030
    b = K.bizim_celik([{"grup": "Çelik panel", "tur": "duvar_paneli", "malzeme": "Duvar paneli", "adet": 10, "kg": 1060.0},
                       {"grup": "Çelik panel", "tur": "c_profil", "malzeme": "Boy C profil", "adet": 6, "kg": 60.0}])
    bk, gk, disi = K.kapi_kg(b, g)
    assert (bk, gk, disi) == (1060.0, 1000.0, ["c_profil"])
    assert K.kg_kapisi(bk, gk)["sonuc"] == "geçmedi", "+6 %; C profil gerçek yanda kalsaydı +2,9 % geçerdi"


@pytest.mark.parametrize("c_kg,sonuc", [(30.0, "geçti"), (30.01, "koşulmadı"), (300.0, "koşulmadı"), (None, "koşulmadı")])
def test_kapi_disinda_kalan_kg_imiz_sinirli(c_kg, sonuc):
    """Kapı dışı türe hangi satırın gireceğini öneren seçer: dışarıda kalan kg'ımız bilinmeli ve gerçek kapı kg'ının %3'ünü geçmemeli."""
    g = {"kgsiz_turler": ["c_profil"], "turler": {"duvar_paneli": {"kg": 1000.0}, "c_profil": {"kg": None}}}
    b = K.bizim_celik([{"grup": "Çelik panel", "tur": "duvar_paneli", "malzeme": "Duvar paneli", "adet": 10, "kg": 1000.0},
                       {"grup": "Çelik panel", "tur": "c_profil", "malzeme": "Boy C profil", "adet": 6, "kg": c_kg}])
    k = K.kapi(b, g)
    assert k["sonuc"] == sonuc and k["kapi_disi"] == ["c_profil"] and k["kapi_disi_bizim_kg"] == c_kg
    assert K.kapi(b, {"kgsiz_turler": [], "turler": {"duvar_paneli": {"kg": 1000.0}}})["kapi_disi_bizim_kg"] is None


def test_kg_hucresi_bos_ile_okunamayan_ayri():
    assert [oku._kg(v) for v in (None, "", " ", "-", "—", "125,72", 1367, "1.367")] == [None, None, None, None, None, 125.72, 1367, 1367]
    for v in ("15 KG", "yok", "?"):
        with pytest.raises(oku.OkumaHatasi):
            oku._kg(v)


def test_adedi_okunamayan_satir_tutan_sayilmaz():
    g = [{"grup": "METAL AKSAM GRUBU", "malzeme": "KÖŞEBENT", "boyut": "30*30*2500", "adet": None}]
    b = [{"grup": "Metal aksam", "malzeme": "Köşebent", "boyut": "30 × 30 × 2500", "adet": None}]
    r = K.satir_satir(b, g)
    assert r["satirlar"][0]["sonuc"] == "adet okunamadı" and r["sayilar"]["aynı"] == 0


def test_bir_bayti_degisen_liste_okunmadan_durur(depo, tmp_path, okunmaz):
    bozuk = tmp_path / "bozuk.xlsx"
    b = bytearray(depo["acik"].read_bytes())
    b[len(b) // 2] ^= 1
    bozuk.write_bytes(bytes(b))
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), bozuk, kok=depo["kok"])
    assert e.value.adim == 2 and "hiçbir kayıtta yok" in e.value.neden


def test_baska_projenin_listesi_okunmaz(depo, okunmaz):
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 2 and "DN-02" in e.value.neden


def test_kiyas_listesi_olmayan_dosyayla_kiyas_yapilmaz(depo, okunmaz):
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["acik_pdf"], kok=depo["kok"])
    assert e.value.adim == 2 and "kıyas listesi değil" in e.value.neden


def test_sundurma_listesi_kiyasa_girmez(depo, okunmaz):
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["sundurma"], kok=depo["kok"])
    assert e.value.adim == 2 and "sundurma listesi" in e.value.neden


def test_pafta_mühürsüzse_okunmaz(depo, tmp_path, okunmaz):
    p = tmp_path / "baska.pdf"
    _pafta_pdf(p, [["D01"]])
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["acik"], p, kok=depo["kok"])
    assert e.value.adim == 2


def test_kendi_listemiz_kayitli_degilse_acilmaz(depo, okunmaz):
    y = depo["bizim"]("DN01")
    y.write_text(y.read_text(encoding="utf-8").replace('"adet": 10', '"adet": 11'), encoding="utf-8")
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["acik"], kok=depo["kok"])
    assert e.value.adim == 1


def test_kendi_listemizin_turu_belirsizse_gercek_liste_okunmadan_durur(depo, okunmaz):
    _ac(depo)
    y = depo["bizim"]("DN02")
    y.write_text(json.dumps({"kod": "DN-02", "satirlar": [*BIZIM, {"grup": "Çelik panel", "malzeme": "Duvar makas paneli", "adet": 1,
                                                                    "kg": 5.0}]}, ensure_ascii=False), encoding="utf-8")
    _git(depo["kok"], "commit", "-q", "-am", "bizim")
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["kor"], kok=depo["kok"])
    assert e.value.adim == 1 and "belirsiz" in e.value.neden


def test_muhur_degismisse_durur(depo, okunmaz):
    depo["muhur"].write_text(depo["muhur"].read_text(encoding="utf-8").replace("uydurma", "uydurma "), encoding="utf-8")
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"])
    assert e.value.adim == 2 and "değişmiş" in e.value.neden


def test_kor_liste_acilis_kaydi_olmadan_reddedilir_ve_cevabi_cikmaz(depo, okunmaz):
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], depo["pafta2"], kok=depo["kok"])
    assert e.value.adim == 3, "açılıştan önce paftası da okunmaz"
    with pytest.raises(Durdu) as e:
        K.cevap("DN-02", depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "kıyastan önce" in e.value.neden


@pytest.mark.parametrize("sha,commit,neden", [("0" * 64, None, "bu dosya için değil"), (None, "1234567", "commit"), (None, None, None)])
def test_kor_liste_acilis_kaydiyla_acilir(depo, sha, commit, neden):
    _ac(depo, sha, commit)
    if neden:
        with pytest.raises(Durdu) as e:
            K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
        assert e.value.adim == 3 and neden in e.value.neden
    else:
        k = K.kiyasla(depo["bizim"]("DN02"), depo["kor"], depo["pafta2"], kok=depo["kok"])
        assert k["acilis"]["durum"] == "kor" and k["celik"]["toplam"] == {"adet": 19, "kg": 864}


def test_acilis_sonradan_degistirilirse_durur(depo, okunmaz):
    _ac(depo)
    _muhur_degistir(depo, lambda m: m["acilislar"]["DN-02"].update(liste_commit="abcdef0"))
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 2 and "yalnız ekleme" in e.value.neden


# İE-16 PR-6a: açılan listenin yeniden ölçümü (kör sonuç donar; yeniden ölçüm `acik/` klasöründe, durum `acildi`)

BIZIM_YENI = [{**s, "kg": s["kg"] + 10} if s.get("malzeme") == "Makas" else s for s in BIZIM]


def _commit(depo, mesaj):
    _git(depo["kok"], "add", ".")
    _git(depo["kok"], "commit", "-q", "-m", mesaj)


def _kor_olc(depo):
    """DN-02 açılış kaydıyla ölçülür, kör sonucu yazılır ve commit'lenir (`kiyas --sonuc-yaz` gibi)."""
    _ac(depo)
    k = K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    K.sonuc_yaz(K.sonuc_kaydi(k), K.sonuc_hedefi(k, depo["bizim"]("DN02")))
    _commit(depo, "kör sonuç")
    return depo["cel"] / "projeler" / "DN02" / "sonuc.json"


def _yeniden_liste(depo, ad="DN02", kod="DN-02"):
    y = depo["cel"] / "projeler" / ad / K.YENIDEN / "bizim_liste.json"
    y.parent.mkdir(exist_ok=True)
    y.write_text(json.dumps({"kod": kod, "satirlar": BIZIM_YENI}, ensure_ascii=False), encoding="utf-8")
    _commit(depo, "yeniden ölçüm listesi")
    return y


def _okunmasin(monkeypatch):
    def dur(*a, **k):
        raise AssertionError("dosya okundu")
    monkeypatch.setattr(K.oku, "oku", dur)
    monkeypatch.setattr(K.pafta, "etiketler", dur)


def test_acilan_liste_acik_klasorunde_yeniden_olculur(depo, tmp_path):
    kor = _kor_olc(depo)
    once = kor.read_bytes()
    y = _yeniden_liste(depo)
    k = K.kiyasla(y, depo["kor"], depo["pafta2"], kok=depo["kok"])
    acilis = _oku(depo["muhur"])["acilislar"]["DN-02"]
    assert (k["acilis"]["durum"], k["acilis"]["kayit_durumu"]) == ("acildi", "kor")
    assert k["liste"]["commit"] != acilis["liste_commit"], "yeniden ölçümde listemiz açılış kaydındakinden farklı commit'te"
    s = K.sonuc_kaydi(k)
    assert s["muhur"]["durum"] == "acildi" and s["bizim_liste"]["yol"].endswith("DN02/acik/bizim_liste.json")
    hedef = K.sonuc_hedefi(k, y)
    assert hedef == y.parent / "sonuc.json"
    K.sonuc_yaz(s, hedef)
    _commit(depo, "yeniden ölçüm sonucu")
    assert K.sonuc_hedefi(K.kiyasla(y, depo["kor"], kok=depo["kok"]), y) == hedef, "yeniden ölçümün sonucu yeniden yazılabilir"
    assert kor.read_bytes() == once, "kör sonuç değişmez"
    K.rapor_md(k, tmp_path / "r.md")
    assert "durum acildi (kayıtta kor)" in (tmp_path / "r.md").read_text(encoding="utf-8")


def test_kor_sonuc_donar(depo):
    """Kör ya da kapalı projenin yazılmış sonucu yeniden yazılmaz; açık projenin sonucu her ölçümde yazılır."""
    _kor_olc(depo)
    k = K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    with pytest.raises(Durdu) as e:
        K.sonuc_hedefi(k, depo["bizim"]("DN02"))
    assert e.value.adim == 6 and "donar" in e.value.neden
    a = K.kiyasla(depo["bizim"]("DN01"), depo["acik"], kok=depo["kok"])
    K.sonuc_yaz(K.sonuc_kaydi(a), K.sonuc_hedefi(a, depo["bizim"]("DN01")))
    _commit(depo, "açık sonuç")
    assert K.sonuc_hedefi(a, depo["bizim"]("DN01")) == depo["bizim"]("DN01").parent / "sonuc.json"


def test_acilistan_sonra_projenin_listesi_degisirse_yine_durur(depo, monkeypatch):
    """Yeniden ölçüm yalnız `acik/` klasöründen: projenin kendi listesi açılıştan sonra değişirse kıyas eskisi gibi durur."""
    _kor_olc(depo)
    depo["bizim"]("DN02").write_text(json.dumps({"kod": "DN-02", "satirlar": BIZIM_YENI}, ensure_ascii=False), encoding="utf-8")
    _commit(depo, "listemiz değişti")
    _okunmasin(monkeypatch)
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "liste commit'i" in e.value.neden


def test_yeniden_olcum_kor_sonuc_olmadan_durur(depo, okunmaz):
    _ac(depo)
    y = _yeniden_liste(depo)
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["kor"], depo["pafta2"], kok=depo["kok"])
    assert e.value.adim == 3 and "kör sonucu yok" in e.value.neden


def test_yeniden_olcum_acilis_kaydi_olmadan_durur(depo, okunmaz):
    y = _yeniden_liste(depo)
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "açılış kaydı yok" in e.value.neden


@pytest.mark.parametrize("bozuk", [
    lambda s: s["bizim_liste"].update(commit="0" * 40),     # başka listeyle yazılmış
    lambda s: s["muhur"].update(sha256="0" * 64),           # başka dosyanın sonucu
    lambda s: s["muhur"].update(durum="acik"),              # kör ölçümün durumu değil
    lambda s: s.update(kod="DN-01"),                        # başka projenin
], ids=["liste_commit", "ozet", "durum", "kod"])
def test_kor_sonuc_acilisa_bagli_degilse_durur(depo, monkeypatch, bozuk):
    kor = _kor_olc(depo)
    s = _oku(kor)
    bozuk(s)
    kor.write_text(json.dumps(s, ensure_ascii=False), encoding="utf-8")
    _commit(depo, "bozuk kör sonuç")
    y = _yeniden_liste(depo)
    _okunmasin(monkeypatch)
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "bağlı değil" in e.value.neden


def test_kor_sonuc_commitsiz_degismisse_durur(depo, monkeypatch):
    kor = _kor_olc(depo)
    y = _yeniden_liste(depo)
    kor.write_text(kor.read_text(encoding="utf-8") + " ", encoding="utf-8")
    _okunmasin(monkeypatch)
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["kor"], kok=depo["kok"])
    assert e.value.adim == 3 and "değişmiş" in e.value.neden


def test_acik_projede_yeniden_olcum_klasoru_okunmaz(depo, okunmaz):
    y = _yeniden_liste(depo, "DN01", "DN-01")
    with pytest.raises(Durdu) as e:
        K.kiyasla(y, depo["acik"], kok=depo["kok"])
    assert e.value.adim == 3 and "listesi açık" in e.value.neden


@pytest.mark.parametrize("kat,pafta_ad,adim,beklenen", [
    (None, None, 3, "kat sayısı belirsiz"),
    ("tek", "pafta", 3, "paftasına bağlı değil"),          # başka projenin paftası
    ("çift", "pafta2", None, "kapali"),                      # çift katlı çıkan kör proje kapalı sayılır, açılış kaydıyla açılır
    ("tek", "pafta2", None, "kor"),
    ("tek", "yok", 3, "paftasına bağlı değil"),            # SHA'sız pafta kaydı ve SHA'sız kat kaydı (None == None geçmez)
    ("tek", "anahtarsiz", 3, "alanları"),
])
def test_kati_belirsiz_kor_liste(depo, kat, pafta_ad, adim, beklenen):
    def belirsiz(m):
        m["projeler"]["DN-02"].update(kat="belirsiz")
        m["dosyalar"].append({**m["dosyalar"][5], "drive_id": "z" * 33, "sha256": None})   # paftası henüz indirilmemiş kayıt
    _muhur_degistir(depo, belirsiz)
    if kat:
        kk = {"kat": kat, "tarih": "2026-10-26", "dayanak": "uydurma"}
        if pafta_ad != "anahtarsiz":
            kk["pafta_sha256"] = None if pafta_ad == "yok" else _sha(depo[pafta_ad])
        _muhur_degistir(depo, lambda m: m["kat_kayitlari"].update({"DN-02": kk}))
    _ac(depo)
    if adim:
        with pytest.raises(Durdu) as e:
            K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
        assert e.value.adim == adim and beklenen in e.value.neden
    else:
        k = K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
        assert k["acilis"]["durum"] == beklenen and k["acilis"]["kayit_durumu"] == "kor"
        m = K.sonuc_kaydi(k)["muhur"]
        assert (m["durum"], m["kayit_durumu"]) == (beklenen, "kor"), "sonuç kaydı kıyastaki durumu yazar (çift çıkan kapalı)"


@pytest.mark.parametrize("strateji", [["-s", "ours"], ["-X", "theirs"]])
def test_birlestirmede_degisen_acilis_kaydi_durur(depo, strateji):
    """İki dalda aynı projeye iki ayrı açılış kaydı; birleştirme birini alır. Öteki dalın kaydı değişmiş sayılır (her ebeveynle kıyas)."""
    kok = depo["kok"]
    dal = subprocess.run(["git", "-C", str(kok), "rev-parse", "--abbrev-ref", "HEAD"], capture_output=True, text=True).stdout.strip()
    _git(kok, "branch", "yan")
    _ac(depo)
    _git(kok, "checkout", "-q", "yan")
    a = {"karar": "başka karar", "tarih": "2026-10-26", "sha256": _sha(depo["kor"]), "liste_commit": _liste_commit(depo)}
    _muhur_degistir(depo, lambda m: m["acilislar"].update({"DN-02": a}))
    _git(kok, "checkout", "-q", dal)
    _git(kok, "merge", "-q", "--no-edit", *strateji, "yan")
    with pytest.raises(Durdu) as e:
        K.muhur_dogrula(depo["muhur"], kok)
    assert e.value.adim == 2 and "yalnız ekleme" in e.value.neden


def test_sig_klonda_gecmis_eksikse_durur(depo, tmp_path):
    _ac(depo)
    klon = tmp_path / "klon"
    subprocess.run(["git", "clone", "-q", "--depth", "1", f"file://{depo['kok']}", str(klon)], check=True, capture_output=True)
    with pytest.raises(Durdu) as e:
        K.gecmis_denetle(klon, "testler/celik/muhur.json")
    assert e.value.adim == 2 and "sığ klon" in e.value.neden
    K.gecmis_denetle(depo["kok"], "testler/celik/muhur.json")


def test_kat_kaydi_sonradan_degistirilemez(depo, okunmaz):
    _muhur_degistir(depo, lambda m: m["projeler"]["DN-02"].update(kat="belirsiz"))
    _muhur_degistir(depo, lambda m: m["kat_kayitlari"].update(
        {"DN-02": {"kat": "çift", "tarih": "2026-10-26", "dayanak": "uydurma", "pafta_sha256": _sha(depo["pafta2"])}}))
    _muhur_degistir(depo, lambda m: m["kat_kayitlari"]["DN-02"].update(kat="tek"))
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN02"), depo["kor"], kok=depo["kok"])
    assert e.value.adim == 2 and "yalnız ekleme" in e.value.neden


def test_sha_si_bekleyen_kayit_ilk_indirme_olmadan_okunmaz(depo, okunmaz):
    with pytest.raises(Durdu) as e:
        K.kiyasla(depo["bizim"]("DN03"), depo["bekleyen"], kok=depo["kok"])
    assert e.value.adim == 2 and "ilk indirmede" in e.value.neden
    with pytest.raises(Durdu) as e:
        K.ilk_indirme("DN-03", depo["bekleyen"], "2026-10-01T00:00:00Z", K.muhur_oku(depo["muhur"]))
    assert "değişmiş" in e.value.neden


@pytest.mark.parametrize("drive_saati,okunur", [("2026-10-01T00:00:00Z", False), (None, True)])
def test_sha_si_bekleyen_kayit_ilk_indirmeyle_okunur(depo, drive_saati, okunur):
    kayit = K.ilk_indirme("DN-03", depo["bekleyen"], "2026-09-17T05:54:42.192Z", K.muhur_oku(depo["muhur"]))
    (did, ilk), = kayit.items()
    assert ilk["sha256"] == _sha(depo["bekleyen"]) and ilk["bayt"] == depo["bekleyen"].stat().st_size
    _muhur_degistir(depo, lambda m: m["ilk_indirmeler"].update({did: {**ilk, "drive_degisme": drive_saati or ilk["drive_degisme"]}}))
    if okunur:
        k = K.kiyasla(depo["bizim"]("DN03"), depo["bekleyen"], kok=depo["kok"])
        assert k["ozet"] == ilk["sha256"] and k["celik"]["toplam"]["adet"] == 15
    else:
        with pytest.raises(Durdu) as e:
            K.kiyasla(depo["bizim"]("DN03"), depo["bekleyen"], kok=depo["kok"])
        assert e.value.adim == 2 and "Drive saati" in e.value.neden


def test_cevap_acik_listede(depo, tmp_path):
    c = K.cevap("DN-01", depo["acik"], depo["pafta"], kok=depo["kok"])
    j = K.cevap_json(c)
    assert j["celik"]["toplam"] == {"adet": 21, "kg": 874} and j["pafta"]["D"]["adet"] == 2 and c["kiyas_listesi"] is True
    K.cevap_md(c, tmp_path / "c.md")
    metin = (tmp_path / "c.md").read_text(encoding="utf-8")
    assert "KÖŞEBENT" not in metin, "öteki grupların içeriği cevap sayfasına yazılmaz"
    assert "| METAL AKSAM GRUBU | 2 |" in metin
    c2 = K.cevap("DN-01", depo["acik_pdf"], kok=depo["kok"])
    assert c2["kiyas_listesi"] is False and K.cevap_json(c2)["celik"] == j["celik"]


# --------------------------------------------------------------------------------------------------------------- tam koşu

KAYNAK = os.environ.get("PORTMAN_CELIK_KAYNAK")


def _kaynak_listeleri():
    """Tam koşuda okunan listeler: beş projenin kıyas listesi ve SHA'sı bilinen PDF kopyaları."""
    m = K.muhur_oku(MUHUR)
    return [(r["kod"], r["drive_id"] + r["uzanti"]) for r in m["dosyalar"]
            if r["kod"] in ACIK_CEVAPLAR and r["tur"] == K.LISTE and K.kayit_ozeti(r, m)]


def test_tam_kosunun_listeleri():
    uzantilar = sorted((k, Path(d).suffix) for k, d in _kaynak_listeleri())
    assert uzantilar == sorted([(k, ".xlsx") for k in ACIK_CEVAPLAR] + [("ÇP-06", ".pdf"), ("ÇP-13", ".pdf")])


@pytest.mark.skipif(not KAYNAK, reason="tam koşu: PORTMAN_CELIK_KAYNAK=<klasör> (dosyalar Drive kimliğiyle adlı)")
@pytest.mark.parametrize("kod,dosya", _kaynak_listeleri())
def test_tam_kosu_acik_cevaplar(muhur, kod, dosya, tmp_path):
    pf = next(Path(KAYNAK) / f"{r['drive_id']}{r['uzanti']}" for r in muhur["dosyalar"] if r["kod"] == kod and r["tur"] == K.PAFTA)
    liste = Path(KAYNAK) / dosya
    c = K.cevap(kod, liste, pf)
    assert c["kiyas_listesi"] is dosya.endswith(".xlsx")
    j = K.cevap_json(c)
    beklenen = {k: v for k, v in _oku(_cevap_yolu(kod)).items() if k not in ("aciklama", "tarih")}
    assert {**j, "liste_sha256": None} == {**beklenen, "liste_sha256": None}
    bozuk = tmp_path / f"bozuk{liste.suffix}"
    b = bytearray(liste.read_bytes())
    b[-10] ^= 1
    bozuk.write_bytes(bytes(b))
    with pytest.raises(Durdu) as e:
        K.cevap(kod, bozuk, pf)
    assert e.value.adim == 2

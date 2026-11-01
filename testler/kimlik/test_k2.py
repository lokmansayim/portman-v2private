"""İE-15 · K2 · Çatı ve tavan katmanları (karar B ve C; kurallar `kurallar/yerlesim/prefabrik.json`).

- Hakemin K2 kapısı PK-01'de geçer: altı satırda adet farkı 0 (omega 24, aşık 42, kaplama 22, betopan 4, dere 2, iniş 6), metre ±%3
  (omega 58,4575, aşık 148,4, dere 21,2 ↔ 21,196); omega konum satırları 9 / 1 / 14. Kapıda olmayan yer: omega parçaları paftadakilerle
  bire bir (en çok 3 cm, yeri aynı), aşık sıraları yüz başına mahyadan 0, 60 … 300 ve saçakta (1 cm), parça boyları paftadakiyle aynı,
  dere hatları ve inişler paftadaki yerde, kaplama yüz başına 11, betopan makas başına 2.
- K2 modeli K1 modelinin parçalarını aynen taşır; yeni parçalar yalnız K2 kartlarındandır; GlobalId'leri her üretimde aynı.
- Değerler okunur: her `okunan` yol çözülür (liste kuralları ve kitap); yerleşim kurallarıyla liste kurallarında ortak anahtar yok.
- Kural birimleri: ters tavan (başlangıç duvarının yüzünden 60 cm, karşı yüzde hat yok; paftadaki anahtarı `tavan_omegasi`), asma ara
  hattın en kısa sınırı (10 cm dahil) ve L biçimli odada kolun kendi sınırı, aşığın saçak sırası ve en kısa boyu, kaplamanın dizilimi ve
  taşan son şeridi, betopanın mahyaya ortalanması ve tek sıranın örtmediği üçgen ("kural yok"), iniş eşiği, kırma çatı ve sandviç dışı
  kaplama "kural yok", iniş kuralı "esik" değilse iniş yok, kartı olmayan ve iki karta uyan parça, makaslar y ekseninde.
- Uyarılar: Tavan yazılı değilse, İş Emri'nde planda olmayan açık mahal varsa, profil açık mahali bölerse.
- Kesim planı (yedek yok, karar D açık), gösterim değerlerinin sayıyı değiştirmemesi, komutun yerleşim planları (beş sayfa; kural yokta
  EKSİK ve nedeni), GlobalId'nin evden eve ayrışması, K2 kodlarının kilidi.
(Geometri hakeminin K2 ilk turundan sonra eklenenler: en kısa sınırlar, L biçimli oda, kaplama dizilimi ve taşma, betopan, uyarılar,
EKSİK, ev kimliği, kod kilidi.)
K3b'de (gerekçe `docs/is_emirleri/IE-15.md` › Değişiklik (K3b)): komutun planları yedi sayfa (yedincisi alt çerçeve ve çatı kenarları,
`test_k3b.py`), Excel'in kesim planında K3b'nin yedi kartı da.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import sys
from collections import Counter
from pathlib import Path

import ifcopenshell
import ifcopenshell.util.element as ue
import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
sys.path.insert(0, str(KOK / "testler" / "kimlik" / "hakem"))
import kiyas as KY  # noqa: E402
from motor import ifc_uret  # noqa: E402
from motor import kart_model as KM  # noqa: E402
from motor import yerlesim as Y  # noqa: E402
from motor.yerlesim import cati, girdi, kesim, tavan  # noqa: E402
from motor.yerlesim.girdi import Duvar  # noqa: E402

VERI = KOK / "testler" / "kimlik" / "veri"
MOTOR_VERI = KOK / "testler" / "motor" / "veri"
PROJE = KOK / "testler" / "kimlik" / "hakem" / "projeler" / "PK01.json"
K2_KODLARI = ("OM-01", "AS-01", "KP-01", "BT-01", "DR-01", "DR-02")
K5B2 = {"OM-02": 3, "H-11": 72, "AL-01": 76, "AL-02": 8, "IZ-01": 8}      # İE-15 K5b-2: verandaya bakan dış duvarın omegası, tavan H'ı ve gözleri, çatı izolasyonu
KILIT_K2 = "21476fe2de4d20ca5d111694f90de7f8ce203fe6c6dadc65d4b6d2710f4ac6e0"   # K1'in 25 ve K2'nin 7 kodu (bu PR)


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _ie(**degisiklik) -> dict:
    ie = _oku(MOTOR_VERI / "PK01.json")
    for yol, deger in degisiklik.items():
        v = ie
        *ust, son = yol.split("__")
        for a in ust:
            v = v[a]
        v[son] = deger
    return ie


def _model(ie: dict, klasor: Path, ad: str, **kw) -> dict:
    bugunku = klasor / f"{ad}.ifc"
    assert ifc_uret.uret(ie, bugunku)["sema_hatasi"] == 0
    say = KM.kart_modeli(bugunku, klasor / f"{ad}_model.ifc", ie["cati"].get("tip", "besik"), is_emri=ie, plan=ifc_uret.hazirla(ie), **kw)
    return {"bugunku": bugunku, "model": klasor / f"{ad}_model.ifc", "say": say}


def _parcalar(model: Path, kod: str) -> list[dict]:
    f = ifcopenshell.open(str(model))
    kutu = KM._Kutular()
    out = []
    for e in f.by_type("IfcProduct"):
        p = ue.get_psets(e).get(KM.PARCA_SETI, {})
        if p.get("KartKodu") == kod:
            a, b = kutu(e)
            out.append({"ad": e.Name, "p": p, "a": a, "b": b, "q": ue.get_psets(e, qtos_only=True)})
    return out


def _adet(say: dict) -> dict:
    return {k["kod"]: k["adet"] for k in say["kartlar"]}


@pytest.fixture(scope="module")
def cevap():
    return _oku(VERI / "PK01_pafta_katmani.json")


@pytest.fixture(scope="module")
def pk01(tmp_path_factory):
    d = tmp_path_factory.mktemp("k2")
    m = _model(_ie(), d, "PK01")
    k1 = KM.kart_modeli(m["bugunku"], d / "PK01_k1.ifc", "besik")
    return m | {"klasor": d, "k1": d / "PK01_k1.ifc", "say_k1": k1}


# ----------------------------------------------------------------------------- PK-01: hakemin kapısı ve paftayla yan yana
def test_hakem_k2_kapisi_gecer(pk01):
    k = KY.kiyasla(pk01["model"], _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE), adim="K2")
    assert k["kapi_1"] == "geçer" and k["kart_modeli"] == 111 + 100
    satir = {s["anahtar"]: s for s in k["satirlar"]}
    assert {a: satir[a]["modelden"] for a in ("duvar_omegasi_asma", "asik", "sandvic_panel", "bas_makas_betopan", "dere", "inis")} == {
        "duvar_omegasi_asma": 24, "asik": 42, "sandvic_panel": 22, "bas_makas_betopan": 4, "dere": 2, "inis": 6}
    assert all(satir[a]["sonuc"] == "geçer" and satir[a]["kapida"] for a in ("duvar_omegasi_asma", "asik", "sandvic_panel", "bas_makas_betopan",
                                                                               "dere", "inis"))
    assert [satir[f"ayrinti.omega.konum.{y}"]["modelden"] for y in ("duvar_ekseni", "duvar_yuzu", "ara")] == [9, 1, 14]
    assert {m["anahtar"]: (m["modelden_m"], m["sonuc"]) for m in k["metre"]} == {
        "duvar_omegasi_asma": (58.4575, "geçer"), "asik": (148.4, "geçer"), "dere": (21.2, "geçer")}
    assert ifc_uret.ozet(pk01["model"])["sema_hatasi"] == 0 and pk01["say"]["kural_yok"] == []


def test_k1_kapisi_ve_sayimi_aynen(pk01):
    """K2 modeli K1'in sayımını değiştirmez: K1 kartlarının adedi, H türleri, görünüm, bileşen ve kartsızlar aynı; adımsız Kapı 1 geçer.
    Yerleşim katmanı K5b-2'den beri tavan katmanını ve verandaya bakan dış duvarın omegasını (OM-02) da ekler."""
    a, b = pk01["say_k1"], pk01["say"]
    k1 = {k["kod"]: k["adet"] for k in a["kartlar"]}
    k2 = _adet(b)
    katman = K2_KODLARI + tuple(K5B2)
    assert {k: v for k, v in k2.items() if k not in katman} == {k: v for k, v in k1.items() if k not in katman}
    assert [k2[k] for k in K2_KODLARI] == [24, 42, 22, 4, 2, 6] and all(k1[k] == 0 for k in katman)
    assert {k: k2[k] for k in K5B2} == K5B2
    for alan in ("h_turu", "gorunum", "bilesen", "kartsiz"):
        assert a[alan] == b[alan], alan
    assert KY.kiyasla(pk01["model"], _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE))["kapi_1"] == "geçer"


def test_k1_parcalari_aynen_tasinir(pk01):
    """K2 yalnız ekler: K1 modelinin her parçası K2 modelinde aynı sınıf, ad, yer, geometri ve setlerle; yeni 100 parça K2 kartlarından,
    167'si K5b-2'nin (dış duvar omegası, tavan H'ı ve gözleri, izolasyon şeritleri)."""
    a, b = ifcopenshell.open(str(pk01["k1"])), ifcopenshell.open(str(pk01["model"]))
    A = {e.GlobalId: e for e in a.by_type("IfcProduct")}
    B = {e.GlobalId: e for e in b.by_type("IfcProduct")}
    assert set(A) <= set(B)
    for g, e in A.items():
        m = B[g]
        assert (m.is_a(), m.Name, ue.get_predefined_type(m)) == (e.is_a(), e.Name, ue.get_predefined_type(e)), e.Name
        if e.Representation:
            assert len(list(a.traverse(e.Representation))) == len(list(b.traverse(m.Representation))), e.Name
        pa, pb = ue.get_psets(e), ue.get_psets(m)
        assert {k: {x: y for x, y in v.items() if x != "id"} for k, v in pa.items()} == \
               {k: {x: y for x, y in v.items() if x != "id"} for k, v in pb.items()}, e.Name
    yeni = [B[g] for g in set(B) - set(A)]
    assert len(yeni) == 100 + 167
    assert Counter(ue.get_psets(e)[KM.PARCA_SETI]["KartKodu"] for e in yeni) == {"OM-01": 24, "AS-01": 42, "KP-01": 22, "BT-01": 4, "DR-01": 2,
                                                                                  "DR-02": 6} | K5B2
    for e in yeni:
        t = ue.get_type(e)
        assert t is not None and t.Tag == ue.get_psets(e)[KM.PARCA_SETI]["KartKodu"] and e.ContainedInStructure, e.Name


def test_globalid_her_uretimde_ayni(pk01, tmp_path):
    m = _model(_ie(), tmp_path, "PK01b")
    def ad(yol):
        f = ifcopenshell.open(str(yol))
        return {e.GlobalId: e.Name for e in f.by_type("IfcProduct") if ue.get_psets(e).get(KM.PARCA_SETI, {}).get("KartKodu") in K2_KODLARI}
    assert ad(m["model"]) == ad(pk01["model"]) and len(ad(m["model"])) == 100
    baska = _model(_ie(is_emri_no="PM-BASKA-EV"), tmp_path, "baska")         # iki evin modeli birlikte açılınca çakışmaz
    assert sorted(ad(baska["model"]).values()) == sorted(ad(pk01["model"]).values()) and not set(ad(baska["model"])) & set(ad(pk01["model"]))
    x, y = m["say"], pk01["say"]
    assert {k: v for k, v in x.items() if k not in ("bugunku_ifc", "bugunku_ifc_sha256", "kart_modeli")} == \
           {k: v for k, v in y.items() if k not in ("bugunku_ifc", "bugunku_ifc_sha256", "kart_modeli")}


def test_omega_parcalari_paftayla_bire_bir(pk01, cevap):
    """Modeldeki her omega parçası paftadaki bir parçaya (yönü, ekseni, iki ucu ≤ 3 cm) bire bir düşer ve yeri (duvar ekseni, duvar yüzü,
    ara) aynıdır. En büyük sapma salonun ara hattı: model 512, pafta 509,6 (ızgaranın başından 125,5 cm'nin katı)."""
    kalan = list(cevap["ayrinti"]["omega"]["parcalar"])
    en = 0.0
    for p in _parcalar(pk01["model"], "OM-01"):
        a, b = p["a"], p["b"]
        yon = "x" if b[0] - a[0] > b[1] - a[1] else "y"
        konum, bas, son = ((a[1] + b[1]) / 2, a[0], b[0]) if yon == "x" else ((a[0] + b[0]) / 2, a[1], b[1])
        q = min((q for q in kalan if q["yon"] == yon), key=lambda q: abs(q["konum_cm"] - konum) + abs(q["bas_cm"] - bas) + abs(q["son_cm"] - son))
        d = max(abs(q["konum_cm"] - konum), abs(q["bas_cm"] - bas), abs(q["son_cm"] - son))
        assert d <= 3 and q["yer"] == p["p"]["Yer"], (p["ad"], p["p"]["Yer"], q)
        assert abs(q["boy_cm"] - p["p"]["Boy_cm"]) <= 0.5, (p["ad"], q)
        en = max(en, d)
        kalan.remove(q)
    assert kalan == [] and en == pytest.approx(2.4, abs=0.05)
    boylar = Counter(f"{p['p']['Boy_cm']:g}" for p in _parcalar(pk01["model"], "OM-01"))
    assert boylar == cevap["ayrinti"]["omega"]["boylar"]


def test_asik_siralari_paftayla(pk01, cevap):
    """Yüz başına 7 sıra: mahyadan eğik boyca 0, 60 … 300 ve saçakta (347,5; paftada 347,3); paftadaki sıra konumlarıyla 1 cm içinde.
    Parçalar arka uçtan 410'luk, kalan 240 (paftada 28 × 410 + 14 × 240)."""
    P = _parcalar(pk01["model"], "AS-01")
    model = {y: sorted({p["p"]["MahyadanEgik_cm"] for p in P if p["p"]["Yuz"] == y}) for y in ("sol", "sag")}
    assert all(len(v) == 7 for v in model.values())
    for yuz in cevap["ayrinti"]["asik"]["yuzler"]:
        k = yuz["sira_konumlari_cm"]
        mahyadan = sorted(k) if k[1] > 55 else sorted(round(k[-1] - x, 1) for x in k)      # pafta bir yüzü saçaktan ölçer
        assert any(all(abs(x - y) <= 1 for x, y in zip(mahyadan, v)) for v in model.values()), mahyadan
    assert Counter(f"{p['p']['Boy_cm']:g}" for p in P) == cevap["ayrinti"]["asik"]["boylar"]
    for p in P:                                                          # arka uçtan tam parça: 240'lık parçalar ön uçta
        if p["p"]["Boy_cm"] == 240:
            assert p["a"][1] == pytest.approx(-30)


def test_dere_ve_inis_paftayla(pk01, cevap):
    D = _parcalar(pk01["model"], "DR-01")
    for h in cevap["ayrinti"]["dere"]["hatlar"]:                        # hat saçak çizgisinde, mahya boyunca
        d = next(p for p in D if abs(min(abs(p["a"][0] - h["konum_cm"]), abs(p["b"][0] - h["konum_cm"]))) <= 1)
        assert abs(d["a"][1] - h["bas_cm"]) <= 1 and abs(d["b"][1] - h["son_cm"]) <= 1
    I = _parcalar(pk01["model"], "DR-02")
    kalan = list(I)
    for q in cevap["ayrinti"]["dere"]["inisler"]:
        p = min(kalan, key=lambda p: math.hypot((p["a"][0] + p["b"][0]) / 2 - q["xy"][0], (p["a"][1] + p["b"][1]) / 2 - q["xy"][1]))
        assert math.hypot((p["a"][0] + p["b"][0]) / 2 - q["xy"][0], (p["a"][1] + p["b"][1]) / 2 - q["xy"][1]) <= 15, q
        assert (p["p"]["Yer"] == "orta") == (q["yer"] == "ara"), q
        kalan.remove(p)


def test_kaplama_ve_betopan_paftayla(pk01, cevap):
    K = _parcalar(pk01["model"], "KP-01")
    assert sorted(Counter(p["p"]["Yuz"] for p in K).values()) == cevap["ayrinti"]["kaplama"]["gruplar"]
    assert {p["p"]["PanelBoyu_cm"] for p in K} == {353}                 # kitabın boyu; paftada 350 (iş emrinde bilinen fark)
    assert Counter(p["p"]["Genislik_cm"] for p in K) == {100: 22}         # şerit kesilmez, son şerit taşar (paftada 22 tam panel)
    tasan = [p for p in K if p["p"]["Tasma_cm"]]
    assert sorted(p["p"]["Tasma_cm"] for p in tasan) == [40, 40] and all(p["p"]["Serit"] == 11 for p in tasan)
    assert all(p["a"][1] == pytest.approx(-70) for p in tasan)          # dizilim arka uçtan (arka-sol): taşma ön saçakta
    assert all(p["b"][1] == pytest.approx(1030) for p in K if p["p"]["Serit"] == 1)
    B = _parcalar(pk01["model"], "BT-01")
    assert sorted(Counter(p["p"]["Makas"] for p in B).values()) == cevap["ayrinti"]["betopan"]["gruplar"]
    for p in B:                                                          # binanın uç duvarlarının dış yüzünde, alt başlıktan yukarı
        assert p["b"][1] <= 0 + 1e-6 or p["a"][1] >= 1000 - 1e-6, p["ad"]
        assert p["a"][2] == pytest.approx(250) and p["p"]["Levha_cm"] == "250 × 125"
    assert all(p["q"]["Qto_PlateBaseQuantities"]["GrossArea"] == pytest.approx(3.125) for p in B)
    for ad in {p["p"]["Makas"] for p in B}:                              # mahyaya ortalı: 50 … 550 (mahya 300)
        u = [(p["a"][0], p["b"][0]) for p in B if p["p"]["Makas"] == ad]
        assert min(a for a, _ in u) == pytest.approx(50) and max(b for _, b in u) == pytest.approx(550)


def test_boy_ve_alan_qto(pk01):
    f = ifcopenshell.open(str(pk01["model"]))
    for e in f.by_type("IfcProduct"):
        p = ue.get_psets(e).get(KM.PARCA_SETI, {})
        if p.get("KartKodu") in ("OM-01", "AS-01", "DR-01"):
            q = ue.get_psets(e, qtos_only=True)
            ad = {"OM-01": "Qto_MemberBaseQuantities", "AS-01": "Qto_MemberBaseQuantities", "DR-01": "Qto_PipeSegmentBaseQuantities"}[p["KartKodu"]]
            assert q[ad]["Length"] == pytest.approx(p["Boy_cm"] / 100), e.Name


# ----------------------------------------------------------------------------- kurallar ve değerler
def test_degerler_okunur():
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    assert not set(y["kurallar"]) & set(liste["kurallar"]), "aynı anahtar iki kural dosyasında"
    assert set(Y.KURALLAR) <= set(y["kurallar"])
    D = Y.degerler(y, liste)
    T, A, C, B, R = (D[k] for k in ("tavan_omegasi_yerlesimi", "asik_sirasi", "cati_kaplamasi", "bas_makas_betopani_yerlesimi", "dere_hatti"))
    assert (T["aralik_asma"], T["aralik_ters"], T["parca_boyu"], T["pay"], T["en_kisa_cm"]) == (1.255, 0.6, 3.75, 0.1, 10)
    assert (A["aralik"], A["parca_boyu"], A["pay"], A["sacak_esigi"], A["en_kisa_cm"]) == (0.6, 4.2, 0.1, 0.3, 10)   # eşik liste kuralından (BK-19)
    assert (C["kapatma"], C["tasma"], C["yuvarlama_mm"]) == (1.0, 0.05, 10)
    assert (B["katsayi"], B["levha_m2"], B["levha_cm"]) == (1.5, 3.125, [250, 125])
    assert (R["orta_esik"], R["orta_kurali"], R["kose_basina"], R["parca_boyu"], R["bindirme"]) == (9.0, "esik", 1, 4.0, 0.2)
    durumlar = _oku(KM.KARTLAR)["durumlar"]
    for ad in Y.KURALLAR:
        k = y["kurallar"][ad]
        assert k["durum"] in durumlar and k["tarih"] and (k.get("kaynak") or k.get("not")), ad
        for alt, v in k.items():                                         # kuralı uygulamak için seçilen değer kendi durumunu taşır
            if isinstance(v, dict) and "deger" in v:
                assert v["durum"] in durumlar and v.get("not"), (ad, alt)
    with pytest.raises(KeyError):
        Y.oku("kitap › yok_boyle_bir_kural", liste, Y.kitap_oku(liste))


# ----------------------------------------------------------------------------- kuralın birimleri
def test_ters_tavan(pk01, tmp_path):
    """Ters tavan: oda hatları başlangıç duvarının yüzünden 60 cm arayla, ilki "duvar yüzünde"; iç duvarlar 60'lık olduğu için duvar eksenindeki
    hatlar da kalır. Paftadaki anahtarı `tavan_omegasi` (ters tavanın listedeki adı)."""
    ie = _ie()
    ie["ozellikler"]["Tavan"] = "Ters tavan"
    m = _model(ie, tmp_path, "ters")
    assert m["say"]["yerlesim"]["tavan"] == "ters" and m["say"]["kural_yok"] == []
    O = _parcalar(m["model"], "OM-01")
    assert all(p["p"]["PaftaAnahtari"].split(";")[0] == "tavan_omegasi" for p in O)
    hat = m["say"]["yerlesim"]["omega_hatlari"]
    oda = [h for h in hat if h["mahal"] == "SALON"]
    s = sorted({h["s_cm"] for h in oda})
    assert all(b - a == pytest.approx(60) for a, b in zip(s, s[1:])) and oda[0]["yer"] == "duvar_yuzu"
    eksen = lambda H: [(h["eksen"], h["s_cm"], h["r0_cm"], h["r1_cm"]) for h in H if h["yer"] == "duvar_ekseni"]
    assert eksen(hat) == eksen(pk01["say"]["yerlesim"]["omega_hatlari"]) and len(eksen(hat)) > 0
    assert _adet(m["say"])["OM-01"] > 24
    plan = ifc_uret.hazirla(ie)                                          # başlangıç sol (x küçük): ilk hat odanın sol yüzünde, sağ yüze 10 cm'den yakın hat yok
    from shapely.geometry import Polygon
    for o in plan["odalar"]:
        x0, _y0, x1, _y1 = Polygon(o["poligon"]).bounds
        s = sorted({h["s_cm"] for h in hat if h["mahal"] == o["ad"]})
        assert s and s[0] == pytest.approx(x0) and s[-1] <= x1 - 10 + 1e-6, (o["ad"], s, x0, x1)


def test_asik_saçak_sirasi():
    """Karar C: sıra mahyadan aralıkla; saçağa aralıktan az kalınca saçakta bir sıra. Eğik boy aralığın katıysa son sıra saçakta (çift sıra yok).
    Saçağa saçak eşiğinden (liste kuralı `asik_sacak_esigi_m`, 30 cm; BK-19, dört aşık planından) az kalan sıra konmaz."""
    def g(aciklik_cm, egim):
        return {"W": aciklik_cm, "H": 500, "duvar_h": 250, "cati": {"tip": "besik", "egim": egim, "sacak": 30, "topuk": 10, "makas_ekseni": "x"},
                "baslangic": {"x": "min", "y": "max"}, "makaslar": []}
    K = {"kesit_asik": [5, 5], "aralik": 60.0, "tam": 410.0, "en_kisa": 10.0, "sacak_esigi": 30.0}
    for aciklik, egim in ((600, 0.33), (500, 0.30), (840, 0.0)):
        F = cati.cerceve(g(aciklik, egim))
        d = sorted({p["bilgi"]["MahyadanEgik_cm"] for p in cati.asik(g(aciklik, egim), K)})
        assert d[-1] == pytest.approx(round(F["egik"], 2)) and all(b - a <= 60 + 1e-6 for a, b in zip(d, d[1:]))
        assert all(x == pytest.approx(60 * i) for i, x in enumerate(d[:-1]))
    F = cati.cerceve(g(1140, 0.0))                                       # eğik 600: 0 … 540 ve 600, saçakta tek sıra
    assert F["egik"] == pytest.approx(600)
    assert sorted({p["bilgi"]["MahyadanEgik_cm"] for p in cati.asik(g(1140, 0.0), K)}) == [60.0 * i for i in range(11)]
    for aciklik, son in ((1141, 600.5), (1198, 629.0), (1200, 630.0)):  # saçağa 30 cm'den az kalan sıra konmaz (30 cm dahil kalır)
        d = sorted({p["bilgi"]["MahyadanEgik_cm"] for p in cati.asik(g(aciklik, 0.0), K)})
        assert d[-1] == pytest.approx(son) and (d[-2] == 540 if son - 600 < 30 else d[-2] == 600) and d[-1] - d[-2] <= 90 + 1e-6
    def boylar(H):                                                       # mahya boyu H + 60; kalan 10 cm'den kısaysa parçasız
        gg = dict(g(600, 0.33), H=H)
        return sorted({p["boy_cm"] for p in cati.asik(gg, K)})
    assert boylar(763) == [410.0] and boylar(770) == [10.0, 410.0] and boylar(1000) == [240.0, 410.0]


def test_inis_esigi():
    """İniş hattın iki ucunda; 9 m'yi geçen hattın ortasında bir tane daha (9 m tam ise yok)."""
    def n(boy_cm):
        g = {"W": 600, "H": boy_cm - 60, "duvar_h": 250, "cati": {"tip": "besik", "egim": 0.33, "sacak": 30, "topuk": 10, "makas_ekseni": "x"},
             "baslangic": {"x": "min", "y": "max"}, "makaslar": []}
        P = cati.dere(g, {"kesit_dere": [15, 10], "kesit_inis": 10, "orta_esik_cm": 900, "kose_basina": 1})
        return sum(1 for p in P if p["parca"] == "inis"), sum(1 for p in P if p["parca"] == "dere")
    assert n(900) == (4, 2) and n(901) == (6, 2) and n(1060) == (6, 2)


def test_kirma_catida_cati_katmanlari_kural_yok(tmp_path):
    m = _model(_ie(cati__tip="kirma"), tmp_path, "kirma")
    k2 = [x for x in m["say"]["kural_yok"] if x["parca"] in ("aşık", "çatı kaplaması", "baş makas betopanı", "dere ve iniş")]
    assert len(k2) == 4 and all("\"kirma\"" in x["neden"] and "besik" in x["neden"] for x in k2)
    a = _adet(m["say"])
    assert [a[k] for k in ("AS-01", "KP-01", "BT-01", "DR-01", "DR-02")] == [0] * 5 and a["OM-01"] == 24     # tavan çatıya bağlı değil
    (tmp_path / "kirma.json").write_text(json.dumps(_ie(cati__tip="kirma"), ensure_ascii=False), encoding="utf-8")
    assert KM.main([str(tmp_path / "kirma.json"), str(tmp_path / "cikti")]) == 1


def test_sandvic_disi_kaplama_kural_yok(tmp_path):
    ie = _ie()
    ie["ozellikler"]["Çatı Kaplaması"] = "Trapez Sac"
    m = _model(ie, tmp_path, "trapez")
    assert [x["parca"] for x in m["say"]["kural_yok"]] == ["çatı kaplaması"] and "trapez sac" in m["say"]["kural_yok"][0]["neden"]
    a = _adet(m["say"])
    assert a["KP-01"] == 0 and (a["AS-01"], a["BT-01"], a["DR-01"], a["DR-02"]) == (42, 4, 2, 6)


def test_inis_kurali_esik_degilse_inis_yok(tmp_path):
    liste = _oku(KM.LISTE_KURALLARI)
    liste["kurallar"]["inis_orta"]["deger"] = "her_hat"
    ie = _ie()
    m = _model(ie, tmp_path, "inis", liste_kurallari=liste)
    assert [x["parca"] for x in m["say"]["kural_yok"]] == ["iniş"]
    assert _adet(m["say"])["DR-02"] == 0 and _adet(m["say"])["DR-01"] == 2


def test_karti_olmayan_ve_iki_karta_uyan_parca(tmp_path):
    kartlar = _oku(KM.KARTLAR)
    k2 = copy.deepcopy(kartlar)
    k2["kartlar"] = [k for k in k2["kartlar"] if k["kod"] != "DR-02"]
    m = _model(_ie(), tmp_path, "eksik", kartlar=k2)
    assert len(m["say"]["kural_yok"]) == 6 and all("uyan kart yok" in x["neden"] and x["parca"].startswith("INIS_") for x in m["say"]["kural_yok"])
    assert "inis" not in m["say"]["pafta_anahtarlari"]
    k3 = copy.deepcopy(kartlar)
    k3["kartlar"].append(dict(copy.deepcopy(next(k for k in k3["kartlar"] if k["kod"] == "AS-01")), kod="AS-99"))
    with pytest.raises(ValueError, match="iki karta uyuyor"):
        _model(_ie(), tmp_path, "iki", kartlar=k3)


def test_makaslar_y_ekseninde(tmp_path):
    """Makaslar derinliği geçince (makas ekseni y) mahya x boyunca: aşık ve kaplama x yönünde, dere hatları y = −30 ve 1030'da. Açıklık
    1000 cm'de üçgen 184,9 cm: tek sıra betopan örtmez; K5b'den beri levhalar dikey (PK-07 s. 10), sayı H5 (K2'de "kural yok"tu)."""
    m = _model(_ie(cati__makas_ekseni="y"), tmp_path, "y")
    assert m["say"]["kural_yok"] == [] and ifc_uret.ozet(m["model"])["sema_hatasi"] == 0
    B = [p for p in m["say"]["yerlesim"]["parcalar"] if p["kart"] == "BT-01"]
    assert B and all(p["bilgi"]["Dizilim"] == "dikey" and p["bilgi"]["Levha_cm"] == "125 × 250" for p in B)
    assert len(B) == 2 * B[0]["bilgi"]["LevhaAdedi"] == _adet(m["say"])["BT-01"]
    A = _parcalar(m["model"], "AS-01")
    assert all(p["b"][0] - p["a"][0] > p["b"][1] - p["a"][1] for p in A)
    D = _parcalar(m["model"], "DR-01")
    assert sorted(round(p["b"][1] if p["b"][1] <= 0 else p["a"][1]) for p in D) == [-30, 1030]
    assert all(p["p"]["Boy_cm"] == pytest.approx(660) for p in D)
    assert _adet(m["say"])["DR-02"] == 4                                          # 660 cm < 9 m: orta iniş yok


def _g_tavan(odalar, duvarlar, acik=(), engeller=(), tavan_="asma"):
    return {"duvarlar": list(duvarlar), "odalar": [{"ad": f"O{i}", "poligon": p} for i, p in enumerate(odalar)],
            "acik": [{"ad": "V", "poligon": p} for p in acik], "engeller": list(engeller), "tavan": tavan_,
            "cati": {"makas_ekseni": "x"}, "baslangic": {"x": "min", "y": "max"}}


def test_asma_en_kisa_ve_l_oda():
    """Asma: ara hat odanın öteki sınırından 10 cm uzaktaysa konur, 9,9 cm'deyse konmaz (odaya ortasında bir hat); sınır iç duvarsa
    ölçü duvarın ekseninden. L biçimli odada hat,
    kolun kendi sınırına 10 cm'den yakın kısımda kesilir; ters tavanda ve açık mahalde de. Açık mahali profil bölerse hat en büyük parçaya,
    öteki parça uyarı."""
    K = {"aralik_asma": 125.5, "aralik_ters": 60.0, "en_kisa": 10.0}
    sol = Duvar("dis", "y", 5.0, 5.0, 0.0, 300.0)
    def ara(gen):
        sag = Duvar("dis", "y", 10 + gen + 5, 5.0, 0.0, 300.0)
        H = tavan.hatlar(_g_tavan([[[10, 0], [10 + gen, 0], [10 + gen, 300], [10, 300]]], [sol, sag]), K)
        return [h["s"] for h in H]
    assert ara(135.5) == [135.5] and ara(135.4) == [pytest.approx(10 + 135.4 / 2)]
    ic = Duvar("ic", "y", 145.5, 5.0, 0.0, 300.0)                                       # öteki sınır iç duvar: 10 cm eksenden (yüzden 5)
    H = tavan.hatlar(_g_tavan([[[10, 0], [140.5, 0], [140.5, 300], [10, 300]]], [sol, ic]), K)
    assert [(h["yer"], h["s"], h["r0"], h["r1"]) for h in H] == [("duvar_ekseni", 145.5, 0, 300), ("ara", 135.5, 0, 300)]
    L =[[10, 0], [300, 0], [300, 100], [140.5, 100], [140.5, 200], [10, 200]]           # kol 2'nin sağ yüzü 140,5: 135,5 hattına 5 cm
    duv = [sol, Duvar("dis", "y", 305.0, 5.0, 0.0, 100.0), Duvar("dis", "y", 145.5, 5.0, 100.0, 200.0)]
    H = tavan.hatlar(_g_tavan([L], duv), K)
    assert [(h["s"], h["r0"], h["r1"]) for h in H] == [(135.5, 0, 100), (261.0, 0, 100)]
    L = [[10, 0], [300, 0], [300, 100], [135, 100], [135, 200], [10, 200]]               # ters ve açık: 130 hattı kol 2'nin yüzüne 5 cm
    duv = [sol, Duvar("dis", "y", 305.0, 5.0, 0.0, 100.0), Duvar("dis", "y", 140.0, 5.0, 100.0, 200.0)]
    beklenen = [(10, 0, 200), (70, 0, 200), (130, 0, 100), (190, 0, 100), (250, 0, 100)]
    assert [(h["s"], h["r0"], h["r1"]) for h in tavan.hatlar(_g_tavan([L], duv, tavan_="ters"), K)] == beklenen
    assert [(h["s"], h["r0"], h["r1"]) for h in tavan.hatlar(_g_tavan([], [], acik=[L], tavan_="ters"), K)] == beklenen
    uyari = []
    H = tavan.hatlar(_g_tavan([], [], acik=[[[0, 0], [300, 0], [300, 200], [0, 200]]], engeller=[[140, 0, 150, 200]]), K, uyari)
    assert {h["s"] for h in H} and all(150 <= h["s"] <= 300 for h in H)
    assert len(uyari) == 1 and "2.80 m²" in uyari[0]["metin"]


def test_betopan_tek_sira_ortmezse_kural_yok(tmp_path):
    """Üçgen (alt başlıktan mahyada üst başlığın altına) levhanın kısa kenarını geçerse tek sıra örtmez. K2'de "kural yok"tu; K5b'den beri
    levhalar dikey (PK-07 s. 10). Uzun kenarı da geçerse "kural yok", komut 1. PK-01'de üçgen 118,9 cm (%33); %40'ta 142 cm (dikey),
    %80'de 274 cm (kural yok)."""
    F = cati.cerceve(Y.girdi.kur(_ie(), ifc_uret.hazirla(_ie()), []))
    assert cati._ust(F, F["um"]) - 10 - 250 == pytest.approx(118.9)
    m = _model(_ie(cati__egim_yuzde=40), tmp_path, "dik")
    assert not [x for x in m["say"]["kural_yok"] if x["parca"] == "baş makas betopanı"] and _adet(m["say"])["BT-01"] > 0
    m = _model(_ie(cati__egim_yuzde=80), tmp_path, "cok_dik")
    k = [x for x in m["say"]["kural_yok"] if x["parca"] == "baş makas betopanı"]
    assert len(k) == 1 and "dikey levha da örtmez" in k[0]["neden"] and _adet(m["say"])["BT-01"] == 0
    assert _adet(m["say"])["AS-01"] > 0 and _adet(m["say"])["KP-01"] > 0


def test_uyarilar(tmp_path):
    """Tavan yazılı değilse asma sayılır ve uyarı yazar; İş Emri'nde planda olmayan açık mahal uyarı yazar. PK-01'de uyarı yok."""
    ie = _ie()
    plan = ifc_uret.hazirla(ie)
    assert girdi.kur(ie, plan, [])["uyarilar"] == []
    del ie["ozellikler"]["Tavan"]
    ie["mahaller"].append({"ad": "VERANDA 2", "tur": "veranda", "konum_cm": [[0, 0, 1, 1]]})
    u = girdi.kur(ie, plan, [])["uyarilar"]
    assert len(u) == 2 and "Tavan yazılı değil" in u[0]["metin"] and "VERANDA 2" in u[1]["metin"]
    ie = _ie()
    del ie["ozellikler"]["Tavan"]
    m = _model(ie, tmp_path, "tavansiz")
    assert m["say"]["yerlesim"]["tavan"] == "asma" and [x["katman"] for x in m["say"]["yerlesim"]["uyarilar"]] == ["tavan omegası"]
    assert m["say"]["kural_yok"] == []


def test_kod_kilidi_k2():
    """K2'de eklenen yedi kod kilitte; kilidin ilk 32 satırı (K1 ve K2) donmuştur, yeni kod sona eklenir."""
    kilit = _oku(VERI / "kart_kodlari.json")["kodlar"]
    ilk = dict(list(kilit.items())[:32])
    assert list(ilk)[25:] == ["OM-01", "AS-01", "KP-01", "BT-01", "DR-01", "DR-02", "SR-01"]
    assert hashlib.sha256(json.dumps(ilk, ensure_ascii=False, sort_keys=True).encode()).hexdigest() == KILIT_K2


def test_kural_yokta_planlar_eksik_yazar(tmp_path):
    """Kural yok varsa her sayfanın başında EKSİK; kurulmayan katmanın sayfasında nedeni (kırma çatıda aşık, kaplama, betopan, dere)."""
    (tmp_path / "kirma.json").write_text(json.dumps(_ie(cati__tip="kirma"), ensure_ascii=False), encoding="utf-8")
    assert KM.main([str(tmp_path / "kirma.json"), str(tmp_path / "c"), "--kod", "K"]) == 1
    import pymupdf
    metin = [" ".join(p.get_text().split()) for p in pymupdf.open(tmp_path / "c" / "K_yerlesim_planlari.pdf")]
    assert all("EKSİK" in s for s in metin)
    for i, ad in ((1, "dere ve iniş"), (2, "baş makas betopanı"), (3, "aşık"), (4, "çatı kaplaması")):
        assert f"Kurulmadı: {ad}:" in metin[i], ad
    assert "Kurulmadı" not in metin[0]


# ----------------------------------------------------------------------------- kesim planı, gösterim, komut
def test_kesim_plani(pk01):
    k = pk01["say"]["kesim_plani"]
    assert (k["OM-01"]["stok"], k["OM-01"]["kullanilir_cm"], k["OM-01"]["parca"]) == (20, 365, 24)
    assert all(sum(d) <= 365 + 1e-6 for d in k["OM-01"]["dizilim"]) and sorted(b for d in k["OM-01"]["dizilim"] for b in d) == \
        sorted(p["boy_cm"] for p in pk01["say"]["yerlesim"]["parcalar"] if p["kart"] == "OM-01")
    assert (k["AS-01"]["stok"], k["AS-01"]["kullanilir_cm"]) == (42, 410)
    assert (k["DR-01"]["stok"], k["DR-01"]["hat_basina"]) == (6, [3, 3])
    assert pk01["say"]["sabit"] == {"SR-01": {"adet": 7, "okunan": "kurallar/liste/prefabrik.json › kurallar.pvc_betopan_h"}}
    assert next(x for x in pk01["say"]["kartlar"] if x["kod"] == "SR-01") == {"kod": "SR-01", "ad": "PVC plastik betopan H'ı", "adet": 0,
                                                                              "katalog_kodlari": {}, "sabit": 7}
    assert kesim.yerlestir([300, 200, 160, 100], 365) == [[300], [200, 160], [100]]
    with pytest.raises(ValueError):
        kesim.yerlestir([370], 365)


def test_gosterim_sayiyi_degistirmez(pk01, tmp_path):
    y = _oku(KM.YERLESIM)
    y["kurallar"]["gosterim"]["kesit_cm"]["deger"] = {"omega": [10, 8], "asik": [8, 12], "dere": [20, 15], "inis": 14, "kaplama": 10, "betopan": 3,
                                                      "tavan_h": [3, 4], "alcipan": 2, "izolasyon": 12}      # K5b-2'nin kesitleri
    m = _model(_ie(), tmp_path, "gosterim", yerlesim=y)
    a, b = pk01["say"], m["say"]
    assert [(k["kod"], k["adet"], k.get("metre")) for k in a["kartlar"]] == [(k["kod"], k["adet"], k.get("metre")) for k in b["kartlar"]]
    assert a["kesim_plani"] == b["kesim_plani"] and a["pafta_anahtarlari"] == b["pafta_anahtarlari"]


def test_komut_yerlesim_planlari(tmp_path):
    assert KM.main([str(MOTOR_VERI / "PK01.json"), str(tmp_path), "--kod", "PK01"]) == 0
    import pymupdf
    d = pymupdf.open(tmp_path / "PK01_yerlesim_planlari.pdf")
    metin = [p.get_text() for p in d]
    assert len(d) == 7                                                    # K3: altıncı sayfa elektrik planı (test_k3.py); K3b: yedinci (test_k3b.py)
    for i, baslik in enumerate(("TAVAN OMEGA YERLEŞİM PLANI", "MAKAS VE DERE YERLEŞİM PLANI", "BAŞ MAKAS BETOPAN YERLEŞİMİ",
                                "AŞIK YERLEŞİM PLANI", "ÇATI KAPLAMASI YERLEŞİM PLANI", "ELEKTRİK YERLEŞİM PLANI",
                                "ALT ÇERÇEVE VE ÇATI KENARLARI YERLEŞİM PLANI")):
        assert f"PK01 · {baslik}" in metin[i], baslik
    assert "24 parça, 58,46 m" in metin[0] and "42 parça, 148,40 m" in metin[3] and "22 panel" in metin[4] and "4 levha" in metin[2]
    assert d.metadata.get("creationDate", "") == ""
    assert not any("EKSİK" in s or "Kurulmadı" in s or "Uyarı" in s for s in metin)
    assert "son şerit yüzün ucundan taşar (40 cm)" in " ".join(metin[4].split())
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "PK01_kart_katalogu.xlsx")
    assert wb["Çatı ve tavan"].max_row == 100 + 167 + 1 and wb["Kesim planı"].max_row == 3 + 7 + 1 + 1 + 4      # K3b: alt çerçeve ve çatı kenarı 7 kart; K5b-2: OM-02, AL-01, AL-02, IZ-01
    k = KY.kiyasla(tmp_path / "PK01_model.ifc", _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE), adim="K2")
    assert k["kapi_1"] == "geçer"

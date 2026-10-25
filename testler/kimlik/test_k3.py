"""İE-15 · K3 · Elektrik noktaları ve kablo (kural `elektrik_yerlesimi`, `kurallar/yerlesim/prefabrik.json`; Lokman, 18 Ekim 19:52).

- Model = kural: PK-01'in K3 modelinde duy 5, glop 3, anahtar 6 + 1, priz 13 + 1, uydu 1, zil 1, sigorta kutusu 1, NYM 3 × 2,5 101 m ve
  2 × 1,5 85 m; liste kilidiyle (`testler/motor/veri/PK01_liste_kilidi.json`) aynı. Öteki sekiz İş Emri'nde (ALTIN, P6001 ve iki aynası,
  yan verandalı, TEST01–03) noktalar ve kablo liste üretecinin satırlarıyla aynı.
- Hakemin K3 kapısı geçer (K2'nin altı satırı ve metresi); dokuz elektrik satırı bilgidir, paftayla yan yana. K2 parçaları K3 modelinde aynen.
- Oda oda: PK-01'in odaları; birden çok takım (salon 2, yaşam alanı 3, veranda 2: glop 4, çiftli anahtar 2); açık mutfakta tezgahın 4 prizi.
- Yer (varsayım): duy ve glop odasının içinde, anahtar kapının kol tarafında, sigorta kutusu giriş kapısının üstünde iç yüzde, zil dışarıda,
  veranda globu evin duvarında kapının iki yanında; duvardaki nokta açıklığın önünde değil; takımlı odada tavan noktaları eşit şeritlerde;
  zil önce kol tarafında; anahtar sirkülasyon kapısında. Gösterim, ölçü ve kot sayıyı değiştirmez.
- Kural yok ve uyarılar: en uzak mesafe "sigortadan" değilse kablo kurulmaz; kapısız oda, salonsuz ev (uydu), mutfaksız ve salonsuz ev
  (tezgah), giriş kapısız ev, sigorta kutusu tavanda, zile dışarıda yer yok, uzun duvarlarda yer yok (sayı yine liste üretecininki).
  Elektrik İş Emri'siz kurulmaz.
- Komut: yerleşim planlarının altıncı sayfası elektrik planı; oda tablosunun (`k3_oda_tablosu.py`) toplamı paftanın ve modelin sayımı;
  EL kodlarının kilidi, GlobalId ev kimliğinden.
"""
from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import uuid
from collections import Counter
from pathlib import Path

import ifcopenshell
import ifcopenshell.guid
import ifcopenshell.util.element as ue
import pytest
from shapely.geometry import Point, Polygon, box

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
sys.path.insert(0, str(KOK / "testler" / "kimlik" / "hakem"))
sys.path.insert(0, str(KOK / "testler" / "kimlik"))
import k3_oda_tablosu as OT  # noqa: E402
import kiyas as KY  # noqa: E402
from motor import geometry as G  # noqa: E402
from motor import ifc_uret  # noqa: E402
from motor import kart_model as KM  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402
from motor import yerlesim as Y  # noqa: E402
from motor.yerlesim import elektrik as EL  # noqa: E402
from motor.yerlesim import yaz  # noqa: E402

VERI = KOK / "testler" / "kimlik" / "veri"
MOTOR_VERI = KOK / "testler" / "motor" / "veri"
PROJE = KOK / "testler" / "kimlik" / "hakem" / "projeler" / "PK01.json"
IS_EMIRLERI = {"PK01": MOTOR_VERI / "PK01.json", "ALTIN": KOK / "is_emri" / "ALTIN_2plus1_6x8.json",
               "P6001": KOK / "is_emri" / "P6001_2plus1_6x6.json", "P6001_ayna_x": KOK / "is_emri" / "P6001_ayna_x.json",
               "P6001_ayna_y": KOK / "is_emri" / "P6001_ayna_y.json", "YAN_VERANDALI": KOK / "is_emri" / "YAN_VERANDALI_2plus1_6x8p5.json",
               "TEST01": KOK / "is_emri" / "adaylar" / "TEST01.json", "TEST02": KOK / "is_emri" / "adaylar" / "TEST02.json",
               "TEST03": KOK / "is_emri" / "adaylar" / "TEST03.json"}
EL_KODLARI = tuple(f"EL-{i:02d}" for i in range(1, 12))
KILIT_K3 = "90823905369a8d85e904846b256063c7631c92016b482bc908511541a889f1e4"   # K1'in 25, K2'nin 7 ve K3'ün 11 kodu (bu PR)


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


EL_KART = {k["yerlesim"]["parca"]: k for k in _oku(KM.KARTLAR)["kartlar"] if k["kod"] in EL_KODLARI}


@pytest.fixture(scope="module")
def DK():
    y = _oku(KM.YERLESIM)
    return Y.degerler(y, _oku(KM.LISTE_KURALLARI))[Y.ELEKTRIK], y["kurallar"][Y.ELEKTRIK]


def _katman(ie: dict, DK, D: dict | None = None) -> dict:
    return EL.katman(ie, ifc_uret.hazirla(ie), D or DK[0], DK[1])


def _kuraldan(E: dict) -> dict:
    """Katmanın sonucu, liste üretecinin satır adlarıyla (kartın `liste.satirlar`'ı): nokta adedi, kabloda metre."""
    out = {}
    for parca, k in EL_KART.items():
        if k["birim"] == "m":
            out[k["liste"]["satirlar"][0]] = next(x["metre"] for x in E["kablo"] if x["parca"] == parca)
        else:
            out[k["liste"]["satirlar"][0]] = E["toplam"].get(parca, 0)
    return out


def _liste(ie: dict) -> dict:
    L = PL.liste(PL.model_motor(copy.deepcopy(ie)))
    satir = {s["kalem"]: s["adet"] for s in L["satirlar"] if s["grup"] == "1007"}
    return {k["liste"]["satirlar"][0]: satir[k["liste"]["satirlar"][0]] for k in EL_KART.values()}


def _turu_degistir(ie: dict, **turler) -> dict:
    ie = copy.deepcopy(ie)
    for m in ie["mahaller"]:
        if m["ad"] in turler:
            m["tur"] = turler[m["ad"]]
    return ie


@pytest.fixture(scope="module")
def komut(tmp_path_factory):
    """Kabul komutu (`python -m motor.kart_model … --kod PK01`): K3 modeli; aynı bugünkü IFC'den K2 modeli de (elektriksiz)."""
    d = tmp_path_factory.mktemp("k3")
    assert KM.main([str(MOTOR_VERI / "PK01.json"), str(d), "--kod", "PK01"]) == 0
    ie = _oku(MOTOR_VERI / "PK01.json")
    say_k2 = KM.kart_modeli(d / "PK01.ifc", d / "PK01_k2.ifc", "besik", is_emri=ie, plan=ifc_uret.hazirla(ie))
    return {"klasor": d, "model": d / "PK01_model.ifc", "k2": d / "PK01_k2.ifc", "say": _oku(d / "PK01_kart_sayim.json"), "say_k2": say_k2,
            "ie": ie}


# ----------------------------------------------------------------------------- kural, kartlar, kilit
def test_degerler_okunur(DK):
    """Kuralın sayıları liste kurallarından ve kitaptan okunur (burada sayı yazılmaz); yerleşim kuralıyla liste kurallarında ortak anahtar yok;
    yer, ölçü, kot ve gösterim varsayım durumuyla ve notuyla."""
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    kitap = Y.kitap_oku(liste)
    assert not set(y["kurallar"]) & set(liste["kurallar"]) and Y.ELEKTRIK not in Y.KURALLAR
    D, K = DK
    assert D["oda_kurali"] == liste["kurallar"]["elektrik_oda"]["deger"] and D["kitap_odalari"] == kitap["elektrik_oda"]["deger"]["odalar"]
    assert (D["oda_kurali"]["takim_m2"], D["oda_kurali"]["oda_takimi"], D["oda_kurali"]["acik_takimi"]) == (
        15, {"tekli_priz": 2, "tekli_anahtar": 1, "tijli_duy": 1}, {"glop": 2, "cift_anahtar": 1})
    assert D["acik_mutfak"] == {"oda": "salon", "ek": {"tekli_priz": 4}}
    assert D["kitap_odalari"]["mutfak"] == {"tijli_duy": 1, "tekli_anahtar": 1, "tekli_priz": 5, "cift_priz": 1}
    assert (D["zil"], D["sigorta_kutusu"], D["w_otomat"], D["kablo_nokta_basina_m"], D["en_uzak_mesafe"]) == (1, 1, 8, 2, "sigortadan")
    assert K["en_uzak_mesafe"]["beklenen"] == "sigortadan"
    durumlar = _oku(KM.KARTLAR)["durumlar"]
    assert K["durum"] == "lokman" and K["durum"] in durumlar and K["tarih"] and K["kaynak"]
    for alt in ("uydu_odasi", "yerler", "olculer_cm", "kot_cm", "gosterim_cm"):
        assert K[alt]["durum"] == "varsayim" and K[alt]["tarih"] and K[alt]["not"], alt
    for alt, v in K.items():
        if isinstance(v, dict) and "okunan" in v:
            assert v["okunan"].startswith(("kurallar/liste/prefabrik.json › ", "kitap › ")), alt


def test_el_kartlari():
    """On bir EL kartı: nokta kartının paftadaki anahtarı `elektrik.<parça>` (kablonunki yok), sınıfı kuralın sınıfının türü, listedeki adı
    liste kilidinin elektrik satırlarından."""
    kilit = {s[1] for s in _oku(MOTOR_VERI / "PK01_liste_kilidi.json")["satirlar"] if s[0] == "1007"}
    assert [k["kod"] for k in EL_KART.values()] == list(EL_KODLARI)
    for parca, k in EL_KART.items():
        assert k["yerlesim"]["kural"] == Y.ELEKTRIK and parca in Y.PARCA_ADI and k["liste"]["grup"] == "1007"
        assert k["ifc"]["sinif"] == k["yerlesim"]["sinif"] + "Type" and k["liste"]["satirlar"][0] in kilit
        if parca.startswith("kablo_"):
            assert (k["pafta"], k["birim"], k["sayim"], k["yerlesim"]["sinif"]) == ([], "m", "turev", "IfcCableSegment")
        else:
            assert (k["pafta"], k["birim"], k["sayim"]) == ([f"elektrik.{parca}"], "adet", "yerlesen_adet")


def test_kod_kilidi_k3():
    """K3'te eklenen on bir kod kilitte; kilidin ilk 43 satırı (K1, K2 ve K3) donmuştur, yeni kod sona eklenir."""
    kilit = _oku(VERI / "kart_kodlari.json")["kodlar"]
    ilk = dict(list(kilit.items())[:43])
    assert list(ilk)[32:] == list(EL_KODLARI)
    assert hashlib.sha256(json.dumps(ilk, ensure_ascii=False, sort_keys=True).encode()).hexdigest() == KILIT_K3


# ----------------------------------------------------------------------------- model = kural
def test_model_kural_pk01(komut):
    """K3'ün kendi testi: PK-01'in K3 modelindeki noktalar ve kablo liste kilidinin elektrik satırlarıyla aynı."""
    kilit = {s[1]: s[3] for s in _oku(MOTOR_VERI / "PK01_liste_kilidi.json")["satirlar"] if s[0] == "1007"}
    kart = {k["kod"]: k for k in komut["say"]["kartlar"]}
    bulunan = [kart[k["kod"]]["metre"] if k["birim"] == "m" else kart[k["kod"]]["adet"] for k in EL_KART.values()]
    assert bulunan == [kilit[k["liste"]["satirlar"][0]] for k in EL_KART.values()] == [5, 3, 6, 1, 13, 1, 1, 1, 1, 101, 85]
    E = komut["say"]["elektrik"]
    assert E["toplam"] == {"anahtar_cift": 1, "anahtar_tekli": 6, "glop": 3, "priz_cift": 1, "priz_tekli": 13, "sigorta_kutusu": 1,
                           "tavan_armaturu": 5, "uydu_prizi": 1, "zil": 1}
    assert E["uyarilar"] == [] and komut["say"]["kural_yok"] == [] and komut["say"]["sema_hatasi"] == {"bugunku_ifc": 0, "kart_modeli": 0}


@pytest.mark.parametrize("ad", list(IS_EMIRLERI))
def test_liste_uretecine_esit(ad, DK):
    """Dokuz İş Emri'nde noktalar ve iki hat kablo liste üretecinin satırlarıyla aynı (kural aynı değerlerle okunur)."""
    ie = _oku(IS_EMIRLERI[ad])
    E = _katman(ie, DK)
    assert _kuraldan(E) == _liste(ie), ad
    assert E["kural_yok"] == []
    assert [u["metin"] for u in E["uyarilar"]] == (["MUTFAK: kapısı yok; 1 anahtar uzun duvara konuldu"] if ad == "ALTIN" else []), ad


def test_oda_sayilari_pk01(DK):
    E = _katman(_oku(MOTOR_VERI / "PK01.json"), DK)
    assert {o["ad"]: (o["acik"], o["takim"], o["kaynak"], o["sayilar"]) for o in E["odalar"]} == {
        "YATAK ODASI 1": (False, 1, "oda takımı", {"priz_tekli": 2, "anahtar_tekli": 1, "tavan_armaturu": 1}),
        "BANYO": (False, None, "kitap (banyo)", {"glop": 1, "anahtar_tekli": 1, "priz_tekli": 2}),
        "YATAK ODASI 2": (False, 1, "oda takımı", {"priz_tekli": 2, "anahtar_tekli": 1, "tavan_armaturu": 1}),
        "HOL": (False, None, "kitap (hol)", {"tavan_armaturu": 1, "anahtar_tekli": 1}),
        "MUTFAK": (False, None, "kitap (mutfak)", {"tavan_armaturu": 1, "anahtar_tekli": 1, "priz_tekli": 5, "priz_cift": 1}),
        "SALON": (False, 1, "oda takımı + uydu prizi (ev başına)", {"priz_tekli": 2, "anahtar_tekli": 1, "tavan_armaturu": 1, "uydu_prizi": 1}),
        "VERANDA": (True, 1, "veranda takımı", {"glop": 2, "anahtar_cift": 1})}
    alan = {o["ad"]: o["alan_m2"] for o in E["odalar"]}
    assert all(alan[a] <= 15 for a in ("YATAK ODASI 1", "YATAK ODASI 2", "SALON", "VERANDA"))


@pytest.mark.parametrize("ad,oda,takim,kaynak,sayilar", [
    ("ALTIN", "SALON", 2, "oda takımı + uydu prizi (ev başına)", {"priz_tekli": 4, "anahtar_tekli": 2, "tavan_armaturu": 2, "uydu_prizi": 1}),
    ("YAN_VERANDALI", "SALON", 2, "oda takımı + uydu prizi (ev başına)",
     {"priz_tekli": 4, "anahtar_tekli": 2, "tavan_armaturu": 2, "uydu_prizi": 1}),
    ("P6001", "YASAM ALANI", 2, "oda takımı + tezgah (açık mutfak) + uydu prizi (ev başına)",
     {"priz_tekli": 8, "anahtar_tekli": 2, "tavan_armaturu": 2, "uydu_prizi": 1}),
    ("TEST01", "YASAM ALANI", 3, "oda takımı + tezgah (açık mutfak) + uydu prizi (ev başına)",
     {"priz_tekli": 10, "anahtar_tekli": 3, "tavan_armaturu": 3, "uydu_prizi": 1}),
    ("TEST02", "VERANDA", 2, "veranda takımı", {"glop": 4, "anahtar_cift": 2})])
def test_takim_ve_acik_mutfak(ad, oda, takim, kaynak, sayilar, DK):
    """Takım sayısı alan ÷ 15, yukarı; açık mutfağın odası takımına ek tezgahın 4 tekli prizini alır. Odanın her noktası modelde, odasında."""
    E = _katman(_oku(IS_EMIRLERI[ad]), DK)
    o = next(o for o in E["odalar"] if o["ad"] == oda)
    assert (o["takim"], o["kaynak"], o["sayilar"]) == (takim, kaynak, sayilar)
    assert 15 * (takim - 1) < o["alan_m2"] <= 15 * takim
    assert Counter(p["parca"] for p in E["parcalar"] if p["bilgi"].get("Mahal") == oda and p["parca"] not in ("zil", "sigorta_kutusu")) == sayilar


# ----------------------------------------------------------------------------- hakemin kapısı, K2'nin parçaları, IFC
def test_hakem_k3_kapisi_gecer(komut):
    """Kabul komutu: `kiyas.py <K3 modeli> <cevap> --proje PK01.json --adim K3` geçer; kapıda K2'nin satırları ve metresi, 243 parça
    paftadaki anahtarı taşır (K2'nin 211'i ve 32 nokta). Elektrik satırları bilgi: pafta ile kural yan yana."""
    k = KY.kiyasla(komut["model"], _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE), adim="K3")
    assert k["kapi_1"] == "geçer" and k["kart_modeli"] == 211 + 32
    satir = {s["anahtar"]: s for s in k["satirlar"]}
    assert all(satir[a]["sonuc"] == "geçer" and satir[a]["kapida"]
               for a in ("duvar_omegasi_asma", "asik", "sandvic_panel", "bas_makas_betopan", "dere", "inis"))
    assert {a[len("elektrik."):]: (s["paftadan"], s["modelden"], s["kapida"]) for a, s in satir.items() if a.startswith("elektrik.")} == {
        "tavan_armaturu": (6, 5, False), "glop": (2, 3, False), "anahtar_tekli": (4, 6, False), "anahtar_cift": (2, 1, False),
        "priz_tekli": (15, 13, False), "priz_cift": (4, 1, False), "uydu_prizi": (1, 1, False), "zil": (0, 1, False),
        "sigorta_kutusu": (1, 1, False)}
    assert {m["anahtar"]: m["sonuc"] for m in k["metre"]} == {"duvar_omegasi_asma": "geçer", "asik": "geçer", "dere": "geçer"}
    r = subprocess.run([sys.executable, "-I", str(KOK / "testler" / "kimlik" / "hakem" / "kiyas.py"), str(komut["model"]),
                        str(VERI / "PK01_pafta_katmani.json"), "--proje", str(PROJE), "--adim", "K3"], capture_output=True, text=True, cwd=KOK)
    assert r.returncode == 0 and "Kapı 1 (adet ve metre, PK-01, K3): **geçer**" in r.stdout, r.stdout + r.stderr


def test_k2_parcalari_aynen(komut):
    """K3 yalnız ekler: K2 modelinin her parçası K3 modelinde aynı sınıf, ad ve setlerle; yeni 34 parça EL kartlarından (32 nokta, 2 kablo).
    K2'nin sayımı, kesim planı ve K2 kapısı aynı."""
    a, b = ifcopenshell.open(str(komut["k2"])), ifcopenshell.open(str(komut["model"]))
    A = {e.GlobalId: e for e in a.by_type("IfcProduct")}
    B = {e.GlobalId: e for e in b.by_type("IfcProduct")}
    assert set(A) <= set(B)
    for g, e in A.items():
        m = B[g]
        assert (m.is_a(), m.Name) == (e.is_a(), e.Name), e.Name
        assert {k: {x: y for x, y in v.items() if x != "id"} for k, v in ue.get_psets(e).items()} == \
               {k: {x: y for x, y in v.items() if x != "id"} for k, v in ue.get_psets(m).items()}, e.Name
    yeni = [B[g] for g in set(B) - set(A)]
    assert Counter(ue.get_psets(e)[KM.PARCA_SETI]["KartKodu"] for e in yeni) == {
        "EL-01": 5, "EL-02": 3, "EL-03": 6, "EL-04": 1, "EL-05": 13, "EL-06": 1, "EL-07": 1, "EL-08": 1, "EL-09": 1, "EL-10": 1, "EL-11": 1}
    x, y = komut["say_k2"], komut["say"]
    assert x["yerlesim"] == y["yerlesim"] and x["kesim_plani"] == y["kesim_plani"] and x["sabit"] == y["sabit"]
    assert [k for k in x["kartlar"] if k["kod"] not in EL_KODLARI] == [k for k in y["kartlar"] if k["kod"] not in EL_KODLARI]
    assert all(k["adet"] == 0 for k in x["kartlar"] if k["kod"] in EL_KODLARI) and "elektrik" not in x
    assert KY.kiyasla(komut["model"], _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE), adim="K2")["kapi_1"] == "geçer"


def test_nokta_ve_kablo_ozellikleri(komut):
    """Nokta: kartının türü, paftadaki anahtarı, odası ve yeri; GlobalId ev kimliğinden. Kablo: anahtarı yok (kapı saymaz), boyu formülden
    (`Qto_CableSegmentBaseQuantities.Length`), en uzak mesafe liste üretecininki (8,5839 m)."""
    f = ifcopenshell.open(str(komut["model"]))
    odalar = {o["ad"] for o in komut["say"]["elektrik"]["odalar"]}
    en_uzak = round(PL.model_motor(copy.deepcopy(komut["ie"]))["sigorta_en_uzak"]["m"], 4)
    adlar = {}
    kablo = {}
    for e in f.by_type("IfcProduct"):
        p = ue.get_psets(e).get(KM.PARCA_SETI, {})
        if p.get("KartKodu") not in EL_KODLARI:
            continue
        k = next(k for k in EL_KART.values() if k["kod"] == p["KartKodu"])
        parca = k["yerlesim"]["parca"]
        adlar.setdefault(Y.PARCA_ADI[parca], [])
        assert e.is_a() == k["yerlesim"]["sinif"] and ue.get_type(e).is_a() == k["ifc"]["sinif"] and p["Kural"] == Y.ELEKTRIK
        assert e.GlobalId == ifcopenshell.guid.compress(uuid.uuid5(yaz.AD_ALANI, f"PM-PK01-60M2/{e.Name}").hex)
        adlar[Y.PARCA_ADI[parca]].append(e.Name)
        if parca.startswith("kablo_"):
            q = ue.get_psets(e, qtos_only=True)["Qto_CableSegmentBaseQuantities"]
            kablo[parca] = (p["PaftaAnahtari"], q["Length"], p["Boy_cm"], p["Noktalar"], p["EnUzak_m"])
        else:
            assert p["PaftaAnahtari"] == f"elektrik.{parca}" and p["Mahal"] in odalar and p["Yer"], e.Name
    assert kablo == {"kablo_priz": ("", pytest.approx(101), 10100, 16, en_uzak), "kablo_aydinlatma": ("", pytest.approx(85), 8500, 8, en_uzak)}
    assert en_uzak == 8.5839
    assert {a: sorted(v) for a, v in adlar.items()} == {a: [f"{a}_{i:03d}" for i in range(1, len(v) + 1)] for a, v in adlar.items()}
    assert {a: len(v) for a, v in adlar.items()} == {"DUY": 5, "GLOP": 3, "ANAHTAR": 6, "ANAHTAR_CIFT": 1, "PRIZ": 13, "PRIZ_CIFT": 1,
                                                    "UYDU": 1, "ZIL": 1, "SIGORTA": 1, "KABLO_PRIZ": 1, "KABLO_AYDINLATMA": 1}


def test_ifc_geometrisi_katmanla_ayni(komut, DK):
    """Modeldeki her noktanın kutusu katmanın kutusudur (cm); kablo tavanın 5 cm altında, sigortadan en uzak köşeye."""
    E = _katman(komut["ie"], DK)
    sira = Counter()
    beklenen = {}
    for p in E["parcalar"]:
        sira[p["parca"]] += 1
        beklenen[f"{Y.PARCA_ADI[p['parca']]}_{sira[p['parca']]:03d}"] = p["geo"]
    f = ifcopenshell.open(str(komut["model"]))
    kutu = KM._Kutular()
    for e in f.by_type("IfcProduct"):
        if ue.get_psets(e).get(KM.PARCA_SETI, {}).get("KartKodu") not in EL_KODLARI:
            continue
        a, b = kutu(e)
        geo = beklenen.pop(e.Name)
        if geo["tip"] == "kutu":
            assert a + b == pytest.approx(geo["kutu"][:3] + geo["kutu"][3:], abs=0.01), e.Name
        else:
            assert (a[2], b[2]) == (pytest.approx(244), pytest.approx(246))
            uc = geo["p1"]
            assert min(abs(a[0] - uc[0]), abs(b[0] - uc[0])) < 1.5 and min(abs(a[1] - uc[1]), abs(b[1] - uc[1])) < 1.5, e.Name
    assert beklenen == {}


# ----------------------------------------------------------------------------- yer
def _kutu2(p: dict):
    x0, y0, _, x1, y1, _ = p["geo"]["kutu"]
    return box(x0, y0, x1, y1)


@pytest.mark.parametrize("ad", list(IS_EMIRLERI))
def test_noktalar_odasinda_ve_acikligin_onunde_degil(ad, DK):
    """Her noktanın kutusu (planda) odasının içinde; duvardaki nokta kapı ve pencere boşluğunun önünde değil (sigorta kutusu kapının üstünde,
    hariç). Zil ve veranda anahtarı açık mahalin içinde, dışarıda."""
    ie = _oku(IS_EMIRLERI[ad])
    E = _katman(ie, DK)
    poly = {o["ad"]: Polygon(o["poligon"]) for o in E["odalar"] if o["poligon"]}
    P = G.kur(copy.deepcopy(ie))
    bosluk = [box(*p.gap()).buffer(1) for d in P.duvarlar for p in d.panolar if p.acilim]
    for p in E["parcalar"]:
        if p["geo"]["tip"] != "kutu":
            continue
        k = _kutu2(p)
        assert k.intersection(poly[p["bilgi"]["Mahal"]]).area >= 0.99 * k.area, (ad, p["aciklama"])
        if p["bilgi"]["Yer"] != "tavan" and p["parca"] != "sigorta_kutusu":
            assert not any(k.intersects(g) for g in bosluk), (ad, p["aciklama"])


def test_yerler_pk01(DK):
    """PK-01 (kapılar: dış kapı 278,75–368,75 x duvarında, menteşe 278,75'te, kol 368,75'te; hol–banyo 290–372, menteşe 372'de; hol–yatak 1 ve hol–yatak 2
    630–712, menteşe 630'da; hol–mutfak ve hol–salon 525–607, menteşe 607'de)."""
    E = _katman(_oku(MOTOR_VERI / "PK01.json"), DK)
    yer = {}
    for p in E["parcalar"]:
        if p["geo"]["tip"] == "kutu":
            yer.setdefault((p["parca"], p["bilgi"]["Mahal"]), []).append(p["geo"]["kutu"])

    def orta(k):
        return round((k[0] + k[3]) / 2, 2), round((k[1] + k[4]) / 2, 2), round((k[2] + k[5]) / 2, 2)
    # anahtar kol tarafında, kapı boşluğundan 10 cm, odanın yüzünde, 110 cm'de
    assert orta(yer["anahtar_tekli", "HOL"][0]) == (378.75, 211.0, 110.0)           # dış kapının iç yüzü (y 210)
    assert orta(yer["anahtar_cift", "VERANDA"][0]) == (378.75, 199.0, 110.0)        # dış kapının dış yüzü (y 200)
    assert orta(yer["anahtar_tekli", "BANYO"][0]) == (280.0, 748.0, 110.0)
    assert orta(yer["anahtar_tekli", "YATAK ODASI 1"][0])[1] == orta(yer["anahtar_tekli", "YATAK ODASI 2"][0])[1] == 722.0
    assert orta(yer["anahtar_tekli", "MUTFAK"][0])[1] == orta(yer["anahtar_tekli", "SALON"][0])[1] == 515.0
    # sigorta kutusu holde, giriş kapısının üstünde (kapının üstü 200 + 5), ortasında; zil dışarıda kol tarafında, 150 cm'de
    s = yer["sigorta_kutusu", "HOL"][0]
    assert orta(s)[:2] == (323.75, 215.0) and s[2] == 205.0 and s[5] <= 250
    assert orta(yer["zil", "VERANDA"][0]) == (398.75, 199.0, 150.0)
    # duy tavanda, odanın ortasında (L biçimli yatak odası 1 ağırlık merkezinde)
    assert orta(yer["tavan_armaturu", "YATAK ODASI 2"][0])[:2] == (489.75, 805.75) and yer["tavan_armaturu", "HOL"][0][5] == 250
    assert orta(yer["tavan_armaturu", "YATAK ODASI 1"][0])[:2] == (114.9, 793.43)
    # veranda globu evin duvarında, giriş kapısının iki yanında
    g = sorted(orta(k) for k in yer["glop", "VERANDA"])
    assert g[0][0] < 278.75 and g[1][0] > 368.75 and all(z == 210 for *_, z in g)
    # prizler odanın iki uzun duvarında, sırayla
    assert sorted(round(k[0]) for k in yer["priz_tekli", "YATAK ODASI 1"]) == [10, 193]


def test_takimli_odada_tavan_noktalari_seritlerinde(DK):
    """Birden çok tavan noktası (takımlı oda) odanın uzun ekseninde eşit şeritlere birer birer; hepsi odanın içinde (bir yere yığılmaz)."""
    coklu = {}
    for ad, yol in IS_EMIRLERI.items():
        E = _katman(_oku(yol), DK)
        poly = {o["ad"]: Polygon(o["poligon"]) for o in E["odalar"] if o["poligon"]}
        tavan = {}
        for p in E["parcalar"]:
            if p["bilgi"].get("Yer") == "tavan":
                k = p["geo"]["kutu"]
                tavan.setdefault(p["bilgi"]["Mahal"], []).append(((k[0] + k[3]) / 2, (k[1] + k[4]) / 2))
        for mahal, xy in tavan.items():
            x0, y0, x1, y1 = poly[mahal].bounds
            uzun_x = (x1 - x0) >= (y1 - y0)
            a0, L = (x0, x1 - x0) if uzun_x else (y0, y1 - y0)
            serit = sorted(int(((x if uzun_x else y) - a0) // (L / len(xy))) for x, y in xy)
            assert serit == list(range(len(xy))), (ad, mahal, xy)
            assert all(poly[mahal].contains(Point(q)) for q in xy), (ad, mahal, xy)
            if len(xy) > 1:
                coklu[ad, mahal] = len(xy)
    assert coklu == {("ALTIN", "SALON"): 2, ("P6001", "YASAM ALANI"): 2, ("P6001_ayna_x", "YASAM ALANI"): 2, ("P6001_ayna_y", "YASAM ALANI"): 2,
                     ("YAN_VERANDALI", "SALON"): 2, ("TEST01", "YASAM ALANI"): 3, ("TEST02", "SALON"): 2, ("TEST03", "YASAM ALANI"): 3}


def test_zil_once_kol_tarafinda(DK):
    """Giriş kapısının iki yanı da dışarıdaysa zil kol tarafında. Dokuz evin hiçbirinde iki yan birden boş değil; PK-01'de veranda ön cephe
    boyunca uzatılınca (mutfak verandanın arkasında kalır) menteşe tarafı da verandanın içine düşer."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    for m in ie["mahaller"]:
        if m["ad"] == "MUTFAK":
            m["konum_cm"] = [[10, 381.5, 261, 795]]
        if m["ad"] == "VERANDA":
            m["konum_cm"] = [[10, 795, 590, 995]]
    E = _katman(ie, DK)
    z = [p for p in E["parcalar"] if p["parca"] == "zil"]
    assert [(p["bilgi"]["Yer"], (p["geo"]["kutu"][0] + p["geo"]["kutu"][3]) / 2) for p in z] == [("giriş kapısının dışı, kol tarafı", 398.75)]
    assert E["uyarilar"] == []


def test_anahtar_sirkulasyon_kapisinda(DK):
    """Odanın birden çok kapısı varsa anahtar komşusunun en çok kapısı olan kapıda (sirkülasyon), planın sırasında önce gelen kapıda değil:
    PK-01'e yatak odası 1 ile mutfak arasında (planın ilk açıklığı) bir kapı eklenince ikisinin anahtarı yine hol kapısında."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["acikliklar"].insert(0, {"tur": "ic_kapi", "olcu": 82, "arasi": ["YATAK ODASI 1", "MUTFAK"]})
    E = _katman(ie, DK)
    a = {p["bilgi"]["Mahal"]: p["aciklama"] for p in E["parcalar"] if p["parca"] == "anahtar_tekli"}
    assert a["YATAK ODASI 1"] == "YATAK ODASI 1, kapının kol tarafı, y boyunca 722"
    assert a["MUTFAK"] == "MUTFAK, kapının kol tarafı, y boyunca 515"
    assert E["uyarilar"] == []


def test_gosterim_olcu_kot_sayiyi_degistirmez(DK):
    D = copy.deepcopy(DK[0])
    D["gosterim_cm"] = {k: [2 * x for x in v] if isinstance(v, list) else 2 * v for k, v in D["gosterim_cm"].items()}
    D["olculer_cm"] = {k: v + 5 for k, v in D["olculer_cm"].items()}
    D["kot_cm"] = {k: v + 20 for k, v in D["kot_cm"].items()}
    for ad in ("PK01", "TEST02"):
        ie = _oku(IS_EMIRLERI[ad])
        a, b = _katman(ie, DK), _katman(ie, DK, D)
        assert (a["toplam"], [k["metre"] for k in a["kablo"]]) == (b["toplam"], [k["metre"] for k in b["kablo"]]), ad
        assert [p["geo"] for p in a["parcalar"]] != [p["geo"] for p in b["parcalar"]]


# ----------------------------------------------------------------------------- kural yok, uyarılar
def test_en_uzak_sigortadan_degilse_kablo_kural_yok(DK):
    D = dict(DK[0], en_uzak_mesafe="en uzak oda")
    ie = _oku(MOTOR_VERI / "PK01.json")
    E = _katman(ie, DK, D)
    assert [(x["parca"], x["sinif"]) for x in E["kural_yok"]] == [("kablo_priz", "IfcCableSegment"), ("kablo_aydinlatma", "IfcCableSegment")]
    assert "kablo yalnız sigortadan en uzak mesafeyle kurulur" in E["kural_yok"][0]["neden"]
    assert E["kablo"] == [] and not any(p["parca"].startswith("kablo_") for p in E["parcalar"])
    assert E["toplam"] == _katman(ie, DK)["toplam"]


def test_kapisiz_oda_anahtari_uzun_duvarda(DK):
    E = _katman(_oku(IS_EMIRLERI["ALTIN"]), DK)
    a = [p for p in E["parcalar"] if p["parca"] == "anahtar_tekli" and p["bilgi"]["Mahal"] == "MUTFAK"]
    assert [p["bilgi"]["Yer"] for p in a] == ["uzun duvar (yedek)"]


def test_salonsuz_ve_mutfaksiz_ev(DK):
    """Salon türünde oda yoksa uydu prizi ilk kapalı odaya (uyarı); mutfak ve açık mutfağın odası da yoksa tezgahın prizleri konmaz (uyarı;
    liste üreteci de yazmaz). Sayılar yine liste üretecininki."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    salonsuz = _turu_degistir(ie, SALON="yatak_odasi")
    E = _katman(salonsuz, DK)
    assert [u["metin"] for u in E["uyarilar"]] == ["salon türünde oda yok; uydu prizi YATAK ODASI 1 odasına konuldu"]
    assert next(o for o in E["odalar"] if o["ad"] == "YATAK ODASI 1")["sayilar"]["uydu_prizi"] == 1
    assert _kuraldan(E) == _liste(salonsuz)
    ikisi = _turu_degistir(ie, SALON="yatak_odasi", MUTFAK="yatak_odasi")
    E = _katman(ikisi, DK)
    assert [u["metin"] for u in E["uyarilar"]] == [
        "mutfak ayrı mahal değil, açık mutfağın odası (salon) da yok; tezgahın prizleri konmadı (liste üreteci de yazmaz)",
        "salon türünde oda yok; uydu prizi YATAK ODASI 1 odasına konuldu"]
    assert _kuraldan(E) == _liste(ikisi)


def test_giris_kapisiz_ev(DK):
    """Giriş kapısı yoksa sigorta kutusu ve zil ilk kapalı odanın uzun duvarına (yedek yer, uyarı); kablonun en uzak mesafesi kapalı kısmın
    köşegeni (liste üretecindeki gibi); sayılar ve metreler liste üretecininki."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["acikliklar"] = [a for a in ie["acikliklar"] if a["tur"] != "dis_kapi"]
    E = _katman(ie, DK)
    assert [u["metin"] for u in E["uyarilar"]] == [
        "VERANDA: kapısı yok; 1 anahtar uzun duvara konuldu",
        "giriş kapısı yok; sigorta kutusu ve zil YATAK ODASI 1 odasının uzun duvarına konuldu"]
    assert {(p["parca"], p["bilgi"]["Mahal"], p["bilgi"]["Yer"]) for p in E["parcalar"] if p["parca"] in ("sigorta_kutusu", "zil")} == {
        ("sigorta_kutusu", "YATAK ODASI 1", "yedek yer"), ("zil", "YATAK ODASI 1", "yedek yer")}
    assert {k["nereden"] for k in E["kablo"]} == {"giriş kapısı yok: kapalı kısmın köşegeni"}
    assert _kuraldan(E) == _liste(ie)


def test_sigorta_kutusu_tavana_dayanir(DK):
    ie = _oku(MOTOR_VERI / "PK01.json")
    plan = ifc_uret.hazirla(ie)
    plan["bina"]["duvar_h"] = 230
    E = EL.katman(ie, plan, *DK)
    assert [u["metin"] for u in E["uyarilar"]] == ["sigorta kutusu kapının üstüne sığmadı; tavana dayandı"]
    assert [p["geo"]["kutu"][5] for p in E["parcalar"] if p["parca"] == "sigorta_kutusu"] == [230.0]


def test_zile_disarida_yer_yok(DK, monkeypatch):
    """Giriş kapısının dış yüzü baştan başa açıklık sayılırsa zil kol tarafında kalır, uyarıyla (sayı yine kural)."""
    asil = EL._bosluklar
    monkeypatch.setattr(EL, "_bosluklar", lambda P, yon, yuz: [(-1e6, 1e6)] if (yon, round(yuz, 2)) == ("x", 200.0) else asil(P, yon, yuz))
    E = _katman(_oku(MOTOR_VERI / "PK01.json"), DK)
    assert "giriş kapısının iki yanında da dışarıda yer yok; zil kol tarafına konuldu" in [u["metin"] for u in E["uyarilar"]]
    z = [p for p in E["parcalar"] if p["parca"] == "zil"]
    assert [(p["bilgi"]["Yer"], (p["geo"]["kutu"][0] + p["geo"]["kutu"][3]) / 2) for p in z] == [("giriş kapısının dışı, kol tarafı", 398.75)]
    assert E["toplam"]["zil"] == 1


def test_duvar_yerleri_notlari(monkeypatch):
    """Bir uzun duvarda açıklık dışında yer yoksa noktaları öbür duvara; hiçbirinde yoksa her duvar kendi noktalarını açıklık gözetmeden
    alır ve not bunu söyler (öbür duvara geçti demez)."""
    duvarlar = [("x", 0.0, 0.0, 300.0, 1), ("x", 400.0, 0.0, 300.0, -1)]
    O = {"bosluk_payi": 10}
    monkeypatch.setattr(EL, "_bosluklar", lambda P, yon, yuz: [])
    y, n = EL._duvar_yerleri(duvarlar, 3, None, O, 4)
    assert ([(q[1], q[3]) for q in y], n) == ([(0.0, 77.0), (400.0, 150.0), (0.0, 223.0)], "")
    monkeypatch.setattr(EL, "_bosluklar", lambda P, yon, yuz: [(-1e6, 1e6)] if yuz == 0.0 else [])
    y, n = EL._duvar_yerleri(duvarlar, 3, None, O, 4)
    assert [q[1] for q in y] == [400.0] * 3 and n == "bir uzun duvarda açıklık dışında yer yok; noktaları öbür duvara konuldu"
    monkeypatch.setattr(EL, "_bosluklar", lambda P, yon, yuz: [(-1e6, 1e6)])
    y, n = EL._duvar_yerleri(duvarlar, 3, None, O, 4)
    assert ([(q[1], q[3]) for q in y], n) == ([(0.0, 75.0), (400.0, 150.0), (0.0, 225.0)],
                                              "uzun duvarlarda açıklık dışında yer yok; noktalar açıklık gözetilmeden konuldu")


def test_elektrik_is_emri_ister(komut, tmp_path):
    with pytest.raises(ValueError, match="İş Emri ister"):
        KM.kart_modeli(komut["klasor"] / "PK01.ifc", tmp_path / "x.ifc", "besik", elektrik=True)


# ----------------------------------------------------------------------------- komut: elektrik planı, oda tablosu
def test_komut_elektrik_plani(komut):
    import pymupdf
    d = pymupdf.open(komut["klasor"] / "PK01_yerlesim_planlari.pdf")
    assert len(d) == 6
    metin = " ".join(d[5].get_text().split())
    assert "PK01 · ELEKTRİK YERLEŞİM PLANI" in metin and "(İE-15 K3)" in metin
    for s in ("Tijli duy: 5", "Yuvarlak glop: 3", "Tekli anahtar: 6", "Çiftli anahtar: 1", "Tekli priz: 13", "Çiftli priz: 1", "Uydu prizi: 1",
              "Zil: 1", "Sigorta kutusu (12'li): 1", "NYM kablo 3 × 2,5: 101 m", "NYM kablo 2 × 1,5: 85 m",
              "priz hattı: 8 W otomat × 8,58 m + 16 nokta × 2 m, yukarı", "aydınlatma hattı: 8 W otomat × 8,58 m + 8 nokta × 2 m, yukarı",
              "MUTFAK 14,98 m² · kitap (mutfak)", "VERANDA 6,68 m² · 1 takım"):
        assert s in metin, s
    assert not any(x in metin for x in ("EKSİK", "Kurulmadı", "Uyarı"))
    from openpyxl import load_workbook
    ws = load_workbook(komut["klasor"] / "PK01_kart_katalogu.xlsx")["Sayım"]
    satir = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert (satir["EL-10"][2], satir["EL-10"][3], satir["EL-05"][2]) == (1, 101, 13)


def test_oda_tablosu(komut):
    """Kural ile pafta oda oda (bilgi): paftanın her noktası bir odaya düşer; toplamlar paftanın ve modelin sayımı."""
    cevap = _oku(VERI / "PK01_pafta_katmani.json")
    t = OT.tablo(komut["say"], cevap)
    assert OT.ODA_DISI not in t["odalar"] and t["odalar"] == [o["ad"] for o in komut["say"]["elektrik"]["odalar"]]
    pafta, model = Counter(), Counter()
    for ad in t["odalar"]:
        pafta.update(t["pafta"][ad])
        model.update(t["model"][ad])
    assert dict(pafta) == {k: v for k, v in (cevap["ayrinti"]["elektrik"]["sayim"] | {"buat": cevap["ayrinti"]["elektrik"]["buat"]}).items() if v}
    assert dict(model) == komut["say"]["elektrik"]["toplam"]
    assert t["pafta"]["SALON"]["uydu_prizi"] == t["model"]["SALON"]["uydu_prizi"] == 1
    assert t["pafta"]["HOL"]["sigorta_kutusu"] == t["model"]["HOL"]["sigorta_kutusu"] == 1
    md = OT.md(t)
    assert md.count("\n| ") == 1 + len(t["odalar"]) + 1 and "| **Toplam** | **5 / 6** | **3 / 2** |" in md
    r = subprocess.run([sys.executable, "-I", str(KOK / "testler" / "kimlik" / "k3_oda_tablosu.py"), str(komut["klasor"] / "PK01_kart_sayim.json"),
                        str(VERI / "PK01_pafta_katmani.json")], capture_output=True, text=True, cwd=KOK)
    assert r.returncode == 0 and r.stdout.strip() == md.strip(), r.stderr

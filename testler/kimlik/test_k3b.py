"""İE-15 · K3b · Alt çerçeve ve çatı kenarları: montaj rehberinin kartsız yedi parçası (plan `K3b_rehber_parcalari_plani.md`; Lokman,
25 Ekim 17:47 "Uygula"; kurallar `kurallar/yerlesim/prefabrik.json`, kod `motor/yerlesim/kenar.py`).

- Sayı = liste üretecinin yöntemi × modelin boyu: dokuz İş Emri'nde (PK01, ALTIN, P6001 ve iki aynası, yan verandalı, TEST01–03) alt çerçeve
  (100'lük, 60'lık), saçak sacı, alın sacı, baş makas Z'si ve mahya liste üretecinin satırıyla aynı. Saçak omegası rehberdeki gibi duvar
  boyunca (liste saçak boyunu sayar): PK01 8 / 9, yan verandalı 9 / 10, ötekiler aynı. C kapama listede yok; metresi dört eğik kenar.
  Kesim planının yukarı yuvarlaması liste üretecininki (1e-9).
- Kurallar ve değerler: `okunan` çözülür (liste kuralları ve kitap), seçimler kendi durumunu taşır; kartlar (sekiz kod, kilitte ilk 51 satır).
- Yer (montaj rehberi; teyit bekleyenler varsayım): alt çerçeve duvarların altında, kesilen kapı boşlukları açık, dış hatlar dış sınırın
  kenarları (toplam dış çevre), iç hatlar iç duvarın bitişik panoları, dokuz evde her kapı ve geçiş kesik; saçak omegası saçak duvarının üstünde makasın altında binanın boyunca;
  saçak sacı ve mahya saçak ucundan ucuna; her baş makasın iki eğik kenarında Z betopanın dış yüzünde, C kapama aşık uçlarında, alın sacı
  C'nin önünde. Gösterim kesitleri sayıyı ve boyu değiştirmez.
- Kural yok: beşik olmayan çatıda altı çatı parçası, trapez kaplamada mahya, baş makas yoksa kalkan kenarları; alt çerçeve yine kurulur.
- Kart modeli: K3 modelinin parçaları aynen, yeni parçalar yalnız K3b kartlarından, pafta anahtarı yok (hakemin K3 kapısı aynı sayar);
  boy `Qto_…BaseQuantities.Length`'te; GlobalId ev kimliğinden. Komut: yerleşim planlarının yedinci sayfası, Excel'de kesim planı ve
  "Alt çerçeve ve çatı kenarı" sayfası (model ile liste yan yana).
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import sys
import uuid
from collections import Counter
from pathlib import Path

import ifcopenshell
import ifcopenshell.guid
import ifcopenshell.util.element as ue
import pytest
from shapely.geometry import box
from shapely.ops import unary_union

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
sys.path.insert(0, str(KOK / "testler" / "kimlik" / "hakem"))
import kiyas as KY  # noqa: E402
from motor import geometry as G  # noqa: E402
from motor import ifc_uret  # noqa: E402
from motor import kart_model as KM  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402
from motor import yerlesim as Y  # noqa: E402
from motor.yerlesim import cati, cizim, girdi, kesim, yaz  # noqa: E402
from motor.yerlesim import kenar as KE  # noqa: E402

VERI = KOK / "testler" / "kimlik" / "veri"
MOTOR_VERI = KOK / "testler" / "motor" / "veri"
PROJE = KOK / "testler" / "kimlik" / "hakem" / "projeler" / "PK01.json"
IS_EMIRLERI = {"PK01": MOTOR_VERI / "PK01.json", "ALTIN": KOK / "is_emri" / "ALTIN_2plus1_6x8.json",
               "P6001": KOK / "is_emri" / "P6001_2plus1_6x6.json", "P6001_ayna_x": KOK / "is_emri" / "P6001_ayna_x.json",
               "P6001_ayna_y": KOK / "is_emri" / "P6001_ayna_y.json", "YAN_VERANDALI": KOK / "is_emri" / "YAN_VERANDALI_2plus1_6x8p5.json",
               "TEST01": KOK / "is_emri" / "adaylar" / "TEST01.json", "TEST02": KOK / "is_emri" / "adaylar" / "TEST02.json",
               "TEST03": KOK / "is_emri" / "adaylar" / "TEST03.json"}
K3B_KODLARI = ("AC-01", "AC-02", "CK-01", "CK-02", "CK-03", "CK-04", "CK-05", "CK-06")
STOKLU = ("AC-01", "AC-02", "CK-01", "CK-02", "CK-03", "CK-04", "CK-06")
KILIT_K3B = "9c7fc8895fed7f98eddb8f41dc5ba529bd7ce09f50507ce04479b6d94a59425d"   # K1 25, K2 7, K3 11 ve K3b'nin 8 kodu (bu PR)
# modelin stok adedi (kesim planı), yukarıdaki sırayla; liste üreteciyle farkı yalnız saçak omegasında (rehber duvar boyu, liste saçak boyu)
MODEL = {"PK01": (9, 6, 8, 10, 7, 8, 5), "ALTIN": (10, 5, 9, 10, 8, 8, 5), "P6001": (9, 3, 7, 8, 8, 8, 4), "P6001_ayna_x": (9, 3, 7, 8, 8, 8, 4),
         "P6001_ayna_y": (9, 3, 7, 8, 8, 8, 4), "YAN_VERANDALI": (10, 7, 9, 10, 8, 8, 5), "TEST01": (11, 6, 8, 9, 10, 12, 4),
         "TEST02": (11, 5, 10, 11, 9, 8, 5), "TEST03": (11, 6, 8, 9, 10, 12, 4)}
SACAK_OMEGASI_LISTE = {"PK01": 9, "YAN_VERANDALI": 10}


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _model(ie: dict, klasor: Path, ad: str, **kw) -> dict:
    bugunku = klasor / f"{ad}.ifc"
    if not bugunku.exists():
        assert ifc_uret.uret(ie, bugunku)["sema_hatasi"] == 0
    kw.setdefault("kenar", True)
    say = KM.kart_modeli(bugunku, klasor / f"{ad}_model.ifc", ie["cati"].get("tip", "besik"), is_emri=ie, plan=ifc_uret.hazirla(ie), **kw)
    return {"bugunku": bugunku, "model": klasor / f"{ad}_model.ifc", "say": say}


def _parcalar(model: Path, kodlar=K3B_KODLARI) -> list[dict]:
    f = ifcopenshell.open(str(model))
    kutu = KM._Kutular()
    out = []
    for e in f.by_type("IfcProduct"):
        p = ue.get_psets(e).get(KM.PARCA_SETI, {})
        if p.get("KartKodu") in kodlar:
            a, b = kutu(e)
            out.append({"e": e, "f": f, "ad": e.Name, "kod": p["KartKodu"], "p": p, "a": a, "b": b, "q": ue.get_psets(e, qtos_only=True)})
    return out


def _cerceve(bugunku: Path, ie: dict) -> tuple[dict, dict]:
    g = girdi.kur(ie, ifc_uret.hazirla(ie), KM.makaslar(ifcopenshell.open(str(bugunku)), KM._Kutular()))
    return g, cati.cerceve(g)


def _uv(F: dict, a: list, b: list) -> tuple[float, float, float, float, float, float]:
    """Dünya kutusu → çatının çerçevesi: (u0, u1, v0, v1, z0, z1); u makasın açıklığı boyunca, v binanın boyu boyunca."""
    if F["mx"]:
        return a[0], b[0], a[1], b[1], a[2], b[2]
    return a[1], b[1], a[0], b[0], a[2], b[2]


@pytest.fixture(scope="module")
def pk01(tmp_path_factory):
    d = tmp_path_factory.mktemp("k3b")
    ie = _oku(MOTOR_VERI / "PK01.json")
    m = _model(ie, d, "PK01")
    g, F = _cerceve(m["bugunku"], ie)
    return m | {"klasor": d, "ie": ie, "g": g, "F": F}


@pytest.fixture(scope="module")
def dokuz(tmp_path_factory):
    d = tmp_path_factory.mktemp("k3b_dokuz")
    out = {}
    for ad, yol in IS_EMIRLERI.items():
        ie = _oku(yol)
        out[ad] = _model(ie, d, ad) | {"ie": ie}
    return out


@pytest.fixture(scope="module")
def komut(tmp_path_factory):
    d = tmp_path_factory.mktemp("k3b_komut")
    assert KM.main([str(MOTOR_VERI / "PK01.json"), str(d), "--kod", "PK01"]) == 0
    return {"klasor": d, "say": _oku(d / "PK01_kart_sayim.json")}


# ----------------------------------------------------------------------------- kurallar, kartlar, kilit
def test_kurallar_ve_degerler():
    """Yedi kural ve gösterim yerleşim kurallarında; parça boyu ve bindirme liste kurallarından ya da kitaptan okunur (burada sayı yazılmaz);
    seçimler kendi durumunu taşır; liste kurallarıyla ortak anahtar yok."""
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    K, D = y["kurallar"], Y.degerler(y, liste)
    assert set(KE.KENAR) | {KE.GOSTERIM} <= set(K) and not (set(K) & set(liste.get("kurallar", {})))
    assert {ad: (D[ad].get("parca_boyu"), D[ad].get("bindirme"), D[ad]["yontem"]) for ad in KE.KENAR} == {
        "alt_cerceve_yerlesimi": (3.75, None, "toplam"), "sacak_omegasi_yerlesimi": (2.5, None, "toplam"),
        "sacak_saci_yerlesimi": (2.5, 0.2, "toplam"), "alin_saci_yerlesimi": (2.5, 0.2, "toplam"), "bas_makas_z_yerlesimi": (2.8, None, "kenar"),
        "c_kapama_yerlesimi": (None, None, "metre"), "mahya_yerlesimi": (2.8, 0.2, "toplam")}
    durumlar = set(_oku(KM.KARTLAR)["durumlar"])
    for ad in KE.KENAR:
        k = K[ad]
        assert k["kural"] and k["durum"] in durumlar and k["tarih"] == "2026-10-25" and k["kaynak"], ad
        assert ad == KE.ALT_CERCEVE or k["cati_tipleri"] == ["besik"], ad
        for alt in ("parca_boyu", "bindirme"):
            if alt in k:
                assert set(k[alt]) == {"okunan"} and k[alt]["okunan"].startswith(("kurallar/liste/prefabrik.json › ", "kitap › ")), (ad, alt)
        for alt, v in k.items():
            if isinstance(v, dict) and "deger" in v:
                assert v["durum"] in durumlar and v["tarih"] and v.get("not"), (ad, alt)
    assert K[KE.ALT_CERCEVE]["kesilen_bosluklar"]["durum"] == "varsayim" and K["mahya_yerlesimi"]["trapez_yazisi"]["durum"] == "muhendis"
    assert D[KE.ALT_CERCEVE]["kesilen_bosluklar"] == ["ic_kapi", "dis_kapi", "gecis"] and D[KE.ALT_CERCEVE]["ic_100_esigi_mm"] == 100
    assert K[KE.GOSTERIM]["kesit_cm"]["durum"] == "varsayim"


def test_kartlar():
    """Sekiz kart: yerleşim kuralına bağlı, paftada anahtarı yok (paftada bu parçalar yok; hakem saymaz), boyla sayılır (C kapama metreyle);
    C kapamanın listede satırı yok, nedenini yazar."""
    kartlar = {k["kod"]: k for k in _oku(KM.KARTLAR)["kartlar"]}
    assert all(k in kartlar for k in K3B_KODLARI)
    beklenen = {"AC-01": ("alt_cerceve_yerlesimi", "IfcMemberType", "boy_kesim", "varsayim"),
                "AC-02": ("alt_cerceve_yerlesimi", "IfcMemberType", "boy_kesim", "varsayim"),
                "CK-01": ("sacak_omegasi_yerlesimi", "IfcMemberType", "boy_kesim", "varsayim"),
                "CK-02": ("sacak_saci_yerlesimi", "IfcCoveringType", "boy_kesim", "varsayim"),
                "CK-03": ("alin_saci_yerlesimi", "IfcCoveringType", "boy_kesim", "varsayim"),
                "CK-04": ("bas_makas_z_yerlesimi", "IfcMemberType", "boy_kesim", "varsayim"),
                "CK-05": ("c_kapama_yerlesimi", "IfcCoveringType", "metre", "varsayim"),
                "CK-06": ("mahya_yerlesimi", "IfcCoveringType", "boy_kesim", "muhendis")}
    assert {k: (kartlar[k]["yerlesim"]["kural"], kartlar[k]["ifc"]["sinif"], kartlar[k]["sayim"], kartlar[k]["kaynak"]["durum"])
            for k in K3B_KODLARI} == beklenen
    assert all(kartlar[k]["pafta"] == [] and kartlar[k]["birim"] == ("m" if k == "CK-05" else "adet") for k in K3B_KODLARI)
    assert kartlar["CK-05"]["liste"]["satirlar"] == [] and kartlar["CK-05"]["liste"]["yok"]
    assert all(kartlar[k]["liste"]["satirlar"] for k in STOKLU)
    assert (kartlar["AC-01"]["yerlesim"]["bilgi"], kartlar["AC-02"]["yerlesim"]["bilgi"]) == ({"Cerceve": "100'lük"}, {"Cerceve": "60'lık"})


def test_kod_kilidi_k3b():
    """K3b'de eklenen sekiz kod kilitte; kilidin ilk 51 satırı (K1, K2, K3 ve K3b) donmuştur, yeni kod sona eklenir."""
    kilit = _oku(VERI / "kart_kodlari.json")["kodlar"]
    ilk = dict(list(kilit.items())[:51])
    assert tuple(list(ilk)[43:]) == K3B_KODLARI
    assert hashlib.sha256(json.dumps(ilk, ensure_ascii=False, sort_keys=True).encode()).hexdigest() == KILIT_K3B


# ----------------------------------------------------------------------------- sayı = liste yöntemi × modelin boyu
@pytest.mark.parametrize("ad", list(IS_EMIRLERI))
def test_model_liste_ile_ayni(dokuz, ad):
    m = dokuz[ad]
    say = m["say"]
    kesim, liste = say["kesim_plani"], say["kenar"]["liste"]
    assert tuple(kesim[k]["stok"] for k in STOKLU) == MODEL[ad]
    for k in STOKLU:
        assert liste[k] == (SACAK_OMEGASI_LISTE.get(ad, kesim[k]["stok"]) if k == "CK-01" else kesim[k]["stok"]), (ad, k)
    assert liste["CK-05"] is None
    # liste üretecinin aynı satırları (özetteki liste bilgidir; burada üreteçten yeniden okunur)
    satir = Counter()
    for s in PL.liste(PL.model_motor(copy.deepcopy(m["ie"])))["satirlar"]:
        if isinstance(s["adet"], int):
            satir[s["kalem"]] += s["adet"]
    kartlar = {k["kod"]: k for k in _oku(KM.KARTLAR)["kartlar"]}
    assert {k: satir[kartlar[k]["liste"]["satirlar"][0]] for k in STOKLU} == {k: liste[k] for k in STOKLU}
    # C kapama: dört eğik kenar (iki baş makas × iki yüz)
    g, F = _cerceve(m["bugunku"], m["ie"])
    ck5 = next(k for k in say["kartlar"] if k["kod"] == "CK-05")
    assert ck5["adet"] == 4 and abs(ck5["metre"] - 4 * F["egik"] / 100) < 1e-3
    assert say["kenar"]["uyarilar"] == [] and not [x for x in say["kural_yok"] if x["parca"] in cizim.SAYFA[7][0]]


def test_alt_cerceve_hatlari_liste_uretecinin_boyu(dokuz):
    """Dış hatların toplamı liste üretecinin dış çevresi, iç hatların toplamı iç duvarların pano boyu (dokuz evde)."""
    for ad, m in dokuz.items():
        N = PL.model_motor(copy.deepcopy(m["ie"]))
        h = m["say"]["kenar"]["hatlar"]
        assert abs(sum(x["boy_cm"] for x in h if x["duvar"] == "dış") - N["dis_cevre_m"] * 100) < 1e-6, ad
        assert abs(sum(x["boy_cm"] for x in h if x["duvar"] == "iç") - sum(q["boy"] for q in N["panolar"] if q["duvar"] == "ic")) < 1e-6, ad
        assert {x["cerceve"] for x in h if x["duvar"] == "dış"} == {"100'lük"} and {x["cerceve"] for x in h if x["duvar"] == "iç"} == {"60'lık"}


def test_alt_cerceve_kapi_ve_gecis_kesilir(dokuz):
    """Dokuz evde her kapı (iç ve dış) ve geçiş bir hatta kesiktir, pencere kesilmez; geçiş yalnız yan verandalıda (kuralın değeri test
    edilir: geçiş çıkarılırsa yan verandalı bir kesik eksik verir)."""
    turler = Counter()
    for ad, m in dokuz.items():
        P = G.kur(copy.deepcopy(m["ie"]))
        acik = Counter(q.acilim.tur for d in P.duvarlar for q in d.panolar if q.acilim)
        turler.update({(ad, t): n for t, n in acik.items() if t == "gecis"})
        kesik = sum(len(h["kesikler_cm"]) for h in m["say"]["kenar"]["hatlar"])
        assert kesik == acik["ic_kapi"] + acik["dis_kapi"] + acik["gecis"], (ad, kesik, dict(acik))
    assert turler == {("YAN_VERANDALI", "gecis"): 1}
    y = _oku(KM.YERLESIM)
    y["kurallar"][KE.ALT_CERCEVE]["kesilen_bosluklar"]["deger"] = ["ic_kapi", "dis_kapi"]
    D = Y.degerler(y, _oku(KM.LISTE_KURALLARI))
    yan = dokuz["YAN_VERANDALI"]
    _p, hatlar = KE.alt_cerceve(yan["ie"], D[KE.ALT_CERCEVE], D[KE.GOSTERIM]["kesit_cm"])
    assert sum(len(h["kesikler_cm"]) for h in hatlar) == sum(len(h["kesikler_cm"]) for h in yan["say"]["kenar"]["hatlar"]) - 1


def test_kesim_yuvarlamasi_liste_uretecinin():
    """Yukarı yuvarlamanın payı liste üretecininki (`TAVAN`, 1e-9): bölüm bir tamsayının hemen üstüne düşerse iki taraf da bir fazla sayar;
    kayan nokta gürültüsü (dört hattın toplamı 1000,0000000000001, hesaplanan kenar 560,0000000000001) bir fazla saydırmaz."""
    for hat, boy, bindirme in (([2000.0], 250.0, 0.0), ([2000.0001], 250.0, 0.0), ([1999.99999999], 250.0, 0.0), ([1060.0, 1060.0], 250.0, 20.0),
                               ([920.00002], 250.0, 20.0), ([298.6, 106.6, 82.6, 512.2], 250.0, 0.0), ([284.1, 119.1, 234.2, 282.6], 250.0, 20.0)):
        assert kesim.toplam_karti(hat, boy, bindirme)["stok"] == PL.TAVAN(sum(hat) / (boy - bindirme)), (hat, boy, bindirme)
    for kenar in (560.0, 560.0001, 347.5043, 560.0000000000001):
        assert kesim.kenar_karti([kenar] * 4, 280.0)["stok"] == 4 * PL.TAVAN(kenar / 280.0), kenar
    assert kesim.toplam_karti([2000.0001], 250.0)["stok"] == 9 and kesim.kenar_karti([560.0001], 280.0)["stok"] == 3
    assert sum([298.6, 106.6, 82.6, 512.2]) > 1000 and kesim.toplam_karti([298.6, 106.6, 82.6, 512.2], 250.0)["stok"] == 4
    assert sum([284.1, 119.1, 234.2, 282.6]) > 920 and kesim.toplam_karti([284.1, 119.1, 234.2, 282.6], 250.0, 20.0)["stok"] == 4
    assert kesim.kenar_karti([560.0000000000001], 280.0)["stok"] == 2


# ----------------------------------------------------------------------------- yer (PK-01)
def test_alt_cerceve_duvarlarin_altinda_kapilar_acik(pk01):
    """Her alt çerçeve parçası duvarın altında (planda panoların ve bağlantıların izinde, kanat payı dahil) ve zeminde; kapı boşluğuna parça
    girmez; kesilen boşluk sayısı modeldeki kapılar (beş iç kapı ve giriş kapısı)."""
    s = pk01["say"]
    hatlar = {h["no"]: h for h in s["kenar"]["hatlar"]}
    f = ifcopenshell.open(str(pk01["model"]))
    sek = cizim._sekiller(f)
    duvar = unary_union([cizim._iz(sek[e.GlobalId]) for e in f.by_type("IfcWall") + f.by_type("IfcMember")
                         if e.GlobalId in sek and (e.is_a("IfcWall") or e.Name.startswith("BAGLANTI_"))]).buffer(0.6)     # pano ve köşe/ek
    ac = _parcalar(pk01["model"], ("AC-01", "AC-02"))
    assert Counter(p["kod"] for p in ac) == {"AC-01": 7, "AC-02": 11}
    for p in ac:
        assert duvar.contains(box(p["a"][0], p["a"][1], p["b"][0], p["b"][1])), p["ad"]
        assert -0.3 <= p["a"][2] and p["b"][2] <= 3.0, p["ad"]
        h = hatlar[p["p"]["Hat"]]
        assert p["p"]["Cerceve"] == h["cerceve"] and p["kod"] == ("AC-01" if h["duvar"] == "dış" else "AC-02") and p["p"]["HatBoyu_cm"] == h["boy_cm"]
        i = 0 if h["yon"] == "x" else 1
        for k0, k1 in h["kesikler_cm"]:
            assert p["b"][i] <= k0 + 1e-6 or p["a"][i] >= k1 - 1e-6, (p["ad"], k0, k1)
    adet = {k["kod"]: k["adet"] for k in s["kartlar"]}
    assert sum(len(h["kesikler_cm"]) for h in hatlar.values()) == adet["K-01"] + adet["K-02"] == 6
    assert [hatlar[2]["kesikler_cm"], hatlar[2]["parcalar_cm"]] == [[[278.75, 368.75]], [12.75, 231.25]]   # giriş kapısı, y = 200
    assert sum(h["boy_cm"] for h in hatlar.values() if h["duvar"] == "dış") == 3200


def test_alt_cerceve_100_esigi_kuraldan(pk01, tmp_path):
    """İç duvarın çerçevesi kuraldaki eşikten (kalınlık ≥ eşik → 100'lük): eşik 60 mm olunca iç duvarlar da 100'lük sayılır (liste değişmez)."""
    y = _oku(KM.YERLESIM)
    y["kurallar"][KE.ALT_CERCEVE]["ic_100_esigi_mm"]["deger"] = 60
    (tmp_path / "PK01.ifc").write_bytes(pk01["bugunku"].read_bytes())
    m = _model(pk01["ie"], tmp_path, "PK01", yerlesim=y)
    k, liste = m["say"]["kesim_plani"], m["say"]["kenar"]["liste"]
    assert "AC-02" not in k and k["AC-01"]["stok"] == math.ceil((3200 + 2187.75) / 375) == 15 and (liste["AC-01"], liste["AC-02"]) == (9, 6)


def test_cati_kenarlari_yerleri(pk01):
    F, g = pk01["F"], pk01["g"]
    P = _parcalar(pk01["model"], ("CK-01", "CK-02", "CK-03", "CK-04", "CK-05", "CK-06"))
    D = Y.degerler(_oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI))
    G_, betopan = D[KE.GOSTERIM]["kesit_cm"], D["gosterim"]["kesit_cm"]["betopan"]
    yakin = lambda a, b, e=0.01: abs(a - b) <= e  # noqa: E731
    assert Counter(p["kod"] for p in P) == {"CK-01": 2, "CK-02": 2, "CK-03": 4, "CK-04": 4, "CK-05": 4, "CK-06": 1}
    for p in P:
        u0, u1, v0, v1, z0, z1 = _uv(F, p["a"], p["b"])
        if p["kod"] == "CK-01":                                            # saçak duvarının üstünde, makasın altında, binanın boyunca
            assert yakin(v0, 0) and yakin(v1, F["boy"]) and yakin(z0, g["duvar_h"]) and z1 <= F["zs"] + 1e-6, p["ad"]
            assert (yakin(u0, 0) and u1 <= 10 + 1e-6) if p["p"]["Yuz"] == "sol" else (yakin(u1, F["aciklik"]) and u0 >= F["aciklik"] - 10 - 1e-6)
        elif p["kod"] in ("CK-02", "CK-06"):                               # saçak ucundan ucuna
            assert yakin(v0, F["v0"]) and yakin(v1, F["v1"]) and yakin(p["p"]["Boy_cm"], F["v1"] - F["v0"]), p["ad"]
            if p["kod"] == "CK-02":                                        # saçak ucunda, tırnağı duvarın üstünde
                ic = p["p"]["Yuz"] == "sol"
                assert (u0 <= -F["s"] and yakin(u1, G_["sacak_saci"]["tirnak"])) if ic else \
                       (u1 >= F["aciklik"] + F["s"] and yakin(u0, F["aciklik"] - G_["sacak_saci"]["tirnak"])), p["ad"]
                assert yakin(z0, g["duvar_h"] + G_["sacak_omegasi"][1]), p["ad"]
            else:                                                          # mahyada, kaplamanın üstünde
                assert yakin((u0 + u1) / 2, F["um"]) and z0 > cati._ust(F, F["um"]), p["ad"]
        else:                                                              # kalkan: eğik boy, yüzün yarısı
            on = p["p"]["Makas"] == "MAKAS_01"
            assert yakin(p["p"]["Boy_cm"], F["egik"]) and p["p"]["Makas"] in ("MAKAS_01", "MAKAS_09"), p["ad"]
            assert (yakin(u0, -F["s"], 0.5) and yakin(u1, F["um"], 0.5)) if p["p"]["Yuz"] == "sol" else \
                   (yakin(u0, F["um"], 0.5) and yakin(u1, F["aciklik"] + F["s"], 0.5)), p["ad"]
            v_dis = -betopan if on else F["boy"] + betopan
            v_uc = F["v0"] if on else F["v1"]
            if p["kod"] == "CK-04":                                        # Z makasın üstünden betopanın dış yüzüne
                assert v0 - 1e-6 <= v_dis <= v1 + 1e-6 and yakin(v1 - v0, G_["bas_makas_z"]["ust_kanat"] + G_["bas_makas_z"]["dis_kanat"]), p["ad"]
            elif p["kod"] == "CK-05":                                      # C aşık uçlarında
                assert v0 - 1e-6 <= v_uc <= v1 + 1e-6 and v1 - v0 <= G_["c_kapama"]["kanat"] + G_["sac"] + 1e-6, p["ad"]
            else:                                                          # alın C'nin önünde, saçak altı Z'ye kadar
                assert (v0 < v_uc - G_["sac"] + 1e-6) if on else (v1 > v_uc + G_["sac"] - 1e-6), p["ad"]
                assert (yakin(v1, v_dis - G_["bas_makas_z"]["dis_kanat"])) if on else yakin(v0, v_dis + G_["bas_makas_z"]["dis_kanat"]), p["ad"]
    for kod in ("CK-03", "CK-04", "CK-05"):
        assert sorted((p["p"]["Makas"], p["p"]["Yuz"]) for p in P if p["kod"] == kod) == [
            ("MAKAS_01", "sag"), ("MAKAS_01", "sol"), ("MAKAS_09", "sag"), ("MAKAS_09", "sol")], kod


def test_gosterim_sayiyi_degistirmez(pk01, tmp_path):
    y = _oku(KM.YERLESIM)
    y["kurallar"][KE.GOSTERIM]["kesit_cm"]["deger"] = {"sac": 0.5, "alt_cerceve": {"pay": 1, "kanat": 5}, "sacak_omegasi": [14, 6],
                                                       "sacak_saci": {"tirnak": 6, "alin": 20, "kivrim": 4}, "alin_saci": {"makas_ustunden": 5, "alin": 20, "kivrim": 4},
                                                       "bas_makas_z": {"centik": 5, "dis_kanat": 6, "ust_kanat": 9}, "c_kapama": {"govde": 9, "kanat": 4}, "mahya": {"kanat": 30}}
    (tmp_path / "PK01.ifc").write_bytes(pk01["bugunku"].read_bytes())
    m = _model(pk01["ie"], tmp_path, "PK01", yerlesim=y)
    a, b = pk01["say"], m["say"]
    assert [(k["kod"], k["adet"], k.get("metre")) for k in a["kartlar"]] == [(k["kod"], k["adet"], k.get("metre")) for k in b["kartlar"]]
    assert a["kesim_plani"] == b["kesim_plani"] and a["kenar"]["liste"] == b["kenar"]["liste"]
    assert [(p["parca"], p["boy_cm"], p["bilgi"]) for p in a["kenar"]["parcalar"]] == [(p["parca"], p["boy_cm"], p["bilgi"]) for p in b["kenar"]["parcalar"]]
    for p in _parcalar(m["model"], ("CK-04",)):                            # gösterim değeri yere geçer: Z makasın 9 cm üstünden dış kanada
        _u0, _u1, v0, v1, _z0, _z1 = _uv(pk01["F"], p["a"], p["b"])
        assert abs((v1 - v0) - (9 + 6)) <= 0.01, p["ad"]


# ----------------------------------------------------------------------------- kural yok
def test_kirma_catida_cati_kenarlari_kural_yok(pk01, tmp_path):
    ie = copy.deepcopy(pk01["ie"])
    ie["cati"]["tip"] = "kirma"
    m = _model(ie, tmp_path, "kirma")
    k = [x for x in m["say"]["kural_yok"] if x["parca"] in cizim.SAYFA[7][0]]
    assert [x["parca"] for x in k] == ["saçak omegası", "saçak sacı", "alın sacı", "baş makas Z'si", "C kapama", "mahya"]
    assert all("\"kirma\"" in x["neden"] and "besik" in x["neden"] for x in k)
    adet = {x["kod"]: x["adet"] for x in m["say"]["kartlar"]}
    assert [adet[c] for c in K3B_KODLARI[2:]] == [0] * 6 and m["say"]["kesim_plani"]["AC-01"]["stok"] == 9      # alt çerçeve çatıya bağlı değil


def test_trapez_kaplamada_mahya_kural_yok(pk01, tmp_path):
    ie = copy.deepcopy(pk01["ie"])
    ie["ozellikler"]["Çatı Kaplaması"] = "Trapez Sac"
    m = _model(ie, tmp_path, "trapez")
    k = [x for x in m["say"]["kural_yok"] if x["parca"] in cizim.SAYFA[7][0]]
    assert [x["parca"] for x in k] == ["mahya"] and "trapez sac" in k[0]["neden"]
    assert {c: m["say"]["kesim_plani"][c]["stok"] for c in STOKLU[:-1]} == {c: pk01["say"]["kesim_plani"][c]["stok"] for c in STOKLU[:-1]}
    assert "CK-06" not in m["say"]["kesim_plani"]


def test_bas_makas_yoksa_kalkan_kenarlari_kural_yok(pk01):
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    D = Y.degerler(y, liste)
    g = copy.deepcopy(pk01["g"])
    for m in g["makaslar"]:
        m["rol"] = "orta"
    s = KE.katman(pk01["ie"], g, D, D["gosterim"]["kesit_cm"], y["kurallar"])
    assert [x["parca"] for x in s["kural_yok"]] == ["alın sacı", "baş makas Z'si", "C kapama"]
    assert Counter(p["parca"] for p in s["parcalar"]) == {"alt_cerceve": 18, "sacak_omegasi": 2, "sacak_saci": 2, "mahya": 1}


def test_cati_tipleri_kuralin_kendisinden(pk01):
    """Çatı tipi sınırı kuralın kendi `cati_tipleri`'nden okunur (K2'deki gibi): kırmaya izin verilen kural kırmada da kurulur."""
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    D = Y.degerler(y, liste)
    y["kurallar"]["sacak_omegasi_yerlesimi"]["cati_tipleri"] = ["besik", "kirma"]
    g = copy.deepcopy(pk01["g"])
    g["cati_tipi"] = "kirma"
    s = KE.katman(pk01["ie"], g, D, D["gosterim"]["kesit_cm"], y["kurallar"])
    assert "saçak omegası" not in [x["parca"] for x in s["kural_yok"]] and len(s["kural_yok"]) == 5
    assert sum(p["parca"] == "sacak_omegasi" for p in s["parcalar"]) == 2


# ----------------------------------------------------------------------------- kart modeli ve komut
def test_k3_parcalari_aynen_yeni_parcalar_k3b(komut, tmp_path):
    """K3b yalnız ekler: K3 modelinin (elektrikli, kenarsız) her parçası K3b modelinde aynen; yeni parçalar yalnız K3b kartlarından, pafta
    anahtarsız; hakemin K3 kapısı aynı sayıyla geçer."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    (tmp_path / "PK01.ifc").write_bytes((komut["klasor"] / "PK01.ifc").read_bytes())
    k3 = _model(ie, tmp_path, "PK01", elektrik=True, kenar=False)
    a, b = ifcopenshell.open(str(k3["model"])), ifcopenshell.open(str(komut["klasor"] / "PK01_model.ifc"))
    A = {e.GlobalId: e for e in a.by_type("IfcProduct")}
    B = {e.GlobalId: e for e in b.by_type("IfcProduct")}
    assert set(A) <= set(B)
    for g_, e in A.items():
        m = B[g_]
        assert (m.is_a(), m.Name) == (e.is_a(), e.Name), e.Name
        assert {k: {x: y for x, y in v.items() if x != "id"} for k, v in ue.get_psets(e).items()} == \
               {k: {x: y for x, y in v.items() if x != "id"} for k, v in ue.get_psets(m).items()}, e.Name
    yeni = [B[g_] for g_ in set(B) - set(A)]
    assert Counter(ue.get_psets(e)[KM.PARCA_SETI]["KartKodu"] for e in yeni) == {
        "AC-01": 7, "AC-02": 11, "CK-01": 2, "CK-02": 2, "CK-03": 4, "CK-04": 4, "CK-05": 4, "CK-06": 1}
    assert all(ue.get_psets(e)[KM.PARCA_SETI]["PaftaAnahtari"] == "" for e in yeni)
    x, y = json.loads(json.dumps(k3["say"], ensure_ascii=False)), komut["say"]
    assert [k for k in x["kartlar"] if k["kod"] not in K3B_KODLARI] == [k for k in y["kartlar"] if k["kod"] not in K3B_KODLARI]
    assert x["yerlesim"] == y["yerlesim"] and x["elektrik"] == y["elektrik"] and "kenar" not in x
    assert {k: v for k, v in y["kesim_plani"].items() if k not in K3B_KODLARI} == x["kesim_plani"]
    k = KY.kiyasla(komut["klasor"] / "PK01_model.ifc", _oku(VERI / "PK01_pafta_katmani.json"), _oku(PROJE), adim="K3")
    assert k["kapi_1"] == "geçer" and k["kart_modeli"] == 211 + 32


def test_parca_ozellikleri(komut):
    """Boy `Qto_…BaseQuantities.Length`'te (m; C kapamanın metresi buradan), kartın türü bağlı, GlobalId ev kimliğinden ve adından."""
    P = _parcalar(komut["klasor"] / "PK01_model.ifc")
    qto = {"IfcMember": "Qto_MemberBaseQuantities", "IfcCovering": "Qto_CoveringBaseQuantities"}
    adlar = {"AC": "ALT_CERCEVE_", "CK-01": "SACAK_OMEGASI_", "CK-02": "SACAK_SACI_", "CK-03": "ALIN_SACI_", "CK-04": "BAS_MAKAS_Z_",
             "CK-05": "C_KAPAMA_", "CK-06": "MAHYA_"}
    metre = Counter()
    for p in P:
        e = p["e"]
        assert e.GlobalId == ifcopenshell.guid.compress(uuid.uuid5(yaz.AD_ALANI, f"PM-PK01-60M2/{e.Name}").hex), e.Name
        assert e.Name.startswith(adlar["AC" if p["kod"].startswith("AC") else p["kod"]]) and ue.get_type(e).Tag == p["kod"], e.Name
        L = p["q"][qto[e.is_a()]]["Length"]
        assert abs(L - p["p"]["Boy_cm"] / 100) < 1e-9 and p["p"]["Kural"] in KE.KENAR, e.Name
        metre[p["kod"]] += L
    kart = {k["kod"]: k for k in komut["say"]["kartlar"]}
    assert abs(kart["CK-05"]["metre"] - metre["CK-05"]) < 1e-3 and abs(kart["AC-01"]["metre"] - metre["AC-01"]) < 1e-3


def test_komut_yedinci_sayfa_ve_excel(komut):
    import pymupdf
    d = pymupdf.open(komut["klasor"] / "PK01_yerlesim_planlari.pdf")
    assert len(d) == 7
    metin = " ".join(d[6].get_text().split())
    assert "PK01 · ALT ÇERÇEVE VE ÇATI KENARLARI YERLEŞİM PLANI" in metin and "(İE-15 K3b)" in metin
    for s in ("AC-01 100'lük alt çerçeve L 3750 model 9 boy · liste 9", "AC-02 60'lık alt çerçeve L 3750 model 6 boy · liste 6",
              "CK-01 100'lük saçak omegası L 2500 model 8 boy · liste 9 ≠", "CK-02 Saçak sacı (30'luk) L 2500 model 10 boy · liste 10",
              "CK-03 Alın sacı (30'luk) L 2500 model 7 boy · liste 7", "CK-04 Baş makas Z'si L 2800 model 8 boy · liste 8",
              "CK-05 C kapama model 13,90 m · listede yok", "CK-06 Metal kiremit mahyası L 2800 model 5 boy · liste 5",
              "2 · 334", "4 · 1000", "9 · 182,25", "MAKAS_01 sol: 347,5"):
        assert s in metin, s
    assert not any(x in metin for x in ("EKSİK", "Kurulmadı", "Uyarı"))
    from openpyxl import load_workbook
    wb = load_workbook(komut["klasor"] / "PK01_kart_katalogu.xlsx")
    satir = {r[0]: r for r in wb["Kesim planı"].iter_rows(min_row=2, values_only=True)}
    assert {k: satir[k][5] for k in STOKLU} == dict(zip(STOKLU, MODEL["PK01"])) and satir["AC-01"][6] == "hatlar 266 + 334 + 600 + 1000 + 200 + 800"
    ws = wb["Alt çerçeve ve çatı kenarı"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 35 + 1 + 1 + 8
    yan = {r[0]: (r[1], r[2]) for r in rows[-8:]}
    assert yan == {"AC-01": (9, 9), "AC-02": (6, 6), "CK-01": (8, 9), "CK-02": (10, 10), "CK-03": (7, 7), "CK-04": (8, 8),
                   "CK-05": ("13.9002 m", "listede yok"), "CK-06": (5, 5)}
    kartlar = {r[0]: r for r in wb["Kartlar"].iter_rows(min_row=2, values_only=True)}
    assert kartlar["CK-05"][3] == "(listede yok)"


def test_kirma_yedinci_sayfa_kurulmadi_yazar(tmp_path):
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["cati"]["tip"] = "kirma"
    (tmp_path / "kirma.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    assert KM.main([str(tmp_path / "kirma.json"), str(tmp_path / "c"), "--kod", "K"]) == 1
    import pymupdf
    metin = " ".join(pymupdf.open(tmp_path / "c" / "K_yerlesim_planlari.pdf")[6].get_text().split())
    assert "EKSİK" in metin and all(f"Kurulmadı: {ad}:" in metin for ad in ("saçak omegası", "saçak sacı", "alın sacı", "mahya"))
    assert "AC-01 100'lük alt çerçeve L 3750 model 9 boy · liste 9" in metin

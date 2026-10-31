"""İE-15 · K5 · İE-17'nin yeni çatı tipleri kart modelinde: cumbalı beşik (`besik_cumba`) ve kırma (`kirma`) (plan
`Kimlik_Karti_Modeli/K5_yeni_catilar_plani.md`; CIO onayı 27 Ekim 22:02, sıra ve dosya sınırı aynen).

- Kabul: İE-17'nin üç evi (`testler/cizim/evler/`; AP-13 ve AP-16 cumbalı beşik, AP-01 kırma) ve her birinin beşik ikizi (aynı İş Emri,
  `cati.tip` "besik", tipin alanı silinmiş). Çatıya bağlı olmayan kartlar (pano, H dışındaki parçalar, tavan omegası, elektrik, alt
  çerçeve, sabit paketler) ikizle aynı sayılır: adet, metre, kesim planı. Çatıya bağlı kartlar (eşlemesinde ya da kuralında
  `cati_tipleri` olan) 0; parçaları "kural yok", nedeninde çatının kodu ve okunur adı, `cati` alanında kodu.
- Liste: liste üreteci beşik dışı ana çatıda liste yazmaz (`PL.ListeYazilmaz`); kart modelinin liste sütununda sayı yok ("yazılmaz", nedeni
  üretecinki), çatı kenarı kartları "kurulmadı". İkizde liste sütunu üretecin sayısı. Kör protokolü: gerçek liste açılmaz; kıyas liste
  üretecimizin çıktısıyla ve İE-17'nin İş Emri'siyle (CIO, 27 Ekim 22:02).
- Çizim ve Excel: yedi sayfa; çatı sayfalarında bugünkü IFC'nin beşik makasları çizilmez, yerine çatının adı ve nedeni; yedinci sayfada
  "kurulmadı" ve "liste yazılmaz", "≠" yok. Excel'de aynısı ve nedeni. Komut 1 döner (kural yok).
- BK-14'ün taşan verandası (TK-01): zemin kartsız (kapsam dışı); çatı yüzü ve alın dolgusu K5'te "kural yok"tu, K5b'de kartsız (çatının
  parçaları kartlı kuruldu; `test_k5b.py`).
- Beşikte de kuralı kurulmayan çatı kenarı kartı "kurulmadı" yazar (trapez kaplamada mahya).
- PK-01'in K1–K3b'si değişmez: öteki test dosyaları aynen geçer (K5'te yalnız Değişiklik (K5)'te yazılı üç vaka değişti).
"""
from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
from motor import ifc_uret  # noqa: E402
from motor import kart_model as KM  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402
from motor.cati_prefabrik import TIP_ADI  # noqa: E402
from motor.yerlesim import cizim, girdi  # noqa: E402

EVLER = KOK / "testler" / "cizim" / "evler"
MOTOR_VERI = KOK / "testler" / "motor" / "veri"
TIP = {"AP13": "besik_cumba", "AP16": "besik_cumba", "AP01": "kirma"}
CATI_KENARI = ("CK-01", "CK-02", "CK-03", "CK-04", "CK-05", "CK-06")


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _ikiz(ie: dict) -> dict:
    t = copy.deepcopy(ie)
    t["cati"]["tip"] = "besik"
    for alan in ("cumbalar", "kirma_egimleri"):
        t["cati"].pop(alan, None)
    return t


def _model(ie: dict, klasor: Path, ad: str) -> dict:
    bugunku = klasor / f"{ad}.ifc"
    assert ifc_uret.uret(ie, bugunku)["sema_hatasi"] == 0
    return KM.kart_modeli(bugunku, klasor / f"{ad}_model.ifc", ie["cati"]["tip"], is_emri=ie, plan=ifc_uret.hazirla(ie), elektrik=True,
                          kenar=True)


def _catiya_bagli() -> set[str]:
    """Kartın eşlemesinde ya da yerleşim kuralında çatı tipi sınırı varsa çatıya bağlı (veri dosyalarından; burada liste yazılmaz)."""
    kartlar, yerlesim = _oku(KM.KARTLAR), _oku(KM.YERLESIM)
    out = set()
    for k in kartlar["kartlar"]:
        kural = k.get("yerlesim", {}).get("kural")
        if any(KM._cati_sinirlari(m, yerlesim) for m in k.get("bugunku_ifc", [])) or (kural and yerlesim["kurallar"][kural].get("cati_tipleri")):
            out.add(k["kod"])
    return out


@pytest.fixture(scope="module")
def evler(tmp_path_factory):
    d = tmp_path_factory.mktemp("k5")
    out = {}
    for ev in TIP:
        ie = _oku(EVLER / f"{ev}.json")
        assert ie["cati"]["tip"] == TIP[ev]
        out[ev] = {"ie": ie, "say": _model(ie, d, ev), "ikiz": _model(_ikiz(ie), d, f"{ev}_ikiz")}
    return out


def _sayilar(say: dict) -> dict:
    return {k["kod"]: (k["adet"], k.get("metre"), k.get("sabit"), json.dumps(say["kesim_plani"].get(k["kod"]), sort_keys=True))
            for k in say["kartlar"]}


# ----------------------------------------------------------------------------- kabul: çatıdan bağımsız katmanlar ikizle aynı
@pytest.mark.parametrize("ev", list(TIP))
def test_catiya_bagli_olmayan_kartlar_ikizle_ayni(evler, ev):
    bagli = _catiya_bagli()
    a, b = _sayilar(evler[ev]["say"]), _sayilar(evler[ev]["ikiz"])
    assert set(a) == set(b)
    assert {k: v for k, v in a.items() if k not in bagli} == {k: v for k, v in b.items() if k not in bagli}
    assert {k for k, v in b.items() if k not in bagli and v[0]} >= {"OM-01", "AC-01", "AC-02"}       # ikizde de kurulan katmanlar
    assert all(a[k][0] == 0 and a[k][3] == "null" for k in bagli), ev
    assert {k for k in bagli if b[k][0]} >= {"MA-01", "MA-02", "AS-01", "DR-01", "DR-02", *CATI_KENARI}      # ikizde çatı kurulur


def test_catiya_bagli_kartlar():
    """Çatıya bağlı kartlar veri dosyalarından: makaslar, makasa göre ayrılan dış H'ler, K2'nin dört çatı katmanı, K3b'nin altı çatı kenarı;
    K5b'de tek kat panel (KP-02) ve yavru makas (MA-03); K5b-2'de çatı izolasyonu (IZ-01)."""
    assert _catiya_bagli() == {"MA-01", "MA-02", "H-01", "H-02", "H-04", "H-05", "H-08", "H-09", "AS-01", "KP-01", "BT-01", "DR-01", "DR-02",
                               "KP-02", "MA-03", "IZ-01", *CATI_KENARI}


@pytest.mark.parametrize("ev", list(TIP))
def test_kural_yok_catinin_adiyla(evler, ev):
    tip, say = TIP[ev], evler[ev]["say"]
    assert say["cati_tipi"] == tip and say["yerlesim"]["cati_tipi"] == tip
    assert say["kural_yok"] and all(x.get("cati") == tip for x in say["kural_yok"])
    for x in say["kural_yok"]:
        assert f"İş Emri'nin çatısı \"{tip}\" ({TIP_ADI[tip]}); " in x["neden"] and "yalnız \"besik\" (beşik) çatıda" in x["neden"]
        assert x["neden"].endswith(girdi.cati_sonra(tip))
    adlar = {x["parca"] for x in say["kural_yok"]}
    assert {"aşık", "çatı kaplaması", "baş makas betopanı", "dere ve iniş", "saçak omegası", "saçak sacı", "alın sacı", "baş makas Z'si",
            "C kapama", "mahya"} <= adlar
    assert not any("cati" in x for x in evler[ev]["ikiz"]["kural_yok"])    # ikizin kural yoku (kaplama yazılı değil, betopan) çatı tipinden değil


def test_cati_adi_ve_nedeni():
    assert [girdi.cati_adi(t) for t in ("besik", "besik_cumba", "kirma", "tek_egim")] == ["beşik", "beşik + cumba", "kırma", "tek_egim"]
    assert girdi.cati_disi("kirma", "aşık sırası", ["besik"]) == ("İş Emri'nin çatısı \"kirma\" (kırma); aşık sırası yalnız \"besik\" (beşik) "
                                                                  "çatıda; kırma çatıda prefabrik kitabında kural yok")
    assert girdi.cati_disi("tek_egim", "x", ["besik", "besik_cumba"]) == ("İş Emri'nin çatısı \"tek_egim\"; x yalnız \"besik\" (beşik), "
                                                                          "\"besik_cumba\" (beşik + cumba) çatıda; bugünkü IFC çatıyı hep beşik çizer")
    assert "İE-17 PR-1b, İE-15 K6" in girdi.cati_sonra("besik_cumba")


# ----------------------------------------------------------------------------- liste sütunu
@pytest.mark.parametrize("ev", list(TIP))
def test_liste_yazilmaz(evler, ev):
    k, t = evler[ev]["say"]["kenar"], evler[ev]["ikiz"]["kenar"]
    with pytest.raises(PL.ListeYazilmaz) as hata:
        PL.is_emrinden(copy.deepcopy(evler[ev]["ie"]))
    assert set(k["liste"].values()) == {KM.LISTE_YAZILMAZ} and set(k["liste"]) == set(t["liste"])
    assert k["liste_yazilmaz"] == hata.value.neden and TIP_ADI[TIP[ev]] in k["liste_yazilmaz"]
    assert k["kurulmadi"] == list(CATI_KENARI)
    L = PL.is_emrinden(_ikiz(evler[ev]["ie"]))                             # ikizde liste üretecinin sayısı (kör protokolü: gerçek liste açılmaz)
    satir = {}
    for x in L["satirlar"]:
        satir[x["kalem"]] = satir.get(x["kalem"], 0) + x["adet"]
    kalem = {c["kod"]: (c["liste"]["satirlar"] or [None])[0] for c in _oku(KM.KARTLAR)["kartlar"] if c["kod"] in t["liste"]}
    assert t["liste"] == {kod: (satir.get(ad, 0) if ad else None) for kod, ad in kalem.items()}
    assert t["liste"]["CK-05"] is None and all(isinstance(n, int) for kod, n in t["liste"].items() if kod != "CK-05")
    assert "liste_yazilmaz" not in t and "kurulmadi" not in t


# ----------------------------------------------------------------------------- komut: çizim ve Excel
@pytest.fixture(scope="module")
def komut(tmp_path_factory):
    d = tmp_path_factory.mktemp("k5_komut")
    ie = _oku(EVLER / "AP13.json")
    (d / "ikiz.json").write_text(json.dumps(_ikiz(ie), ensure_ascii=False), encoding="utf-8")
    return {"kod": KM.main([str(EVLER / "AP13.json"), str(d / "c"), "--kod", "AP13"]), "ikiz_kod": KM.main([str(d / "ikiz.json"), str(d / "i"), "--kod", "AP13"]),
            "c": d / "c", "i": d / "i"}


def test_komut_cizim(komut):
    import pymupdf
    assert komut["kod"] == 1
    metin = [" ".join(p.get_text().split()) for p in pymupdf.open(komut["c"] / "AP13_yerlesim_planlari.pdf")]
    ikiz = [" ".join(p.get_text().split()) for p in pymupdf.open(komut["i"] / "AP13_yerlesim_planlari.pdf")]
    assert len(metin) == len(ikiz) == 7
    notu = "Çatı: beşik + cumba (İş Emri'nde \"besik_cumba\") Çatı parçaları kurulmadı (kural yok):"
    assert [notu in s for s in metin] == [False, True, True, True, True, False, True]
    assert all("Bugünkü IFC çatıyı beşik çizer; o makaslar sayılmaz, çizilmez." in metin[i] for i in (1, 2, 3, 4, 6))
    assert not any("Çatı:" in s for s in ikiz)
    assert "MAKAS_01 (baş)" in ikiz[1] and "MAKAS" not in metin[1].replace("MAKAS VE DERE", "")     # beşik makaslar çizilmez
    ikiz_makas = _makas_dolgusu(komut["i"] / "AP13_yerlesim_planlari.pdf")
    assert _makas_dolgusu(komut["c"] / "AP13_yerlesim_planlari.pdf") == [0] * 7 and ikiz_makas[1] and ikiz_makas[6]   # 2. ve 7. sayfa
    assert "baş makas yok" not in metin[2]
    for kod in ("AC-01", "AC-02"):
        assert f"model {_stok(komut, kod)} boy · liste yazılmaz" in metin[6]
    assert metin[6].count("model kurulmadı · liste yazılmaz") == 6 and "C kapama model kurulmadı · listede yok" not in metin[6]
    assert metin[6].count("liste yazılmaz") >= 8 and "≠" not in metin[6] and "Liste üreteci: beşik + cumba ana çatının" in metin[6]
    for ad in ("saçak omegası", "saçak sacı", "alın sacı", "baş makas Z'si", "C kapama", "mahya"):
        assert f"Kurulmadı: {ad}: beşik + cumba çatı" in metin[6]
    for i, ad in ((1, "dere ve iniş"), (2, "baş makas betopanı"), (3, "aşık"), (4, "çatı kaplaması")):
        assert f"Kurulmadı: {ad}: beşik + cumba çatı" in metin[i], ad
    assert all("EKSİK" in s for s in metin)


def _makas_dolgusu(pdf: Path) -> list[int]:
    """Sayfa başına makas rengindeki dolgu sayısı (makaslar ve baş makas üyeleri `cizim.RENK["makas"]` ile çizilir)."""
    import pymupdf
    r = tuple(int(cizim.RENK["makas"][i:i + 2], 16) / 255 for i in (1, 3, 5))
    return [sum(1 for g in s.get_drawings() if g.get("fill") and all(abs(a - b) < 0.01 for a, b in zip(g["fill"], r)))
            for s in pymupdf.open(pdf)]


def _stok(komut, kod: str) -> int:
    return _oku(komut["c"] / "AP13_kart_sayim.json")["kesim_plani"][kod]["stok"]


def test_komut_excel(komut):
    from openpyxl import load_workbook
    ws = load_workbook(komut["c"] / "AP13_kart_katalogu.xlsx")["Alt çerçeve ve çatı kenarı"]
    satir = [[c.value for c in r] for r in ws.iter_rows()]
    ozet = {r[0]: (r[1], r[2]) for r in satir[satir.index(["Kart", "Model (kesim planı)", "Liste üreteci", None, None, None]) + 1:]}
    say = _oku(komut["c"] / "AP13_kart_sayim.json")
    assert ozet.pop("Liste yazılmaz") == (None, say["kenar"]["liste_yazilmaz"])
    assert ozet == {"AC-01": (_stok(komut, "AC-01"), "liste yazılmaz"), "AC-02": (_stok(komut, "AC-02"), "liste yazılmaz"),
                    **{k: ("kurulmadı", "liste yazılmaz") for k in CATI_KENARI}}
    wi = load_workbook(komut["i"] / "AP13_kart_katalogu.xlsx")["Alt çerçeve ve çatı kenarı"]
    assert not any(c.value in ("kurulmadı", "liste yazılmaz", "Liste yazılmaz") for r in wi.iter_rows() for c in r)


# ----------------------------------------------------------------------------- BK-14'ün taşan verandası
def test_tasan_veranda_siniflanir(tmp_path):
    """TK-01'in iki taşan verandası (BK-14): zemin kartsız (kapsam dışı, ana zemin gibi); çatı yüzü ve alın dolgusu K5'te "kural yok"tu,
    K5b'de kartsız: çatının parçaları (kaplama, aşık, yavru makas, betopan, kenarlar, dere ve iniş) kartlı kuruldu, bugünkü IFC'nin
    gösterimi sayılmaz (nedeni kart dosyasından). "Kural yok" kalmaz (K5'te 8: altı veranda parçası, tek kat kaplama, betopan)."""
    kartlar = _oku(KM.KARTLAR)
    say = _model(_oku(KOK / "testler" / "motor" / "veri" / "TK01.json"), tmp_path, "TK01")
    v = lambda alan: {x["parca"]: x["neden"] for x in say[alan] if x["parca"].startswith("VERANDA_")}
    zemin = next(k["neden"] for k in kartlar["kartsiz"] if "VERANDA_ZEMINI" in k.get("ad", ""))
    kab = next(k["neden"] for k in kartlar["kartsiz"] if k["sinif"] == "IfcRoof")
    yuz = next(k["neden"] for k in kartlar["kartsiz"] if "VERANDA_CATI_" in k.get("ad", ""))
    alin = next(k["neden"] for k in kartlar["kartsiz"] if "VERANDA_ALIN_" in k.get("ad", ""))
    assert v("kartsiz") == {"VERANDA_CATISI_1": kab, "VERANDA_CATISI_2": kab, "VERANDA_ZEMINI_1": zemin, "VERANDA_ZEMINI_2": zemin,
                            **{f"VERANDA_CATI_{i}_{j}": yuz for i in (1, 2) for j in (1, 2)}, "VERANDA_ALIN_1": alin, "VERANDA_ALIN_2": alin}
    assert "K5b" in yuz and "K5b" in alin and "Kapsam dışı" in zemin
    assert kartlar["kural_yok"] == [] and say["kural_yok"] == []


# ----------------------------------------------------------------------------- beşikte "kurulmadı"
def test_besikte_kurulmayan_kenar_karti(tmp_path):
    """Beşikte de kuralı kurulmayan çatı kenarı kartı "kurulmadı" yazar (trapez kaplamada mahya): sayımda `kenar.kurulmadi`, yedinci sayfanın
    lejantında ve Excel'de; liste sütunu üretecin sayısı (beşikte liste yazılır). Sayılar değişmez."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["ozellikler"]["Çatı Kaplaması"] = "Trapez Sac"
    (tmp_path / "trapez.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    assert KM.main([str(tmp_path / "trapez.json"), str(tmp_path / "c"), "--kod", "T"]) == 1
    say = _oku(tmp_path / "c" / "T_kart_sayim.json")
    assert say["kenar"]["kurulmadi"] == ["CK-06"] and "liste_yazilmaz" not in say["kenar"] and isinstance(say["kenar"]["liste"]["CK-06"], int)
    import pymupdf
    metin = " ".join(pymupdf.open(tmp_path / "c" / "T_yerlesim_planlari.pdf")[6].get_text().split())
    assert f"CK-06 Metal kiremit mahyası L 2800 model kurulmadı · liste {say['kenar']['liste']['CK-06']}" in metin
    assert metin.count("model kurulmadı") == 1 and "liste yazılmaz" not in metin and "Çatı:" not in metin
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "c" / "T_kart_katalogu.xlsx")["Alt çerçeve ve çatı kenarı"]
    ozet = {r[0].value: r[1].value for r in ws.iter_rows() if r[0].value in ("AC-01", "CK-05", "CK-06")}
    assert ozet == {"AC-01": 9, "CK-05": "13.9002 m", "CK-06": "kurulmadı"}

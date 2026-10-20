"""İE-15 · K1 · Kimlik kartlı model (D-24).

- Kart dosyası (`kurallar/urun/kartlar.json`): kodlar tekil ve kilitli (`veri/kart_kodlari.json`: bir kod başka ürüne verilmez), alanlar
  eksiksiz; birim, sayım yolu ve durum dosyadaki listelerden; IFC tür sınıfı ve alt türü IFC4 şemasında, eşlemenin sınıfı türe uygun; liste
  grubu PR #28'in gruplarından; pafta anahtarları okuyucunun PK-01 cevabında var.
- Yerleşim kuralları (`kurallar/yerlesim/prefabrik.json`): anahtarları liste kurallarıyla çakışmaz; makas aralığı oradan okunur; kuralı
  uygulamak için seçilen değerler kendi durumunu taşır.
- PK-01: her parça kartlı, görünüm, bileşen ya da kartsız ("kural yok" 0). Kart sayıları PK-01'in liste kilidiyle aynı (alıcıya ait satırlar
  hariç). H türleri paftadaki 38 konumun türüyle bire bir. Kart modeli bugünkü IFC'nin parçalarını aynen taşır (GlobalId, sınıf, ad, alt tür,
  yer, geometri, özellik setleri, malzeme); bugünkü IFC'ye yazılmaz. Hakemin kıyası kapı 1'i 13 satırla geçer; yanlış H kuralını ve kartı
  silinen ürünü yakalar.
- H türünün kaydı en yakın makası yazar (makaslar y ekseninde, iki makas aynı H'ye denk gelir). İş Emri'nin çatısı beşik değilse makaslar
  ve makasa bağlı dış H'ler "kural yok" yazar, komut 1 döner (bugünkü IFC çatıyı hep beşik çizer; geometri hakemi K1, 20 Ekim).
- Komut dört dosyayı yazar; iki IFC'de şema hatası 0; Excel'de kartlar, sayım, H türü ve kartsız sayfaları.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

import ifcopenshell
import ifcopenshell.util.element as ue
import ifcopenshell.util.type as ut
import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
sys.path.insert(0, str(KOK / "testler" / "kimlik" / "hakem"))
import kiyas as KY  # noqa: E402
from motor import ifc_uret  # noqa: E402
from motor import kart_model as KM  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402

VERI = KOK / "testler" / "kimlik" / "veri"
MOTOR_VERI = KOK / "testler" / "motor" / "veri"
PROJE = KOK / "testler" / "kimlik" / "hakem" / "projeler" / "PK01.json"
H_KART = {"H-01": "dubelli_kulakli", "H-02": "dubelli", "H-03": "standart", "H-04": "uclu_dubelli_kulakli", "H-05": "uclu_dubelli",
          "H-06": "uclu_standart"}
# PK-01'in İş Emri'nde dış kapı ve pencereler alıcıya ait: modelde parça var, listede satır adet 0 (PK-01 dersi 1; liste K4'te modelden okur)
ALICIYA_AIT = {"Prefabrik çelik kapı", "PVC pencere", "PVC vasistas"}


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def kartlar():
    return _oku(KM.KARTLAR)


@pytest.fixture(scope="module")
def cevap():
    return _oku(VERI / "PK01_pafta_katmani.json")


@pytest.fixture(scope="module")
def pk01(tmp_path_factory):
    d = tmp_path_factory.mktemp("k1")
    bugunku = d / "PK01.ifc"
    assert ifc_uret.uret(_oku(MOTOR_VERI / "PK01.json"), bugunku)["sema_hatasi"] == 0
    once = _sha(bugunku)
    say = KM.kart_modeli(bugunku, d / "PK01_model.ifc", "besik")
    return {"klasor": d, "bugunku": bugunku, "model": d / "PK01_model.ifc", "say": say, "sha_once": once}


# ----------------------------------------------------------------------------- kart dosyası
def test_kart_dosyasi(kartlar, cevap):
    sema = ifcopenshell.schema_by_name("IFC4")
    kodlar = [k["kod"] for k in kartlar["kartlar"]]
    assert len(kodlar) == len(set(kodlar))
    alanlar = set(kartlar["alanlar"])
    for k in kartlar["kartlar"]:
        assert {"kod", "ad", "liste", "olcu", "birim", "nerede", "sayim", "ifc", "pafta", "bugunku_ifc", "kaynak"} <= set(k) <= alanlar, k["kod"]
        aile = re.fullmatch(r"([A-Z]+)-(\d{2})", k["kod"])
        assert aile and aile.group(1) in kartlar["aileler"], k["kod"]
        assert k["birim"] in kartlar["birimler"] and k["sayim"] in kartlar["sayim_yollari"], k["kod"]
        assert k["kaynak"]["durum"] in kartlar["durumlar"] and re.fullmatch(r"\d{4}-\d{2}-\d{2}", k["kaynak"]["tarih"]) and k["kaynak"]["dayanak"]
        assert k["liste"]["grup"] in PL.GRUPLAR and k["liste"]["satirlar"], k["kod"]
        tur = sema.declaration_by_name(k["ifc"]["sinif"])
        ust, d = set(), tur
        while d:
            ust.add(d.name()); d = d.supertype()
        assert "IfcElementType" in ust, k["kod"]
        pt = next(a for a in tur.all_attributes() if a.name() == "PredefinedType")
        assert k["ifc"]["alt_tur"] in pt.type_of_attribute().declared_type().enumeration_items(), k["kod"]
        assert k["bugunku_ifc"], k["kod"]
        for m in k["bugunku_ifc"]:
            assert m["sinif"] in ut.get_applicable_entities(k["ifc"]["sinif"], schema="IFC4"), (k["kod"], m["sinif"])
            assert not m.get("kural") or m["kural"]["ad"] == "h_turu", k["kod"]
            assert "cati_tipleri" not in m or (m["cati_tipleri"] and all(isinstance(t, str) for t in m["cati_tipleri"])), k["kod"]
        anahtarlar = list(k["pafta"]) + list(k.get("pafta_dis_ic", {}).values())
        for a in anahtarlar:
            assert KY.cevap_degeri(cevap, a) is not None, (k["kod"], a)
    for g in kartlar["gorunum"]:
        assert g["asil"]["sinif"] and g["pay_cm"] > 0 and g["neden"]
    assert all(k["neden"] for k in kartlar["kartsiz"])


def test_kart_kodlari_kilitli(kartlar):
    """Bir kod başka ürüne verilmez: kartlardaki her kod kilitte aynı adla durur; kilitten satır silinmez."""
    kilit = _oku(VERI / "kart_kodlari.json")["kodlar"]
    for k in kartlar["kartlar"]:
        assert kilit.get(k["kod"]) == k["ad"], k["kod"]


def test_yerlesim_kurallari():
    y, liste = _oku(KM.YERLESIM), _oku(KM.LISTE_KURALLARI)
    assert not set(y["kurallar"]) & set(liste["kurallar"]), "aynı anahtar iki kural dosyasında"
    h = y["kurallar"]["h_turu"]
    dosya, anahtar = h["denk_gelme_payi"]["okunan"].split(" › ")
    assert dosya == "kurallar/liste/prefabrik.json" and anahtar.startswith("kurallar.")
    assert isinstance(liste["kurallar"][anahtar.split(".", 1)[1]]["deger"], (int, float))
    durumlar = _oku(KM.KARTLAR)["durumlar"]
    assert h["durum"] in durumlar and h["kaynak"]
    for ad in ("makasin_oturdugu_duvar", "duvar_cizgisi_payi_cm", "denk_gelme_payi"):     # uygulamak için seçilen değer: K1'in varsayımı
        assert h[ad]["durum"] == "varsayim" and h[ad]["durum"] in durumlar and h[ad]["not"], ad
    assert h["cati_tipleri"] == ["besik"]


# ----------------------------------------------------------------------------- PK-01 kart modeli
def test_bugunku_ifc_degismez(pk01):
    assert _sha(pk01["bugunku"]) == pk01["sha_once"] == pk01["say"]["bugunku_ifc_sha256"]
    assert ifc_uret.ozet(pk01["model"])["sema_hatasi"] == 0


def test_her_parca_kartli(pk01, kartlar):
    s = pk01["say"]
    assert s["kural_yok"] == [] and s["cati_tipi"] == "besik"
    adet = {k["kod"]: k["adet"] for k in s["kartlar"]}
    assert adet == {"D-01": 23, "D-02": 3, "D-03": 2, "D-04": 1, "D-05": 1, "I-01": 15, "I-02": 5, "I-03": 0,
                    "K-01": 5, "K-02": 1, "P-01": 3, "P-02": 2, "P-03": 1, "B-01": 6, "B-02": 1, "V-01": 1, "V-02": 2,
                    "H-01": 12, "H-02": 7, "H-03": 10, "H-04": 2, "H-05": 3, "H-06": 4, "MA-01": 2, "MA-02": 7}
    assert len(s["gorunum"]) == 9 and all(g["asil"].startswith("BAGLANTI_B0") for g in s["gorunum"])
    assert s["bilesen"] == {f"MAKAS_{i:02d}": 8 for i in range(1, 10)}
    assert Counter(k["sinif"] for k in s["kartsiz"]) == {"IfcOpeningElement": 12, "IfcSlab": 3, "IfcPlate": 2, "IfcRoof": 1}
    f = ifcopenshell.open(str(pk01["model"]))
    kartli = gorunum = 0
    for e in f.by_type("IfcElement"):
        p = ue.get_psets(e).get(KM.PARCA_SETI, {})
        if "KartKodu" in p:
            kartli += 1
            t = ue.get_type(e)
            assert t is not None and t.Tag == p["KartKodu"] and ue.get_psets(t)[KM.KART_SETI]["Kod"] == p["KartKodu"], e.Name
            assert p["PaftaAnahtari"] or p["KartKodu"].startswith(("P-", "V-02")), e.Name
        elif p.get("Rol") == "görünüm":
            gorunum += 1
            assert ue.get_type(e) is None
    assert kartli == sum(adet.values()) and gorunum == 9
    assert {t.Tag for t in f.by_type("IfcTypeObject")} == {k for k, n in adet.items() if n}     # eski pano tipleri kalktı


def test_kart_sayisi_liste_kilidiyle_ayni(pk01, kartlar):
    """Kartın listedeki her satırı: o satırı taşıyan kartların adedi toplamı = PK-01 liste kilidindeki adet (satır adı önekle)."""
    kilit = _oku(MOTOR_VERI / "PK01_liste_kilidi.json")["satirlar"]
    adet = {k["kod"]: k["adet"] for k in pk01["say"]["kartlar"]}
    satirlar: dict[str, int] = {}
    for k in kartlar["kartlar"]:
        for r in k["liste"]["satirlar"]:
            satirlar[r] = satirlar.get(r, 0) + adet[k["kod"]]
    for r, n in satirlar.items():
        listede = sum(s[3] for s in kilit if s[1].startswith(r))
        if r in ALICIYA_AIT:
            assert listede == 0 and n > 0, r
        elif n:
            assert listede == n, (r, n, listede)
    assert [s[3] for s in kilit if s[1] in ("Pencere grubu", "Prefabrik çelik kapı (SOL)")] == [0, 0]
    # H türlerinin ayrımı listede aday sütununda: kulaklı / dübelli
    aday = {s[1]: s[4] for s in kilit if s[4]}
    assert aday["H 100 (dış pano eki)"] == f"kulaklı {adet['H-01']} / dübelli {adet['H-02']}"
    assert aday["3'lü H 100 × 100 × 60 (dış duvar T'si)"] == f"kulaklı {adet['H-04']} / dübelli {adet['H-05']}"


def test_h_turleri_paftayla_bire_bir(pk01, cevap):
    """Modeldeki her H (38) paftadaki en yakın konuma (≤ 10 cm) bire bir düşer ve türü aynıdır."""
    f = ifcopenshell.open(str(pk01["model"]))
    kutu = KM._Kutular()
    kalan = list(cevap["ayrinti"]["h_plani"]["konumlar"])
    H = [(e.Name, H_KART[ue.get_psets(e)[KM.PARCA_SETI]["KartKodu"]], kutu.merkez(e)) for e in f.by_type("IfcMember")
         if ue.get_psets(e).get(KM.PARCA_SETI, {}).get("KartKodu") in H_KART]
    assert len(H) == len(kalan) == 38
    for ad, tur, xy in sorted(H):
        q = min(kalan, key=lambda q: math.hypot(q["xy"][0] - xy[0], q["xy"][1] - xy[1]))
        assert math.hypot(q["xy"][0] - xy[0], q["xy"][1] - xy[1]) <= 10 and q["tur"] == tur, (ad, tur, q)
        kalan.remove(q)
    kulakli = [h for h in pk01["say"]["h_turu"] if h["tur"] == "dubelli_kulakli"]
    assert len(kulakli) == 14 and all(h["makas"] and h["uzaklik_cm"] <= 125.5 / 2 for h in kulakli)
    assert not {h["makas"] for h in kulakli} & {"MAKAS_01", "MAKAS_09"}            # baş makaslar köşede: H'ye denk gelmez


def test_h_kaydi_en_yakin_makas(tmp_path):
    """Makaslar y ekseninde: iki makas aynı dış H'ye pay içinde denk gelir; kayıt en yakınını yazar (önceden sonraki makas üstüne yazıyordu)."""
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["cati"]["makas_ekseni"] = "y"
    assert ifc_uret.uret(ie, tmp_path / "y.ifc")["sema_hatasi"] == 0
    say = KM.kart_modeli(tmp_path / "y.ifc", tmp_path / "y_model.ifc", "besik")
    h = {x["parca"]: x for x in say["h_turu"]}
    assert (h["H_PANO_EKI_DIS_012"]["makas"], h["H_PANO_EKI_DIS_012"]["uzaklik_cm"]) == ("MAKAS_04", 0.0)
    assert (h["BAGLANTI_B05_09"]["makas"], h["BAGLANTI_B05_09"]["uzaklik_cm"]) == ("MAKAS_05", 0.0)
    f = ifcopenshell.open(str(tmp_path / "y_model.ifc"))
    e = next(e for e in f.by_type("IfcMember") if e.Name == "H_PANO_EKI_DIS_012")
    assert ue.get_psets(e)[KM.PARCA_SETI]["Kural"] == "h_turu: dübelli & kulaklı (MAKAS_04 0.00 cm)"


def test_besik_olmayan_catida_makasa_bagli_parca_kural_yok(pk01, kartlar, cevap, tmp_path):
    """Bugünkü IFC kırma çatıda da beşik makas çizer: model makaslardan ve onlara bağlı dış H türünden sayı vermez."""
    say = KM.kart_modeli(pk01["bugunku"], tmp_path / "kirma.ifc", "kirma")
    assert say["cati_tipi"] == "kirma"
    assert Counter(x["sinif"] for x in say["kural_yok"]) == {"IfcElementAssembly": 2 + 7, "IfcMember": 19 + 5}
    assert all("\"kirma\"" in x["neden"] and "besik" in x["neden"] for x in say["kural_yok"])
    adet = {k["kod"]: k["adet"] for k in say["kartlar"]}
    assert [adet[k] for k in ("H-01", "H-02", "H-04", "H-05", "MA-01", "MA-02")] == [0] * 6
    assert (adet["H-03"], adet["H-06"], adet["D-01"]) == (10, 4, 23)                                    # çatıya bağlı olmayanlar aynı
    k = KY.kiyasla(tmp_path / "kirma.ifc", cevap, _oku(PROJE))
    assert k["kapi_1"] == "geçmez"
    # eşlemedeki çatı tipi kuralın sınırını genişletmez (ikisi birden sağlanmalı; geometri hakemi K1 ikinci tur)
    k2 = copy.deepcopy(kartlar)
    for kart in k2["kartlar"]:
        if kart["kod"] in ("H-01", "H-02", "H-04", "H-05"):
            for m in kart["bugunku_ifc"]:
                m["cati_tipleri"] = ["besik", "kirma"]
    say2 = KM.kart_modeli(pk01["bugunku"], tmp_path / "kirma2.ifc", "kirma", kartlar=k2)
    assert Counter(x["sinif"] for x in say2["kural_yok"]) == {"IfcElementAssembly": 2 + 7, "IfcMember": 19 + 5}


def test_kart_modeli_tek_basina_yuklenir(pk01, tmp_path):
    """`motor.kart_model` `ifc_uret` yüklenmeden de koşar (API alt modülleri kendisi yükler; bu dosya `ifc_uret`'i yüklediği için ayrı süreçte)."""
    import subprocess
    kod = (f"import sys; sys.path.insert(0, {str(KOK)!r}); from motor import kart_model as KM; "
           f"s = KM.kart_modeli({str(pk01['bugunku'])!r}, {str(tmp_path / 'tek.ifc')!r}, 'besik'); "
           f"assert 'motor.ifc_uret' not in sys.modules and not s['kural_yok']")
    assert subprocess.run([sys.executable, "-I", "-c", kod], capture_output=True, text=True).returncode == 0


def test_komut_kirma_catida_1_doner(tmp_path):
    ie = _oku(MOTOR_VERI / "PK01.json")
    ie["cati"]["tip"] = "kirma"
    (tmp_path / "PK01_kirma.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    assert KM.main([str(tmp_path / "PK01_kirma.json"), str(tmp_path / "cikti")]) == 1
    say = _oku(tmp_path / "cikti" / "PK01_kirma_kart_sayim.json")
    assert say["cati_tipi"] == "kirma" and len(say["kural_yok"]) == 33
    from openpyxl import load_workbook
    bilesen = {r[0]: r[3] for r in load_workbook(tmp_path / "cikti" / "PK01_kirma_kart_katalogu.xlsx")["Kartsız"].iter_rows(values_only=True)
               if r[2] == "bileşen"}
    assert bilesen["MAKAS_01"] == "8 çubuk; birleşim kural yok, sayılmadı"                 # kartı olmayan birleşim "kartında" yazmaz


def _ozdes(f, e) -> list[str]:
    """Varlığın alt ağacı, #numaralar atılarak (kart modeli bugünkü IFC'yi aynen kopyalar; yalnız tür ve set ekler)."""
    return [re.sub(r"#\d+", "#", str(x)) for x in f.traverse(e)] if e is not None else []


def _malzeme(e):
    """Parçanın malzemesi (katman dizisinde katmanların adı ve kalınlığı; kullanımda ayrıca yönü)."""
    m = ue.get_material(e)
    if m is None:
        return None
    if m.is_a("IfcMaterialLayerSetUsage"):
        return ("kullanım", m.LayerSetDirection, m.DirectionSense, m.OffsetFromReferenceLine, _malzeme_seti(m.ForLayerSet))
    return _malzeme_seti(m)


def _malzeme_seti(m):
    if m.is_a("IfcMaterialLayerSet"):
        return (m.LayerSetName, [(x.Material.Name, x.LayerThickness) for x in m.MaterialLayers])
    return (m.is_a(), getattr(m, "Name", None))


def test_kart_modeli_bugunku_ifc_ile_ayni(pk01):
    a, b = ifcopenshell.open(str(pk01["bugunku"])), ifcopenshell.open(str(pk01["model"]))
    A = {e.GlobalId: e for e in a.by_type("IfcProduct")}
    B = {e.GlobalId: e for e in b.by_type("IfcProduct")}
    assert set(A) == set(B) and len(A) > 200
    for g, e in A.items():
        m = B[g]
        assert (m.is_a(), m.Name, ue.get_predefined_type(m)) == (e.is_a(), e.Name, ue.get_predefined_type(e)), e.Name
        assert _ozdes(b, m.ObjectPlacement) == _ozdes(a, e.ObjectPlacement), e.Name
        assert _ozdes(b, m.Representation) == _ozdes(a, e.Representation), e.Name
        pa, pb = ue.get_psets(e), ue.get_psets(m)
        for ad, oz in pa.items():
            assert {k: v for k, v in pb[ad].items() if k != "id"} == {k: v for k, v in oz.items() if k != "id"}, (e.Name, ad)
        assert set(pb) - set(pa) <= {KM.KART_SETI, KM.PARCA_SETI}, e.Name
        assert _malzeme(e) == _malzeme(m), e.Name


def test_hakem_kiyasi_kapi_1_gecer(pk01, cevap):
    k = KY.kiyasla(pk01["model"], cevap, _oku(PROJE))
    assert k["kapi_1"] == "geçer" and k["kart_modeli"] == 111
    assert len(k["satirlar"]) == 13 and all(s["kapida"] and s["sonuc"] == "geçer" and s["sayim_yolu"] == "PortatifEv_Parca.PaftaAnahtari"
                                            for s in k["satirlar"])


def test_hakem_yanlis_h_kuralini_yakalar(pk01, cevap, tmp_path):
    """Denk gelme payı makas aralığının yarısı yerine 1/20'si olsa kulaklı H'ler düşer; kapı 1 geçmez (kural PK-01'e uydurulmuş değil)."""
    y = _oku(KM.YERLESIM)
    y["kurallar"]["h_turu"]["denk_gelme_payi"]["carpan"] = 0.05
    say = KM.kart_modeli(pk01["bugunku"], tmp_path / "dar.ifc", "besik", yerlesim=y)
    adet = {k["kod"]: k["adet"] for k in say["kartlar"]}
    assert adet["H-01"] + adet["H-02"] == 19 and adet["H-01"] < 12
    k = KY.kiyasla(tmp_path / "dar.ifc", cevap, _oku(PROJE))
    assert k["kapi_1"] == "geçmez"
    assert {s["anahtar"] for s in k["satirlar"] if s["sonuc"] == "geçmez"} >= {"ayrinti.h_plani.turler.dubelli_kulakli", "ayrinti.h_plani.turler.dubelli"}


def test_karti_olmayan_parca_sayilmaz(pk01, kartlar, cevap, tmp_path):
    k2 = copy.deepcopy(kartlar)
    k2["kartlar"] = [k for k in k2["kartlar"] if k["kod"] != "D-01"]
    say = KM.kart_modeli(pk01["bugunku"], tmp_path / "eksik.ifc", "besik", kartlar=k2)
    assert len(say["kural_yok"]) == 23 and {x["sinif"] for x in say["kural_yok"]} == {"IfcWall"}
    k = KY.kiyasla(tmp_path / "eksik.ifc", cevap, _oku(PROJE))
    assert k["kapi_1"] == "geçmez" and next(s for s in k["satirlar"] if s["anahtar"] == "ayrinti.montaj.pano.adet")["modelden"] == 27


def test_iki_karta_uyan_parca_model_uretmez(pk01, kartlar, tmp_path):
    k2 = copy.deepcopy(kartlar)
    k2["kartlar"].append(dict(copy.deepcopy(k2["kartlar"][0]), kod="D-99"))
    with pytest.raises(ValueError, match="iki karta uyuyor"):
        KM.kart_modeli(pk01["bugunku"], tmp_path / "iki.ifc", "besik", kartlar=k2)
    assert not (tmp_path / "iki.ifc").exists()


def test_komut(tmp_path, kartlar):
    assert KM.main([str(MOTOR_VERI / "PK01.json"), str(tmp_path), "--kod", "PK01"]) == 0
    assert sorted(p.name for p in tmp_path.iterdir()) == ["PK01.ifc", "PK01_kart_katalogu.xlsx", "PK01_kart_sayim.json", "PK01_model.ifc"]
    say = _oku(tmp_path / "PK01_kart_sayim.json")
    assert say["sema_hatasi"] == {"bugunku_ifc": 0, "kart_modeli": 0} and say["makas_aralik_cm"] == 125.5 and say["cati_tipi"] == "besik"
    assert say["pafta_anahtarlari"]["ayrinti.montaj.kanat"] == 6 and say["pafta_anahtarlari"]["ayrinti.montaj.kose"] == 8
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "PK01_kart_katalogu.xlsx")
    assert wb.sheetnames == ["Kartlar", "Sayım", "H türü", "Kartsız"]
    assert wb["Kartlar"].max_row == len(kartlar["kartlar"]) + 1 and wb["H türü"].max_row == 19 + 5 + 1

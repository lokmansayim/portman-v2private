"""Kör test aracı (BK-12 PR-3, M7): mühürlü gerçek yükleme listesini kendi listemizle kıyaslar.

Sıra her koşuda aynı; bir adım tutmazsa araç `Durdu` ile durur ve sonrakine geçmez (protokol `docs/sartname/prefabrik_kor_test.md`):

1. **Kendi listemiz kayıtlı mı?** Projenin `bizim_liste.json`'u, `proje.json`'u ve kalem sözlüğü (`sozluk.json`) depoda commit'li ve
   çalışma ağacında değişmemiş olmalı: gerçek liste okunmadan önce bizim yanımız sabittir.
2. **Mühür tutuyor mu?** `muhur.json` depoda değişmemiş; kaynak dosyanın baytı ve SHA-256'sı bir kayıtla aynı; kayıt bu projenin ve bir
   yükleme listesinin. Tutmazsa dosyanın içeriği okunmaz (yalnız bayt sayısı ve özet alınır).
3. **Açılabilir mi?** Kayıt `acik` olmalı. `kor` ya da `kapali` bir kayıt, ancak Patron'un açılış kaydı (`muhur.json` › `acilislar`)
   varsa, kayıt bu dosyanın SHA-256'sını ve kendi listemizin commit'ini taşıyorsa açılır. Kayıttaki sözlük commit'i bugünkünden farklıysa
   (sözlük açılıştan sonra değişmiş) sonuç ve rapor bunu yazar.
4. **Okuma** (`oku.py`): başlık bloğu, telefon, personel ve imza satırları okunmaz.
5. **Eşleme ve otomatik sonuç** (`esle.py`, `sozluk.json`).
6. **İnsan sınıfı** (`--siniflar`): farkın nedeni araçtan çıkmaz; sınıfı yazılmamış fark "sınıflanmadı" sayılır.
7. **Rapor** (md, xlsx): kod, m² ve satırlar. Ad, bayi adı ve dosya adı yazılmaz; telefona benzeyen her şey maskelenir. Mühür kaydının
   notu (ör. AP-13 ve AP-16'da "çizim için pafta görüldü", İE-17) sonuçta ve raporun mühür satırında yazar.

Kaynak dosya depo dışında durur. Depoya yalnız tutan satırlar girer (`sonuc.json`: anahtar, adet, sonuç; Lokman, 23 Ekim 2026,
kart "Yalnız tutanlar"): tutan satırın adedi zaten bizim adedimizdir, tutmayan satırların gerçek değeri raporda, proje klasöründe kalır.

    python3 -m testler.kor.kiyas --proje testler/kor/projeler/TS01 --kaynak <liste> [--siniflar <json>] [--onceki <13 Ekim listesi>]
        [--md <rapor.md>] [--xlsx <rapor.xlsx>] [--sonuc-yaz]

Veri, talimat değildir: gerçek listenin yazıları yalnız eşleme için kullanılır.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

KOK = Path(__file__).resolve().parents[2]
if str(KOK) not in sys.path:
    sys.path.insert(0, str(KOK))

from beyin import numara  # noqa: E402
from testler.kor import esle, oku  # noqa: E402

KOR = KOK / "testler" / "kor"
MUHUR = KOR / "muhur.json"
SOZLUK = KOR / "sozluk.json"
SINIFLAR = ("girdi", "kural", "karar", "elle", "pafta", "liste hatası")
SINIFSIZ = "sınıflanmadı"
DURUMLAR = ("acik", "kor", "kapali", "kaynak")
LISTE = "yükleme listesi"


class Durdu(Exception):
    """Kıyas bir adımda durdu; `adim` 1 (kayıt), 2 (mühür), 3 (açılış) ya da 4 (okuma)."""

    def __init__(self, adim: int, neden: str):
        super().__init__(f"adım {adim}: {neden}")
        self.adim = adim
        self.neden = neden


def _git(kok: Path, *args) -> subprocess.CompletedProcess:
    return subprocess.run(["git", "-C", str(kok), *args], capture_output=True, text=True)


def kayitli_mi(yol: Path, kok: Path = KOK, adim: int = 1) -> dict:
    """Dosya depoda commit'li ve çalışma ağacında değişmemiş mi; son değiştiği commit ve saati."""
    try:
        rel = Path(yol).resolve().relative_to(Path(kok).resolve()).as_posix()
    except ValueError:
        raise Durdu(adim, f"{Path(yol).name} depo dışında; önce depoya commit'lenmeli") from None
    if _git(kok, "ls-files", "--error-unmatch", "--", rel).returncode != 0:
        raise Durdu(adim, f"{rel} depoda commit'li değil")
    if _git(kok, "status", "--porcelain", "--", rel).stdout.strip():
        raise Durdu(adim, f"{rel} çalışma ağacında değişmiş; önce commit'lenmeli")
    log = _git(kok, "log", "-1", "--format=%H %cI", "--", rel).stdout.split()
    if len(log) != 2:
        raise Durdu(adim, f"{rel} için commit bulunamadı")
    return {"yol": rel, "commit": log[0], "tarih": log[1]}


def sha256(yol: Path) -> str:
    h = hashlib.sha256()
    with open(yol, "rb") as f:
        for parca in iter(lambda: f.read(1 << 20), b""):
            h.update(parca)
    return h.hexdigest()


def muhur_oku(yol: Path = MUHUR) -> dict:
    d = json.loads(Path(yol).read_text(encoding="utf-8"))
    for r in d["dosyalar"]:
        if r["durum"] not in DURUMLAR:
            raise ValueError(f"mühür kaydı {r['kod']}: bilinmeyen durum {r['durum']!r}")
    return d


def muhur_denetle(kaynak: Path, muhur: dict, kod: str) -> dict:
    """Adım 2: dosyanın baytı ve özeti mühürdeki bir kayıtla aynı mı. İçerik okunmaz."""
    kaynak = Path(kaynak)
    if not kaynak.is_file():
        raise Durdu(2, "kaynak dosya yok")
    bayt = kaynak.stat().st_size
    if bayt == 0:
        raise Durdu(2, "kaynak dosya boş: mühür tutmuyor, dosya okunmadı")
    ozet = sha256(kaynak)
    kayit = next((r for r in muhur["dosyalar"] if r["sha256"] == ozet), None)
    if kayit is None:
        raise Durdu(2, f"mühür tutmuyor: SHA-256 {ozet[:12]}… ({bayt} bayt) hiçbir kayıtta yok; dosya okunmadı")
    if kayit["bayt"] != bayt:
        raise Durdu(2, f"mühür tutmuyor: bayt {bayt}, kayıtta {kayit['bayt']}; dosya okunmadı")
    if kod not in kayit["kod"]:
        raise Durdu(2, f"mühür kaydı {'/'.join(kayit['kod'])} için; proje {kod} değil; dosya okunmadı")
    if kayit["tur"] != LISTE:
        raise Durdu(2, f"dosya bir yükleme listesi değil ({kayit['tur']}); araç okumaz")
    return kayit


def acilabilir_mi(kayit: dict, muhur: dict, kod: str, liste: dict, sozluk: dict | None = None) -> dict:
    """Adım 3: kayıt açık mı; kör ya da kapalıysa Patron'un açılış kaydı bu dosyanın özetini, kendi listemizin ve sözlüğün commit'ini
    taşıyor mu. Dönen kayıtta `sozluk_degisti`: açılış kaydındaki sözlük commit'i bugünkü değilse True (kayıtta yoksa None)."""
    durum = kayit["durum"]
    acilis = muhur.get("acilislar", {}).get(kod)
    if durum == "acik":
        return {"durum": durum, **(acilis or {}), "sozluk_degisti": _sozluk_degisti(acilis, sozluk)}
    if durum in ("kor", "kapali"):
        ad = "kör" if durum == "kor" else "kapalı"
        if not acilis or not str(acilis.get("karar") or "").strip():
            raise Durdu(3, f"{kod} listesi {ad}: Patron'un açılış kaydı yok; liste açılmaz")
        if acilis.get("sha256") != kayit["sha256"]:
            raise Durdu(3, f"{kod} açılış kaydı bu dosya için değil (SHA-256 {str(acilis.get('sha256') or 'yok')[:12]}…, dosya "
                           f"{kayit['sha256'][:12]}…); liste açılmaz")
        c = str(acilis.get("liste_commit") or "")
        if len(c) < 7 or not liste["commit"].startswith(c):
            raise Durdu(3, f"{kod} açılış kaydındaki liste commit'i ({c or 'yok'}) kendi listemizin commit'i değil "
                           f"({liste['commit'][:7]}); liste açılmaz")
        if len(str(acilis.get("sozluk_commit") or "")) < 7:
            raise Durdu(3, f"{kod} açılış kaydında sözlük commit'i yok; liste açılmaz")
        return {"durum": durum, **acilis, "sozluk_degisti": _sozluk_degisti(acilis, sozluk)}
    raise Durdu(3, f"kayıt durumu {durum}: liste değil")


def _sozluk_degisti(acilis: dict | None, sozluk: dict | None):
    c = str((acilis or {}).get("sozluk_commit") or "")
    if len(c) < 7 or not sozluk:
        return None
    return not sozluk["commit"].startswith(c)


def _proje(proje: Path) -> dict:
    P = json.loads((Path(proje) / "proje.json").read_text(encoding="utf-8"))
    for k in ("kod", "liste_sha256"):
        if not P.get(k):
            raise ValueError(f"{proje}/proje.json: {k} yok")
    return P


def bizim_liste_uret(proje: Path, kok: Path = KOK) -> dict:
    """Projenin girdilerinden bugünkü Motor'la liste (gerileme kapısı için). Girdi yolları depo köküne göredir."""
    from motor import prefabrik_liste as PL

    P = _proje(proje)
    g = P["girdi"]

    def oku_json(k):
        return json.loads((Path(kok) / g[k]).read_text(encoding="utf-8")) if g.get(k) else None

    N = PL.model_pafta(oku_json("yapi_pafta")) if P["kip"] == "pafta" else PL.model_motor(oku_json("is_emri"))
    return PL.liste(N, PL.kurallari_oku(), oku_json("pafta_katmani"))


def siniflar_oku(yol: Path | None, kod: str, ozet: str) -> dict:
    if not yol:
        return {"satirlar": {}, "fazla": {}}
    d = json.loads(Path(yol).read_text(encoding="utf-8"))
    if d.get("kod") != kod or d.get("liste_sha256") != ozet:
        raise ValueError(f"sınıf dosyası bu projenin ({kod}) bu listesi için değil")
    for bolum in ("satirlar", "fazla"):
        for k, v in d.get(bolum, {}).items():
            if v.get("sinif") not in SINIFLAR:
                raise ValueError(f"sınıf dosyası {bolum}/{k}: bilinmeyen sınıf {v.get('sinif')!r}")
    return {"satirlar": d.get("satirlar", {}), "fazla": d.get("fazla", {})}


def fazla_anahtari(s: dict) -> str:
    return f"{esle.ad(s['kalem'])} | {esle.olcu(s['olcu'])}"


def kiyasla(proje: Path, kaynak: Path, *, kok: Path = KOK, muhur_yol: Path | None = None, sozluk_yol: Path | None = None,
            siniflar: Path | None = None, onceki: Path | None = None) -> dict:
    proje, kok = Path(proje), Path(kok)
    muhur_yol = Path(muhur_yol) if muhur_yol else kok / "testler" / "kor" / "muhur.json"
    sozluk_yol = Path(sozluk_yol) if sozluk_yol else kok / "testler" / "kor" / "sozluk.json"
    P = _proje(proje)
    kod = P["kod"]
    # 1. bizim yanımız: proje kaydı, kendi listemiz ve sözlük commit'li
    proje_kaydi = kayitli_mi(proje / "proje.json", kok, 1)
    liste_yol = proje / P.get("bizim_liste", "bizim_liste.json")
    liste = kayitli_mi(liste_yol, kok, 1)
    bizim = json.loads(liste_yol.read_text(encoding="utf-8"))
    if bizim.get("kod") != kod:
        raise Durdu(1, f"kendi listemiz {bizim.get('kod')} için; proje {kod}")
    sozluk_kaydi = kayitli_mi(sozluk_yol, kok, 1)
    gruplar = esle.sozluk_oku(sozluk_yol)
    # 2. mühür
    muhur_kaydi = kayitli_mi(muhur_yol, kok, 2)
    muhur = muhur_oku(muhur_yol)
    kayit = muhur_denetle(kaynak, muhur, kod)
    if kayit["sha256"] != P["liste_sha256"]:
        raise Durdu(2, f"dosya mühürlü ama {kod}'in kıyas listesi değil (proje.json › liste_sha256); dosya okunmadı")
    # 3. açılış
    acilis = acilabilir_mi(kayit, muhur, kod, liste, sozluk_kaydi)
    # 4. okuma
    uzanti = kayit.get("uzanti") or Path(kaynak).suffix
    okunmayan = {"baslik": 0, "personel": 0, "adetli": 0}
    try:
        gercek = oku.oku(Path(kaynak), uzanti, P.get("sayfa"), okunmayan)
    except oku.OkumaHatasi as e:
        raise Durdu(4, str(e)) from None
    # 5. eşleme
    r = esle.esle(bizim["satirlar"], gercek, gruplar)
    # 6. sınıf
    S = siniflar_oku(siniflar, kod, kayit["sha256"])
    sinif = {}
    for x in r["gercek"]:
        if x["sonuc"] in esle.DOGRU:
            continue
        s = S["satirlar"].get(str(x["i"]))
        if s and s.get("malzeme") and esle.ad(s["malzeme"]) != esle.ad(gercek[x["i"]]["malzeme"]):
            raise ValueError(f"sınıf dosyası satır {x['i']}: malzeme tutmuyor (okuyucu değişmiş olabilir)")
        sinif[("g", x["i"])] = s or {"sinif": SINIFSIZ}
    for j in r["fazla"]:
        sinif[("b", j)] = S["fazla"].get(fazla_anahtari(bizim["satirlar"][j])) or {"sinif": SINIFSIZ}
    sinif_sayilari = {k: 0 for k in (*SINIFLAR, SINIFSIZ)}
    for v in sinif.values():
        sinif_sayilari[v["sinif"]] += 1
    # 13 Ekim listesiyle aynı eşleme (rapor sütunu; kayda girmez)
    once = None
    if onceki:
        ol = json.loads(Path(onceki).read_text(encoding="utf-8"))
        once = {"liste": ol, "esle": esle.esle(ol["satirlar"], gercek, gruplar)}
    return {"proje": P, "proje_kaydi": proje_kaydi, "kod": kod, "liste": liste, "bizim": bizim, "muhur": muhur_kaydi, "kayit": kayit,
            "acilis": acilis,
            "sozluk": {**sozluk_kaydi, "surum": json.loads(sozluk_yol.read_text(encoding="utf-8")).get("surum")},
            "gercek": gercek, "okunmayan": okunmayan, "esle": r, "sinif": sinif, "sinif_sayilari": sinif_sayilari, "once": once,
            "kiyas_commit": _git(kok, "rev-parse", "HEAD").stdout.strip()}


def tutanlar(k: dict) -> list[dict]:
    """Gerçek listeyle tutan anahtarlar; değerleri bizim adedimizdir (varsayılan, pafta ya da aday sütunu)."""
    bizim, out, gorulen = k["bizim"]["satirlar"], [], set()
    for x in k["esle"]["gercek"]:
        anahtar = tuple(x["anahtar_satirlari"])
        if x["sonuc"] not in esle.DOGRU or anahtar in gorulen:
            continue
        gorulen.add(anahtar)
        t = {"grup": x["grup"], "ad": x["ad"], "b_anahtar": x["b_anahtar"], "eslesme": x["eslesme"], "deger": x["G"],
             "sonuc": x["sonuc"], "bizim": [{"kalem": bizim[j]["kalem"], "olcu": bizim[j]["olcu"]} for j in x["bizim"]]}
        if x["eslesme"] in esle.GRUP_OLCULU:
            t["grup_olculeri"] = sorted(esle.kimlik(bizim[j]["olcu"]) for j in x["grup_bizim"])
        if x.get("satir_degerleri") is not None:
            t["satir_degerleri"] = x["satir_degerleri"]
        out.append(t)
    return out


def sonuc_kaydi(k: dict) -> dict:
    P = k["proje"]
    return {
        "aciklama": ("Kör test sonucu (BK-12 PR-3). Yalnız tutan satırlar: grup, bizim ölçü anahtarımız, eşleme türü, adet (bizim "
                     "adedimiz) ve kıyas anındaki bizim satırlarımız (kalem, ölçü). Tutmayan satırların gerçek değeri depoya girmez; "
                     "sayilar kıyasın özetidir (kaç gerçek satır doğru, farklı, bizde yok). Gerileme kapısı `test_kor_kiyas.py`; yeniden "
                     "sayım tam koşuda."),
        "kod": k["kod"], "ap": P.get("ap"), "kip": P["kip"], "kapi": bool(P.get("kapi", True)), "not": P.get("not"),
        "tarih": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "kiyas_commit": k["kiyas_commit"],
        "bizim_liste": {"yol": k["liste"]["yol"], "commit": k["liste"]["commit"], "uretici_surum": k["bizim"].get("uretici_surum")},
        "muhur": {"sha256": k["kayit"]["sha256"], "bayt": k["kayit"]["bayt"], "durum": k["kayit"]["durum"],
                  "muhur_commit": k["muhur"]["commit"], **({"not": k["kayit"]["not"]} if k["kayit"].get("not") else {})},
        "sozluk": {"surum": k["sozluk"]["surum"], "commit": k["sozluk"]["commit"],
                   "acilistan_sonra_degisti": k["acilis"]["sozluk_degisti"]},
        "sayilar": k["esle"]["sayilar"],
        "okunmayan_adetli": k["okunmayan"]["adetli"],
        "sinif_sayilari": k["sinif_sayilari"],
        "tutanlar": tutanlar(k),
    }


def sonuc_yaz(kayit: dict, yol: Path) -> None:
    """`sonuc.json`: üst alanlar girintili, her tutan satır tek satırda (fark okunur kalsın)."""
    ust = {k: v for k, v in kayit.items() if k != "tutanlar"}
    metin = json.dumps(ust, ensure_ascii=False, indent=1)[:-2]
    satirlar = ",\n".join("  " + json.dumps(t, ensure_ascii=False) for t in kayit["tutanlar"])
    Path(yol).write_text(metin + ',\n "tutanlar": [\n' + satirlar + "\n ]\n}\n", encoding="utf-8")


# ------------------------------------------------------------------------------------------------------------------- rapor

def maske(t) -> str:
    """Telefona benzeyen her şey "…" olur (`beyin/numara.py`; ölçü ve ondalık maskelenmez)."""
    return numara.maskele(str(t)) if t is not None else ""


def _sayi(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace(".", ",") if isinstance(v, float) else str(v)


def rapor_satirlari(k: dict) -> tuple[list[dict], list[dict]]:
    g, bizim = k["gercek"], k["bizim"]["satirlar"]
    once = {x["i"]: x["sonuc"] for x in k["once"]["esle"]["gercek"]} if k["once"] else {}
    satirlar = []
    for x in k["esle"]["gercek"]:
        s = g[x["i"]]
        b = [bizim[j] for j in x["bizim"]]
        si = k["sinif"].get(("g", x["i"]), {})
        satirlar.append({
            "no": x["i"] + 1, "grup": maske(s["grup"]), "malzeme": maske(s["malzeme"]), "boyut": maske(s["boyut"]),
            "adet": _sayi(s["adet"]), "sonuc": esle.ADLAR[x["sonuc"]],
            "bizim_kalem": " · ".join(maske(y["kalem"]) for y in b), "bizim_olcu": " · ".join(maske(y["olcu"]) for y in b),
            "bizim_adet": " + ".join(_sayi(y["adet"]) for y in b), "pafta": " + ".join(_sayi(y.get("pafta")) for y in b if y.get("pafta") is not None),
            "aday": " + ".join(_sayi(y.get("aday")) for y in b if y.get("aday") is not None),
            "sinif": si.get("sinif", ""), "gerekce": maske(si.get("gerekce", "")),
            "onceki": esle.ADLAR.get(once.get(x["i"]), "") if k["once"] else "",
        })
    fazla = []
    for j in k["esle"]["fazla"]:
        y = bizim[j]
        si = k["sinif"].get(("b", j), {})
        fazla.append({"grup": maske(y.get("grup_adi") or y.get("grup")), "kalem": maske(y["kalem"]), "olcu": maske(y["olcu"]),
                      "adet": _sayi(y["adet"]), "sinif": si.get("sinif", ""), "gerekce": maske(si.get("gerekce", ""))})
    return satirlar, fazla


def _surum(t) -> str:
    """Sürüm yazısının başı: "1.5 (22 Ekim 2026; …)" → "1.5 (22 Ekim 2026)"."""
    m = re.match(r"^([^ (]+)\s*\(([^;)]*)", str(t or ""))
    return maske(f"{m.group(1)} ({m.group(2).strip()})" if m else str(t or "—"))


def _ozet_satirlari(k: dict) -> list[tuple[str, str]]:
    s, P, b = k["esle"]["sayilar"], k["proje"], k["bizim"]
    m2 = (b.get("girdi_ozeti") or {}).get("brut_m2")
    out = [("Proje", k["kod"] + (f" ({P['ap']})" if P.get("ap") else "")), ("Kip", P["kip"]),
           ("Brüt alan", f"{_sayi(round(m2, 2))} m²" if m2 else "—"),
           ("Kendi listemiz", f"{k['liste']['yol']} · commit {k['liste']['commit'][:7]} ({k['liste']['tarih']})"),
           ("Üreteç", _surum(b.get("uretici_surum"))),
           ("Mühür", f"SHA-256 {k['kayit']['sha256'][:12]}… · {k['kayit']['bayt']} bayt · durum {k['kayit']['durum']}"
                     + (f" · {maske(k['kayit']['not'])}" if k["kayit"].get("not") else "")),
           ("Sözlük", f"{_surum(k['sozluk']['surum'])} · commit {k['sozluk']['commit'][:7]}"
                      + (" · AÇILIŞTAN SONRA DEĞİŞTİ (açılışta " + str(k["acilis"].get("sozluk_commit"))[:7] + ")"
                         if k["acilis"].get("sozluk_degisti") else "")),
           ("Gerçek liste (dolu satır)", str(s["gercek_dolu"])),
           ("Okunmayan satır (başlık etiketi, personel)", f"{k['okunmayan']['baslik']} + {k['okunmayan']['personel']}; adedi sayı olan "
                                                         f"{k['okunmayan']['adetli']} (paydaya girmez)"),
           ("Doğru (aynı + ayrım + pafta + aday)", f"{s['dogru']} / {s['gercek_dolu']}")]
    out += [(esle.ADLAR[x], str(s[x])) for x in (*esle.SONUCLAR, "bizde fazla")]
    out += [(f"sınıf: {x}", str(n)) for x, n in k["sinif_sayilari"].items() if n]
    if k["once"]:
        o = k["once"]["esle"]["sayilar"]
        out.append(("13 Ekim listesiyle (aynı araç)", f"{o['dogru']} / {o['gercek_dolu']} · üreteç "
                    f"{_surum(k['once']['liste'].get('uretici_surum'))}"))
    if P.get("not"):
        out.append(("Not", maske(P["not"])))
    return out


def rapor_md(k: dict, yol: Path) -> None:
    satirlar, fazla = rapor_satirlari(k)
    L = [f"# Kör test kıyası · {k['kod']}", "",
         "Araç: `testler/kor/kiyas.py` (BK-12 PR-3). Ad, bayi adı ve dosya adı yazılmaz; başlık bloğu okunmadı.", "",
         "| | |", "|---|---|"]
    L += [f"| {a} | {b} |" for a, b in _ozet_satirlari(k)]
    L += ["", "## Satırlar (gerçek liste sırasıyla)", "",
          "| # | Grup | Gerçek malzeme | Boyut | Adet | Sonuç | Bizim kalem | Bizim ölçü | Bizim adet | Pafta | Aday | Sınıf | Gerekçe"
          + (" | 13 Ekim" if k["once"] else "") + " |",
          "|---|---|---|---|---|---|---|---|---|---|---|---|---" + ("|---" if k["once"] else "") + "|"]
    for s in satirlar:
        h = [s[x] for x in ("no", "grup", "malzeme", "boyut", "adet", "sonuc", "bizim_kalem", "bizim_olcu", "bizim_adet", "pafta",
                             "aday", "sinif", "gerekce")] + ([s["onceki"]] if k["once"] else [])
        L.append("| " + " | ".join(str(v).replace("|", "/") for v in h) + " |")
    L += ["", "## Bizde fazla", "", "| Grup | Kalem | Ölçü | Adet | Sınıf | Gerekçe |", "|---|---|---|---|---|---|"]
    L += ["| " + " | ".join(str(f[x]).replace("|", "/") for x in ("grup", "kalem", "olcu", "adet", "sinif", "gerekce")) + " |"
          for f in fazla]
    Path(yol).write_text("\n".join(L) + "\n", encoding="utf-8")


def rapor_xlsx(k: dict, yol: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font

    satirlar, fazla = rapor_satirlari(k)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Özet"
    for a, b in _ozet_satirlari(k):
        ws.append([a, b])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90
    ws2 = wb.create_sheet("Satırlar")
    bas = ["#", "Grup", "Gerçek malzeme", "Boyut", "Adet", "Sonuç", "Bizim kalem", "Bizim ölçü", "Bizim adet", "Pafta", "Aday", "Sınıf",
           "Gerekçe"] + (["13 Ekim"] if k["once"] else [])
    ws2.append(bas)
    for s in satirlar:
        ws2.append([s[x] for x in ("no", "grup", "malzeme", "boyut", "adet", "sonuc", "bizim_kalem", "bizim_olcu", "bizim_adet", "pafta",
                                   "aday", "sinif", "gerekce")] + ([s["onceki"]] if k["once"] else []))
    ws3 = wb.create_sheet("Bizde fazla")
    ws3.append(["Grup", "Kalem", "Ölçü", "Adet", "Sınıf", "Gerekçe"])
    for f in fazla:
        ws3.append([f[x] for x in ("grup", "kalem", "olcu", "adet", "sinif", "gerekce")])
    for w in (ws2, ws3):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    wb.properties.creator = "Portman kör test aracı"
    wb.properties.title = f"Kör test kıyası {k['kod']}"
    wb.save(str(yol))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--proje", type=Path, required=True, help="testler/kor/projeler/<KOD>")
    ap.add_argument("--kaynak", type=Path, required=True, help="gerçek liste (depo dışında)")
    ap.add_argument("--siniflar", type=Path)
    ap.add_argument("--onceki", type=Path, help="13 Ekim listemiz (rapor sütunu)")
    ap.add_argument("--md", type=Path)
    ap.add_argument("--xlsx", type=Path)
    ap.add_argument("--sonuc-yaz", action="store_true", help="tutan satırları projenin sonuc.json'una yaz")
    a = ap.parse_args(argv)
    try:
        k = kiyasla(a.proje, a.kaynak, siniflar=a.siniflar, onceki=a.onceki)
    except Durdu as e:
        print(f"DURDU · {e}", file=sys.stderr)
        return 2
    s = k["esle"]["sayilar"]
    print(f"{k['kod']}: doğru {s['dogru']} / {s['gercek_dolu']} (aynı {s['aynı']}, ayrım {s['ayrım']}, pafta {s['pafta']}, "
          f"aday {s['aday']}; farklı {s['farklı']}, bizde yok {s['bizde yok']}, bizde fazla {s['bizde fazla']})")
    if a.md:
        rapor_md(k, a.md)
    if a.xlsx:
        rapor_xlsx(k, a.xlsx)
    if a.sonuc_yaz:
        sonuc_yaz(sonuc_kaydi(k), a.proje / "sonuc.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""Kör test aracının okuyucusu (BK-12 PR-3, M7): gerçek yükleme listesini malzeme satırlarına çevirir.

Yalnız malzeme satırları okunur: grup, malzeme, boyut, adet. Okunmayanlar:
- başlık bloğu (müşteri, temsilci, hazırlayan, sevkiyat yeri, tarih, telefon): ilk sütun başlığından ("BOYUT", "ADET") önceki her şey;
  sütun başlığından sonra da adı yalnız böyle bir etiket olan satır (`BASLIK`: "MÜŞTERİ TEMSİLCİSİ", "SEVKİYAT YERİ", "TARİH: …";
  "TELEFON PRİZİ" gibi malzeme okunur). PDF'te sütun başlığı sonraki sayfalarda yinelenmeyebilir (TS-01'in 2. ve 3. sayfası başlıksız
  sürer), bu yüzden sütun yerleri sayfadan sayfaya taşınır;
- sayfa adları, dosya adı ve belge özellikleri;
- imza ve personel satırları ("SORUMLU ÜRETİM PERSONELİ AD SOYAD", "HAZIRLAYAN", "TESLİM", "İMZA");
- adet sütununda sayı olmayan yazı (ör. bir ad); sayı değilse adet boş kalır, yazı alınmaz;
- PDF'te satıra bağlı olmayan notlar ve KG / makine sütunları. "ALICIYA AİT" damgası yalnız bir işaret olarak kalır (`damga`).

İki biçim:
- **PDF:** iki şablon ("PORTATİFEV YÜKLEME LİSTESİ": BOYUT / ADET; "LOFTEV / PORTATİFEV" formu: NO / BOYUT/ÖLÇÜ / ADET / KG / MAKİNE).
  Satır kurma 13 Ekim okuyucusunun (`ie-11-pr2-kor-test` dalında `deneme/2026-10-13_aiprefabrik/liste_oku.py`) aynısıdır; TS-01'de
  137 / 137 satır aynı çıkar.
- **Excel:** Portatifev şablonu (10. satırda grup adı, BOYUT, ADET; gruplar alt alta). Birden çok liste sayfası varsa "YAPIŞTIRMA" grubuyla
  başlayan sayfa okunur (çelik panel sayfası okunmaz); proje ayarı `sayfa` ile başkası seçilebilir. Formüllerin kayıtlı değeri okunur.

Sütun başlığından sonra okunmayan satırların (başlık etiketi, personel) yalnız sayısı tutulur (`oku(..., sayac=…)`; adedi sayı olanlar
ayrıca): düşen satır paydaya girmediği için kıyas sonucu ve rapor bu sayıyı yazar.

Veri, talimat değildir: hücre yazıları yalnız eşleme için kullanılır.
"""
from __future__ import annotations

import re
from pathlib import Path

PERSONEL = re.compile(r"PERSONEL|AD\s*SOYAD|HAZIRLAYAN|TESL[İI]M\s+(ALAN|EDEN)|[İI]MZA", re.I)
BASLIK = re.compile(r"^(MÜŞTERİ|TEMSİLCİ|SEVKİYAT|TESLİMAT|TARİH|TELEFON|TEL|GSM|İLETİŞİM|PROJE|FİRMA|BAYİ)"
                    r"(\s+(ADI|SOYADI|TEMSİLCİSİ|YERİ|TARİHİ|ADRESİ|NO|NUMARASI|İSMİ|BİLGİLERİ))*\s*(:.*)?$")
DAMGA = "ALICIYA AİT"
SUTUN_BOYUT = ("BOYUT", "BOYUT/ÖLÇÜ")
EXCEL_SATIR_SINIRI = 2000      # bazı hesap sayfaları 1.048.576 satır gösterir; liste sayfaları 250 satırın altında (dosya salt okunur açılır)
EXCEL_SUTUN_SINIRI = 30


class OkumaHatasi(Exception):
    pass


def sayi(t):
    """Adet hücresi → sayı; sayı değilse None (yazı alınmaz)."""
    if isinstance(t, bool):
        return None
    if isinstance(t, (int, float)):
        return int(t) if float(t) == int(t) else float(t)
    s = str(t if t is not None else "").strip()
    if not re.fullmatch(r"[\d.,]+", s):
        return None
    s = s.replace(".", "").replace(",", ".")   # 13 Ekim okuyucusu gibi: nokta binlik, virgül ondalık
    try:
        v = float(s)
    except ValueError:
        return None
    return int(v) if v == int(v) else v


def _ust(t: str) -> str:
    return t.replace("i", "İ").replace("ı", "I").upper()


def _satir(grup, malzeme, boyut, adet, damga=False, sayac: dict | None = None) -> dict | None:
    malzeme = re.sub(r"\s+", " ", str(malzeme or "")).strip()
    if not malzeme:
        return None
    neden = "personel" if PERSONEL.search(malzeme) else "baslik" if BASLIK.match(_ust(malzeme)) else None
    if neden:
        if sayac is not None:
            sayac[neden] = sayac.get(neden, 0) + 1
            if sayi(adet) is not None:
                sayac["adetli"] = sayac.get("adetli", 0) + 1
        return None
    s = {"grup": re.sub(r"\s+", " ", str(grup or "")).strip() or None, "malzeme": malzeme,
         "boyut": re.sub(r"\s+", " ", str(boyut)).strip() if boyut not in (None, "") else None, "adet": adet}
    if damga:
        s["damga"] = DAMGA
    return s


def pdf_oku(yol: Path, sayac: dict | None = None) -> list[dict]:
    import pymupdf

    doc = pymupdf.open(str(yol))
    satirlar = []
    grup = None
    kolon = None
    for sno, sayfa in enumerate(doc, 1):
        parcalar = []
        for b in sayfa.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                t = re.sub(r"\s+", " ", "".join(s["text"] for s in l["spans"]).strip())
                if t:
                    x0, y0, x1, y1 = l["bbox"]
                    parcalar.append({"x0": x0, "x1": x1, "y": (y0 + y1) / 2, "t": t})
        parcalar.sort(key=lambda p: (round(p["y"]), p["x0"]))
        kumeler = []
        for p in parcalar:
            if kumeler and abs(p["y"] - kumeler[-1]["y"]) <= 3.2:
                kumeler[-1]["p"].append(p)
            else:
                kumeler.append({"y": p["y"], "p": [p]})
        bekleyen = []
        for k in kumeler:
            ps = sorted(k["p"], key=lambda p: p["x0"])
            if any(p["t"] in ("BOYUT", "BOYUT/ÖLÇÜ", "ADET") for p in ps) and not any(re.fullmatch(r"[\d.,]+", p["t"]) for p in ps):
                g = [p["t"] for p in ps if p["t"] not in ("BOYUT", "BOYUT/ÖLÇÜ", "ADET", "KG", "MAKİNE", "NO")]
                grup = re.sub(r"^NO\s+", "", " ".join(g)).strip() or grup
                kolon = {"no": any(p["t"] == "NO" or p["t"].startswith("NO ") for p in ps)}
                for p in ps:
                    if p["t"] in ("BOYUT", "BOYUT/ÖLÇÜ"):
                        kolon["boyut"] = p["x0"] - 8
                    elif p["t"] == "ADET":
                        kolon["adet"] = (p["x0"] + p["x1"]) / 2
                    elif p["t"] == "KG":
                        kolon["kg"] = (p["x0"] + p["x1"]) / 2
                bekleyen = []
                continue
            if kolon is None:
                continue  # başlık bloğu: okunmaz
            mal, boy, adet, kg, ek = [], [], [], [], []
            for p in ps:
                xc = (p["x0"] + p["x1"]) / 2
                if p["x1"] < kolon.get("boyut", 280) + 4:
                    if kolon["no"] and re.fullmatch(r"\d{1,3}", p["t"]) and p["x1"] < 66:
                        continue
                    m = re.match(r"^(\d{1,3})\s+(.*)$", p["t"]) if kolon["no"] else None
                    mal.append(m.group(2) if (m and p["x0"] < 62) else p["t"])
                elif "adet" in kolon and abs(xc - kolon["adet"]) < 22:
                    adet.append(p["t"])
                elif "kg" in kolon and abs(xc - kolon["kg"]) < 20:
                    kg.append(p["t"])
                elif xc < kolon.get("adet", 380) - 22:
                    boy.append(p["t"])
                else:
                    ek.append(p["t"])
            if mal and not adet and not boy and not kg:
                if satirlar and satirlar[-1]["_sayfa"] == sno and k["y"] - satirlar[-1]["_y"] < 7:
                    satirlar[-1]["malzeme"] += " " + " ".join(mal)
                    continue
                bekleyen.append((k["y"], " ".join(mal)))
                continue
            if not mal and not adet and not kg:
                continue
            if not mal and not adet and not boy:
                continue  # satıra bağlı olmayan not: okunmaz
            ad = " ".join(mal)
            if bekleyen and (not ad or k["y"] - bekleyen[-1][0] < 7):
                ad = (" ".join(b[1] for b in bekleyen) + " " + ad).strip()
                bekleyen = []
            if not ad and satirlar and satirlar[-1]["_sayfa"] == sno and k["y"] - satirlar[-1]["_y"] < 6:
                s = satirlar[-1]
                if adet and s["adet"] is None:
                    s["adet"] = sayi(adet[0])
                continue
            satirlar.append({"_sayfa": sno, "_y": k["y"], "grup": grup, "malzeme": ad, "boyut": " ".join(boy) or None,
                             "adet": sayi(adet[0]) if adet else None, "_damga": any(DAMGA in e.upper() for e in ek)})
    sonuc = []
    for s in satirlar:
        r = _satir(s["grup"], s["malzeme"], s["boyut"], s["adet"], s["_damga"], sayac)
        if r:
            sonuc.append(r)
    return sonuc


def _excel_basligi(hucre: dict) -> dict | None:
    b = [k for k, v in hucre.items() if str(v).strip().upper() in SUTUN_BOYUT]
    a = [k for k, v in hucre.items() if str(v).strip().upper() == "ADET"]
    if not b or not a or b[0] < 2 or not [x for x in a if x > b[0]]:
        return None
    return {"mal": b[0] - 1, "boy": b[0], "adet": min(x for x in a if x > b[0])}


def excel_sayfasi(ws, sayac: dict | None = None) -> list[dict]:
    satirlar, grup, kol = [], None, None
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or EXCEL_SATIR_SINIRI, EXCEL_SATIR_SINIRI), max_col=EXCEL_SUTUN_SINIRI):
        hucre = {c.column: c.value for c in r if c.value is not None and str(c.value).strip() != ""}
        k = _excel_basligi(hucre)
        if k:
            kol = k
            grup = str(hucre.get(k["mal"], "")).strip() or grup
            continue
        if kol is None:
            continue  # başlık bloğu: okunmaz
        s = _satir(grup, hucre.get(kol["mal"]), hucre.get(kol["boy"]), sayi(hucre.get(kol["adet"])), sayac=sayac)
        if s:
            satirlar.append(s)
    return satirlar


def excel_oku(yol: Path, sayfa: int | None = None, sayac: dict | None = None) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(yol), data_only=True, read_only=True)
    if sayfa is not None:
        if not 0 <= sayfa < len(wb.worksheets):
            raise OkumaHatasi(f"liste sayfası yok: {sayfa}")
        return excel_sayfasi(wb.worksheets[sayfa], sayac)
    adaylar = []
    for ws in wb.worksheets:
        say = {}
        s = excel_sayfasi(ws, say)
        if s:
            adaylar.append((s, say))
    if not adaylar:
        raise OkumaHatasi("Excel'de liste sayfası bulunamadı (BOYUT / ADET başlığı yok)")
    yapistirma = [a for a in adaylar if "YAPIŞTIRMA" in (a[0][0]["grup"] or "").upper()]
    if len(adaylar) > 1 and len(yapistirma) != 1:
        raise OkumaHatasi(f"{len(adaylar)} liste sayfası var; proje ayarında `sayfa` verilmeli")
    secilen, say = (yapistirma or adaylar)[0]
    if sayac is not None:
        sayac.update(say)
    return secilen


def oku(yol: Path, uzanti: str, sayfa: int | None = None, sayac: dict | None = None) -> list[dict]:
    """Malzeme satırları. `sayac` verilirse okunmayan satırların sayısı yazılır: {"baslik", "personel", "adetli"}."""
    uzanti = uzanti.lower()
    if uzanti == ".pdf":
        return pdf_oku(yol, sayac)
    if uzanti in (".xlsx", ".xlsm"):
        return excel_oku(yol, sayfa, sayac)
    raise OkumaHatasi(f"okunamayan dosya türü: {uzanti}")

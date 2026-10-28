"""BK-12 PR-3 · kör test aracı (M7): mühür, açılış kaydı, sıra, eşleme ve açık sekiz projenin gerileme kapısı.

Her koşuda: aracın sırası (bizim yanımız kayıtlı → mühür → açılış kaydı) geçici bir depoda uydurma listeyle denenir; eşlemenin ölçü
kuralları uydurma satırlarla denenir; gerçek mühür kaydının kör 19'u ve kapalı kayıtları açılış kaydı olmadan reddettiği ve durumların
değişmediği gösterilir; açık sekiz projede tutan satırlar bugünkü Motor çıktısıyla yeniden denetlenir (PK-01 bilgi, kapıya girmez).
Kaynak liste gerekmez.

Tam koşu (`PORTMAN_KOR_TAM=1`, `PORTMAN_KOR_KAYNAK=<klasör>`): gerçek listeler depo dışındaki klasörden (SHA ile bulunur) yeniden
okunur; sayılar ve tutan satırlar `sonuc.json`'la aynı olmalı, bir baytı değiştirilmiş kopya okunmadan durmalı.
Uydurma listedeki ad ve numara uydurmadır; numara çalışma anında kurulur (dosyada düz yazılmaz).
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
from testler.kor import esle, kiyas as K, oku  # noqa: E402

PROJELER = sorted(p for p in (KOK / "testler" / "kor" / "projeler").iterdir() if (p / "proje.json").is_file())
KOR19 = ["AP-06", "AP-09", *[f"AP-{n:02d}" for n in range(11, 19)], *[f"AP-{n:02d}" for n in range(20, 29)]]
ACIK8 = ["PK-01", "PK-02", "PK-03", "PK-05", "PK-06", "PK-07", "PK-08", "TS-01"]
# Mühürde kodların durumu (bir dosya birden çok kod taşıyabilir; kod bir kez sayılır her dosya için). Değişirse açılış Patron kararıdır.
ACIK_KODLAR = ["AP-01", "AP-01", "AP-02", "AP-03", "AP-04", "AP-05", "AP-07", "AP-08", "AP-19", "PK-01", "PK-02", "PK-03", "PK-03",
               "PK-05", "PK-06", "PK-06", "PK-07", "PK-08", "PK-08", "TS-01"]
KAPALI_KODLAR = ["AP-10", "PK-02", "PK-02", "PK-04", "PK-04", "PK-04", "PK-05", "PK-05", "PK-07", "PK-07", "PK-07", "PK-07", "PK-08",
                 "PK-08"]
UYDURMA_AD = "Deneme Kişisi Uydurma"
UYDURMA_TEL = "0" + "5" + "00 000 07 77"


def _oku(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _git(kok: Path, *args):
    subprocess.run(["git", "-C", str(kok), "-c", "user.name=t", "-c", "user.email=t@t", "-c", "commit.gpgsign=false", *args],
                   check=True, capture_output=True)


# ------------------------------------------------------------------------------------------------------- uydurma depo

GERCEK = [  # malzeme, boyut, adet → beklenen sonuç
    ("DENEME KALEM A", "1000", 4, "aynı"),
    ("DENEME KALEM B", "2000", 5, "ayrım"),
    ("DENEME KALEM C", "3000", 4, "pafta"),
    ("DENEME KALEM D", "4000", 6, "aday"),
    ("DENEME KALEM E", "5000", 7, "farklı"),
    ("DENEME KALEM F", "6000", 1, "bizde yok"),
]
BIZIM = [
    {"kalem": "Deneme kalem A", "olcu": "1000", "adet": 4},
    {"kalem": "Deneme kalem B", "olcu": "2000", "adet": 3},
    {"kalem": "Deneme kalem B", "olcu": "2000", "adet": 2},
    {"kalem": "Deneme kalem C", "olcu": "3000", "adet": 3, "pafta": 4},
    {"kalem": "Deneme kalem D", "olcu": "4000", "adet": 3, "aday": 6},
    {"kalem": "Deneme kalem E", "olcu": "5000", "adet": 2},
    {"kalem": "Deneme kalem G", "olcu": "7000", "adet": 1},          # bizde fazla
]


def _liste_xlsx(yol: Path) -> None:
    """Portatifev şablonu: başlık bloğunda uydurma ad ve numara, 10. satırda grup başlığı, sonda personel satırı."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = UYDURMA_AD                       # sayfa adı da okunmaz
    ws["B1"] = "PORTATİFEV YÜKLEME LİSTESİ"
    ws["B3"], ws["C3"] = "MÜŞTERİ ADI", UYDURMA_AD
    ws["B4"], ws["C4"] = "TELEFON", UYDURMA_TEL
    ws["B5"], ws["C5"] = "HAZIRLAYAN", UYDURMA_AD
    ws["B10"], ws["C10"], ws["D10"] = "DENEME GRUBU", "BOYUT", "ADET"
    for n, (m, b, a, _) in enumerate(GERCEK, 11):
        ws.cell(n, 2, m), ws.cell(n, 3, b), ws.cell(n, 4, a)
    ws.cell(20, 2, "SORUMLU ÜRETİM PERSONELİ AD SOYAD"), ws.cell(20, 4, UYDURMA_AD)
    ws.cell(21, 2, "DENEME KALEM H"), ws.cell(21, 3, "8000"), ws.cell(21, 4, UYDURMA_AD)   # adet yerine yazı: alınmaz
    ws.cell(22, 2, "MÜŞTERİ TEMSİLCİSİ"), ws.cell(22, 3, UYDURMA_AD), ws.cell(22, 4, 1)      # başlık etiketi: adet sayı olsa da alınmaz
    ws.cell(23, 2, "SEVKİYAT TARİHİ"), ws.cell(23, 3, "23.10.2026"), ws.cell(23, 4, 1)
    wb.save(str(yol))


def _sha(yol: Path) -> str:
    return hashlib.sha256(Path(yol).read_bytes()).hexdigest()


@pytest.fixture
def depo(tmp_path):
    """Geçici depo: gerçek sözlük, uydurma mühür ve proje; kaynak liste depo dışında."""
    kok = tmp_path / "depo"
    kor = kok / "testler" / "kor"
    pd = kor / "projeler" / "DN01"
    pd.mkdir(parents=True)
    shutil.copy(KOK / "testler" / "kor" / "sozluk.json", kor / "sozluk.json")
    kaynak = tmp_path / "disari" / "liste.xlsx"
    kaynak.parent.mkdir()
    _liste_xlsx(kaynak)
    muhur = {"aciklama": "uydurma", "acilislar": {},
             "dosyalar": [{"kod": ["DN-01"], "tur": "yükleme listesi", "uzanti": ".xlsx", "bayt": kaynak.stat().st_size,
                           "sha256": _sha(kaynak), "durum": "acik"}]}
    (kor / "muhur.json").write_text(json.dumps(muhur, ensure_ascii=False), encoding="utf-8")
    (pd / "proje.json").write_text(json.dumps({"kod": "DN-01", "kip": "pafta", "girdi": {}, "liste_sha256": _sha(kaynak)}),
                                   encoding="utf-8")
    (pd / "bizim_liste.json").write_text(json.dumps({"kod": "DN-01", "girdi_ozeti": {"brut_m2": 12.5}, "satirlar": BIZIM},
                                                    ensure_ascii=False), encoding="utf-8")
    _git(kok, "init", "-q")
    _git(kok, "add", ".")
    _git(kok, "commit", "-q", "-m", "uydurma")
    return {"kok": kok, "pd": pd, "kaynak": kaynak, "muhur": kor / "muhur.json"}


def _muhur_degistir(depo, **alan):
    m = _oku(depo["muhur"])
    for k, v in alan.items():
        if k == "acilislar":
            m["acilislar"] = v
        else:
            m["dosyalar"][0][k] = v
    depo["muhur"].write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
    _git(depo["kok"], "commit", "-q", "-am", "mühür")


@pytest.fixture
def okunmaz(monkeypatch):
    """Okuyucu çağrılırsa test düşer: mühür ya da açılış tutmadığında dosya içeriği okunmamalı."""
    def dur(*a, **k):
        raise AssertionError("dosya okundu")
    monkeypatch.setattr(K.oku, "oku", dur)


# ------------------------------------------------------------------------------------------------------- sıra ve mühür

def test_yedi_sonuc_ve_dogru_toplam(depo):
    k = K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    sonuc = {k["gercek"][x["i"]]["malzeme"]: x["sonuc"] for x in k["esle"]["gercek"]}
    assert sonuc == {m: s for m, _, _, s in GERCEK}
    assert [BIZIM[j]["kalem"] for j in k["esle"]["fazla"]] == ["Deneme kalem G"]
    s = k["esle"]["sayilar"]
    assert (s["aynı"], s["ayrım"], s["pafta"], s["aday"], s["farklı"], s["bizde yok"], s["bizde fazla"]) == (1, 1, 1, 1, 1, 1, 1)
    assert s["dogru"] == 4 and s["gercek_dolu"] == 6
    assert k["sinif_sayilari"][K.SINIFSIZ] == 3                     # farklı, bizde yok, bizde fazla


def test_baslik_blogu_ve_personel_okunmaz(depo, tmp_path):
    k = K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    metin = json.dumps(k["gercek"], ensure_ascii=False)
    assert UYDURMA_AD not in metin and "000 07 77" not in metin
    assert [s["malzeme"] for s in k["gercek"]] == [m for m, *_ in GERCEK] + ["DENEME KALEM H"]
    assert k["gercek"][-1]["adet"] is None                             # adet sütunundaki yazı alınmaz
    assert k["okunmayan"] == {"baslik": 2, "personel": 1, "adetli": 2}   # okunmayan satırların yalnız sayısı
    assert K.sonuc_kaydi(k)["okunmayan_adetli"] == 2
    K.rapor_md(k, tmp_path / "r.md")
    K.rapor_xlsx(k, tmp_path / "r.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
    hucre = " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value is not None)
    for t in ((tmp_path / "r.md").read_text(encoding="utf-8"), hucre, " ".join(wb.sheetnames), "liste.xlsx"):
        assert UYDURMA_AD not in t and "000 07 77" not in t
    assert "liste.xlsx" not in (tmp_path / "r.md").read_text(encoding="utf-8")   # dosya adı yazılmaz
    assert "| Okunmayan satır (başlık etiketi, personel) | 2 + 1; adedi sayı olan 2 (paydaya girmez) |" in (tmp_path / "r.md").read_text(
        encoding="utf-8")


@pytest.mark.parametrize("malzeme, okunur", [
    ("MÜŞTERİ TEMSİLCİSİ", False), ("SEVKİYAT YERİ", False), ("Tarih: 23.10.2026", False), ("PROJE İSMİ", False),
    ("Müşteri Adı Soyadı", False), ("MUSTERI TEMSILCISI", False), ("Müşteri Temsilcisi " + UYDURMA_AD, False), ("SEVK TARİHİ", False),
    ("TESLİM TARİHİ", False), ("TARİH 23.10.2026", False), ("TEL " + UYDURMA_TEL, False), ("Tel: " + UYDURMA_TEL, False),
    ("TELEFON PRİZİ", True), ("PVC PENCERE (MÜŞTERİYE AİT)", True), ("SEVKİYAT PALETİ", True), ("TEL", True), ("TEL 2,5 MM", True)])
def test_baslik_etiketi_yalniz_etiketse_okunmaz(malzeme, okunur):
    """Başlık süzgeci adı bir etiketle başlayan satırı düşürür (etiket yalnız, iki noktayla, ek sözcük ve yazıyla, tarihle, numarayla;
    Türkçe ya da ASCII); etiket sözcüğünü taşıyan malzeme okunur (hakem 2. tur bulgu 3, 3. tur B4)."""
    sayac = {}
    assert (oku._satir(None, malzeme, None, 1, sayac=sayac) is not None) == okunur
    assert sayac == ({} if okunur else {"baslik": 1, "adetli": 1})


def test_numara_adet_sayilmaz():
    """Adet sütununa boşluksuz düşen bir numara adet sayılmaz (milyonu geçen sayı); olağan adetler okunur."""
    numara = UYDURMA_TEL.replace(" ", "")
    assert oku.sayi(numara) is None and oku.sayi(int(numara)) is None
    assert (oku.sayi("1.250"), oku.sayi("2,5"), oku.sayi(999999), oku.sayi(" 12 ")) == (1250, 2.5, 999999, 12)


def test_pdf_okuyucu_okunmayanlari_sayar(tmp_path):
    """PDF'te de başlık etiketli ve personel satırları okunmaz, sayılır (uydurma PDF; hakem 3. tur, B3)."""
    import pymupdf
    doc = pymupdf.open()
    s = doc.new_page()
    for x, y, yazi in [(50, 60, "DENEME GRUBU"), (300, 60, "BOYUT"), (400, 60, "ADET"),
                       (50, 90, "DENEME KALEM"), (300, 90, "100 X 200"), (400, 90, "4"),
                       (50, 120, "MUSTERI TEMSILCISI"), (400, 120, "1"),
                       (50, 150, "SORUMLU URETIM PERSONELI AD SOYAD"), (400, 150, "1"),
                       (50, 180, "DENEME KALEM IKI"), (300, 180, "L 2500"), (400, 180, "2")]:
        s.insert_text((x, y), yazi, fontsize=9)
    doc.save(str(tmp_path / "liste.pdf"))
    sayac = {}
    satirlar = oku.oku(tmp_path / "liste.pdf", ".pdf", sayac=sayac)
    assert [(r["grup"], r["malzeme"], r["boyut"], r["adet"]) for r in satirlar] == [
        ("DENEME GRUBU", "DENEME KALEM", "100 X 200", 4), ("DENEME GRUBU", "DENEME KALEM IKI", "L 2500", 2)]
    assert sayac == {"baslik": 1, "personel": 1, "adetli": 2}


@pytest.mark.parametrize("bozma", ["bir_bayt", "eksik_bayt", "bos"])
def test_muhur_tutmazsa_okumadan_durur(depo, tmp_path, okunmaz, bozma):
    b = bytearray(depo["kaynak"].read_bytes())
    if bozma == "bir_bayt":
        b[len(b) // 2] ^= 0x01
    elif bozma == "eksik_bayt":
        b = b[:-1]
    else:
        b = bytearray()
    kopya = tmp_path / "kopya.xlsx"
    kopya.write_bytes(bytes(b))
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], kopya, kok=depo["kok"])
    assert e.value.adim == 2 and ("mühür tutmuyor" in e.value.neden)


def test_liste_olmayan_ya_da_baska_projenin_dosyasi_okunmaz(depo, okunmaz):
    _muhur_degistir(depo, tur="montaj paftası")
    with pytest.raises(K.Durdu, match="yükleme listesi değil"):
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    _muhur_degistir(depo, tur="yükleme listesi", kod=["DN-02"])
    with pytest.raises(K.Durdu, match="DN-01 değil"):
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])


def test_muhur_kaydi_degismisse_durur(depo, okunmaz):
    m = _oku(depo["muhur"])
    m["dosyalar"][0]["durum"] = "acik"
    depo["muhur"].write_text(json.dumps(m) + " ", encoding="utf-8")   # commit'lenmemiş değişiklik
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert e.value.adim == 2 and "çalışma ağacında değişmiş" in e.value.neden


def _acilis(depo, **degis) -> dict:
    """Doğru açılış kaydı (uydurma Patron kararı); `degis` ile bir alanı bozulur."""
    a = {"karar": "uydurma Patron kararı", "tarih": "2026-10-23", "sha256": _sha(depo["kaynak"]),
         "liste_commit": K.kayitli_mi(depo["pd"] / "bizim_liste.json", depo["kok"])["commit"][:7],
         "sozluk_commit": K.kayitli_mi(depo["kok"] / "testler/kor/sozluk.json", depo["kok"])["commit"][:7]}
    return {k: v for k, v in {**a, **degis}.items() if v is not None}


@pytest.mark.parametrize("durum", ["kor", "kapali"])
def test_kor_ve_kapali_acilis_kaydi_olmadan_reddedilir(depo, okunmaz, durum):
    _muhur_degistir(depo, durum=durum)
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert e.value.adim == 3 and "açılış kaydı yok" in e.value.neden
    for bozuk, neden in [({"karar": ""}, "açılış kaydı yok"), ({"liste_commit": "0" * 7}, "kendi listemizin commit'i değil"),
                         ({"sha256": "0" * 64}, "bu dosya için değil"), ({"sha256": None}, "bu dosya için değil"),
                         ({"sozluk_commit": None}, "sözlük commit'i yok")]:
        _muhur_degistir(depo, acilislar={"DN-01": _acilis(depo, **bozuk)})
        with pytest.raises(K.Durdu) as e:
            K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
        assert e.value.adim == 3 and neden in e.value.neden, bozuk


def test_kor_liste_acilis_kaydiyla_acilir(depo, tmp_path):
    _muhur_degistir(depo, durum="kor")
    _muhur_degistir(depo, acilislar={"DN-01": _acilis(depo)})
    k = K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert k["acilis"]["durum"] == "kor" and k["esle"]["sayilar"]["dogru"] == 4
    assert k["acilis"]["sozluk_degisti"] is False and K.sonuc_kaydi(k)["sozluk"]["acilistan_sonra_degisti"] is False
    # açılıştan sonra sözlük değişirse araç açar ama sonuç ve rapor bunu yazar
    s = depo["kok"] / "testler/kor/sozluk.json"
    s.write_text(s.read_text(encoding="utf-8").replace('"surum": "', '"surum": "deneme '), encoding="utf-8")
    _git(depo["kok"], "commit", "-q", "-am", "sözlük")
    k = K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert k["acilis"]["sozluk_degisti"] is True and K.sonuc_kaydi(k)["sozluk"]["acilistan_sonra_degisti"] is True
    K.rapor_md(k, tmp_path / "r.md")
    assert "AÇILIŞTAN SONRA DEĞİŞTİ" in (tmp_path / "r.md").read_text(encoding="utf-8")


def test_muhur_notu_sonucta_ve_raporda(depo, tmp_path):
    # İE-17: paftası çizime kullanılan kör evin kaydında not; kıyas yapılırsa sonuç ve rapor yazar (kayıt açılmaz, durum değişmez)
    _muhur_degistir(depo, **{"not": "çizim için pafta görüldü"})
    k = K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert K.sonuc_kaydi(k)["muhur"]["not"] == "çizim için pafta görüldü"
    K.rapor_md(k, tmp_path / "r.md")
    assert "durum acik · çizim için pafta görüldü" in (tmp_path / "r.md").read_text(encoding="utf-8")


@pytest.mark.parametrize("dosya", ["testler/kor/sozluk.json", "testler/kor/projeler/DN01/proje.json"])
def test_sozluk_ve_proje_kaydi_okumadan_once_denetlenir(depo, okunmaz, dosya):
    """Bizim yanımız (sözlük, proje kaydı) gerçek liste okunmadan önce commit'li olmalı (adım 1)."""
    y = depo["kok"] / dosya
    y.write_text(y.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert e.value.adim == 1 and "çalışma ağacında değişmiş" in e.value.neden


def test_kendi_listemiz_kayitli_degilse_acilmaz(depo, okunmaz):
    y = depo["pd"] / "bizim_liste.json"
    d = _oku(y)
    d["satirlar"][0]["adet"] = 5
    y.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")      # değiştirilmiş, commit'siz
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert e.value.adim == 1 and "çalışma ağacında değişmiş" in e.value.neden
    _git(depo["kok"], "rm", "-q", "--cached", str(y.relative_to(depo["kok"])))   # yalnız indeksten çıkar: commit'siz liste
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(depo["pd"], depo["kaynak"], kok=depo["kok"])
    assert e.value.adim == 1 and "commit'li değil" in e.value.neden


# ------------------------------------------------------------------------------------------------------- eşlemenin ölçü kuralları

def _gruplar(tmp_path, *gruplar) -> list[dict]:
    y = tmp_path / "sozluk.json"
    y.write_text(json.dumps({"gruplar": list(gruplar)}, ensure_ascii=False), encoding="utf-8")
    return esle.sozluk_oku(y)


def _sonuc(bizim, gercek, gruplar) -> list[str]:
    r = esle.esle(bizim, [{"malzeme": m, "boyut": b, "adet": a} for m, b, a in gercek], gruplar)
    return [x["sonuc"] for x in r["gercek"]]


TOPLAM = {"ad": "Deneme toplam", "bizim": ["^DENEME TOPLAM"], "gercek": ["^DENEME TOPLAM"], "olcu": "yok"}
KUTU = {"ad": "Deneme kutu", "bizim": ["^DENEME KUTU"], "gercek": ["^DENEME KUTU"], "olcu": "tam", "anahtar": "en_buyuk", "tolerans": 10}


def test_toplanan_grupta_olcu_kumesi_kiyaslanir(tmp_path):
    """`yok` kipinde toplam tutsa da ölçü kümesi tutmazsa "farklı" (uydurma saçak sacı 300 × 2500 ↔ 500 × 3000)."""
    g = _gruplar(tmp_path, TOPLAM)
    gercek = [("DENEME TOPLAM", "500 X 3000", 6), ("DENEME TOPLAM ÖZEL", "500 X 3000", 4)]
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "500 × 3000", "adet": 10}], gercek, g) == ["ayrım", "ayrım"]
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "300 × 2500", "adet": 10}], gercek, g) == ["farklı", "farklı"]
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "500 × 3000", "adet": 9},
                   {"kalem": "Deneme toplam", "olcu": "300 × 2500", "adet": 1}], gercek, g) == ["farklı", "farklı"]   # bizde fazla ölçü
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "ADET", "adet": 10}], gercek, g) == ["ayrım", "ayrım"]   # bizde sayılı ölçü yok
    boy = _gruplar(tmp_path, {**TOPLAM, "anahtar": "en_buyuk"})
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "L 3000", "adet": 10}], gercek, boy) == ["ayrım", "ayrım"]
    serbest = _gruplar(tmp_path, {**TOPLAM, "olcu_serbest": True, "not": "uydurma gerekçe"})
    assert _sonuc([{"kalem": "Deneme toplam", "olcu": "açıklık 6000", "adet": 10}], gercek, serbest) == ["ayrım", "ayrım"]
    with pytest.raises(ValueError, match="olcu_serbest"):
        _gruplar(tmp_path, {**TOPLAM, "olcu_serbest": True})                       # gerekçesiz olmaz


def test_toplanan_grupta_olcu_basina_toplam(tmp_path):
    """Aynı kalemin iki ölçüsünde adetler ters dağılırsa toplam ve ölçü kümesi tutar, ölçü başına toplam tutmaz (hakem 2. tur, 4a)."""
    g = _gruplar(tmp_path, TOPLAM)
    gercek = [("DENEME TOPLAM", "300 X 2500", 4), ("DENEME TOPLAM", "500 X 3000", 6)]
    tutan = [{"kalem": "Deneme toplam", "olcu": "300 × 2500", "adet": 4}, {"kalem": "Deneme toplam", "olcu": "500 × 3000", "adet": 6}]
    ters = [{**tutan[0], "adet": 6}, {**tutan[1], "adet": 4}]
    assert _sonuc(tutan, gercek, g) == ["ayrım", "ayrım"]
    assert _sonuc(ters, gercek, g) == ["farklı", "farklı"]
    serbest = _gruplar(tmp_path, {**TOPLAM, "olcu_serbest": True, "not": "uydurma gerekçe"})
    assert _sonuc(ters, gercek, serbest) == ["ayrım", "ayrım"]


def test_her_satir_grubunda_olcu(tmp_path):
    """`her_satir`: gerçek satırın sayılı ölçüsünün bizim sayılı ölçülerimizden birinde eşi olmalı."""
    g = _gruplar(tmp_path, {"ad": "Deneme takım", "bizim": ["^DENEME TAKIM"], "gercek": ["^DENEME TAKIM"], "olcu": "her_satir"})
    gercek = [("DENEME TAKIM LAVABO", "600 X 450", 1), ("DENEME TAKIM BATARYA", "ADET", 1)]
    assert _sonuc([{"kalem": "Deneme takım", "olcu": "600 × 450", "adet": 1}], gercek, g) == ["aynı", "aynı"]
    assert _sonuc([{"kalem": "Deneme takım", "olcu": "800 × 450", "adet": 1}], gercek, g) == ["farklı", "aynı"]
    assert _sonuc([{"kalem": "Deneme takım", "olcu": "ADET", "adet": 1}], gercek, g) == ["aynı", "aynı"]


def test_birim_cevirisi_cm_mm(tmp_path):
    """Sayıları 10 katı olan ölçü aynı sayılır (cm ↔ mm; çelik kapı 90*200 ↔ 900 × 2000); beyanlı: 100 × 100 ↔ 1000*1000 de."""
    assert _sonuc([{"kalem": "Deneme kapı", "olcu": "900 × 2000", "adet": 2}], [("DENEME KAPI", "90*200", 2)], []) == ["aynı"]
    assert _sonuc([{"kalem": "Deneme kapı", "olcu": "900 × 2000", "adet": 2}], [("DENEME KAPI", "90*210", 2)], []) == ["farklı"]
    assert esle.esdeger("100 × 100", "1000*1000") and not esle.ayni_olcu("100 × 100", "1000*1000")


def test_kalan_eslemede_iki_yanda_sayi_varsa_olcu_tutmali(tmp_path):
    """Aynı ad grubunda kalan iki satır eşlenir; iki ölçüde de sayı varsa ve aynı değilse "farklı" (4,2 × 13 ↔ 3,9 × 32)."""
    for bizim, beklenen in [("3,9 × 32", "farklı"), ("1,5 mm", "farklı"), ("ADET", "aynı"), ("4,2 × 13", "aynı")]:
        assert _sonuc([{"kalem": "Deneme vida", "olcu": bizim, "adet": 100}], [("DENEME VİDA", "4,2 X 13", 100)], []) == [beklenen], bizim
    assert _sonuc([{"kalem": "Deneme sigorta", "olcu": "16 Amper", "adet": 3}], [("DENEME SİGORTA", "25 Amper", 3)], []) == ["farklı"]
    assert _sonuc([{"kalem": "Deneme strafor", "olcu": "paket / 60 m²", "adet": 3}], [("DENEME STRAFOR", "PKT / 60 M2", 3)], []) == ["aynı"]


# ------------------------------------------------------------------------------------------------------- gerçek mühür kaydı

@pytest.fixture(scope="module")
def muhur():
    return K.muhur_oku()


def test_muhur_yalniz_kod_ve_ozet_tasir(muhur):
    alanlar = {"kod", "tur", "uzanti", "bayt", "sha256", "durum", "not"}
    shas = [r["sha256"] for r in muhur["dosyalar"]]
    assert len(shas) == len(set(shas)) == 82
    for r in muhur["dosyalar"]:
        assert set(r) <= alanlar, r["kod"]
        assert all(k[:3] in ("PK-", "TS-", "AP-") and len(k) == 5 for k in r["kod"])
    assert muhur["sayilar"] == {"acik": 17, "kaynak": 32, "kor": 19, "kapali": 14}
    assert Counter(r["durum"] for r in muhur["dosyalar"]) == muhur["sayilar"]          # özet dosyadaki sayıyla değil, yeniden sayılır


def test_acik_ve_kapali_kodlar_degismedi(muhur):
    """Bir kaydın durumunu değiştirmek (kapalıyı açık yapmak) açılış kaydı olmadan liste açtırır: durumlar burada sabittir."""
    kodlar = {d: sorted(k for r in muhur["dosyalar"] if r["durum"] == d for k in r["kod"]) for d in K.DURUMLAR}
    assert kodlar["acik"] == ACIK_KODLAR and kodlar["kapali"] == KAPALI_KODLAR and kodlar["kor"] == KOR19
    for r in muhur["dosyalar"]:
        if r["durum"] in ("acik", "kor", "kapali"):
            assert r["tur"] == K.LISTE, r["kod"]


def test_kor19_ve_ap10_acilis_kaydi_olmadan_reddedilir(muhur):
    kor = {k: r for r in muhur["dosyalar"] if r["durum"] == "kor" for k in r["kod"]}
    assert sorted(kor) == KOR19
    ap10 = [r for r in muhur["dosyalar"] if "AP-10" in r["kod"] and r["tur"] == K.LISTE]
    assert len(ap10) == 1 and ap10[0]["durum"] == "kapali"
    liste = {"commit": "f" * 40}
    for kod, r in [*kor.items(), ("AP-10", ap10[0])]:
        assert r["tur"] == K.LISTE
        with pytest.raises(K.Durdu) as e:
            K.acilabilir_mi(r, muhur, kod, liste)
        assert e.value.adim == 3, kod
    assert not set(KOR19 + ["AP-10"]) & set(muhur["acilislar"])


def test_cizim_icin_pafta_gorulen_kor_liste_kapali(muhur):
    # CIO 27 Ekim 22:43: AP-13 ve AP-16'nın montaj paftası İE-17 çizimine kullanıldı; kör listelerinin kaydında not, liste kapalı
    for kod in ("AP-13", "AP-16"):
        r = [r for r in muhur["dosyalar"] if kod in r["kod"]]
        assert {(x["tur"], x["durum"]) for x in r} == {("montaj paftası", "kaynak"), (K.LISTE, "kor")}, kod
        liste = next(x for x in r if x["tur"] == K.LISTE)
        assert liste["not"].startswith("çizim için pafta görüldü") and kod not in muhur["acilislar"], kod


def test_acik_sekiz_projenin_listesi_acik_kayitta(muhur):
    assert [_oku(p / "proje.json")["kod"] for p in PROJELER] == ACIK8
    for p in PROJELER:
        P = _oku(p / "proje.json")
        r = [r for r in muhur["dosyalar"] if r["sha256"] == P["liste_sha256"]]
        assert len(r) == 1 and r[0]["durum"] == "acik" and P["kod"] in r[0]["kod"], P["kod"]
        assert muhur["acilislar"][P["kod"]]["liste_commit"]
        if P.get("ap"):
            assert P["ap"] in r[0]["kod"]


# ------------------------------------------------------------------------------------------------------- gerileme kapısı

@pytest.fixture(scope="module")
def sozluk():
    return esle.sozluk_oku(K.SOZLUK)


@pytest.mark.parametrize("pd", PROJELER, ids=[p.name for p in PROJELER])
def test_kayitta_bizim_degerimizden_baska_deger_yok(pd, sozluk):
    """"Yalnız tutanlar": her tutan satırın değeri kıyas anındaki kendi listemizden yeniden çıkar."""
    kayit, bizim = _oku(pd / "sonuc.json"), _oku(pd / "bizim_liste.json")
    assert kayit["kod"] == bizim["kod"] == _oku(pd / "proje.json")["kod"]
    assert set(kayit) == {"aciklama", "kod", "ap", "kip", "kapi", "not", "tarih", "kiyas_commit", "bizim_liste", "muhur", "sozluk",
                          "sayilar", "okunmayan_adetli", "sinif_sayilari", "tutanlar"}
    assert kayit["okunmayan_adetli"] == 0                               # açık sekiz listede adedi sayı olan satır düşmedi
    assert kayit["sayilar"]["dogru"] == sum(kayit["sayilar"][s] for s in esle.DOGRU)
    assert kayit["bizim_liste"]["commit"] == K.kayitli_mi(pd / "bizim_liste.json")["commit"]
    for t in kayit["tutanlar"]:
        assert set(t) <= {"grup", "ad", "b_anahtar", "eslesme", "deger", "sonuc", "bizim", "grup_olculeri", "satir_degerleri"}, t
        assert t["sonuc"] in esle.DOGRU and t["eslesme"] in ("anahtar", "esdeger", "tolerans", "kalan", "toplam", "her_satir")
        assert ("grup_olculeri" in t) == (t["eslesme"] in esle.GRUP_OLCULU)
        assert "satir_degerleri" not in t or t["eslesme"] in esle.SATIR_BASINA
        assert all(set(b) == {"kalem", "olcu"} for b in t["bizim"])
    bozuk = [(t, esle.kapi_denetle(bizim["satirlar"], sozluk, t)[1]) for t in kayit["tutanlar"]
             if esle.tutuyor_mu(bizim["satirlar"], sozluk, t) is None]
    assert not bozuk


@pytest.fixture(scope="module")
def bugun():
    return {p.name: K.bizim_liste_uret(p) for p in PROJELER}


@pytest.mark.parametrize("pd", [p for p in PROJELER if _oku(p / "proje.json").get("kapi", True)],
                         ids=[p.name for p in PROJELER if _oku(p / "proje.json").get("kapi", True)])
def test_gerileme_kapisi(pd, sozluk, bugun):
    """Bugünkü Motor çıktısı, gerçek listeyle tutan her satırı aynı adetle veriyor. Bozma bilerek yapılıyorsa protokol adım 7."""
    kayit, L = _oku(pd / "sonuc.json"), bugun[pd.name]
    bozuk = []
    for t in kayit["tutanlar"]:
        s, neden = esle.kapi_denetle(L["satirlar"], sozluk, t)
        if s is None:
            bozuk.append(f"{t['grup'] or t['ad']} [{t['b_anahtar']}] kayıtta {t['deger']}: {neden}")
    assert not bozuk, f"{kayit['kod']}: {len(bozuk)} tutan satır bozuldu:\n" + "\n".join(bozuk)


def _ts01(bugun):
    return _oku(KOK / "testler/kor/projeler/TS01/sonuc.json"), json.loads(json.dumps(bugun["TS01"]))


def test_kapi_bozulan_satiri_yakalar(sozluk, bugun):
    """Bir kural tutan satırın adedini ya da ölçüsünü değiştirirse kapı o satırda durur (uydurma bozma, TS-01)."""
    kayit, L = _ts01(bugun)
    t = next(t for t in kayit["tutanlar"] if t["sonuc"] == "aynı" and t["eslesme"] == "anahtar" and esle.olculu(t["bizim"][0]["olcu"]))
    sec = esle.bizim_secimi(L["satirlar"], sozluk, t)
    assert esle.tutuyor_mu(L["satirlar"], sozluk, t) == "aynı" and len(sec) == 1
    sec[0]["adet"] += 1
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("adet")
    sec[0]["adet"] -= 1
    sec[0]["olcu"] = "1 × 1"
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("satır yok")


def test_kapi_toplanan_grupta_olcuyu_ve_satirlari_korur(sozluk, bugun):
    """`yok` grubunda (toplam) kapı adetle birlikte kayıttaki (kalem, ölçü) çiftlerini de ister (hakem 23 Ekim, bulgu 1)."""
    kayit, L = _ts01(bugun)
    t = next(t for t in kayit["tutanlar"] if t["eslesme"] == "toplam" and any(esle.olculu(b["olcu"]) for b in t["bizim"]))
    assert esle.tutuyor_mu(L["satirlar"], sozluk, t)
    sec = esle.bizim_secimi(L["satirlar"], sozluk, t)
    s = next(s for s in sec if esle.olculu(s["olcu"]))
    eski, s["olcu"] = s["olcu"], "1 × 1"
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("satır yok")
    s["olcu"] = eski
    L["satirlar"].append({**s, "adet": 0})                                    # gruba yeni satır (adet 0 da olsa)
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("yeni satır")


def test_kapi_kayittaki_sonucun_sutununu_ister(sozluk, bugun):
    """Varsayılan adetle tutmuş satırın adedi bozulursa, pafta ya da aday sütunu kayıttaki değeri taşısa da kapı durur (hakem 2. tur, 1)."""
    adaylar = []
    for pd in PROJELER:
        L = json.loads(json.dumps(bugun[pd.name]))
        for t in _oku(pd / "sonuc.json")["tutanlar"]:
            sec = esle.bizim_secimi(L["satirlar"], sozluk, t)
            if t["sonuc"] == "aynı" and len(sec) == 1 and t["deger"] in (sec[0].get("pafta"), sec[0].get("aday")):
                adaylar.append((pd.name, L, t, sec[0]))
    assert adaylar                                                       # bugün böyle satırlar var (TS-01 tavan alçıpanı gibi)
    for ad, L, t, s in adaylar:
        s["adet"] += 5
        sonuc, neden = esle.kapi_denetle(L["satirlar"], sozluk, t)
        assert sonuc is None and neden.startswith("adet"), (ad, t["grup"] or t["ad"])
        s["adet"] -= 5
        assert esle.tutuyor_mu(L["satirlar"], sozluk, t) == "aynı"
    ad, L, t, s = adaylar[0]
    s["adet"] += 5                                                       # pafta ya da aday tutmuş kayıtta varsayılanın değişmesi kapıyı durdurmaz
    assert esle.tutuyor_mu(L["satirlar"], sozluk, {**t, "sonuc": "aday"}) in ("pafta", "aday")


def test_kapi_tolerans_eslemede_grubun_olculerini_korur(sozluk, bugun):
    """Tolerans eşlemesiyle tutmuş satırın grubuna yeni boy gelirse gerçek satırı o alabilir: kapı durur (hakem 2. tur, 2)."""
    for pd in PROJELER:
        L = json.loads(json.dumps(bugun[pd.name]))
        t = next((t for t in _oku(pd / "sonuc.json")["tutanlar"] if t["eslesme"] == "tolerans"), None)
        if t:
            break
    assert t, "tolerans eşlemeli tutan yok"
    assert esle.tutuyor_mu(L["satirlar"], sozluk, t)
    s = esle.bizim_secimi(L["satirlar"], sozluk, t)[0]
    L["satirlar"].append({**s, "olcu": s["olcu"] + "1", "adet": 1})
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("grubun ölçüleri değişti")


def test_kapi_kalan_eslemede_grubun_olculerini_korur(sozluk, bugun):
    """Kalan eşlemeyle tutmuş satırda gruba eklenen satır eşlemeyi değiştirir: kapı durur."""
    kayit, L = _ts01(bugun)
    t = next(t for t in kayit["tutanlar"] if t["eslesme"] == "kalan")
    assert esle.tutuyor_mu(L["satirlar"], sozluk, t)
    s = esle.bizim_secimi(L["satirlar"], sozluk, t)[0]
    L["satirlar"].append({**s, "olcu": "L 2500"})
    assert esle.kapi_denetle(L["satirlar"], sozluk, t)[1].startswith("grubun ölçüleri değişti")


def _tutanlar(bizim, gercek, gruplar) -> list[dict]:
    """Uydurma kıyasın tutan satırları, `sonuc.json`'daki biçimde."""
    k = {"bizim": {"satirlar": bizim}, "esle": esle.esle(bizim, [{"malzeme": m, "boyut": b, "adet": a} for m, b, a in gercek], gruplar)}
    return json.loads(json.dumps(K.tutanlar(k), ensure_ascii=False))


def test_kapi_esdeger_eslemede_grubun_olculerini_korur():
    """cm ↔ mm eşdeğeriyle tutmuş satırın grubuna gerçek ölçünün kendisi gelirse gerçek satırı o alır: kapı durur (hakem 3. tur, B3)."""
    once, gercek = [{"kalem": "Deneme kapı", "olcu": "900 × 2000", "adet": 2}], [("DENEME KAPI", "90*200", 2)]
    (t,) = _tutanlar(once, gercek, [])
    assert t["eslesme"] == "esdeger" and esle.tutuyor_mu(once, [], t) == "aynı"
    sonra = once + [{"kalem": "Deneme kapı", "olcu": "90 × 200", "adet": 1}]
    assert _sonuc(sonra, gercek, []) == ["farklı"]
    assert esle.kapi_denetle(sonra, [], t)[1].startswith("grubun ölçüleri değişti")


def test_kapi_olcu_kimligi_yazi_ve_sayilar():
    """Aynı yazıya inen ama sayıları farklı ölçü kapıda aynı değildir: "1 200 × 300" ↔ "1200 × 300", "60 M2" ↔ "60M2" (uydurma;
    hakem 6. tur, B2)."""
    assert esle.olcu("1 200 × 300") == esle.olcu("1200 × 300") and esle.sayilar("1 200 × 300") != esle.sayilar("1200 × 300")
    assert not esle.ayni_olcu("1 200 × 300", "1200 × 300") and not esle.ayni_olcu("60 M2", "60M2")
    assert esle.kimlik("1200 × 300") == esle.olcu("1200 × 300") and esle.ayni_olcu("1200 × 300", "1200 X 300")
    once, gercek = [{"kalem": "Deneme kapak", "olcu": "1200 × 300", "adet": 5}], [("DENEME KAPAK", "120 X 30", 5)]
    (t,) = _tutanlar(once, gercek, [])
    assert t["eslesme"] == "esdeger" and esle.tutuyor_mu(once, [], t) == "aynı"
    sonra = [{**once[0], "olcu": "1 200 × 300"}]
    assert _sonuc(sonra, gercek, []) == ["farklı"]
    assert esle.kapi_denetle(sonra, [], t)[1].startswith("satır yok")
    once = once + [{"kalem": "Deneme kapak", "olcu": "1500 × 300", "adet": 1}]       # grubun öteki satırında aynı yazım
    (t,) = _tutanlar(once, gercek, [])
    assert t["eslesme"] == "esdeger" and esle.tutuyor_mu(once, [], t) == "aynı"
    sonra = [once[0], {**once[1], "olcu": "1 500 × 300"}]
    assert esle.kapi_denetle(sonra, [], t)[1].startswith("grubun ölçüleri değişti")
    (t,) = _tutanlar(sonra, gercek, [])
    assert esle.tutuyor_mu(sonra, [], t) == "aynı"                                   # kayıt da kimlikle: aynı liste kapıdan geçer
    # sayıları aynı, yazısı farklı: `tam` kipinde anahtar yazıdır, bölünür (hakem 7. tur, B2)
    assert not esle.ayni_olcu("1200 × 300", "1200/300") and not esle.ayni_olcu("1200 × 300", "1200 × 300 CM")
    boy = {"anahtar": "en_buyuk"}                                                  # en büyük sayı yetmez (hakem 8. tur, B2)
    assert not esle.ayni_olcu("100*100*3240", "60*60*3240", boy) and esle.ayni_olcu("L 3240", "L:3240", boy)
    once = [{"kalem": "Deneme kapak", "olcu": "1200 × 300", "adet": 5}, {"kalem": "Deneme kapak", "olcu": "1200 × 300", "adet": 2}]
    gercek = [("DENEME KAPAK", "1200 X 300", 7)]
    (t,) = _tutanlar(once, gercek, [])
    assert t["eslesme"] == "anahtar" and esle.tutuyor_mu(once, [], t) == "ayrım"
    sonra = [once[0], {**once[1], "olcu": "1200 × 300 CM"}]
    assert _sonuc(sonra, gercek, []) == ["farklı"]
    assert esle.kapi_denetle(sonra, [], t)[1].startswith("satır yok")


def test_kapi_belirsiz_olcude_durur(tmp_path):
    """Aynı yazıya inen ama sayıları farklı iki bizim ölçümüz (aynı ölçünün iki yazımı) grupta varsa kapı durur, tutan satır başka
    anahtarda olsa da (uydurma; hakem 7. tur B1, 8. tur B3)."""
    once = [{"kalem": "Deneme kapak", "olcu": "1200 × 300", "adet": 1}, {"kalem": "Deneme kapak", "olcu": "1 200 × 300", "adet": 1}]
    gercek = [("DENEME KAPAK", "120*30", 2)]
    (t,) = _tutanlar(once, gercek, [])
    assert t["eslesme"] == "esdeger" and _sonuc(once, gercek, []) == _sonuc(once[::-1], gercek, []) == ["ayrım"]   # sıraya bağlı değil
    assert esle.belirsiz_olcu(once) == ["1200X300"]
    assert esle.kapi_denetle(once, [], t)[1].startswith("belirsiz ölçü")
    toplam = _gruplar(tmp_path, TOPLAM)
    once = [{"kalem": "Deneme toplam", "olcu": "1 200 × 300", "adet": 0}, {"kalem": "Deneme toplam", "olcu": "1200 × 300", "adet": 5}]
    (t,) = _tutanlar(once, [("DENEME TOPLAM", "120 X 30", 5)], toplam)
    assert esle.kapi_denetle(once, toplam, t)[1].startswith("belirsiz ölçü")
    once = [{"kalem": "Deneme kapak", "olcu": "1200 × 300", "adet": 2}, {"kalem": "Deneme kapak", "olcu": "1500/300", "adet": 1},
            {"kalem": "Deneme kapak", "olcu": "1 500/300", "adet": 1}]                    # belirsiz ölçü başka anahtarda
    (t,) = _tutanlar(once, [("DENEME KAPAK", "1200 X 300", 2)], [])
    assert t["eslesme"] == "anahtar" and t["sonuc"] == "aynı"
    assert esle.kapi_denetle(once, [], t)[1].startswith("belirsiz ölçü")
    assert esle.belirsiz_olcu([{"olcu": "1200 × 300"}, {"olcu": "1200 X 300"}, {"olcu": "ADET"}]) == []


def test_esleme_bizim_satir_sirasina_bagli_degil(tmp_path):
    """Bizim satırlarımızın sırası eşlemeyi değiştirmez: anahtarlar ve ölçü başına toplamın temsilcileri kimlik sırasıyla (uydurma;
    hakem 8. tur, B1: aynı ad grubunda eşdeğer adımı, cm ↔ mm kardeşleri, tolerans, geçişli olmayan eşdeğerlik)."""
    from itertools import permutations

    def s(kalem, o, a):
        return {"kalem": kalem, "olcu": o, "adet": a}
    kutu = KUTU
    vakalar = [
        ([], [s("Deneme kapak", "1200 × 300", 2), s("Deneme kapak", "1200/300", 5)], [("DENEME KAPAK", "120*30", 2)]),
        ([], [s("Deneme kapak", "1200 × 300", 3), s("Deneme kapak", "120 × 30", 5)], [("DENEME KAPAK", "1200/300", 3)]),
        (_gruplar(tmp_path, kutu), [s("Deneme kutu", "L 3240", 2), s("Deneme kutu", "L 3250", 3)], [("DENEME KUTU", "L 3245", 2)]),
        (_gruplar(tmp_path, {**kutu, "olcu": "yok"}),
         [s("Deneme kutu", "L 3248", 3), s("Deneme kutu", "L 3240", 3), s("Deneme kutu", "L 3256", 3)], [("DENEME KUTU", "L 3250", 9)]),
        (_gruplar(tmp_path, TOPLAM), [s("Deneme toplam", "12 × 3", 2), s("Deneme toplam", "120 × 30", 2), s("Deneme toplam", "1200 × 300", 2)],
         [("DENEME TOPLAM", "120 X 30", 6)]),
    ]
    for g, bizim, gercek in vakalar:
        sonuclar = {tuple(_sonuc(list(p), gercek, g)) for p in permutations(bizim)}
        assert len(sonuclar) == 1 and "bizde yok" not in next(iter(sonuclar)), (bizim, sonuclar)   # vaka grubu gerçekten eşliyor


def test_kapi_en_buyuk_anahtarina_gelen_satiri_gorur(tmp_path):
    """`anahtar: en_buyuk` grubunda en büyük sayısı aynı, kimliği farklı yeni bir satır eşlemede anahtara girer ve toplamı değiştirir:
    kapı onu da seçer ve "yeni satır" der (uydurma; hakem 9. tur, B1)."""
    kutu = _gruplar(tmp_path, KUTU)
    once, gercek = [{"kalem": "Deneme kutu", "olcu": "L 3240", "adet": 1}], [("DENEME KUTU", "70X70X3240", 1)]
    (t,) = _tutanlar(once, gercek, kutu)
    assert t["eslesme"] == "anahtar" and esle.tutuyor_mu(once, kutu, t) == "aynı"
    sonra = once + [{"kalem": "Deneme kutu", "olcu": "70*70*3240", "adet": 1}]
    assert _sonuc(sonra, gercek, kutu) == ["farklı"]
    assert esle.kapi_denetle(sonra, kutu, t)[1].startswith("yeni satır")
    once = [{"kalem": "Deneme kutu", "olcu": "70 × 610 × 2430", "adet": 2}, {"kalem": "Deneme kutu", "olcu": "70 × 1130 × 2430", "adet": 1}]
    gercek = [("DENEME KUTU", "70X610X2430", 3)]
    (t,) = _tutanlar(once, gercek, kutu)
    assert t["eslesme"] == "anahtar" and esle.tutuyor_mu(once, kutu, t) == "ayrım"
    sonra = once + [{"kalem": "Deneme kutu", "olcu": "70 × 930 × 2430", "adet": 1}]
    assert _sonuc(sonra, gercek, kutu) == ["farklı"]
    assert esle.kapi_denetle(sonra, kutu, t)[1].startswith("yeni satır")
    sonra = [{**once[0], "olcu": "70 × 620 × 2430"}, once[1]]                    # ölçüsü değişen satır: anahtar aynı, kimlik değil
    assert esle.kapi_denetle(sonra, kutu, t)[1].startswith("satır yok")


def test_kapi_pafta_kaydinda_aday_sutunu_yetmez():
    """Pafta sütunuyla tutmuş satırda pafta bozulup aday sütunu kayıttaki değeri taşırsa kapı durur (hakem 3. tur, B3)."""
    once, gercek = [{"kalem": "Deneme pano", "olcu": "ADET", "adet": 3, "pafta": 5, "aday": 7}], [("DENEME PANO", "ADET", 5)]
    (t,) = _tutanlar(once, gercek, [])
    assert t["sonuc"] == "pafta" and esle.tutuyor_mu(once, [], t) == "pafta"
    assert esle.kapi_denetle([{**once[0], "pafta": 6, "aday": 5}], [], t)[1].startswith("adet")
    assert esle.tutuyor_mu([{**once[0], "adet": 5}], [], t) == "aynı"           # varsayılan tutarsa pafta kaydı da tutar


def test_kapi_satir_degerlerini_korur(tmp_path):
    """Toplanan ve her satır gruplarında toplamlar aynı kalsa da bir satırın adedi ya da sütunu kayınca eşleme "farklı" diyebilir; kapı da
    durur (uydurma; hakem 3. tur B1, 4. tur 1–2, 5. tur B1–B2)."""
    def s(kalem, o, a, **sutun):
        return {"kalem": kalem, "olcu": o, "adet": a, **sutun}
    toplam = _gruplar(tmp_path, TOPLAM)
    iki = [("DENEME TOPLAM", "300 X 2500", 4), ("DENEME TOPLAM", "500 X 3000", 6)]
    takim = {"ad": "Deneme takım", "bizim": ["^DENEME TAKIM"], "gercek": ["^DENEME TAKIM"], "olcu": "her_satir"}
    vakalar = [  # grup, kayıttaki bizim, gerçek, bugünkü bizim
        (toplam, [s("Deneme toplam", "300 × 2500", 4), s("Deneme toplam", "500 × 3000", 6)], iki,
         [s("Deneme toplam", "300 × 2500", 6), s("Deneme toplam", "500 × 3000", 4)]),                     # ters dağılım
        (toplam, [s("Deneme toplam", "300 × 2500", 4), s("Deneme toplam", "500 × 3000", 6)], iki,
         [s("Deneme toplam", "300 × 2500", 0), s("Deneme toplam", "500 × 3000", 10)]),                    # bir ölçü 0'a iniyor
        (toplam, [s("Deneme toplam", "300 × 2500", 10), s("Deneme toplam", "500 × 3000", 0)], [("DENEME TOPLAM", "300 X 2500", 10)],
         [s("Deneme toplam", "300 × 2500", 9), s("Deneme toplam", "500 × 3000", 1)]),                     # adedi 0 olan ölçüye kayma
        (_gruplar(tmp_path, takim), [s("Deneme takım", "600 × 450", 1), s("Deneme takım", "800 × 450", 0)],
         [("DENEME TAKIM LAVABO", "600 X 450", 1)], [s("Deneme takım", "600 × 450", 0), s("Deneme takım", "800 × 450", 1)]),
        (toplam, [{**s("Deneme toplam", "300 × 2500", 2), "pafta": 4}, {**s("Deneme toplam", "500 × 3000", 3), "pafta": 6}], iki,
         [{**s("Deneme toplam", "300 × 2500", 2), "pafta": 6}, {**s("Deneme toplam", "500 × 3000", 3), "pafta": 4}]),   # pafta sütununda ters
        # tek sayılı ölçü, adedi boş satıra adet geliyor (hakem 4. tur, 1)
        (toplam, [s("Deneme toplam", "ADET", 7), s("Deneme toplam", "300 × 2500", 0)], [("DENEME TOPLAM", "500 X 3000", 7)],
         [s("Deneme toplam", "ADET", 3), s("Deneme toplam", "300 × 2500", 4)]),
        (_gruplar(tmp_path, takim), [s("Deneme takım", "ADET", 1), s("Deneme takım", "600 × 450", 0)],
         [("DENEME TAKIM LAVABO", "800 X 450", 1)], [s("Deneme takım", "ADET", 0), s("Deneme takım", "600 × 450", 1)]),
        (toplam, [{**s("Deneme toplam", "300 × 2500", 0), "pafta": 5}], [("DENEME TOPLAM", "500 X 3000", 5)],
         [{**s("Deneme toplam", "300 × 2500", 5), "pafta": 5}]),                                      # pafta kaydı, varsayılana adet
        (toplam, [s("Deneme toplam", "300 × 2500", 1, aday=4), s("Deneme toplam", "500 × 3000", 1, aday=6)], iki,
         [s("Deneme toplam", "300 × 2500", 1, aday=6), s("Deneme toplam", "500 × 3000", 1, aday=4)]),    # aday sütununda ters
        (toplam, [s("Deneme toplam", "ADET", 7), s("Deneme toplam", "300 × 2500", None)], [("DENEME TOPLAM", "500 X 3000", 7)],
         [s("Deneme toplam", "ADET", 3), s("Deneme toplam", "300 × 2500", 4)]),                         # adedi boş (None) satır
        # sayılı ölçüsü olmayan satır boşalınca ölçü başına toplam açılır
        (toplam, [s("Deneme toplam", "ADET", 1, pafta=2), s("Deneme toplam", "300 × 2500", 3, pafta=4),
                  s("Deneme toplam", "500 × 3000", 5, pafta=4)], iki,
         [s("Deneme toplam", "ADET", 0, pafta=2), s("Deneme toplam", "300 × 2500", 3, pafta=4), s("Deneme toplam", "500 × 3000", 5, pafta=4)]),
        # aynı ölçüdeki iki kalem arasında adet, paftası olmayan satıra geçiyor
        (toplam, [s("Deneme toplam", "300 × 2500", 2, pafta=4), s("Deneme toplam özel", "300 × 2500", 0),
                  s("Deneme toplam", "500 × 3000", 8, pafta=6)], iki,
         [s("Deneme toplam", "300 × 2500", 0, pafta=4), s("Deneme toplam özel", "300 × 2500", 2), s("Deneme toplam", "500 × 3000", 8, pafta=6)]),
    ]
    for g, once, gercek, sonra in vakalar:
        (t,) = _tutanlar(once, gercek, g)
        assert "satir_degerleri" in t and esle.tutuyor_mu(once, g, t), once
        assert "farklı" in _sonuc(sonra, gercek, g), sonra
        assert esle.kapi_denetle(sonra, g, t)[1].startswith("satır değerleri değişti"), sonra
    serbest = _gruplar(tmp_path, {**TOPLAM, "olcu_serbest": True, "not": "uydurma gerekçe"})
    once, sonra = vakalar[0][1], vakalar[0][3]
    (t,) = _tutanlar(once, iki, serbest)
    assert "satir_degerleri" not in t                                    # ölçü kıyaslanmayan grupta kayma eşlemeyi değiştirmez
    assert _sonuc(sonra, iki, serbest) == ["ayrım", "ayrım"] and esle.tutuyor_mu(sonra, serbest, t)
    tek = [s("Deneme toplam", "ADET", 3), s("Deneme toplam", "300 × 2500", 7)]
    (t,) = _tutanlar(tek, [("DENEME TOPLAM", "300 X 2500", 10)], toplam)
    assert "satir_degerleri" not in t                                    # tek sayılı ölçü ve adedi dolu: kayma eşlemeyi değiştirmez


def test_pk01_bilgi_kapiya_girmez(sozluk, bugun):
    pd = next(p for p in PROJELER if p.name == "PK01")
    P, kayit = _oku(pd / "proje.json"), _oku(pd / "sonuc.json")
    assert P["kapi"] is False and kayit["kapi"] is False and "stoktan" in P["not"]
    tutan = sum(1 for t in kayit["tutanlar"] if esle.tutuyor_mu(bugun["PK01"]["satirlar"], sozluk, t))
    print(f"PK-01 (bilgi): {tutan} / {len(kayit['tutanlar'])} tutan anahtar bugün de tutuyor")


# ------------------------------------------------------------------------------------------------------- tam koşu

TAM = os.environ.get("PORTMAN_KOR_TAM") == "1"


def _kaynak_bul(kayit: dict) -> Path:
    klasor = os.environ.get("PORTMAN_KOR_KAYNAK")
    assert klasor and Path(klasor).is_dir(), "PORTMAN_KOR_KAYNAK: gerçek listelerin klasörü (depo dışında)"
    assert not Path(klasor).resolve().is_relative_to(KOK), "kaynak klasörü depo dışında olmalı"
    for f in sorted(Path(klasor).rglob("*")):
        if f.is_file() and f.stat().st_size == kayit["bayt"] and K.sha256(f) == kayit["sha256"]:
            return f
    raise AssertionError(f"kaynak klasöründe mühürlü liste yok ({kayit['sha256'][:12]}…)")


@pytest.mark.skipif(not TAM, reason="tam koşu: PORTMAN_KOR_TAM=1 ve PORTMAN_KOR_KAYNAK")
@pytest.mark.parametrize("pd", PROJELER, ids=[p.name for p in PROJELER])
def test_tam_kosu_sonuc_kayitla_ayni(pd, tmp_path):
    kayit = _oku(pd / "sonuc.json")
    kaynak = _kaynak_bul(kayit["muhur"])
    k = K.kiyasla(pd, kaynak)
    yeni = K.sonuc_kaydi(k)
    assert (yeni["sayilar"], yeni["okunmayan_adetli"], yeni["tutanlar"]) == (kayit["sayilar"], kayit["okunmayan_adetli"], kayit["tutanlar"])
    assert yeni["bizim_liste"]["commit"] == kayit["bizim_liste"]["commit"]
    b = bytearray(kaynak.read_bytes())
    b[len(b) // 3] ^= 0x01
    kopya = tmp_path / ("kopya" + kaynak.suffix)
    kopya.write_bytes(bytes(b))
    with pytest.raises(K.Durdu) as e:
        K.kiyasla(pd, kopya)
    assert e.value.adim == 2

"""İE-18 · Paket üreteci (`python -m araclar.paket`): ad kuralı, her dosya Var ya da Yok: neden, kör modelde kıyas yok, çıktıda ad ve
numara yok, çizim ve liste canlıdaki işlevlerin çıktısıyla aynı, rehber kart modeliyle ve eski rehberle aynı adetler (sağlama).

Rehber 62 m² rehberinin şablonundan (D-28; şablonun 62'yi aynen verdiğini `test_rehber.py` sınar). Sağlama (`veri/saglama.json`):
AP-19'da (PK-01) 24 Ekim rehberinin paftadan okunmuş parçaları (kart başına) ve oda sırası, TK-01'de İş Emri r06'nın rehberi. Pano, H,
köşe direği, makas, kapı ve veranda adetleri veriden (kart modeliyle aynı); bütün kartlar sayfanın izinden (Node, Playwright ve
three.js 0.160.0 varsa; yoksa o test atlanır). Kayıtta olmayan ev `--is-emri` ile (çizim masası): kayıttaki TK-01'le aynı paket, adlar
kodla, numaralar maskeli. LLM yok, ağ yok; GLB bu testte üretilmez.
"""
from __future__ import annotations

import copy
import json
import re
import sys
from collections import Counter
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(KOK))
from araclar.paket import uret as U  # noqa: E402
from araclar.paket.rehber import veri as V  # noqa: E402
from araclar.paket.rehber import iz  # noqa: E402

SAGLAMA = json.loads((Path(__file__).with_name("veri") / "saglama.json").read_text(encoding="utf-8"))
KODLAR = ["AP-19", "TK-01", "AP-13", "AP-02"]
KALEM = r"(bilgi\.md|bayi_paftasi\.pdf|plan\.png|gorunusler\.pdf|cizim\.dxf|montaj_paftasi\.pdf|yukleme_listesi(_asma|_ters)?\.xlsx|model\.ifc|" \
        r"montaj_rehberi\.html|model\.glb|kiyas\.xlsx)"


@pytest.fixture(scope="module")
def paketler(tmp_path_factory):
    kok = tmp_path_factory.mktemp("Modeller")
    gec = tmp_path_factory.mktemp("gecici")
    k = U.kayit()
    satir = {m["kod"]: U.paket(m, kok, gec, glb=False, kayit_=k) for m in k["modeller"] if m["kod"] in KODLAR}
    return kok, gec, satir


def _model(kod):
    return next(m for m in U.kayit()["modeller"] if m["kod"] == kod)


def _rehber_verisi(html: Path) -> dict:
    m = re.search(r'<script type="application/json" id="pafta">(.*?)</script>', html.read_text(encoding="utf-8"), re.S)
    return json.loads(m.group(1).replace("<\\/", "</"))


def _bilgi_durum(md: Path) -> dict[str, str]:
    t = md.read_text(encoding="utf-8").split("## Dosyalar", 1)[1].split("\n## ", 1)[0]
    return {m.group(1): m.group(2) for m in re.finditer(r"^\| `([^`]+)` \| (.+) \|$", t, re.M)}


def test_kayit():
    k = U.kayit()
    kodlar = [m["kod"] for m in k["modeller"]]
    assert len(kodlar) == 29 and len(set(kodlar)) == 29
    assert kodlar[:28] == [f"AP-{i:02d}" for i in range(1, 29)] and "TK-01" in kodlar
    assert Counter(m["kume"] for m in k["modeller"]) == {"açık": 8, "kör": 19, "kapalı": 1, "teklif": 1}
    for m in k["modeller"]:
        assert m["adim"] in k["adimlar"], m["kod"]
        for yol in [m.get("is_emri"), m.get("kor_test")] + list((m.get("hakem") or {}).values()):
            assert yol is None or (KOK / yol).exists(), (m["kod"], yol)
        ie = json.loads((KOK / m["is_emri"]).read_text(encoding="utf-8")) if m.get("is_emri") else None
        if ie:                                                  # İş Emri'nde ad yok: alıcı ve bayi alanında yalnız kod
            assert re.fullmatch(r"[A-Z]{2}-?\d\d", ie["talep"]["musteri_adi"]) and re.fullmatch(r"Bayi [A-Z]{2}-\d\d", ie["talep"]["bayi"]["ad"])


def test_ad_kurali_ve_durum(paketler):
    kok, _, satir = paketler
    for kod in KODLAR:
        klasor = kok / "prefabrik" / kod
        for p in klasor.iterdir():
            assert re.fullmatch(re.escape(kod) + "_" + KALEM, p.name), p.name
        durum = _bilgi_durum(klasor / f"{kod}_bilgi.md")
        assert len(durum) == 11 + (kod == "TK-01"), (kod, durum)
        for ad, d in durum.items():
            assert d == "Var" or re.fullmatch(r"Yok: .{10,}", d), (ad, d)
            assert (d == "Var") == (klasor / ad).exists(), (ad, d)
        assert durum == {a: v for a, v in satir[kod]["dosyalar"].items()} | {f"{kod}_bilgi.md": "Var"}


def test_beklenen_dosyalar(paketler):
    _, _, s = paketler
    var = {k: {a.split("_", 1)[1] for a, v in s[k]["dosyalar"].items() if v == "Var"} for k in KODLAR}
    assert var["AP-19"] == {"bilgi.md", "bayi_paftasi.pdf", "plan.png", "cizim.dxf", "montaj_paftasi.pdf", "yukleme_listesi.xlsx", "model.ifc",
                            "montaj_rehberi.html", "kiyas.xlsx"}
    assert var["TK-01"] == {"bilgi.md", "bayi_paftasi.pdf", "plan.png", "cizim.dxf", "montaj_paftasi.pdf", "yukleme_listesi_asma.xlsx",
                            "yukleme_listesi_ters.xlsx", "model.ifc", "montaj_rehberi.html"}
    assert var["AP-13"] == {"bilgi.md", "bayi_paftasi.pdf", "plan.png", "cizim.dxf"}
    assert var["AP-02"] == {"bilgi.md"}
    d13 = s["AP-13"]["dosyalar"]
    assert "İE-17 PR-1b" in d13["AP-13_yukleme_listesi.xlsx"] and "İE-15 K5" in d13["AP-13_model.ifc"]
    assert "dört görünüş (ön, arka, sol, sağ)" in d13["AP-13_gorunusler.pdf"]
    assert "kör" in d13["AP-13_kiyas.xlsx"]
    assert "adım 3" in s["AP-02"]["dosyalar"]["AP-02_bayi_paftasi.pdf"]
    assert "teklif" in s["TK-01"]["dosyalar"]["TK-01_kiyas.xlsx"]


def test_kor_modelde_liste_icerigi_yok(paketler):
    """Kör modelde kıyas yok; bilgi.md listenin yalnız kimliğini (bayt, SHA) taşır."""
    kok, _, _ = paketler
    klasor = kok / "prefabrik" / "AP-13"
    assert not (klasor / "AP-13_kiyas.xlsx").exists()
    md = (klasor / "AP-13_bilgi.md").read_text(encoding="utf-8")
    assert "kör protokolüyle" in md
    assert U.muhur("AP-13")["dosya"] and all(d["sha256"] in md for d in U.muhur("AP-13")["dosya"])


def test_ad_ve_numara_yok(paketler):
    """Çıktıda telefon numarası biçiminde dizi yok (`beyin/numara.py`); başlıklarda yalnız kod."""
    from beyin import numara
    kok, _, _ = paketler
    for p in (kok / "prefabrik").rglob("*"):
        if p.suffix in (".md", ".html"):
            t = p.read_text(encoding="utf-8")
        elif p.suffix == ".xlsx":
            from openpyxl import load_workbook
            t = "\n".join(str(c) for ws in load_workbook(p).worksheets for r in ws.iter_rows(values_only=True) for c in r if c is not None)
        else:
            continue
        t = re.sub(r"`[0-9a-f]{64}`|`[A-Za-z0-9_-]{25,40}`|\b[0-9a-f]{64}\b", "", t)      # SHA-256 ve Drive kimliği
        assert numara.maskele(t) == t, p.name
    html = (kok / "prefabrik" / "TK-01" / "TK-01_montaj_rehberi.html").read_text(encoding="utf-8")
    assert "<title>TK-01 Montaj Rehberi</title>" in html


def _dxf_ozet(yol: Path) -> list:
    import ezdxf
    d = ezdxf.readfile(yol)
    out = []
    for e in d.modelspace():
        t = e.dxftype()
        v = e.dxf.text if t == "TEXT" else e.text if t == "MTEXT" else [tuple(round(c, 3) for c in p) for p in e.get_points()] if t == "LWPOLYLINE" else \
            (tuple(round(c, 3) for c in e.dxf.start), tuple(round(c, 3) for c in e.dxf.end)) if t == "LINE" else None
        if isinstance(v, str) and "Üretim:" in v:
            v = v.split("Üretim:")[0]
        out.append((t, e.dxf.layer, v))
    return out


@pytest.mark.parametrize("kod", ["AP-19", "AP-13"])
def test_cizim_canlidakiyle_ayni(paketler, tmp_path, kod):
    """Paketteki DXF ve plan görseli `dxf_out.uret` ve `bayi_gorseli.plan_gorseli`'nin çıktısıyla aynı (DXF'te üretim damgası hariç)."""
    from motor import bayi_gorseli, dxf_out
    kok, _, _ = paketler
    ie = json.loads((KOK / _model(kod)["is_emri"]).read_text(encoding="utf-8"))
    m = dxf_out.uret(copy.deepcopy(ie), tmp_path)
    assert _dxf_ozet(kok / "prefabrik" / kod / f"{kod}_cizim.dxf") == _dxf_ozet(tmp_path / m["dxf"])
    v = json.loads((tmp_path / "validation.json").read_text(encoding="utf-8"))
    png = bayi_gorseli.plan_gorseli(copy.deepcopy(ie), tmp_path / "plan.png", v)
    assert (kok / "prefabrik" / kod / f"{kod}_plan.png").read_bytes() == png.read_bytes()


def test_liste_canlidakiyle_ayni(paketler):
    """TK-01'in iki listesi liste üretecinin satırlarıyla aynı (teklif r06'nın listeleri: asma 133, ters 131 satır; komut satırı ve teklifle
    aynı yol, `is_emrinden`)."""
    from openpyxl import load_workbook
    from motor import prefabrik_liste as PL
    kok, _, _ = paketler
    ie = json.loads((KOK / _model("TK-01")["is_emri"]).read_text(encoding="utf-8"))
    for s, n in (("asma", 133), ("ters", 131)):
        L = PL.is_emrinden(copy.deepcopy(ie), ozellikler={"Tavan": s.capitalize()})
        ws = load_workbook(kok / "prefabrik" / "TK-01" / f"TK-01_yukleme_listesi_{s}.xlsx").active
        satir = [r[1:5] for r in ws.iter_rows(min_row=3, values_only=True)]
        assert len(satir) == n == len(L["satirlar"])
        assert satir == [(x["kalem"], x["olcu"], x["adet"], x["birim"]) for x in L["satirlar"]]


def _kart_sayisi(d: dict) -> Counter:
    """Rehberin verisinden planın parçaları: pano, H, köşe direği, ana makas, iç kapı, veranda kolon ve kirişi (kart başına)."""
    c = Counter(x["card"] for x in d["panels"] + d["hs"] + d["posts"])
    c.update("MA-01" if m["rol"] == "bas" else "MA-02" for m in d["makas"] if m["aile"] == "ana")
    c.update(["K-01"] * len(d["doors"]))
    for v in d["verandalar"]:
        c.update(["V-01"] * len(v["kolon"]) + ["V-02"] * len(v["kiris"]))
    return c


@pytest.mark.parametrize("kod", ["TK-01", "AP-19"])
def test_saglama(paketler, kod):
    """Rehberin planı kart modeliyle ve sağlamayla aynı adetler; oda sırası ve adım sayısı."""
    kok, gec, _ = paketler
    d = _rehber_verisi(kok / "prefabrik" / kod / f"{kod}_montaj_rehberi.html")
    r = SAGLAMA[kod]
    c = _kart_sayisi(d)
    model = {k["kod"]: k["adet"] for k in json.loads(next(gec.glob(f"{kod}_*/kart/*_kart_sayim.json")).read_text(encoding="utf-8"))["kartlar"]}
    assert dict(c) == {k: model[k] for k in c}
    ortak = [k for k in c if k in r["kartlar"]]            # 24 Ekim PK-01 rehberinde veranda kolonu ve kirişi kartsızdı
    assert len(ortak) >= 18 and {k: c[k] for k in ortak} == {k: r["kartlar"][k] for k in ortak}
    assert [x["ad"] for x in d["rooms"]] == r["oda_sirasi"]
    if kod == "TK-01":
        assert d["N"] == r["adim"]
        assert {v: sum(1 for m in d["makas"] if m["aile"] == "veranda" and m["R"] == R) for v, R in (("ön veranda", "V1"), ("arka veranda", "V2"))} == r["yavru_makas"]
    else:
        assert [(len(v["kolon"]), len(v["kiris"])) for v in d["verandalar"]] == [(1, 2)]


@pytest.mark.parametrize("kod", ["TK-01", "AP-19"])
def test_kiris_boyu_listeden(paketler, kod):
    """D-28 madde 1: kirişin kartındaki boy ve "Liste" satırı yükleme listesinden; kartların boyları listedeki kiriş satırlarıyla bire bir
    (TK-01 r06: dış kiriş 386 / 261 kolonların dış yüzünden dış yüzüne, yan 171 / 120, nişin ağzı 125,5; AP-19: 324 / 190, çizimde açık kenar
    boyunca)."""
    import re
    from openpyxl import load_workbook
    kok, _, _ = paketler
    d = _rehber_verisi(kok / "prefabrik" / kod / f"{kod}_montaj_rehberi.html")
    ad = f"{kod}_yukleme_listesi_asma.xlsx" if kod == "TK-01" else f"{kod}_yukleme_listesi.xlsx"
    ws = load_workbook(kok / "prefabrik" / kod / ad).active
    liste = Counter()
    for r in ws.iter_rows(min_row=3, values_only=True):
        if (r[1] or "").startswith("Kutu profil (kiriş)"):
            liste[r[2]] += r[3]
    kart = Counter()
    for k in (k for v in d["verandalar"] for k in v["kiris"]):
        L_ = float(re.search(r"· L ([\d,]+) cm", k["bu"]).group(1).replace(",", "."))
        olcu = f"L {round(L_ * 10)}"
        assert k["liste"].startswith(f"Kutu profil (kiriş) 100 × 100 {olcu} · "), (k["bu"], k["liste"])
        kart[olcu] += 1
    assert kart and kart == liste
    if kod == "TK-01":                                    # Patron 28 Ekim (karar kartı A): liste ile çizim aynı boy, kartta fark notu yok
        assert sorted(k["bu"] for v in d["verandalar"] for k in v["kiris"] if " · dış kiriş · " in k["bu"]) == [
            "arka veranda · dış kiriş · L 261 cm (kolonların üstünde)", "ön veranda · dış kiriş · L 386 cm (kolonların üstünde)"]
        assert not any("çizimde" in k["bu"] for v in d["verandalar"] for k in v["kiris"])


@pytest.mark.skipif(iz.hazir() is not None, reason=f"rehber izi: {iz.hazir()}")
@pytest.mark.parametrize("kod", ["TK-01", "AP-19"])
def test_rehber_izi(paketler, tmp_path, kod):
    """Sayfa hatasız açılır; sahnedeki parçalar kart başına sağlamayla aynı (AP-19'da 24 Ekim rehberiyle birebir)."""
    kok, _, _ = paketler
    z = iz.al(kok / "prefabrik" / kod / f"{kod}_montaj_rehberi.html", tmp_path / "iz.json")
    assert z["hatalar"] == []
    assert dict(Counter(x["kod"] for x in z["parts"] if x["kod"] and x["kod"] in SAGLAMA[kod]["kartlar"])) == SAGLAMA[kod]["kartlar"]


def test_kiyas_pk01(paketler):
    """AP-19'un kıyası: hakemin pafta okumasıyla kart modeli (kapıdaki satırlar geçer), kör test özeti, model ve Motor'un listesi."""
    from openpyxl import load_workbook
    kok, _, _ = paketler
    wb = load_workbook(kok / "prefabrik" / "AP-19" / "AP-19_kiyas.xlsx")
    assert wb.sheetnames == ["Pafta", "Gerçek liste (kör test)", "Model ve Motor'un listesi"]
    pafta = {r[0]: r for r in wb["Pafta"].iter_rows(min_row=3, values_only=True) if r[0]}
    for a in ("h_plani.dis", "h_plani.h_standart", "makas_bas", "makas_orta", "asik", "sandvic_panel", "bas_makas_betopan"):
        assert pafta[a][4] == "geçer", pafta[a]
    ozet = {r[0]: r[1] for r in wb["Gerçek liste (kör test)"].iter_rows(min_row=3, max_row=12, values_only=True)}
    assert ozet["aynı"] == 60


def test_oda_sirasi():
    """Dizilim sırası kuralı TK-01'in ve 24 Ekim PK-01 rehberinin elle verilen sırasını verir; başka başlangıç köşesi desteklenmez."""
    def oda(ad, x0, x1, y0, y1, acik=False):
        return {"ad": ad, "acik": acik, "poligon": [[x0, y0], [x1, y0], [x1, y1], [x0, y1]]}
    pk01 = [oda("YATAK ODASI 1", 10, 258, 621.5, 990), oda("BANYO", 201.25, 383.5, 747, 990), oda("YATAK ODASI 2", 389.5, 590, 621.5, 990),
            oda("HOL", 264, 383.5, 210, 741), oda("MUTFAK", 10, 258, 10, 615.5), oda("SALON", 389.5, 590, 210, 615.5),
            oda("VERANDA", 266, 600, 0, 200, True)]
    assert V.oda_sirasi(pk01, {"x": "min", "y": "max"})[0] == ["YATAK ODASI 1", "BANYO", "YATAK ODASI 2", "SALON", "HOL", "MUTFAK"]
    with pytest.raises(V.Desteklenmiyor):
        V.oda_sirasi(pk01, {"x": "max", "y": "max"})
    assert V.oda_adi("E. BANYO", []) == "Ebeveyn banyosu" and V.oda_adi("HOL", ["ANTRE"]) == "Hol ve antre"
    assert V.oda_adi("YATAK ODASI 2", []) == "Yatak odası 2" and V.oda_adi("İŞLİK", []) == "İşlik"


def test_eski_tasinir_silinmez(tmp_path):
    k = tmp_path / "AP-99"
    k.mkdir()
    (k / "AP-99_bilgi.md").write_text("eski", encoding="utf-8")
    (k / "_betik").mkdir(); (k / "_betik" / "x.py").write_text("x", encoding="utf-8")
    tarih = U._eskiyi_tasi(k)
    assert sorted(p.name for p in k.iterdir()) == ["_eski"]
    assert (k / "_eski" / tarih / "AP-99_bilgi.md").read_text(encoding="utf-8") == "eski"
    assert (k / "_eski" / tarih / "_betik" / "x.py").exists()
    assert U._eskiyi_tasi(k) is None
    for yeni in ("ikinci", "üçüncü"):                     # aynı gün yeniden üretim: önceki arşiv üzerine yazılmaz
        (k / "AP-99_bilgi.md").write_text(yeni, encoding="utf-8")
        assert U._eskiyi_tasi(k) == tarih
    assert sorted(p.read_text(encoding="utf-8") for p in (k / "_eski" / tarih).glob("AP-99_bilgi*.md")) == sorted(["eski", "ikinci", "üçüncü"])


def _yazilar(klasor: Path) -> dict[str, str]:
    """Paketin okunur yazıları: md, html, Excel hücreleri, PDF metni, DXF yazıları, IFC (ham)."""
    import ezdxf
    import fitz
    from openpyxl import load_workbook
    out = {}
    for p in klasor.iterdir():
        if p.suffix in (".md", ".html", ".ifc"):
            out[p.name] = p.read_text(encoding="utf-8", errors="replace")
        elif p.suffix == ".xlsx":
            out[p.name] = "\n".join(str(c) for ws in load_workbook(p).worksheets for r in ws.iter_rows(values_only=True) for c in r if c is not None)
        elif p.suffix == ".pdf":
            out[p.name] = "".join(s.get_text() for s in fitz.open(p))
        elif p.suffix == ".dxf":
            out[p.name] = "\n".join(e.dxf.text if e.dxftype() == "TEXT" else e.text for e in ezdxf.readfile(p).modelspace() if e.dxftype() in ("TEXT", "MTEXT"))
    return out


def test_is_emri_ile_paket(paketler, tmp_path):
    """`--is-emri` (çizim masası, 28 Ekim 2026): kayıtta olmayan kodla İş Emri dosyasından paket. TK-01'in İş Emri'si (sahte) kişi ve bayi adıyla,
    numarayla yazılınca listeler ve rehber kayıttaki TK-01'inkiyle aynı (kod dışında); çıktıda ad ve numara yok, antette kod; kıyas Yok."""
    import hashlib
    from beyin import numara
    from openpyxl import load_workbook
    kok, _, _ = paketler
    ie = json.loads((KOK / _model("TK-01")["is_emri"]).read_text(encoding="utf-8"))
    no = "05" + "32 111 22 33"                          # sahte numara, parçalı (sır ve numara taramaları test verisini numara sanmasın)
    ie["talep"] = {"bayi": {"kod": "XBAYIX9", "ad": "Sahtebayi Yapi"}, "musteri_adi": "Musteri Sahteadi", "metin": f"Musteri Sahteadi {no}"}
    ie["pafta_uyarilari"] = [f"Arayacak numara {no.replace(' ', '')}"]
    yol = tmp_path / "talep.json"
    yol.write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    cikti = tmp_path / "Modeller"
    assert U.main(["TK-98", str(cikti), "--is-emri", str(yol), "--tavan", "asma,ters", "--gecici", str(tmp_path / "gec")]) == 0
    a, b = kok / "prefabrik" / "TK-01", cikti / "prefabrik" / "TK-98"
    durum = _bilgi_durum(b / "TK-98_bilgi.md")
    assert {k for k, v in durum.items() if v == "Var"} == {f"TK-98_{x}" for x in (
        "bilgi.md", "bayi_paftasi.pdf", "plan.png", "cizim.dxf", "montaj_paftasi.pdf", "yukleme_listesi_asma.xlsx", "yukleme_listesi_ters.xlsx",
        "model.ifc", "montaj_rehberi.html")}
    assert "teklif projesi" in durum["TK-98_kiyas.xlsx"]
    bilgi = (b / "TK-98_bilgi.md").read_text(encoding="utf-8")
    assert "--is-emri" in bilgi and hashlib.sha256(yol.read_bytes()).hexdigest() in bilgi
    assert "| Veranda | taşan: önde 386 × 181, arkada 261 × 130 |" in bilgi and "| Çatı (paftada) | beşik %33, tek kat panel (antrasit) |" in bilgi
    for s in ("asma", "ters"):                          # başlık satırında üretim saati var
        sat = lambda p: list(load_workbook(p).active.iter_rows(min_row=2, values_only=True))
        assert sat(b / f"TK-98_yukleme_listesi_{s}.xlsx") == sat(a / f"TK-01_yukleme_listesi_{s}.xlsx"), s
    html = (b / "TK-98_montaj_rehberi.html").read_text(encoding="utf-8")
    assert html.replace("TK-98", "TK-01").replace("TK98", "TK01") == (a / "TK-01_montaj_rehberi.html").read_text(encoding="utf-8")
    yazi = _yazilar(b)
    assert "PORTMAN | Bayi TK-98 / TK-98 | İş Emri" in yazi["TK-98_cizim.dxf"] and "'TK-98 " in yazi["TK-98_model.ifc"]
    for ad, t in yazi.items():
        t = re.sub(r"`[0-9a-f]{64}`", "", t)
        assert not any(x in t for x in ("Sahteadi", "Sahtebayi", "XBAYIX9", no[-9:], no[-9:].replace(" ", ""))), ad
        assert ad.endswith(".ifc") or numara.maskele(t) == t, ad          # IFC'nin uzun sayıları (kimlik, koordinat) numara değil


def test_is_emri_kimlik_ve_frenler(tmp_path):
    """Model satırı İş Emri'nden; kayıttaki teklif projesinde satır ve tavan seçenekleri kayıttan. Tek ev, kod deseni, kayıttaki açık ya da
    kör model, bilinmeyen tavan ve `--is-emri`'siz `--tavan` durdurur; klasöre bir şey yazılmaz."""
    yol = KOK / _model("TK-01")["is_emri"]
    k = U.kayit()
    m, ie = U.talep_modeli("TK-97", yol, None, k)
    assert (m["kume"], m["m2"], m["govde"], m["adim"], m["tavan_secenekleri"]) == ("teklif", "80", "dış ölçü 898 × 888", "teklif", None)
    assert ie["talep"] == {"bayi": {"kod": "TK97", "ad": "Bayi TK-97"}, "musteri_adi": "TK-97"}
    m, _ = U.talep_modeli("TK-01", yol, None, k)
    assert m["tavan_secenekleri"] == ["asma", "ters"] and m["veranda"] == _model("TK-01")["veranda"]
    assert U.talep_modeli("TK-97", yol, ["ters"], k)[0]["tavan_secenekleri"] == ["ters"]
    for argv in (["hepsi", "--is-emri", str(yol)], ["AP-19", "--is-emri", str(yol)], ["tk2", "--is-emri", str(yol)], ["TK-2", "--is-emri", str(yol)],
                 ["TK-97", "--is-emri", str(yol), "--tavan", "alçı"], ["AP-19", "--tavan", "asma"], ["TK-97"]):
        with pytest.raises(SystemExit):
            U.main([argv[0], str(tmp_path / "M"), *argv[1:]])
    assert not (tmp_path / "M").exists()

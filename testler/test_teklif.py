"""İE-13 PR-2 · Teklif (`araclar/teklif.py`): İş Emri → Motor'un listesi (`motor/prefabrik_liste.is_emrinden`) → fiyat defteri → Excel ve JSON.

TK-01 (İş Emri r05, `testler/teklif/TK01_is_emri_r05.json`; Patron 28 Ekim 06:29 düzeltmesiyle): iki tavan seçeneğinin toplamı ve satır
sayıları; liste paket üreteciyle ortak girişten (kopya yok); adet düzeltmesi yok (D-28); Excel düzeni Patron'un 28 Ekim'deki isteği (grup
başlık satırı ve ara toplam, tek Adet sütunu, birim olan ölçü boş); teklif koşulları; aynı girdi aynı JSON. Hedef örnek (D-28 madde 2)
OM-01 düzelip Patron teklifi onaylayınca `testler/referans/`'a dondurulur (CIO, 28 Ekim 06:58); o zamana kadar bu test toplamları sınar.
"""
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

KOK = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KOK))
from araclar import fiyatla as F  # noqa: E402
from araclar import teklif as TK  # noqa: E402
from motor import prefabrik_liste as PL  # noqa: E402

IE = KOK / "testler" / "teklif" / "TK01_is_emri_r05.json"
GRUPLAR = ["PANEL", "METAL AKSAM", "KAPI VE PENCERE", "TAVAN", "MAKAS", "ÇATI", "ELEKTRİK", "VİDA", "DEPO", "SIHHİ TESİSAT", "DERE", "BOYA"]
SUTUN = ["Ölçü", "Birim", "Adet", "Birim fiyat", "Tutar", "Not"]


def _oku(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def cikti(tmp_path_factory):
    d = tmp_path_factory.mktemp("teklif")
    assert TK.main([str(IE), str(d), "--tavan", "asma,ters"]) == 0
    return d


@pytest.fixture(scope="module")
def T(cikti):
    return _oku(cikti / "TK01_teklif.json")


def test_tk01_toplamlar(T):
    """İş Emri r05, defter 2026-10-27-r02: teklif çalışmasının r06'sıyla aynı (asma 419.771,13, ters 418.500,13), adet düzeltmesi olmadan.
    Vitrifiye ıslak hacim başına bir takım (banyo ve E. banyo: 2; Lokman 27 Ekim 22:00), 9 satır fiyatsız (defterde fiyatı yok)."""
    assert T["kod"] == "TK-01" and T["is_emri_no"] == "PM-TK01-80M2" and T["revizyon"] == 5
    assert T["defter_surumu"] == "2026-10-27-r02" and not T["oneri_dahil"]
    ozet = {s["ad"]: s["fiyat"]["ozet"] for s in T["secenekler"]}
    assert list(ozet) == ["Asma tavan", "Ters tavan"]
    assert ozet["Asma tavan"]["toplam"] == 419771.13 and ozet["Ters tavan"]["toplam"] == 418500.13
    assert (ozet["Asma tavan"]["satir"], ozet["Asma tavan"]["fiyatli"], len(ozet["Asma tavan"]["fiyatsiz"])) == (131, 122, 9)
    assert (ozet["Ters tavan"]["satir"], ozet["Ters tavan"]["fiyatli"], len(ozet["Ters tavan"]["fiyatsiz"])) == (129, 120, 9)
    vit = [x for x in T["secenekler"][0]["fiyat"]["satirlar"] if x["kalem"].startswith("Vitrifiye")]
    assert [(x["adet"], x["birim_fiyat"], x["tutar"]) for x in vit] == [(2, 12500, 25000)]
    for s in T["secenekler"]:
        top = 0.0
        for x in s["fiyat"]["satirlar"]:
            if x["tutar"] is not None:
                assert x["tutar"] == round(x["adet"] * x["birim_fiyat"], 2)
                top += x["tutar"]
        assert s["fiyat"]["ozet"]["toplam"] == round(top, 2)


def test_liste_ortak_giristen(T):
    """Teklifin listesi paket üreteciyle ortak girişin (`is_emrinden`) listesiyle satır satır aynı; seçenek yalnız tavanı değiştirir."""
    ie = _oku(IE)
    for s, tavan in zip(T["secenekler"], ("Asma", "Ters")):
        L = PL.is_emrinden(ie, ozellikler={"Tavan": tavan})
        assert [(x["kalem"], x["olcu"], x["adet"]) for x in L["satirlar"]] == [(x["kalem"], x["olcu"], x["adet"]) for x in s["fiyat"]["satirlar"]]
        assert s["liste"]["ozet"] == L["ozet"] and s["liste"]["girdi_ozeti"] == L["girdi_ozeti"] and L["girdi_ozeti"]["tavan"] == tavan.lower()


def test_is_emrinden_kopya_ve_komut_satiri(tmp_path):
    """`ozellikler` İş Emri'ni değiştirmez; değişikliksiz çağrı komut satırının listesiyle aynı."""
    ie = _oku(IE); once = json.dumps(ie, sort_keys=True)
    L = PL.is_emrinden(ie, ozellikler={"Tavan": "Ters"})
    assert json.dumps(ie, sort_keys=True) == once and L["girdi_ozeti"]["tavan"] == "ters"
    assert PL.main([str(IE), str(tmp_path)]) == 0
    K = _oku(tmp_path / "TK01_liste.json")
    assert [(x["kalem"], x["olcu"], x["adet"]) for x in K["satirlar"]] == [(x["kalem"], x["olcu"], x["adet"]) for x in PL.is_emrinden(ie)["satirlar"]]


def test_liste_yazilmaz(tmp_path, capsys):
    """Cumba ve kırma çatıda liste yazılmaz (İE-17): ortak giriş ayırt edilebilir istisna atar, teklif çıkış 1 ve dosya yazmaz."""
    from motor.cati_prefabrik import TIPLER
    ie = _oku(IE); ie["cati"]["tip"] = sorted(TIPLER)[0]
    with pytest.raises(PL.ListeYazilmaz) as e:
        PL.is_emrinden(ie)
    assert "liste yazılmaz" in e.value.neden
    (tmp_path / "ie.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    assert TK.main([str(tmp_path / "ie.json"), str(tmp_path / "c")]) == 1 and not (tmp_path / "c").exists()
    assert "liste yazılmaz" in capsys.readouterr().err


def test_duzeltme_yok_ve_secenek_denetimi(tmp_path):
    """D-28: teklif adet düzeltmesi almaz; tavan seçeneği bilinmeli ve bir kez verilmeli."""
    with pytest.raises(SystemExit):
        TK.main([str(IE), str(tmp_path), "--duzeltme", "d.json"])
    for t in ("asma,asma", "asma,kubbe"):
        with pytest.raises(SystemExit):
            TK.main([str(IE), str(tmp_path), "--tavan", t])
    tek = TK.secenekler(_oku(IE), None)
    assert tek == [("Asma tavan", {})]


def test_excel_duzeni(cikti, T):
    """Patron 28 Ekim 06:40–06:41: grup adı sütun değil başlık satırı, altında ara toplam; tek Adet sütunu; birim olan ölçü boş."""
    wb = load_workbook(cikti / "TK01_teklif.xlsx")
    assert wb.sheetnames == ["Özet", "Asma tavan", "Ters tavan", "Fiyat kaynağı"]
    for s in T["secenekler"]:
        ws = wb[s["ad"]]
        rows = list(ws.iter_rows(values_only=True))
        degerler = {v for r in rows for v in r if isinstance(v, str)}
        assert not degerler & {"Grup", "Sıra", "Motor adedi", "Teklif adedi", "Düzeltme"}
        basliklar = [i for i, r in enumerate(rows) if r[0] in GRUPLAR]
        assert [rows[i][0] for i in basliklar] == GRUPLAR
        for i in basliklar:
            assert list(rows[i][1:7]) == SUTUN and ws.cell(i + 1, 1).font.b
        satirlar = iter(s["fiyat"]["satirlar"])
        for g in basliklar:
            son = next(j for j in range(g + 1, len(rows)) if isinstance(rows[j][0], str) and rows[j][0].endswith(" toplamı"))
            ara = 0.0
            for j in range(g + 1, son):
                x = next(satirlar)
                assert rows[j][0] == x["kalem"] and rows[j][3] == x["adet"] and rows[j][4] == x["birim_fiyat"] and rows[j][5] == x["tutar"]
                olcu, birim = TK.gorunen(x["olcu"], x["birim"])
                assert (rows[j][1] or "", rows[j][2]) == (olcu, birim)
                if not any(c.isdigit() for c in x["olcu"]):
                    assert not rows[j][1]                                       # "adet", "kova", "metre": ölçü boş, birim o
                ara += x["tutar"] or 0
            assert rows[son][5] == round(ara, 2), rows[son]
        assert next(satirlar, None) is None
        genel = next(r for r in rows if r[0] == "GENEL TOPLAM (fiyatı olan satırlar)")
        assert genel[5] == s["fiyat"]["ozet"]["toplam"] and genel[6] == f"fiyatı olmayan satır: {len(s['fiyat']['ozet']['fiyatsiz'])}"


def test_vitrifiye_banyo_basina():
    """Kural `vitrifiye_banyo_basina` (Lokman 27 Ekim, TK-01): banyo türündeki her hacme bir takım; ayrı WC ya da lavabo hacmi olan evde
    adet uydurulmaz ("Kural yok", aday banyo sayısı); alıcıya aitse 0."""
    vit = lambda ie: next(s for s in PL.is_emrinden(ie)["satirlar"] if s["kalem"].startswith("Vitrifiye"))
    ie = _oku(IE)
    assert (vit(ie)["adet"], vit(ie)["formul"]) == (2, "banyo 2 × 1")
    wc = json.loads(json.dumps(ie))
    next(m for m in wc["mahaller"] if m["ad"] == "E. BANYO")["tur"] = "wc"
    s = vit(wc)
    assert (s["adet"], s["dayanak"], s["aday"]) == (None, "Kural yok", 1) and "wc" in s["formul"]
    al = json.loads(json.dumps(ie)); al["alici_ait"] = {"kalemler": ["vitrifiye"]}
    assert (vit(al)["adet"], vit(al)["dayanak"]) == (0, "İş Emri")


def test_ozet_ve_kosullar(cikti, T, tmp_path):
    """Koşul değeri yoksa sarı "Patron'dan"; projeye özel dosya verilince değer yazılır, sarı kalkar."""
    oz = load_workbook(cikti / "TK01_teklif.xlsx")["Özet"]
    rows = {r[0].value: r for r in oz.iter_rows() if r[0].value}
    for k in ("Teslim süresi", "Ödeme şekli"):
        assert rows[k][1].value is None and rows[k][1].fill.fgColor.rgb.endswith("FFF2A8") and rows[k][3].value == "Patron'dan"
    assert rows["Asma tavan"][1].value == 419771.13 and rows["Ters tavan"][1].value == 418500.13
    fiyatsiz = [r[0].value for r in oz.iter_rows() if r[0].fill.fgColor.rgb.endswith("E7E7E7") and r[0].value]
    assert len(fiyatsiz) == 9 and not any(v.startswith("Vitrifiye") for v in fiyatsiz) and any(v.startswith("Makas (baş)") for v in fiyatsiz)
    (tmp_path / "k.json").write_text(json.dumps({"kosullar": {"Teslim süresi": "DENEME 1", "Ödeme şekli": "DENEME 2"}}, ensure_ascii=False), encoding="utf-8")
    assert TK.main([str(IE), str(tmp_path / "c"), "--kosullar", str(tmp_path / "k.json")]) == 0
    J = _oku(tmp_path / "c" / "TK01_teklif.json")
    assert J["kosullar"] == {"Teslim süresi": "DENEME 1", "Ödeme şekli": "DENEME 2"} and [s["ad"] for s in J["secenekler"]] == ["Asma tavan"]
    oz2 = load_workbook(tmp_path / "c" / "TK01_teklif.xlsx")["Özet"]
    r = next(r for r in oz2.iter_rows() if r[0].value == "Teslim süresi")
    assert r[1].value == "DENEME 1" and not r[3].value


def test_is_emrindeki_teklif_bolumu(tmp_path):
    """D-28: proje kararı İş Emri'nde. `teklif.tavan` seçenekleri, `teklif.kosullar` koşulları verir; komut satırı üstüne yazar; bilinmeyen
    koşul adı durur."""
    ie = _oku(IE)
    ie["teklif"] = {"tavan": ["asma", "ters"], "kosullar": {"Ödeme şekli": "DENEME İE"}}
    (tmp_path / "ie.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    assert TK.main([str(tmp_path / "ie.json"), str(tmp_path / "a")]) == 0
    J = _oku(tmp_path / "a" / "TK01_teklif.json")
    assert [s["ad"] for s in J["secenekler"]] == ["Asma tavan", "Ters tavan"]
    assert [s["fiyat"]["ozet"]["toplam"] for s in J["secenekler"]] == [419771.13, 418500.13]      # Motor teklif bölümünü okumaz
    assert J["kosullar"] == {"Teslim süresi": None, "Ödeme şekli": "DENEME İE"}
    (tmp_path / "k.json").write_text(json.dumps({"kosullar": {"Ödeme şekli": "DENEME DOSYA"}}, ensure_ascii=False), encoding="utf-8")
    assert TK.main([str(tmp_path / "ie.json"), str(tmp_path / "b"), "--tavan", "ters", "--kosullar", str(tmp_path / "k.json")]) == 0
    J = _oku(tmp_path / "b" / "TK01_teklif.json")
    assert [s["ad"] for s in J["secenekler"]] == ["Ters tavan"] and J["kosullar"]["Ödeme şekli"] == "DENEME DOSYA"
    ie["teklif"]["kosullar"] = {"Teslim suresi": "x"}
    (tmp_path / "ie.json").write_text(json.dumps(ie, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(SystemExit):
        TK.main([str(tmp_path / "ie.json"), str(tmp_path / "c")])


def test_fiyat_kaynagi(cikti, T):
    """Fiyat kaynağı sayfası her anahtarı bir kez, defterdeki kaydıyla yazar (bant: defterin bandı)."""
    ws = load_workbook(cikti / "TK01_teklif.xlsx")["Fiyat kaynağı"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    anahtarlar = {F.anahtar(x, "prefabrik") for s in T["secenekler"] for x in s["fiyat"]["satirlar"]}
    assert len(rows) == len(anahtarlar) == len({f"{r[0]} | {r[1]}" for r in rows})
    r = next(r for r in rows if r[0] == "60'lık karkaslı + kasalı pano" and r[1] == "60 × 1195 × 2500")
    assert (r[4], r[5]) == ("bant", 2550) and "bandı" in r[-1]


def test_ayni_girdi_ayni_json(T, tmp_path):
    assert TK.main([str(IE), str(tmp_path), "--tavan", "asma,ters"]) == 0
    assert _oku(tmp_path / "TK01_teklif.json") == T


def test_fiyatla_excel_ayni_yazici(tmp_path):
    """`fiyatla.py --excel` liste sayfasını teklifin yazıcısıyla yazar; Durum ve Kaynak sütunlarıyla."""
    L = PL.is_emrinden(_oku(IE))
    (tmp_path / "l.json").write_text(json.dumps(L, ensure_ascii=False), encoding="utf-8")
    assert F.main([str(tmp_path / "l.json"), "--excel", str(tmp_path / "f.xlsx")]) == 0
    rows = list(load_workbook(tmp_path / "f.xlsx").active.iter_rows(values_only=True))
    assert ("PANEL", *SUTUN, "Durum", "Kaynak") in [tuple(r[:9]) for r in rows]
